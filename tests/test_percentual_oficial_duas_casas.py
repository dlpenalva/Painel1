"""REGRA PETREA — percentual oficial do ciclo fechado em 2 casas.

Caso real (C1/ICTI): variacao bruta 4,052187881255853% -> percentual oficial
4,05% -> fator proprio oficial 1,0405. Depois do fechamento, nenhum calculo
financeiro pode usar o fator bruto 1,0405218788125585; a variacao bruta
sobrevive apenas como memoria tecnica da apuracao.

Multiciclo: cada ciclo e fechado individualmente e o acumulado e o PRODUTO dos
fatores oficiais — nunca soma de percentuais nem composicao de fatores brutos.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from _coleta_oficial import _percentual, gerar_coleta_oficial_preenchida  # noqa: E402
from _coleta_reajuste import _percentual_ciclo  # noqa: E402
from _leitor_masterfile_v10 import ler_masterfile_v10  # noqa: E402
from _memoria_calculo import (  # noqa: E402
    TIPO_VARIACAO_BRUTA,
    ler_memoria_calculo,
    normalizar_memoria_calculo,
)
from _objeto_processo_reajuste import _fatores_acumulados_por_percentual  # noqa: E402
from _reajuste_utils import (  # noqa: E402
    APLICAR_VARIACAO_NEGATIVA,
    NEUTRALIZAR_VARIACAO_NEGATIVA,
    fator_acumulado_oficial,
    fator_oficial,
    fechar_percentual_oficial,
    resolver_tratamento_variacao_negativa,
)
from _templates_documentos import _extrair_dados  # noqa: E402
from test_reajuste_negativo import PAGINA_MULTIPLA, PAGINA_SIMPLES  # noqa: E402

RAW_C1 = 0.04052187881255853
FATOR_BRUTO_C1 = 1.0405218788125585
RAW_C2 = 0.04496
RAW_C3 = 0.03814

# (item, VU C0, VU esperado com 4,05%, VU antigo com o fator bruto)
ITENS_CASO_REAL = (
    ("1", 1043304.26, 1085558.08, 1085580.91),
    ("2", 737983.34, 767871.67, 767887.81),
    ("3", 31551.30, 32829.13, 32829.82),
    ("10", 196942.70, 204918.88, 204923.19),
    ("11", 24813.13, 25818.06, 25818.60),
    ("12", 31551.30, 32829.13, 32829.82),
)


def _round2(valor: float) -> float:
    """ROUND(x, 2) do template/aplicacao (meio para cima)."""
    return float(Decimal(repr(valor)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


# ------------------------------------------------------------------ unitarios

def test_fechamento_do_percentual_do_caso_real():
    assert fechar_percentual_oficial(RAW_C1) == 0.0405
    assert fator_oficial(RAW_C1) == 1.0405


def test_fator_oficial_nao_e_o_fator_bruto():
    assert 1.0 + RAW_C1 == FATOR_BRUTO_C1
    assert fator_oficial(RAW_C1) != FATOR_BRUTO_C1
    resolvido = resolver_tratamento_variacao_negativa(RAW_C1)
    assert resolvido["percentual_indice"] == RAW_C1  # memoria bruta preservada
    assert resolvido["percentual_aplicado"] == 0.0405
    assert resolvido["fator"] == 1.0405


@pytest.mark.parametrize(
    ("bruto", "oficial"),
    (
        (0.04045, 0.0405),        # meio para cima em pontos percentuais
        (0.040449999, 0.0404),
        (0.0375, 0.0375),         # ja oficial: idempotente
        (0.0, 0.0),
        (-0.00004, 0.0),          # sem -0,00%
        (-0.020349, -0.0203),
        (0.04496, 0.045),
        (0.03814, 0.0381),
    ),
)
def test_regra_decimal_meio_para_cima_e_idempotente(bruto, oficial):
    assert fechar_percentual_oficial(bruto) == oficial
    assert fechar_percentual_oficial(fechar_percentual_oficial(bruto)) == oficial
    assert str(fator_oficial(bruto)) == str(Decimal(1) + Decimal(repr(oficial)))


def test_ausencia_permanece_ausencia():
    assert fechar_percentual_oficial(None) is None
    assert fator_oficial(None) is None
    assert fator_acumulado_oficial([RAW_C1, None]) is None


def test_negativo_preserva_os_tres_estados_e_fecha_depois():
    pendente = resolver_tratamento_variacao_negativa(-0.020349)
    assert pendente["pendente"] is True
    assert pendente["percentual_aplicado"] is None
    assert pendente["fator"] is None

    neutro = resolver_tratamento_variacao_negativa(
        -0.020349, NEUTRALIZAR_VARIACAO_NEGATIVA
    )
    assert neutro["percentual_aplicado"] == 0.0
    assert neutro["fator"] == 1.0

    aplicado = resolver_tratamento_variacao_negativa(
        -0.020349, APLICAR_VARIACAO_NEGATIVA
    )
    assert aplicado["percentual_indice"] == -0.020349
    assert aplicado["percentual_aplicado"] == -0.0203
    assert aplicado["fator"] == 0.9797


def test_multiciclo_fecha_cada_ciclo_e_compoe_fatores_oficiais():
    acumulado = fator_acumulado_oficial([RAW_C1, RAW_C2, RAW_C3])
    assert acumulado == 1.0405 * 1.045 * 1.0381


def test_multiciclo_proibe_soma_de_percentuais():
    acumulado = fator_acumulado_oficial([RAW_C1, RAW_C2, RAW_C3])
    soma = 1.0 + 0.0405 + 0.045 + 0.0381
    assert abs(acumulado - soma) > 1e-3


def test_multiciclo_proibe_fatores_brutos_e_acumulado_fechado():
    acumulado = fator_acumulado_oficial([RAW_C1, RAW_C2, RAW_C3])
    brutos = (1 + RAW_C1) * (1 + RAW_C2) * (1 + RAW_C3)
    assert abs(acumulado - brutos) > 1e-6
    # fechar o ACUMULADO e reaplicar tambem e proibido
    acumulado_fechado = fator_oficial(brutos - 1)
    assert abs(acumulado - acumulado_fechado) > 1e-6


@pytest.mark.parametrize(("item", "vu_c0", "esperado", "antigo"), ITENS_CASO_REAL)
def test_vu_c1_do_caso_real_com_4_05(item, vu_c0, esperado, antigo):
    fator = resolver_tratamento_variacao_negativa(RAW_C1)["fator"]
    assert _round2(vu_c0 * fator) == esperado
    assert _round2(vu_c0 * FATOR_BRUTO_C1) == antigo  # o valor que sai de cena
    assert esperado != antigo


# ------------------------------------------------------- fronteira payload/XLS

def test_fronteira_do_xls_fecha_payload_so_com_variacao_bruta():
    assert _percentual({"percentual_indice": RAW_C1}) == 0.0405
    assert _percentual({"fator": FATOR_BRUTO_C1}) == 0.0405
    assert _percentual({"percentual_aplicado": 4.052187881255853}) == 0.0405
    assert _percentual_ciclo({"percentual_indice": RAW_C1}) == 0.0405
    assert _percentual_ciclo({"fator": FATOR_BRUTO_C1}) == 0.0405


def test_consumidor_do_objeto_compoe_sobre_percentuais_oficiais():
    fatores = _fatores_acumulados_por_percentual([
        {"ciclo": "C1", "indice_percentual": 0.0405},
        {"ciclo": "C2", "indice_percentual": 0.045},
    ])
    assert fatores["C1"] == 1.0405
    assert fatores["C2"] == pytest.approx(1.0405 * 1.045, abs=1e-15)


# --------------------------------------------------------------- memoria

def _res_icti(variacao: float) -> dict:
    import pandas as pd

    return {
        "variacao": variacao,
        "var": variacao,
        "metodo": "ICTI",
        "dados": pd.DataFrame({
            "data": [pd.Timestamp("2024-08-01")],
            "valor": [variacao * 100],
            "fator_mensal": [1 + variacao],
            "fator_acumulado_progressivo": [1 + variacao],
        }),
    }


def test_memoria_preserva_bruto_e_resultado_e_oficial():
    oficial = resolver_tratamento_variacao_negativa(RAW_C1)
    memoria = normalizar_memoria_calculo(
        _res_icti(RAW_C1), oficial["fator"], oficial["percentual_aplicado"]
    )
    assert [m["tipo"] for m in memoria] == ["MES", TIPO_VARIACAO_BRUTA, "RESULTADO"]
    bruta, resultado = memoria[-2], memoria[-1]
    assert bruta["variacao_final"] == RAW_C1
    assert bruta["fator_acumulado"] == FATOR_BRUTO_C1
    assert "não alimenta" in bruta["metodo_fonte"]
    assert resultado["variacao_final"] == 0.0405
    assert resultado["fator_acumulado"] == 1.0405
    assert [m["ordem"] for m in memoria] == [1, 2, 3]


def test_memoria_sem_linha_bruta_quando_nada_foi_fechado_ou_ha_tratamento():
    # ja oficial: nada a declarar
    memoria = normalizar_memoria_calculo(_res_icti(0.05), 1.05, 0.05)
    assert [m["tipo"] for m in memoria] == ["MES", "RESULTADO"]
    # precluso/neutralizado: o 0,00% nao e fechamento do bruto
    memoria = normalizar_memoria_calculo(_res_icti(RAW_C1), 1.0, 0.0)
    assert [m["tipo"] for m in memoria] == ["MES", "RESULTADO"]


# --------------------------------------------------------------- payloads

def _ciclo(nome: str, ano: int, bruto: float) -> dict:
    resolvido = resolver_tratamento_variacao_negativa(bruto)
    return {
        "ciclo": nome,
        "data_base": f"01/08/{ano}",
        "periodo_inicio": f"01/08/{ano}",
        "periodo_fim": f"31/07/{ano + 1}",
        "financeiro_inicio": f"01/08/{ano + 1}",
        "percentual_indice": bruto,
        "percentual_aplicado": resolvido["percentual_aplicado"],
        "variacao": resolvido["percentual_aplicado"],
        "fator": resolvido["fator"],
        "situacao": "TEMPESTIVO",
        "situacao_aplicada": "TEMPESTIVO",
        "memoria_calculo": normalizar_memoria_calculo(
            _res_icti(bruto), resolvido["fator"], resolvido["percentual_aplicado"]
        ),
    }


def _dados(*brutos: float) -> dict:
    ciclos = [
        _ciclo(f"C{i}", 2024 + i - 1, bruto) for i, bruto in enumerate(brutos, start=1)
    ]
    acumulado = fator_acumulado_oficial(brutos)
    fator = 1.0
    for ciclo in ciclos:
        fator *= ciclo["fator"]
        ciclo["fator_acumulado"] = fator
    return {
        "origem": "Reajuste Simples" if len(brutos) == 1 else "Reajuste Represado",
        "indice": "ICTI",
        "data_base_original": "01/08/2024",
        "fator": acumulado,
        "fator_acumulado": acumulado,
        "variacao": acumulado - 1,
        "variacao_acumulada": acumulado - 1,
        "ciclos": ciclos,
    }


@pytest.fixture(scope="module")
def coleta_c1() -> bytes:
    return gerar_coleta_oficial_preenchida(_dados(RAW_C1))


@pytest.fixture(scope="module")
def coleta_multiciclo() -> bytes:
    return gerar_coleta_oficial_preenchida(_dados(RAW_C1, RAW_C2, RAW_C3))


def test_xlsx_parametros_recebe_percentual_oficial(coleta_c1):
    wb = load_workbook(io.BytesIO(coleta_c1))
    par = wb["parametros"]
    assert par["E3"].value == 0.0405
    # a cadeia do XLS continua 100% formula sobre parametros!E
    assert par["F3"].value.startswith("=") and "(1+E3)" in par["F3"].value
    assert par["D12"].value == '=IF(C12="","",1+C12)'
    assert wb["historico_VU"]["K3"].value.startswith("=")


def test_xlsx_legado_so_com_bruto_tambem_sai_oficial():
    dados = _dados(RAW_C1)
    for chave in ("percentual_aplicado", "variacao", "fator", "fator_acumulado"):
        dados["ciclos"][0].pop(chave)
    wb = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(dados)))
    assert wb["parametros"]["E3"].value == 0.0405


def test_xlsx_memoria_demonstra_bruto_e_oficial(coleta_c1):
    memoria = ler_memoria_calculo(load_workbook(io.BytesIO(coleta_c1))["parametros"])
    tipos = [r["tipo"] for r in memoria["C1"]]
    assert tipos[-2:] == [TIPO_VARIACAO_BRUTA, "RESULTADO"]
    assert memoria["C1"][-2]["variacao_final"] == RAW_C1
    assert memoria["C1"][-1]["variacao_final"] == 0.0405
    assert memoria["C1"][-1]["fator_acumulado"] == 1.0405


def test_xlsx_multiciclo_grava_cada_ciclo_fechado(coleta_multiciclo):
    par = load_workbook(io.BytesIO(coleta_multiciclo))["parametros"]
    assert [par[f"E{r}"].value for r in (3, 4, 5)] == [0.0405, 0.045, 0.0381]


def _com_pc(conteudo: bytes, data_pc: datetime, valor: float) -> bytes:
    wb = load_workbook(io.BytesIO(conteudo))
    pc = wb["itens_PC"]
    pc["A2"], pc["B2"], pc["D2"], pc["G2"] = "PC-0001", data_pc, valor, "Sim"
    wb["CONTROLE"]["B1"] = "PCs"
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _pc_lido(conteudo: bytes) -> dict:
    leitura = ler_masterfile_v10(conteudo)
    return next(
        p for p in leitura["itens_pc_v10"]["itens"] if p.get("numero_pc") == "PC-0001"
    )


def test_leitor_e_pc_usam_percentual_oficial(coleta_c1):
    leitura = ler_masterfile_v10(coleta_c1)
    c1 = leitura["parametros_v10"]["por_ciclo"]["C1"]
    assert c1["percentual_reajuste"] == 0.0405
    pc = _pc_lido(_com_pc(coleta_c1, datetime(2025, 9, 15), 737983.34))
    assert pc["ciclo"] == "C1"
    assert pc["fator_acumulado"] == 1.0405
    assert pc["valor_atualizado"] == 767871.67


def test_pc_multiciclo_infere_produto_dos_fatores_oficiais(coleta_multiciclo):
    pc = _pc_lido(_com_pc(coleta_multiciclo, datetime(2026, 9, 15), 1000.0))
    assert pc["ciclo"] == "C2"
    assert pc["fator_acumulado"] == pytest.approx(1.0405 * 1.045, abs=1e-12)
    assert pc["valor_atualizado"] == _round2(1000.0 * 1.0405 * 1.045)
    assert pc["valor_atualizado"] != _round2(1000.0 * (1 + RAW_C1) * (1 + RAW_C2))


def test_documentos_recebem_o_mesmo_percentual_oficial(coleta_c1):
    dados = _extrair_dados(ler_masterfile_v10(coleta_c1), None)
    ciclo = next(c for c in dados["ciclos"] if c["ciclo"] == "C1")
    assert ciclo["percentual_reajuste"] == 0.0405


# ------------------------------------------------------ calculadoras (origem)

def _wrapper_bruto(pagina: Path, por_ano: dict[int, float]) -> str:
    return f'''
import runpy
import sys

import pandas as pd

sys.path.insert(0, {str(ROOT)!r})
import _indice_utils

POR_ANO = {por_ano!r}

def _ist_bruto(data_inicio, *args, **kwargs):
    inicio = pd.Timestamp(data_inicio).replace(day=1)
    fim = inicio + pd.DateOffset(months=12)
    variacao = POR_ANO[inicio.year]
    return {{
        "variacao": variacao, "var": variacao,
        "i_ini": 100.0, "i_fim": 100.0 * (1.0 + variacao),
        "d_ini": inicio, "d_fim": fim, "p_ini": inicio, "p_fim": fim,
        "metodo": "IST determinístico de teste", "serie": "IST-TESTE",
        "dados": pd.DataFrame({{"data": [inicio, fim],
                               "indice": [100.0, 100.0 * (1.0 + variacao)]}}),
    }}

_original = _indice_utils.calcular_ist_numero_indice
try:
    _indice_utils.calcular_ist_numero_indice = _ist_bruto
    runpy.run_path({str(pagina)!r}, run_name="__main__")
finally:
    _indice_utils.calcular_ist_numero_indice = _original
'''


def _e_da_coleta(at) -> list:
    """parametros!E3:E6 da Coleta gerada a partir do estado da Calculadora."""
    assert at.get("download_button"), "botao da Coleta ausente"
    conteudo = at.session_state["dados_admissibilidade"]
    par = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(conteudo)))["parametros"]
    return [par[f"E{r}"].value for r in range(3, 7)]


def test_calculadora_simples_fecha_na_origem():
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_string(
        _wrapper_bruto(PAGINA_SIMPLES, {2023: RAW_C1, 2024: RAW_C1}),
        default_timeout=180,
    ).run()
    at.date_input[0].set_value(date(2023, 2, 1))
    at.run()
    at.date_input[1].set_value(date(2024, 2, 1))
    at.run()
    next(b for b in at.button if "Processar" in str(b.label)).click()
    at.run()

    adm = at.session_state["dados_admissibilidade"]
    ciclo = adm["ciclos"][0]
    assert ciclo["percentual_indice"] == RAW_C1
    assert ciclo["percentual_aplicado"] == 0.0405
    assert ciclo["variacao"] == 0.0405
    assert ciclo["fator"] == 1.0405
    assert adm["fator"] == adm["fator_acumulado"] == 1.0405
    tipos = [r["tipo"] for r in ciclo["memoria_calculo"]]
    assert tipos[-2:] == [TIPO_VARIACAO_BRUTA, "RESULTADO"]
    assert _e_da_coleta(at)[0] == 0.0405


def test_calculadora_multiciclo_fecha_cada_ciclo_antes_de_compor():
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_string(
        _wrapper_bruto(PAGINA_MULTIPLA, {2022: RAW_C1, 2023: RAW_C2, 2024: RAW_C3}),
        default_timeout=180,
    ).run()
    at.selectbox(key="rep_ciclo_final_analise").select("C3")
    at.run()
    next(b for b in at.button if "Processar" in str(b.label)).click()
    at.run()

    adm = at.session_state["dados_admissibilidade"]
    ciclos = {c["ciclo"]: c for c in adm["ciclos"]}
    assert [ciclos[n]["percentual_aplicado"] for n in ("C1", "C2", "C3")] == [
        0.0405, 0.045, 0.0381
    ]
    assert [ciclos[n]["fator"] for n in ("C1", "C2", "C3")] == [1.0405, 1.045, 1.0381]
    assert [ciclos[n]["percentual_indice"] for n in ("C1", "C2", "C3")] == [
        RAW_C1, RAW_C2, RAW_C3
    ]
    oficial = 1.0405 * 1.045 * 1.0381
    assert ciclos["C3"]["fator_acumulado"] == pytest.approx(oficial, abs=1e-15)
    assert adm["fator_acumulado"] == pytest.approx(oficial, abs=1e-15)
    bruto = (1 + RAW_C1) * (1 + RAW_C2) * (1 + RAW_C3)
    assert abs(adm["fator_acumulado"] - bruto) > 1e-6
    assert _e_da_coleta(at)[:3] == [0.0405, 0.045, 0.0381]


# ------------------------------------------------ XLSX real recalculado (Excel)

@pytest.fixture(scope="module")
def excel():
    client = pytest.importorskip("win32com.client")
    pythoncom = pytest.importorskip("pythoncom")
    pythoncom.CoInitialize()
    app = client.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    yield app
    app.Quit()
    del app
    pythoncom.CoUninitialize()


def _cenario_excel(conteudo: bytes, destino: Path) -> Path:
    wb = load_workbook(io.BytesIO(conteudo))
    rem = wb["itens_Remanesc"]
    for linha, (item, vu_c0, _esperado, _antigo) in enumerate(ITENS_CASO_REAL, start=2):
        rem[f"A{linha}"], rem[f"B{linha}"], rem[f"C{linha}"] = item, 10, vu_c0
        rem[f"E{linha}"] = 4
    fin = wb["financeiro"]
    for linha in range(2, 80):
        competencia = fin[f"A{linha}"].value
        if isinstance(competencia, datetime) and competencia <= datetime(2026, 7, 1):
            fin[f"C{linha}"] = 187345.67
    wb["CONTROLE"]["B1"] = "Financeiro"
    wb.save(destino)
    return destino


@pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar Excel COM",
)
def test_xlsx_real_recalculado_propaga_4_05_por_toda_a_cadeia(
    excel, coleta_c1, tmp_path
):
    caminho = _cenario_excel(coleta_c1, tmp_path / "c1_oficial.xlsx")
    book = excel.Workbooks.Open(str(caminho), UpdateLinks=0, ReadOnly=True, CorruptLoad=0)
    try:
        excel.CalculateFullRebuild()

        def v(aba, celula):
            return book.Worksheets(aba).Range(celula).Value

        assert v("parametros", "E3") == 0.0405
        assert v("parametros", "F3") == 1.0405
        assert v("parametros", "D12") == 1.0405
        assert v("RESULTADOS", "H5") == 1.0405
        for linha, (item, _vu, esperado, antigo) in enumerate(ITENS_CASO_REAL, start=2):
            assert str(v("historico_VU", f"A{linha}")) == item
            assert v("historico_VU", f"D{linha}") == esperado        # VU_C1
            assert v("historico_VU", f"D{linha}") != antigo
            assert v("itens_RC", f"E{linha + 1}") == esperado        # mesmo VU canonico
            assert v("itens_Remanesc", f"F{linha}") == _round2(4 * esperado)
        # retroativo C1 (financeiro) e remanescente atualizado com 4,05%
        assert float(v("RESULTADOS", "D17")) == _round2(12 * _round2(187345.67 * 0.0405))
        assert float(v("RESULTADOS", "C27")) == _round2(
            sum(_round2(4 * esperado) for *_x, esperado, _a in ITENS_CASO_REAL)
        )
    finally:
        book.Close(False)


# =================================================== revisao independente (P1/P2)

from _coleta_oficial import normalizar_dados_calculadora  # noqa: E402
from _coleta_reajuste import _montar_ciclos  # noqa: E402
from _reajuste_utils import (  # noqa: E402
    fator_oficial_de_fator,
    percentual_contexto_oficial,
    percentual_oficial_do_payload,
)


def _dados_c2_com_c1_so_no_contexto() -> dict:
    dados = _dados(RAW_C1, RAW_C2)
    c2 = dados["ciclos"][1]
    dados["ciclos"] = [c2]
    dados["contexto_contratual_anterior"] = {
        "ultimo_ciclo_concedido": "C1",
        "percentual_ja_aplicado_pct": RAW_C1 * 100,  # 4,052187881255853 (pontos)
    }
    return dados


def test_contexto_percentual_historico_e_fechado():
    contexto = {"percentual_ja_aplicado_pct": 4.052187881255853}
    assert percentual_contexto_oficial(contexto) == 0.0405
    assert percentual_contexto_oficial({"percentual_ja_aplicado_pct": RAW_C1}) == 0.0405
    assert percentual_contexto_oficial({}) is None


def test_contexto_c1_entra_oficial_no_montador_da_coleta_legada():
    ciclos, _alvos, _alertas = _montar_ciclos(_dados_c2_com_c1_so_no_contexto())
    por_nome = {c["nome"]: c for c in ciclos}
    assert por_nome["C1"]["percentual"] == 0.0405
    assert por_nome["C2"]["percentual"] == 0.045


def test_contexto_c1_nao_se_perde_na_coleta_oficial():
    normalizado = normalizar_dados_calculadora(_dados_c2_com_c1_so_no_contexto())
    por_nome = {c["ciclo"]: c for c in normalizado["ciclos"]}
    assert por_nome["C1"]["percentual"] == 0.0405
    assert por_nome["C2"]["percentual"] == 0.045

    conteudo = gerar_coleta_oficial_preenchida(_dados_c2_com_c1_so_no_contexto())
    par = load_workbook(io.BytesIO(conteudo))["parametros"]
    assert par["E3"].value == 0.0405
    assert par["E4"].value == 0.045
    por_ciclo = ler_masterfile_v10(conteudo)["parametros_v10"]["por_ciclo"]
    assert por_ciclo["C1"]["percentual_reajuste"] == 0.0405
    assert por_ciclo["C2"]["percentual_reajuste"] == 0.045
    from _reajuste_utils import cadeia_fatores_oficiais

    cadeia = cadeia_fatores_oficiais([0.0, 0.0405, 0.045])
    assert cadeia[2] == pytest.approx(fator_oficial(RAW_C1) * fator_oficial(RAW_C2), abs=1e-15)


PAYLOAD_NEGATIVO_PENDENTE = {
    "percentual_aplicado": None,
    "percentual_indice": -0.020349,
    "tratamento_negativo": None,
}


def test_negativo_pendente_permanece_fail_closed_nas_fronteiras():
    assert percentual_oficial_do_payload(dict(PAYLOAD_NEGATIVO_PENDENTE)) is None
    assert _percentual(dict(PAYLOAD_NEGATIVO_PENDENTE)) is None
    assert _percentual_ciclo(dict(PAYLOAD_NEGATIVO_PENDENTE)) is None
    # payload legado sem a chave percentual_aplicado: bruto negativo sem
    # tratamento tambem e pendente
    legado = {"percentual_indice": -0.020349}
    assert _percentual(legado) is None
    assert _percentual_ciclo(legado) is None
    # com tratamento aprovado, os estados existentes sao preservados
    aplicar = {**legado, "tratamento_ciclo_negativo": APLICAR_VARIACAO_NEGATIVA}
    neutralizar = {**legado, "tratamento_ciclo_negativo": NEUTRALIZAR_VARIACAO_NEGATIVA}
    assert _percentual(aplicar) == -0.0203
    assert _percentual(neutralizar) == 0.0


def test_negativo_pendente_nao_vira_percentual_na_coleta():
    dados = _dados(RAW_C1)
    dados["ciclos"][0].update(PAYLOAD_NEGATIVO_PENDENTE)
    par = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(dados)))["parametros"]
    assert par["E3"].value in (None, "")


def test_fator_de_fator_usa_decimal_e_meio_para_cima():
    assert fator_oficial_de_fator(1.04025) == 1.0403
    assert fator_oficial_de_fator(FATOR_BRUTO_C1) == 1.0405
    assert _percentual({"fator": 1.04025}) == 0.0403


# ------------------------------------------------------------- XLS legado

def _coleta_legada_bruta(conteudo: bytes) -> bytes:
    """Coleta 'antiga' ja calculada pelo Excel: E e F com precisao bruta."""
    wb = load_workbook(io.BytesIO(conteudo))
    par = wb["parametros"]
    par["E3"] = RAW_C1
    par["F2"] = 1.0
    par["F3"] = FATOR_BRUTO_C1
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def test_xls_legado_bruto_e_oficializado_na_leitura(coleta_c1):
    legado = _coleta_legada_bruta(coleta_c1)
    leitura = ler_masterfile_v10(legado)
    c1 = leitura["parametros_v10"]["por_ciclo"]["C1"]
    assert c1["percentual_reajuste"] == 0.0405
    assert c1["fator_acumulado"] == 1.0405
    # memoria tecnica preservada
    assert c1["percentual_reajuste_bruto"] == RAW_C1
    assert c1["fator_acumulado_bruto"] == pytest.approx(FATOR_BRUTO_C1, abs=1e-15)
    assert any("precisao bruta" in a for a in leitura["parametros_v10"]["alertas"])
    # o arquivo fisico nao e regravado
    par = load_workbook(io.BytesIO(legado))["parametros"]
    assert par["E3"].value == pytest.approx(RAW_C1, abs=1e-16)
    assert par["F3"].value == pytest.approx(FATOR_BRUTO_C1, abs=1e-15)


def test_xls_legado_multiciclo_reconstroi_acumulado_pelos_oficiais(coleta_multiciclo):
    wb = load_workbook(io.BytesIO(coleta_multiciclo))
    par = wb["parametros"]
    par["E3"], par["E4"], par["E5"] = RAW_C1, RAW_C2, RAW_C3
    par["F2"] = 1.0
    par["F3"] = 1 + RAW_C1
    par["F4"] = (1 + RAW_C1) * (1 + RAW_C2)
    par["F5"] = (1 + RAW_C1) * (1 + RAW_C2) * (1 + RAW_C3)
    saida = io.BytesIO()
    wb.save(saida)
    por_ciclo = ler_masterfile_v10(saida.getvalue())["parametros_v10"]["por_ciclo"]
    assert por_ciclo["C2"]["fator_acumulado"] == pytest.approx(1.0405 * 1.045, abs=1e-15)
    oficial_c3 = 1.0405 * 1.045 * 1.0381
    assert por_ciclo["C3"]["fator_acumulado"] == pytest.approx(oficial_c3, abs=1e-15)
    # nao e o fator bruto acumulado apenas arredondado
    bruto_c3 = (1 + RAW_C1) * (1 + RAW_C2) * (1 + RAW_C3)
    assert abs(por_ciclo["C3"]["fator_acumulado"] - fator_oficial_de_fator(bruto_c3)) > 1e-6


def test_xls_legado_pc_documentos_e_objeto_recebem_oficial(coleta_c1):
    legado = _com_pc(_coleta_legada_bruta(coleta_c1), datetime(2025, 9, 15), 737983.34)
    leitura = ler_masterfile_v10(legado)
    pc = next(p for p in leitura["itens_pc_v10"]["itens"] if p.get("numero_pc") == "PC-0001")
    assert pc["fator_acumulado"] == 1.0405
    assert pc["valor_atualizado"] == 767871.67
    dados = _extrair_dados(leitura, None)
    assert next(c for c in dados["ciclos"] if c["ciclo"] == "C1")["percentual_reajuste"] == 0.0405
    objeto = leitura.get("objeto_processo") or {}
    indice = (objeto.get("resultados") or {}).get("indice_acumulado") or {}
    if indice.get("fator_acumulado") is not None:
        assert indice["fator_acumulado"] == 1.0405


def test_xls_legado_adaptador_valor_global_usa_cadeia_oficial():
    from openpyxl import Workbook

    from _coleta_reajuste_documentos import _cadeia_fator_acumulado

    ws = Workbook().active
    ws["E3"], ws["E4"] = RAW_C1, RAW_C2
    cadeia = _cadeia_fator_acumulado(ws)
    assert cadeia[2] == 1.0
    assert cadeia[3] == 1.0405
    assert cadeia[4] == pytest.approx(1.0405 * 1.045, abs=1e-15)
    assert 5 not in cadeia  # E5 ausente: a cadeia para, como a formula


# ------------------------------------------------ Valor Global / Adequacao

def _funcao_da_pagina(pagina: str, nome: str, globais: dict):
    import ast

    fonte = (ROOT / "pages" / pagina).read_text(encoding="utf-8")
    arvore = ast.parse(fonte)
    no = next(n for n in arvore.body if isinstance(n, ast.FunctionDef) and n.name == nome)
    espaco = dict(globais)
    exec(compile(ast.Module(body=[no], type_ignores=[]), f"<{pagina}>", "exec"), espaco)
    return espaco[nome], fonte


def test_valor_global_fator_operacional_usa_round_half_up_canonico():
    fator_operacional, fonte = _funcao_da_pagina(
        "03_Valor_Global.py", "fator_operacional",
        {"fator_oficial_de_fator": fator_oficial_de_fator},
    )
    assert fator_operacional(1.04025) == 1.0403          # 4,025% -> 4,03%
    assert fator_operacional(FATOR_BRUTO_C1) == 1.0405
    assert round(1.04025, 4) == 1.0402                   # o comportamento antigo
    assert "round(float(valor), 4)" not in fonte
    assert "fator = 1 + perc_legacy" not in fonte


def test_adequacao_ajuste_manual_passa_pelo_helper_canonico():
    from _adequacao_ui import percentual_e_fator_da_adequacao

    assert percentual_e_fator_da_adequacao(0.1, False, "4,052187881255853%") == (0.0405, 1.0405)
    assert percentual_e_fator_da_adequacao(0.1, True, "4,025%") == (0.0403, 1.0403)
    # campo intacto: o percentual canonico importado da apuracao segue exato
    assert percentual_e_fator_da_adequacao(0.0289, True, "2,89%") == (0.0289, 1.0289)
    fonte = (ROOT / "pages" / "12_Adequacao_Orcamentaria.py").read_text(encoding="utf-8")
    assert "fator_reajuste = 1 + percentual_reajuste" not in fonte
    assert "parse_moeda_br(percentual_txt) / 100" not in fonte


# ------------------------------------------------------------- Garantia / DOU

def test_garantia_nao_forma_percentual_nem_fator_de_reajuste():
    for arquivo in (ROOT / "_garantia_calculo.py", ROOT / "pages" / "05_Garantia.py"):
        fonte = arquivo.read_text(encoding="utf-8")
        assert "_indice_utils" not in fonte
        assert "fator_oficial" not in fonte and "fator_acumulado" not in fonte
        assert "1 + percentual" not in fonte and "1.0 + percentual" not in fonte


def test_dou_so_apresenta_percentual_e_fator_da_cadeia_oficial():
    fonte = (ROOT / "pages" / "13_DOU.py").read_text(encoding="utf-8")
    assert "_indice_utils" not in fonte
    assert "1 + percentual" not in fonte and "1.0 + percentual" not in fonte
    assert "1 + pct" not in fonte and "1.0 + pct" not in fonte
    # o texto do DOU le 'Percentual aplicado' / 'Fator acumulado' do resultado
    assert '"Percentual aplicado"' in fonte and '"Fator acumulado"' in fonte


# ======================================== Coleta Financeiro antiga (bloqueio)

from _politica_entrega_segura import MENSAGEM_COLETA_PRECISAO_ANTERIOR  # noqa: E402


def _coleta_com_metodo(conteudo: bytes, modo: str, *, bruta: bool) -> bytes:
    wb = load_workbook(io.BytesIO(conteudo))
    wb["CONTROLE"]["B1"] = modo
    if bruta:
        par = wb["parametros"]
        par["E3"], par["F2"], par["F3"] = RAW_C1, 1.0, FATOR_BRUTO_C1
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def test_coleta_financeiro_antiga_e_lida_mas_formalizacao_bloqueada(coleta_c1):
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    antiga = _coleta_com_metodo(coleta_c1, "Financeiro", bruta=True)
    # leitura/upload permitidos: o runtime nao rejeita o arquivo
    resultado, diagnostico = processar_coleta_oficial_runtime(antiga)
    # diagnostico da precisao bruta exibido
    assert any("precisao bruta" in aviso for aviso in diagnostico.get("avisos") or [])
    # formalizacao bloqueada com a orientacao para regenerar
    assert resultado["formalizacao_bloqueada"] is True
    assert MENSAGEM_COLETA_PRECISAO_ANTERIOR in resultado["bloqueios_formalizacao"]
    assert "Regere a Coleta e faça novo upload" in MENSAGEM_COLETA_PRECISAO_ANTERIOR
    formalizacao = resultado["resultado_consolidado"]["formalizacao"]
    assert formalizacao["bloqueada"] is True and formalizacao["status"] == "BLOQUEADA"
    assert MENSAGEM_COLETA_PRECISAO_ANTERIOR in resultado["resultado_consolidado"]["bloqueios"]
    # o arquivo fisico continua intacto (nada e regravado)
    assert load_workbook(io.BytesIO(antiga))["parametros"]["E3"].value == pytest.approx(RAW_C1, abs=1e-16)


def test_coleta_financeiro_pela_regra_vigente_nao_recebe_o_bloqueio(coleta_c1):
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, diagnostico = processar_coleta_oficial_runtime(
        _coleta_com_metodo(coleta_c1, "Financeiro", bruta=False)
    )
    assert MENSAGEM_COLETA_PRECISAO_ANTERIOR not in resultado["bloqueios_formalizacao"]
    assert not any("precisao bruta" in a for a in diagnostico.get("avisos") or [])


def test_metodo_pc_mantem_a_regra_existente(coleta_c1):
    """No PC o efeito ja decorre da reconciliacao XLS x Python; nada muda."""
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, _diagnostico = processar_coleta_oficial_runtime(
        _coleta_com_metodo(coleta_c1, "PCs", bruta=True)
    )
    assert MENSAGEM_COLETA_PRECISAO_ANTERIOR not in resultado["bloqueios_formalizacao"]


def test_golden_financeiro_real_antigo_bloqueia_sem_recalcular_o_vta():
    """Coleta Financeiro REAL (C3 com 2,8899...% bruto), recalculada no Excel."""
    import test_baseline_resultados_goldens as goldens

    arquivo = goldens.GOLDENS["financeiro_multiciclo_validado"]
    if not arquivo.exists():
        pytest.skip(f"golden externo ausente: {arquivo}")
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, diagnostico = processar_coleta_oficial_runtime(arquivo.read_bytes())
    assert any("precisao bruta" in a for a in diagnostico.get("avisos") or [])
    assert resultado["bloqueios_formalizacao"] == [MENSAGEM_COLETA_PRECISAO_ANTERIOR]
    consolidado = resultado["resultado_consolidado"]
    assert consolidado["formalizacao"]["status"] == "BLOQUEADA"
    assert consolidado["formalizacao"]["mensagem"] == MENSAGEM_COLETA_PRECISAO_ANTERIOR
    # o VTA gravado no XLS antigo NAO e recalculado silenciosamente
    web = goldens._fotografar("financeiro_multiciclo_validado")["web"]
    assert web["vta_oficial"] == goldens.VTA_FINANCEIRO_HOMOLOGADO


# ============================== revisao final: 3 bypasses P1 (docs/negativo/fator)

DOCS_FORMALIZADORES = ("sumario_executivo", "despacho_saneador", "termo_apostila")


# ---------------------------------------------------- P1.1 documentos formalizadores

def test_p1_coleta_antiga_nao_libera_nenhum_documento_formalizador(coleta_c1):
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, diagnostico = processar_coleta_oficial_runtime(
        _coleta_com_metodo(coleta_c1, "Financeiro", bruta=True)
    )
    # leitura permitida e diagnostico presente
    assert any("precisao bruta" in a for a in diagnostico.get("avisos") or [])
    assert resultado["formalizacao_bloqueada"] is True
    documentos = resultado["capacidades"]["documentos"]
    for chave in DOCS_FORMALIZADORES:
        assert documentos[chave]["habilitado"] is False, chave
        assert documentos[chave]["motivo"] == MENSAGEM_COLETA_PRECISAO_ANTERIOR, chave
    # a pagina exibe a mesma orientacao canonica junto aos documentos
    fonte = (ROOT / "pages" / "03_Valor_Global.py").read_text(encoding="utf-8")
    assert "st.warning(MENSAGEM_COLETA_PRECISAO_ANTERIOR)" in fonte


def test_p1_coleta_vigente_mantem_os_tres_documentos_disponiveis(coleta_c1):
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, _diagnostico = processar_coleta_oficial_runtime(
        _coleta_com_metodo(coleta_c1, "Financeiro", bruta=False)
    )
    documentos = resultado["capacidades"]["documentos"]
    for chave in DOCS_FORMALIZADORES:
        assert documentos[chave]["habilitado"] is True, chave


def test_p1_golden_financeiro_real_antigo_sem_documentos_formalizadores():
    import test_baseline_resultados_goldens as goldens

    arquivo = goldens.GOLDENS["financeiro_multiciclo_validado"]
    if not arquivo.exists():
        pytest.skip(f"golden externo ausente: {arquivo}")
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, _diagnostico = processar_coleta_oficial_runtime(arquivo.read_bytes())
    documentos = resultado["capacidades"]["documentos"]
    for chave in DOCS_FORMALIZADORES:
        assert documentos[chave]["habilitado"] is False, chave
        assert documentos[chave]["motivo"] == MENSAGEM_COLETA_PRECISAO_ANTERIOR


# ------------------------- Valor Global: negativo preservado como no main

def _funcoes_valor_global() -> dict:
    """Funcoes/constantes de modulo da pagina, sem executar a interface."""
    import ast

    caminho = (ROOT / "pages" / "03_Valor_Global.py").resolve()
    fonte = caminho.read_text(encoding="utf-8")
    arvore = ast.parse(fonte)
    corpo = [
        n for n in arvore.body
        if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef))
        or (isinstance(n, ast.Assign) and "st." not in ast.get_source_segment(fonte, n))
    ]
    espaco = {"__name__": "valor_global_funcoes", "__file__": str(caminho)}
    exec(compile(ast.Module(body=corpo, type_ignores=[]), "<valor_global>", "exec"), espaco)
    return espaco


def _padronizar(linha_c1: dict, *, situacao_c1: str = "TEMPESTIVO"):
    import pandas as pd

    funcoes = _funcoes_valor_global()
    df = pd.DataFrame([
        {"Ciclo": "C1", "Situação": situacao_c1,
         "Percentual apurado pelo índice": -0.020349, **linha_c1},
        {"Ciclo": "C2", "Situação": "TEMPESTIVO",
         "Percentual apurado pelo índice": 0.05, "Percentual aplicado": 0.05,
         "Tratamento ciclo negativo": ""},
    ])
    ciclos = funcoes["padronizar_ciclos"](df)
    return funcoes, {r["Ciclo"]: r for r in ciclos.to_dict("records")}


def _assert_negativo_como_no_main(ciclos):
    """Valor Global (regra preexistente do main): ciclo negativo sem acordo
    negocial entra com 0,00% / fator 1,0000; o ciclo seguinte compoe
    normalmente a partir dele."""
    assert ciclos["C1"]["Percentual aplicado"] == 0.0
    assert ciclos["C1"]["Variação"] == 0.0
    assert ciclos["C1"]["Fator"] == 1.0
    assert ciclos["C1"]["Fator acumulado"] == 1.0
    assert ciclos["C2"]["Fator"] == 1.05
    assert ciclos["C2"]["Fator acumulado"] == 1.05


def test_valor_global_negativo_sem_decisao_mantem_regra_do_main():
    _funcoes, ciclos = _padronizar(
        {"Percentual aplicado": None, "Tratamento ciclo negativo": None}
    )
    _assert_negativo_como_no_main(ciclos)
    assert ciclos["C1"]["Situação aplicada"].endswith("CICLO NEGATIVO (APLICADO 0,00%)")


def test_valor_global_negativo_neutralizado_mantem_regra_do_main():
    _funcoes, ciclos = _padronizar(
        {"Percentual aplicado": 0.0, "Tratamento ciclo negativo": NEUTRALIZAR_VARIACAO_NEGATIVA}
    )
    _assert_negativo_como_no_main(ciclos)


def test_valor_global_negativo_aplicado_mantem_regra_do_main():
    _funcoes, ciclos = _padronizar(
        {"Percentual aplicado": -0.020349, "Tratamento ciclo negativo": APLICAR_VARIACAO_NEGATIVA}
    )
    _assert_negativo_como_no_main(ciclos)


def test_valor_global_precluso_sem_pedido_mantem_regra_do_main():
    _funcoes, ciclos = _padronizar(
        {"Percentual aplicado": None, "Tratamento ciclo negativo": None},
        situacao_c1="PRECLUSO | SEM PEDIDO",
    )
    _assert_negativo_como_no_main(ciclos)


def test_valor_global_ciclo_positivo_usa_fator_oficial_de_duas_casas():
    import pandas as pd

    funcoes = _funcoes_valor_global()
    df = pd.DataFrame([
        {"Ciclo": "C1", "Situação": "TEMPESTIVO",
         "Percentual apurado pelo índice": RAW_C1, "Percentual aplicado": RAW_C1,
         "Tratamento ciclo negativo": ""},
    ])
    c1 = funcoes["padronizar_ciclos"](df).to_dict("records")[0]
    assert c1["Fator"] == 1.0405
    assert c1["Percentual aplicado"] == 0.0405
    assert c1["Fator acumulado"] == 1.0405


# ------------------------------------- P1.3 fator acumulado legado sem percentual

FATOR_LEGADO_ISOLADO = 1.040521878812559


def _coleta_so_com_fator(conteudo: bytes) -> bytes:
    wb = load_workbook(io.BytesIO(conteudo))
    par = wb["parametros"]
    par["E3"] = None                      # percentual ausente
    par["F2"] = 1.0
    par["F3"] = FATOR_LEGADO_ISOLADO      # so o fator acumulado bruto
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def test_p1_fator_legado_sem_percentual_nao_vira_canonico(coleta_c1):
    leitura = ler_masterfile_v10(_coleta_so_com_fator(coleta_c1))
    c1 = leitura["parametros_v10"]["por_ciclo"]["C1"]
    assert c1["percentual_reajuste"] is None
    assert c1["fator_acumulado"] is None
    assert c1["fator_acumulado_bruto"] == pytest.approx(FATOR_LEGADO_ISOLADO, abs=1e-15)
    assert any("sem percentual do ciclo suficiente" in a
               for a in leitura["parametros_v10"]["alertas"])


def test_p1_fator_legado_sem_percentual_nao_chega_aos_consumidores(coleta_c1):
    def _contem_fator_bruto(valor) -> bool:
        if isinstance(valor, float):
            return abs(valor - FATOR_LEGADO_ISOLADO) < 1e-12
        if isinstance(valor, dict):
            return any(
                _contem_fator_bruto(v) for k, v in valor.items()
                if not str(k).endswith("_bruto")
            )
        if isinstance(valor, (list, tuple)):
            return any(_contem_fator_bruto(v) for v in valor)
        return False

    leitura = ler_masterfile_v10(_coleta_so_com_fator(coleta_c1))
    # objeto do processo, VTA sombra, potencial, composicao e PCs nunca veem o
    # fator bruto fora de campo *_bruto
    for chave in ("objeto_processo", "vta_sombra", "potencial_futuro",
                  "composicao_vta", "itens_pc_v10"):
        assert not _contem_fator_bruto(leitura.get(chave)), chave
    assert not _contem_fator_bruto(leitura["parametros_v10"]["por_ciclo"])
    dados = _extrair_dados(leitura, None)
    assert not _contem_fator_bruto(dados)
    c1 = next(c for c in dados["ciclos"] if c["ciclo"] == "C1")
    assert c1["percentual_reajuste"] is None
