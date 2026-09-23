"""Ajustes documentais consolidados do TERMO DE APOSTILA.

1  Considerando fixo da Ata da 1869ª Reunião Ordinária (independe do campo
   manual `deliberacao_institucional`, que fica so para deliberacao adicional).
2  Nova redacao sobre dupla contagem do retroativo (secao 3).
3  Considerando das duas casas decimais + itens 1.2 e 1.3+ abaixo do Quadro 1
   (numeracao pertence ao Termo; o Despacho Saneador nao a recebe).
4  Linha Total do Quadro 3 cita as referencias das parcelas presentes.

Sao ajustes de APRESENTACAO: nenhum valor, percentual, fator, VTA ou regra de
negocio e recalculado ou alterado aqui.
"""
from __future__ import annotations

import re
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from _sumario_executivo import TEXTO_PERCENTUAL_FATOR  # noqa: E402
from _templates_documentos import (  # noqa: E402
    TEXTO_CONSIDERANDO_ATA_1869,
    TEXTO_CONSIDERANDO_DUAS_CASAS,
    _extrair_dados,
    _composicao_didatica_vta,
    _ta_retroativo_do_metodo,
    _ta_secao3_composicao_vta,
    _vta_texto_doc,
    formatar_moeda,
    gerar_despacho_saneador,
    gerar_modelo_branco_termo,
    gerar_termo_apostila,
)

from test_sumario_executivo import (  # noqa: E402
    leitura_multiciclo_pc,
    leitura_simples_financeiro,
)
from test_templates_documentos import (  # noqa: E402
    CAMPOS_SANEADOR,
    CAMPOS_TERMO,
    _leitura_consumidos_bruta,
    _leitura_financeiro,
    _leitura_pc_sem_potencial,
)

ATA_1869 = (
    "A deliberação da Diretoria Executiva da Telebras, consignada na Ata da "
    "1869ª Reunião Ordinária, de 13 de janeiro de 2026, que revogou a "
    "suspensão anteriormente imposta à tramitação dos reajustes contratuais e "
    "restabeleceu a normalidade do respectivo processamento;"
)
DUAS_CASAS_CONSIDERANDO = (
    "A adoção, em cada ciclo de reajuste, do respectivo percentual apurado com "
    "duas casas decimais;"
)
ITEM_1_2 = (
    "1.2. Para o cálculo, o percentual de reajuste de cada ciclo é considerado "
    "com duas casas decimais, sendo o fator correspondente aplicado aos valores "
    "unitários."
)
FRASE_PERDA = "Em razão da data do pedido, os efeitos financeiros do reajuste"


# ---------------------------------------------------------------------------
# Auxiliares
# ---------------------------------------------------------------------------

def _termo(leitura: dict, **campos) -> bytes:
    return gerar_termo_apostila(leitura, campos_manuais=dict(CAMPOS_TERMO, **campos))


def _paragrafos(docx_bytes: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(docx_bytes)).paragraphs]


def _texto(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.extend(c.text for c in linha.cells)
    return "\n".join(partes)


def _considerandos(docx_bytes: bytes) -> list[str]:
    textos = _paragrafos(docx_bytes)
    inicio = textos.index("CONSIDERANDO:") + 1
    saida = []
    for texto in textos[inicio:]:
        if not texto.strip():
            break
        saida.append(texto)
    return saida


def _corpo_ordenado(docx_bytes: bytes) -> list[tuple[str, str]]:
    """Corpo do DOCX na ordem real: ('p', texto) ou ('t', 'celula | celula')."""
    doc = Document(BytesIO(docx_bytes))
    saida: list[tuple[str, str]] = []
    for filho in doc.element.body.iterchildren():
        etiqueta = filho.tag.rsplit("}", 1)[-1]
        if etiqueta == "p":
            texto = "".join(t.text or "" for t in filho.iter() if t.tag.endswith("}t"))
            saida.append(("p", texto))
        elif etiqueta == "tbl":
            primeira = filho.find(".//{*}tr")
            celulas = [
                "".join(t.text or "" for t in c.iter() if t.tag.endswith("}t"))
                for c in primeira.findall("{*}tc")
            ]
            saida.append(("t", " | ".join(celulas)))
    return saida


def _quadro3(docx_bytes: bytes) -> list[list[str]]:
    doc = Document(BytesIO(docx_bytes))
    tabela = next(
        t for t in doc.tables
        if [c.text for c in t.rows[0].cells] == ["Ref.", "Descrição", "Valor"]
    )
    return [[c.text for c in linha.cells] for linha in tabela.rows[1:]]


def _multiciclo_com_duas_perdas() -> dict:
    # C1 abre em 01/05/2024 e passa a produzir efeito so em 01/07/2024
    # (perde maio e junho); C2 ja perde maio, junho e julho de 2025.
    leitura = leitura_multiciclo_pc()
    leitura["parametros_v10"]["por_ciclo"]["C1"]["inicio_efeito_financeiro"] = (
        date(2024, 7, 1)
    )
    return leitura


def _leitura_sem_perda() -> dict:
    leitura = leitura_simples_financeiro()
    leitura["parametros_v10"]["por_ciclo"]["C1"]["inicio_efeito_financeiro"] = (
        date(2025, 2, 1)
    )
    return leitura


def _dados_quadro3(**extra) -> dict:
    base = {"_modo_branco": False, "metodo": "financeiro"}
    base.update(extra)
    return base


def _quadro3_sintetico(dados: dict) -> list[list[str]]:
    doc = Document()
    _ta_secao3_composicao_vta(doc, dados)
    tabela = next(
        t for t in doc.tables
        if [c.text for c in t.rows[0].cells] == ["Ref.", "Descrição", "Valor"]
    )
    return [[c.text for c in linha.cells] for linha in tabela.rows[1:]]


# ---------------------------------------------------------------------------
# ALTERACAO 1 — Ata da 1869ª RO (testes 1-5)
# ---------------------------------------------------------------------------

def test_ata_1869_texto_fonte_e_integral():
    assert TEXTO_CONSIDERANDO_ATA_1869 == ATA_1869
    assert TEXTO_CONSIDERANDO_DUAS_CASAS == DUAS_CASAS_CONSIDERANDO


def test_ata_1869_aparece_uma_vez_com_data_e_redacao_integral():
    texto = _texto(_termo(leitura_multiciclo_pc()))
    assert texto.count("1869ª Reunião Ordinária") == 1                     # 1
    assert "13 de janeiro de 2026" in texto                                 # 2
    assert (                                                                # 3
        "que revogou a suspensão anteriormente imposta à tramitação dos "
        "reajustes contratuais e restabeleceu a normalidade do respectivo "
        "processamento"
    ) in texto


def test_ata_1869_vem_logo_apos_a_clausula_e_antes_da_solicitacao():
    considerandos = _considerandos(_termo(leitura_multiciclo_pc()))
    assert "Cláusula Oitava" in considerandos[0]
    assert considerandos[1] == "2. " + ATA_1869
    assert "solicitação" in considerandos[2]
    assert considerandos[2].startswith("3. ")


def test_ata_1869_nao_depende_de_deliberacao_institucional():
    # 4 — sem o campo manual, a Ata continua presente.
    assert "deliberacao_institucional" not in CAMPOS_TERMO
    sem_campo = _considerandos(_termo(leitura_multiciclo_pc()))
    assert sum("1869ª Reunião Ordinária" in c for c in sem_campo) == 1
    # e o considerando adicional nao aparece "por engano".
    assert not any("Ata nº 1" in c for c in sem_campo)


def test_deliberacao_adicional_nao_substitui_nem_duplica_a_ata():
    # 5 — deliberacao adicional distinta: convive com a Ata, sem substitui-la.
    adicional = "A deliberação da Diretoria Executiva registrada na Ata nº 1"
    considerandos = _considerandos(
        _termo(leitura_multiciclo_pc(), deliberacao_institucional=adicional)
    )
    assert sum("1869ª Reunião Ordinária" in c for c in considerandos) == 1
    assert considerandos[1] == "2. " + ATA_1869
    assert sum("Ata nº 1" in c for c in considerandos) == 1
    assert considerandos[-1].startswith(f"{len(considerandos)}. ")
    for numero, paragrafo in enumerate(considerandos, start=1):
        assert paragrafo.startswith(f"{numero}. ")

    # Campo preenchido com a propria Ata: nao gera segundo considerando.
    repetida = _considerandos(
        _termo(leitura_multiciclo_pc(), deliberacao_institucional=ATA_1869)
    )
    assert sum("1869ª Reunião Ordinária" in c for c in repetida) == 1
    assert len(repetida) == len(_considerandos(_termo(leitura_multiciclo_pc())))

    # Deliberacao adicional que apenas cita o numero 1869 (nao a Ata) e mantida.
    outra = _considerandos(_termo(
        leitura_multiciclo_pc(),
        deliberacao_institucional="A deliberação constante do Processo 1869/2026",
    ))
    assert sum("Processo 1869/2026" in c for c in outra) == 1
    assert len(outra) == len(repetida) + 1


def test_ata_1869_tambem_integra_o_modelo_em_branco():
    considerandos = _considerandos(gerar_modelo_branco_termo())
    assert considerandos[1] == "2. " + ATA_1869
    assert sum("1869ª Reunião Ordinária" in c for c in considerandos) == 1


# ---------------------------------------------------------------------------
# ALTERACAO 2 — dupla contagem do retroativo (testes 6-12)
# ---------------------------------------------------------------------------

def _paragrafo_dupla_contagem(docx_bytes: bytes) -> str:
    return next(
        p for p in _paragrafos(docx_bytes)
        if p.startswith("3.") and "dupla contagem" in p
    )


def test_redacao_antiga_da_dupla_contagem_nao_existe_mais():
    for leitura in (_leitura_financeiro(), _leitura_pc_sem_potencial(),
                    _leitura_consumidos_bruta()):
        texto = _texto(_termo(leitura))
        assert "não é somado como parcela autônoma" not in texto           # 6
        assert "Sua inclusão adicional representaria dupla contagem" not in texto  # 7


def test_nova_redacao_da_dupla_contagem_financeiro_e_pc():
    for leitura in (_leitura_financeiro(), _leitura_pc_sem_potencial()):
        docx = _termo(leitura)
        dados = _extrair_dados(leitura, None)
        retro = _ta_retroativo_do_metodo(dados)
        assert retro is not None and round(retro, 2)
        par = _paragrafo_dupla_contagem(docx)
        # 11 — valor vem da cadeia canonica ja selecionada, formatado.
        assert par.endswith(
            f"O retroativo reconhecido de {formatar_moeda(retro)} já está "
            "incorporado ao valor da execução atualizada considerado na "
            "composição do Quadro 3. Por essa razão, ele não é somado "
            "novamente como parcela separada, pois isso resultaria em dupla "
            "contagem do mesmo valor."
        )
        assert (                                                             # 8
            "já está incorporado ao valor da execução atualizada considerado "
            "na composição do Quadro 3"
        ) in par
        assert "não é somado novamente como parcela separada" in par        # 9
        assert "isso resultaria em dupla contagem do mesmo valor" in par    # 10


def test_nova_redacao_da_dupla_contagem_itens_consumidos_nao_diz_reconhecido():
    leitura = _leitura_consumidos_bruta()
    docx = _termo(leitura)
    dados = _extrair_dados(leitura, None)
    retro = _ta_retroativo_do_metodo(dados)
    par = _paragrafo_dupla_contagem(docx)
    assert par.endswith(
        f"O retroativo de {formatar_moeda(retro)} já está incorporado ao valor "
        "da execução atualizada considerado na composição do Quadro 3. Por "
        "essa razão, ele não é somado novamente como parcela separada, pois "
        "isso resultaria em dupla contagem do mesmo valor."
    )
    assert "retroativo reconhecido" not in par                              # 12
    assert "retroativo reconhecido" not in _texto(docx)


def test_dupla_contagem_nao_hardcoda_valor():
    # O valor acompanha a cadeia canonica: fixtures distintas, valores distintos.
    valores = set()
    for leitura in (_leitura_financeiro(), _leitura_pc_sem_potencial(),
                    _leitura_consumidos_bruta()):
        par = _paragrafo_dupla_contagem(_termo(leitura))
        valores.add(re.search(r"R\$ [\d.]+,\d{2}", par).group(0))
    assert len(valores) == 3


# ---------------------------------------------------------------------------
# ALTERACAO 3.1 — considerando das duas casas decimais (testes 13-14)
# ---------------------------------------------------------------------------

def test_considerando_duas_casas_aparece_uma_unica_vez():
    for docx in (_termo(leitura_multiciclo_pc()), _termo(_leitura_financeiro()),
                 gerar_modelo_branco_termo()):
        texto = _texto(docx)
        assert texto.count(DUAS_CASAS_CONSIDERANDO) == 1                    # 13


def test_numeracao_dos_considerandos_segue_automatica_e_sequencial():
    for docx in (_termo(leitura_multiciclo_pc()),
                 _termo(leitura_multiciclo_pc(),
                        deliberacao_institucional="Deliberação adicional X",
                        instrumentos_posteriores="Termo Aditivo nº 3")):
        considerandos = _considerandos(docx)
        for numero, paragrafo in enumerate(considerandos, start=1):         # 14
            assert paragrafo.startswith(f"{numero}. ")
        # O numero do novo considerando nao e fixo: acompanha os itens.
        idx = next(i for i, c in enumerate(considerandos)
                   if DUAS_CASAS_CONSIDERANDO in c)
        assert considerandos[idx].startswith(f"{idx + 1}. ")
    # Sem adicionais, o considerando das duas casas e o ultimo.
    base = _considerandos(_termo(leitura_multiciclo_pc()))
    assert base[-1] == f"{len(base)}. " + DUAS_CASAS_CONSIDERANDO
    # Com adicionais, eles vem DEPOIS dos itens fixos.
    com = _considerandos(_termo(
        leitura_multiciclo_pc(), deliberacao_institucional="Deliberação adicional X"
    ))
    assert DUAS_CASAS_CONSIDERANDO in com[-2]
    assert "Deliberação adicional X" in com[-1]


# ---------------------------------------------------------------------------
# ALTERACAO 3.2 — Secao 1: itens 1.2 e 1.3+ abaixo do Quadro 1 (testes 15-24)
# ---------------------------------------------------------------------------

def _secao1(docx_bytes: bytes) -> list[tuple[str, str]]:
    corpo = _corpo_ordenado(docx_bytes)
    ini = next(i for i, (t, x) in enumerate(corpo)
               if t == "p" and x == "1. Dos reajustes concedidos")
    fim = next(i for i, (t, x) in enumerate(corpo)
               if t == "p" and x.startswith("2. Da apuração"))
    return [(t, x) for t, x in corpo[ini:fim] if x.strip()]


def _itens_secao1(secao: list[tuple[str, str]]) -> list[str]:
    return [x for t, x in secao if t == "p" and re.match(r"1\.\d+\. ", x)]


def test_item_1_2_vem_logo_abaixo_do_quadro_1_com_texto_exato():
    secao = _secao1(_termo(leitura_simples_financeiro()))
    assert secao[0] == ("p", "1. Dos reajustes concedidos")
    assert secao[1][1].startswith("1.1. Ao Contrato nº ")
    assert secao[2] == ("p", "Quadro 1 — Síntese dos reajustes concedidos")
    assert secao[3][0] == "t"                                               # tabela
    assert secao[4] == ("p", ITEM_1_2)                                      # 15, 16
    assert secao[4][1] == "1.2. " + TEXTO_PERCENTUAL_FATOR  # fonte unica reaproveitada


def test_perda_de_efeitos_vira_item_1_3_com_dados_automaticos_um_ciclo():
    secao = _secao1(_termo(leitura_simples_financeiro()))
    assert secao[5] == ("p", (
        "1.3. Em razão da data do pedido, os efeitos financeiros do reajuste "
        "deste ciclo iniciam-se em 04/2025, não alcançando as competências de "
        "fevereiro e março de 2025."                     # 17, 18, 19, 20
    ))
    assert len(secao) == 6
    assert "do ciclo C1" not in secao[5][1]


def test_perda_em_ciclo_nomeado_quando_ha_varios_ciclos():
    itens = _itens_secao1(_secao1(_termo(leitura_multiciclo_pc())))
    assert [x[:4] for x in itens] == ["1.1.", "1.2.", "1.3."]
    assert itens[1] == ITEM_1_2
    assert itens[2] == (
        "1.3. Em razão da data do pedido, os efeitos financeiros do reajuste "
        "do ciclo C2 iniciam-se em 08/2025, não alcançando as competências de "
        "maio, junho e julho de 2025."                                     # 21
    )


def test_multiplas_perdas_recebem_numeracao_sequencial():
    itens = _itens_secao1(_secao1(_termo(_multiciclo_com_duas_perdas())))
    assert [x[:4] for x in itens] == ["1.1.", "1.2.", "1.3.", "1.4."]       # 22
    assert "do ciclo C1 iniciam-se em 07/2024" in itens[2]
    assert "as competências de maio e junho de 2024" in itens[2]
    assert "do ciclo C2 iniciam-se em 08/2025" in itens[3]
    assert "deste ciclo" not in " ".join(itens)


def test_sem_perda_de_efeitos_nao_ha_item_1_3_artificial():
    docx = _termo(_leitura_sem_perda())
    secao = _secao1(docx)
    itens = _itens_secao1(secao)
    assert [x[:4] for x in itens] == ["1.1.", "1.2."]                       # 23
    assert FRASE_PERDA not in _texto(docx)
    assert secao[-1] == ("p", ITEM_1_2)


def test_despacho_saneador_nao_recebe_a_numeracao_do_termo():
    for leitura in (leitura_simples_financeiro(), leitura_multiciclo_pc(),
                    _multiciclo_com_duas_perdas()):
        docx = gerar_despacho_saneador(leitura, campos_manuais=CAMPOS_SANEADOR)
        paragrafos = [p for p in _paragrafos(docx) if p.strip()]
        perdas = [p for p in paragrafos if FRASE_PERDA in p]
        assert perdas
        assert all(p.startswith(FRASE_PERDA) for p in perdas)               # 24
        assert not any(re.match(r"1\.[234]\. ", p) for p in paragrafos)
        assert not any(p.startswith("1.2. Para o cálculo") for p in paragrafos)
        # texto do percentual/fator segue sem numeracao propria do Termo
        assert any(p.startswith("Para o cálculo, o percentual") for p in paragrafos)


def test_termo_e_saneador_mantem_a_mesma_frase_de_perda_sem_o_prefixo():
    leitura = _multiciclo_com_duas_perdas()
    sane = [p for p in _paragrafos(gerar_despacho_saneador(
        leitura, campos_manuais=CAMPOS_SANEADOR)) if p.startswith(FRASE_PERDA)]
    termo = [p[len("1.x. "):] for p in _paragrafos(_termo(leitura))
             if re.match(r"1\.[34]\. " + re.escape(FRASE_PERDA), p)]
    assert sane == termo and len(termo) == 2


def test_modelo_em_branco_nao_afirma_item_1_2_nem_perdas():
    texto = _texto(gerar_modelo_branco_termo())
    assert "1.2. Para o cálculo" not in texto
    assert FRASE_PERDA not in texto
    assert not re.search(r"^1\.[3-9]\. ", texto, flags=re.M)


# ---------------------------------------------------------------------------
# ALTERACAO 4 — Quadro 3: Total (A + B + ...) (testes 25-30)
# ---------------------------------------------------------------------------

def test_total_do_quadro_3_com_duas_parcelas():
    for leitura in (_leitura_financeiro(), _leitura_consumidos_bruta()):
        linhas = _quadro3(_termo(leitura))
        assert [l[0] for l in linhas] == ["A", "B", "Total (A + B)"]        # 25
        assert linhas[-1][1] == "Valor Total Atualizado do Contrato"
        dados = _extrair_dados(leitura, None)
        assert linhas[-1][2] == _vta_texto_doc(dados)                        # 30


def test_total_do_quadro_3_com_tres_parcelas():
    dados = _dados_quadro3(
        vta=150.0, vta_execucao_atualizada=100.0,
        vta_saldo_remanescente_atualizado=40.0, vta_retroativo_potencial=10.0,
        vta_tem_parcela_potencial=True,
    )
    assert len(_composicao_didatica_vta(dados)) == 3
    linhas = _quadro3_sintetico(dados)
    assert [l[0] for l in linhas] == ["A", "B", "C", "Total (A + B + C)"]   # 26
    assert linhas[-1][2] == _vta_texto_doc(dados) == formatar_moeda(150.0)   # 30
    assert [l[2] for l in linhas[:-1]] == [
        formatar_moeda(100.0), formatar_moeda(40.0), formatar_moeda(10.0)
    ]


def test_total_do_quadro_3_expande_dinamicamente_com_mais_parcelas():
    def parcela(ciclo, valor, fonte=""):
        return {"ciclo": ciclo, "valor_atualizado": valor, "fonte_parcela": fonte}

    parcelas = [
        parcela("C0", 100.0), parcela("C1", 50.0), parcela("C2", 25.0),
        parcela("", 10.0, "saldo remanescente"),
    ]
    dados = _dados_quadro3(vta=185.0, parcelas_vta=parcelas)
    linhas = _quadro3_sintetico(dados)
    assert [l[0] for l in linhas] == ["A", "B", "C", "D", "Total (A + B + C + D)"]  # 27
    parcelas.append(parcela("C3", 5.0))
    linhas5 = _quadro3_sintetico(_dados_quadro3(vta=190.0, parcelas_vta=parcelas))
    assert linhas5[-1][0] == "Total (A + B + C + D + E)"


def test_total_usa_exatamente_as_referencias_das_linhas_presentes():
    for leitura in (_leitura_financeiro(), leitura_simples_financeiro()):
        linhas = _quadro3(_termo(leitura))
        refs = [l[0] for l in linhas[:-1]]
        assert linhas[-1][0] == "Total (" + " + ".join(refs) + ")"          # 28
        # 29 — nenhuma referencia inexistente
        citadas = re.findall(r"[A-Z]", linhas[-1][0].replace("Total", ""))
        assert set(citadas) <= set(refs)

    dados = _dados_quadro3(
        vta=140.0, vta_execucao_atualizada=100.0,
        vta_saldo_remanescente_atualizado=40.0,
    )
    linhas = _quadro3_sintetico(dados)
    assert linhas[-1][0] == "Total (A + B)"
    assert "C" not in linhas[-1][0].replace("Total", "")                    # 29


def test_quadro_3_sem_parcelas_mantem_total_simples():
    linhas = _quadro3_sintetico(_dados_quadro3(vta=10.0))
    assert [l[0] for l in linhas] == ["Total"]


# ---------------------------------------------------------------------------
# DOCX real: ordem estrutural (validacao final)
# ---------------------------------------------------------------------------

def test_docx_real_ordem_considerandos_secao_1_e_secao_3():
    docx = _termo(_leitura_financeiro())
    corpo = [x for _, x in _corpo_ordenado(docx) if x.strip()]

    def pos(prefixo):
        return next(i for i, x in enumerate(corpo) if x.startswith(prefixo))

    # CONSIDERANDOS: clausula -> Ata -> ... -> duas casas
    assert pos("1. Cláusula Oitava") < pos("2. A deliberação da Diretoria") \
        < pos("3. A solicitação")
    assert pos("3. A solicitação") < pos("11. A adoção") \
        < pos("FORMALIZA-SE O PRESENTE TERMO")
    # SECAO 1: 1.1 -> Quadro 1 -> tabela -> 1.2 -> 1.3+
    assert pos("1.1. ") < pos("Quadro 1 —") < pos("Ref. | Ciclo") \
        < pos("1.2. Para o cálculo") < pos("1.3. Em razão") \
        < pos("2. Da apuração")
    # SECAO 3: Quadro 3 -> tabela -> paragrafo da dupla contagem
    assert pos("3. Da composição") < pos("Quadro 3 —") < pos("Ref. | Descrição") \
        < pos("3.2. O retroativo reconhecido de")
