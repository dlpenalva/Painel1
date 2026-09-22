"""Paragrafo documental sobre o percentual de 2 casas e o fator do ciclo.

Tarefa exclusivamente documental: Sumario Executivo e Despacho Saneador
trazem o exemplo com o ULTIMO ciclo computado (dados reais da apuracao);
o Termo de Apostila traz apenas a frase geral. Nenhum valor e recalculado.
"""
from __future__ import annotations

import copy
import re
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from _leitor_masterfile_v10 import _oficializar_parametros_v10  # noqa: E402
from _sumario_executivo import (  # noqa: E402
    TEXTO_PERCENTUAL_FATOR,
    gerar_sumario_executivo_pdf,
    montar_dados_sumario_executivo,
    texto_percentual_fator,
)
from _templates_documentos import (  # noqa: E402
    _extrair_dados,
    gerar_despacho_saneador,
    gerar_termo_apostila,
)
from test_sumario_executivo import (  # noqa: E402
    leitura_multiciclo_pc,
    leitura_simples_financeiro,
)
from test_templates_documentos import CAMPOS_SANEADOR, CAMPOS_TERMO  # noqa: E402

TRECHO_UNICO = "duas casas decimais, sendo o fator correspondente"


def _texto_docx(conteudo: bytes) -> str:
    doc = Document(BytesIO(conteudo))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.extend(celula.text for celula in linha.cells)
    return "\n".join(partes)


def _texto_sumario(dados) -> str:
    """Texto do PDF a partir dos flowables efetivamente renderizados.

    O ambiente nao tem extrator de PDF (fitz/pypdf); captura-se a historia
    entregue ao reportlab e confirma-se que o PDF foi gerado.
    """
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table

    capturado: list = []
    original = SimpleDocTemplate.build

    def _build(self, flowables, *args, **kwargs):
        capturado.extend(list(flowables))
        return original(self, flowables, *args, **kwargs)

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(SimpleDocTemplate, "build", _build)
        pdf = gerar_sumario_executivo_pdf(dados)
    assert pdf.startswith(b"%PDF")

    partes: list[str] = []

    def _coletar(obj) -> None:
        if isinstance(obj, Paragraph):
            partes.append(obj.getPlainText())
        elif isinstance(obj, Table):
            for linha in obj._cellvalues:
                for celula in linha:
                    _coletar(celula)
        elif isinstance(obj, (list, tuple)):
            for item in obj:
                _coletar(item)
        elif isinstance(obj, str):
            partes.append(obj)

    _coletar(capturado)
    return re.sub(r"[ \t]+", " ", "\n".join(partes))


def _leitura_tres_ciclos():
    """C1, C2 e C3 computados; C4 com percentual, mas FORA da apuracao."""
    leitura = leitura_multiciclo_pc()
    por_ciclo = leitura["parametros_v10"]["por_ciclo"]
    por_ciclo["C3"].update(
        percentual_reajuste=0.0381, fator_acumulado=1.100249,
        computar_nesta_apuracao="Sim", situacao="Computado",
    )
    por_ciclo["C4"].update(percentual_reajuste=0.0999, computar_nesta_apuracao="Não")
    leitura["controle"]["ciclo_vigente"] = "C3"
    return leitura


def _leitura_percentual_bruto_oficializado():
    """C1 com percentual bruto 4,052187...% passado pela fronteira oficial."""
    leitura = leitura_simples_financeiro()
    parametros = leitura["parametros_v10"]
    parametros["por_ciclo"]["C1"]["percentual_reajuste"] = 0.04052187881
    parametros["por_ciclo"]["C1"]["fator_acumulado"] = 1.04052187881
    parametros["alertas"] = []
    _oficializar_parametros_v10(parametros, fator_proprio_e_percentual=False)
    return leitura


def _gerar_tres(leitura):
    sumario = _texto_sumario(
        montar_dados_sumario_executivo(copy.deepcopy(leitura))
    )
    saneador = _texto_docx(gerar_despacho_saneador(
        copy.deepcopy(leitura), campos_manuais=CAMPOS_SANEADOR
    ))
    apostila = _texto_docx(gerar_termo_apostila(
        copy.deepcopy(leitura), campos_manuais=CAMPOS_TERMO
    ))
    return sumario, saneador, apostila


def test_um_ciclo_usa_o_proprio_ciclo():
    dados = _extrair_dados(leitura_simples_financeiro(), None)
    assert texto_percentual_fator(dados["ciclos"], com_exemplo=True) == (
        TEXTO_PERCENTUAL_FATOR
        + " No ciclo C1, por exemplo, o percentual de 5,25% corresponde ao "
          "fator 1,0525."
    )


def test_tres_ciclos_usa_o_ultimo_ciclo_aplicavel():
    dados = _extrair_dados(_leitura_tres_ciclos(), None)
    texto = texto_percentual_fator(dados["ciclos"], com_exemplo=True)
    assert texto.endswith(
        "No ciclo C3, por exemplo, o percentual de 3,81% corresponde ao "
        "fator 1,0381."
    )
    assert "C4" not in texto and "9,99%" not in texto


def test_percentual_bruto_oficializado_mostra_duas_casas_e_fator_oficial():
    leitura = _leitura_percentual_bruto_oficializado()
    assert leitura["parametros_v10"]["por_ciclo"]["C1"]["percentual_reajuste"] == 0.0405
    sumario, saneador, _ = _gerar_tres(leitura)
    frase = "No ciclo C1, por exemplo, o percentual de 4,05% corresponde ao fator 1,0405."
    assert frase in sumario
    assert frase in saneador
    assert "4,0521" not in sumario + saneador


def test_percentual_nao_oficializado_nao_gera_exemplo():
    """Precisao bruta nunca e fechada no documento: so a frase geral."""
    leitura = leitura_simples_financeiro()
    leitura["parametros_v10"]["por_ciclo"]["C1"]["percentual_reajuste"] = 0.04052187881
    dados = _extrair_dados(leitura, None)
    assert texto_percentual_fator(dados["ciclos"], com_exemplo=True) == TEXTO_PERCENTUAL_FATOR


def test_sem_ciclo_valido_apresenta_so_a_frase_geral():
    leitura = leitura_simples_financeiro()
    leitura["parametros_v10"]["por_ciclo"]["C1"]["percentual_reajuste"] = None
    dados = _extrair_dados(leitura, None)
    assert texto_percentual_fator(dados["ciclos"], com_exemplo=True) == TEXTO_PERCENTUAL_FATOR
    assert texto_percentual_fator([], com_exemplo=True) == TEXTO_PERCENTUAL_FATOR


def test_sumario_e_saneador_com_exemplo_apostila_sem_exemplo_uma_vez_cada():
    sumario, saneador, apostila = _gerar_tres(_leitura_tres_ciclos())
    exemplo = "No ciclo C3, por exemplo, o percentual de 3,81% corresponde ao fator 1,0381."
    assert exemplo in sumario
    assert exemplo in saneador
    assert TEXTO_PERCENTUAL_FATOR in apostila
    assert "por exemplo, o percentual" not in apostila
    for texto in (sumario, saneador, apostila):
        assert texto.count(TRECHO_UNICO) == 1


@pytest.mark.parametrize(
    "fabrica", [leitura_simples_financeiro, _leitura_tres_ciclos]
)
def test_valores_da_apuracao_nao_mudam(fabrica):
    leitura = fabrica()
    antes = copy.deepcopy(leitura)
    dados_antes = _extrair_dados(copy.deepcopy(leitura), None)
    ciclos = copy.deepcopy(dados_antes["ciclos"])
    texto_percentual_fator(ciclos, com_exemplo=True)
    assert ciclos == dados_antes["ciclos"]
    gerar_despacho_saneador(leitura, campos_manuais=CAMPOS_SANEADOR)
    gerar_termo_apostila(leitura, campos_manuais=CAMPOS_TERMO)
    assert leitura == antes
    assert _extrair_dados(copy.deepcopy(leitura), None) == dados_antes
