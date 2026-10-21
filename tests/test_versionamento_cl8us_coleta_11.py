from __future__ import annotations

import io
import re
from contextlib import nullcontext
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

import _ui_utils
from _coleta_oficial import (
    TEMPLATE_COLETA_OFICIAL,
    gerar_coleta_oficial_preenchida,
    obter_coleta_oficial_bytes,
    registrar_versoes_coleta,
)
from _leitor_masterfile_v10 import ler_masterfile_v10
from _versao import (
    CL8US_VERSION,
    COLETA_COMPATIBILIDADE_ANTERIORES,
    COLETA_VERSION,
    COLETA_VERSOES_ACEITAS,
)


def _dados_calculadora() -> dict:
    return {
        "origem": "Reajuste Simples",
        "indice": "IST",
        "data_base_original": "01/01/2023",
        "ciclos": [{
            "ciclo": "C1",
            "data_base": "01/01/2023",
            "data_pedido": "01/01/2024",
            "situacao": "TEMPESTIVO",
            "percentual_aplicado": 0.10,
            "financeiro_inicio": "01/01/2024",
        }],
    }


def _formulas(wb) -> dict[tuple[str, str], str]:
    return {
        (ws.title, cell.coordinate): cell.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }


def _conteudo_exceto_marcadores(wb) -> dict[tuple[str, str], tuple]:
    ignorar = {("CONTROLE", ref) for ref in ("A16", "B16", "A17", "B17")}
    return {
        (ws.title, cell.coordinate): (
            cell.value,
            cell.data_type,
            cell.style_id,
            cell.number_format,
            cell.protection.locked,
        )
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if (ws.title, cell.coordinate) not in ignorar
    }


def test_versoes_publicas_e_politica_preparada() -> None:
    # Quadro "Execucao sem efeito financeiro" na RESULTADOS: mudanca ESTRUTURAL do
    # XLS (Cl8us 11.4, Modelo de Coleta 11.1); as duas versoes sao independentes e
    # a Coleta 11.0 continua aceita (ver COLETA_VERSOES_ACEITAS).
    # RESULTADOS executiva + RESULTADOS_DETALHE: mudanca ESTRUTURAL do XLS
    # (Cl8us 11.5, Modelo de Coleta 11.2); 11.0 e 11.1 seguem aceitas.
    # PR #174 = Cl8us 11.6 (so app); PR #175 = Cl8us 11.7 / Coleta 11.3 (UX
    # da Coleta). Bumps aplicados pelo hotfix de versionamento obrigatorio.
    # Cl8us 11.8 / Coleta 11.4: fator dos aditivos, ciclo em execucao pela
    # data de corte e "Acrescimo - novo item" (formulas do XLS).
    # Cl8us 11.9: checkbox "Nao houve pedido" oculto em analise de um ciclo (so app).
    # Cl8us 12.0: pagina de ciclo unico sem o checkbox "Nao houve pedido" (so app).
    # Cl8us 12.1 / Coleta 11.5: aviso de historico necessario + orientacao de
    # itens novos por aditivo (XLS).
    # Cl8us 12.2: hotfix do carimbo de publicacao; Coleta permanece 11.5.
    # Cl8us 12.3 / Coleta 11.6: base economica do VU em aditivos (XLS).
    # Cl8us 12.4 / Coleta 11.7: historico_VU segue a base economica (XLS).
    # Cl8us 12.5 / Coleta 11.8: ajustes da Coleta em uso (VTA com a posicao
    # atual, PC_PAGO Nao/Não, PC precluso, ajustes visuais) (XLS).
    # Cl8us 12.6 / Coleta 11.9: B2 automatico sem dropdown residual e
    # parametros!U:W ocultas (XLS).
    # Cl8us 12.7 / Coleta 12.0: itens_PC com filtro, PC_PAGO normalizado e
    # coerencia metodo x fonte.
    # Cl8us 12.8 / Coleta 12.1: RESULTADOS com o card RETRO. RECONHECIDO +
    # POTENCIAL e a SITUACAO DA APURACAO em tres estados (XLS).
    # Cl8us 12.9 / Coleta 12.2: acabamento visual de itens_PC, aditivos,
    # CONTROLE e CICLO_EM_EXECUCAO (XLS).
    # Cl8us 12.10 / Coleta 12.3: CONTROLE simplificada e carimbo ao fim da sidebar.
    # Cl8us 12.11 / Coleta 12.4: precisao financeira canonica e saldo remanescente.
    # Cl8us 12.12 / Coleta 12.5: saldo remanescente logo apos a observacao final.
    # Cl8us 12.13 / Coleta 12.6: retroativo zero e referencias temporais do C1.
    # Cl8us 12.14: Painel Executivo com o VTA canonico (so aplicativo, #195).
    assert CL8US_VERSION == "12.14"
    assert COLETA_VERSION == "12.6"
    assert COLETA_VERSOES_ACEITAS == (
        "11.0", "11.1", "11.2", "11.3", "11.4", "11.5", "11.6", "11.7", "11.8", "11.9",
        "12.0", "12.1", "12.2", "12.3", "12.4", "12.5", "12.6",
    )
    assert CL8US_VERSION != COLETA_VERSION
    assert COLETA_COMPATIBILIDADE_ANTERIORES == 2
    assert re.fullmatch(r"\d{2}\.\d+", CL8US_VERSION)
    assert re.fullmatch(r"\d{2}\.\d", COLETA_VERSION)


def test_sidebar_exibe_as_duas_versoes_da_fonte_unica(monkeypatch) -> None:
    legendas: list[str] = []
    containers: list[str] = []
    monkeypatch.setattr(_ui_utils.st, "markdown", lambda *_a, **_k: None)
    monkeypatch.setattr(_ui_utils.st, "caption", legendas.append)
    monkeypatch.setattr(
        _ui_utils.st,
        "container",
        lambda *, key: containers.append(key) or nullcontext(),
    )
    monkeypatch.setattr(_ui_utils, "atualizado_em", lambda: "30/09/2026 12:00")

    _ui_utils.render_versao_sidebar()

    assert legendas == [
        "Última atualização publicada em 30/09/2026 12:00",
        "Cl8us 12.14",
        "Modelo de Coleta 12.6",
    ]
    assert containers == ["cl8us_version_sidebar"]


def test_carimbo_tem_ordem_final_escopada_a_sidebar() -> None:
    app = (Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8")
    assert ".st-key-cl8us_version_sidebar" in app
    assert 'data-testid="stLayoutWrapper"' in app
    assert "order: 9999" in app
    trecho = app[app.index(".st-key-cl8us_version_sidebar"):
                 app.index(".cl8us-page-header")]
    assert "position: fixed" not in trecho
    assert "position: sticky" not in trecho


def test_area_de_versao_estava_livre_no_template_oficial() -> None:
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    ws = wb["CONTROLE"]
    assert all(ws[coordenada].value is None for coordenada in ("A16", "B16", "A17", "B17"))
    assert not any(
        faixa.min_row <= linha <= faixa.max_row
        and faixa.min_col <= coluna <= faixa.max_col
        for faixa in ws.merged_cells.ranges
        for linha in (16, 17)
        for coluna in (1, 2)
    )
    assert not ws.tables
    assert not ws._images
    assert not ws._charts
    for nome in wb.defined_names.values():
        try:
            destinos = list(nome.destinations)
        except Exception:
            continue
        for aba, referencia in destinos:
            if aba != "CONTROLE":
                continue
            min_col, min_row, max_col, max_row = range_boundaries(
                referencia.replace("$", "")
            )
            assert max_row < 16 or min_row > 17 or max_col < 1 or min_col > 2
    assert not any(
        faixa.min_row <= linha <= faixa.max_row
        and faixa.min_col <= coluna <= faixa.max_col
        for validacao in ws.data_validations.dataValidation
        for faixa in validacao.ranges.ranges
        for linha in (16, 17)
        for coluna in (1, 2)
    )


def test_registro_de_versao_nao_altera_formulas() -> None:
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    formulas_antes = _formulas(wb)
    conteudo_antes = _conteudo_exceto_marcadores(wb)
    nomes_antes = {
        nome: definido.attr_text for nome, definido in wb.defined_names.items()
    }

    registrar_versoes_coleta(wb)

    assert _formulas(wb) == formulas_antes
    assert _conteudo_exceto_marcadores(wb) == conteudo_antes
    assert {
        nome: definido.attr_text for nome, definido in wb.defined_names.items()
    } == nomes_antes
    ws = wb["CONTROLE"]
    assert ws["A16"].value == "Modelo de Coleta"
    assert ws["B16"].value == COLETA_VERSION
    assert ws["A17"].value == "Gerado pelo Cl8us"
    assert ws["B17"].value == CL8US_VERSION


def test_registro_de_versao_nunca_sobrescreve_conteudo() -> None:
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    wb["CONTROLE"]["A16"] = "Conteudo preexistente"

    with pytest.raises(ValueError, match="A16"):
        registrar_versoes_coleta(wb)

    assert wb["CONTROLE"]["A16"].value == "Conteudo preexistente"


def test_coleta_nova_em_branco_e_preenchida_recebem_os_marcadores() -> None:
    for conteudo in (
        obter_coleta_oficial_bytes(),
        gerar_coleta_oficial_preenchida(_dados_calculadora()),
    ):
        ws = load_workbook(io.BytesIO(conteudo), data_only=False)["CONTROLE"]
        assert ws["A16"].value == "Modelo de Coleta"
        assert ws["B16"].value == COLETA_VERSION
        assert ws["A17"].value == "Gerado pelo Cl8us"
        assert ws["B17"].value == CL8US_VERSION


def test_leitor_recupera_versao_explicita_da_coleta_nova() -> None:
    leitura = ler_masterfile_v10(obter_coleta_oficial_bytes())
    assert leitura["coleta_version"] == COLETA_VERSION
    assert leitura["versao_detectada"] == "v10-rc"


def test_arquivo_anterior_sem_marcador_mantem_comportamento_atual() -> None:
    leitura = ler_masterfile_v10(TEMPLATE_COLETA_OFICIAL.read_bytes())
    assert leitura["ok"] is True
    assert leitura["erro"] == ""
    assert leitura["versao_detectada"] == "v10-rc"
    assert "coleta_version" not in leitura


@pytest.mark.parametrize("marcador", ["11.2", "11.3"])
def test_coleta_11_2_e_11_3_com_detalhe_sao_aceitas(marcador) -> None:
    """11.3 (PR #175) so muda apresentacao: arquivo 11.2 segue aceito e ambos
    exigem RESULTADOS_DETALHE (fail-closed desde a 11.2)."""
    from _compatibilidade_coleta import (
        LINHAGEM_COLETA_11,
        detectar_linhagem_coleta,
        exige_resultados_detalhe,
    )

    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()), data_only=False)
    wb["CONTROLE"]["B16"] = marcador
    deteccao = detectar_linhagem_coleta(wb)
    assert deteccao["codigo"] == LINHAGEM_COLETA_11
    assert deteccao["marcador_publico"] == marcador
    assert exige_resultados_detalhe(wb)
