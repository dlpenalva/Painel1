"""CONSUMO-GLOSA-1 — ajuste opcional da execucao por valor pago / glosa.

Cobre os treze testes focais exigidos pela tarefa:

  1. baseline sem ajuste (identico ao comportamento homologado);
  2. valor pago menor;       3. glosa equivalente;
  4. zero;                   5. pago maior  -> REVISAR;
  6. glosa maior -> REVISAR; 7. dois ciclos (so o ajustado muda);
  8. retroativo;             9. VTA;
 10. quantidades iguais;    11. remanescente igual;
 12. Financeiro igual;      13. PC igual.

Mais: leitura compativel com XLS antigo (sem o bloco), VAZIO != ZERO, a
cascata de validacao completa e a estrutura do bloco no template oficial.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook  # noqa: E402

from _leitor_masterfile_v10 import (  # noqa: E402
    _ler_ajustes_execucao_consumidos,
    _mapear_colunas_por_cabecalho,
)
from _objeto_processo_reajuste import _montar_memoria_por_ciclo  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"

# Caso controlado, com os numeros do enunciado:
#   valor calculado da execucao em C1 = 1000 x 100,00 = 100.000,00
#   fator acumulado de C1             = 1,08
#   execucao atualizada sem ajuste    = 108.000,00
#   remanescente = (1500 - 1000) x 100,00 = 50.000,00 -> 54.000,00 atualizado
FATOR_C1 = 1.08
FATOR_C2 = 1.1664

_ITEM = {
    "item": "I1",
    "qtd_contratada": 1500,
    "vu_original": 100.0,
    "qtd_total": 1000,
    "check": "OK",
    "consumos": {
        "C0": {"qtd": None, "valor": None},
        "C1": {"qtd": 1000, "valor": 108000.0},
        "C2": {"qtd": None, "valor": None},
        "C3": {"qtd": None, "valor": None},
        "C4": {"qtd": None, "valor": None},
    },
}


def _item_dois_ciclos():
    """Mesmo item, com consumo tambem em C2 (500 x 100 = 50.000 calculado)."""
    item = dict(_ITEM)
    item["qtd_total"] = 1500
    item["consumos"] = dict(_ITEM["consumos"])
    item["consumos"]["C2"] = {"qtd": 500, "valor": 58320.0}  # 50.000 x 1,1664
    return item


def _leitura(
    itens=None, ajustes=None, fatores=None, financeiro=None,
    percentuais=None, computar=None,
):
    fatores = fatores or {"C0": 1.0, "C1": FATOR_C1}
    percentuais = percentuais or {"C1": 0.08}
    computar = computar or {}
    por_ciclo = {}
    for ciclo in ("C0", "C1", "C2", "C3", "C4"):
        por_ciclo[ciclo] = {
            "fator_acumulado": fatores.get(ciclo),
            "percentual_reajuste": percentuais.get(ciclo),
            "computar_nesta_apuracao": computar.get(ciclo, "Sim"),
        }
    bloco = {"itens": list(itens if itens is not None else [_ITEM])}
    if ajustes is not None:
        bloco["ajustes_execucao"] = ajustes
    leitura = {
        "controle": {"modo": "d", "ciclo_vigente": "C1"},
        "parametros_v10": {"por_ciclo": por_ciclo},
        "itens_consumidos_v10": bloco,
    }
    if financeiro is not None:
        leitura["vta_sombra"] = {"parcelas_computadas": financeiro}
    return leitura


def _ajuste(tipo, valor):
    """Bloco cru como o leitor entrega, a partir dos dois campos manuais."""
    return {
        "ciclo": "?",
        "tipo_bruto": tipo,
        "tipo": {"valor pago": "Valor pago", "glosa": "Glosa"}.get(
            str(tipo).strip().lower()
        ),
        "valor_bruto": valor,
        "valor_vazio": valor is None or str(valor).strip() == "",
    }


def _memoria(**kwargs):
    return _montar_memoria_por_ciclo(_leitura(**kwargs), {}, [], {})


def _ciclo(memoria, nome):
    return next(c for c in memoria["ciclos"] if c["ciclo"] == nome)


def _consumidos(memoria, nome="C1"):
    return _ciclo(memoria, nome)["retroativo"]["consumidos"]


def _conferencia(memoria, metodo="consumidos"):
    return next(
        c for c in memoria["conferencias_metodologicas"] if c["metodo"] == metodo
    )


# ---------------------------------------------------------------------------
# 1. Baseline — sem nenhum campo novo preenchido, nada muda
# ---------------------------------------------------------------------------

def test_1_baseline_sem_ajuste_preserva_comportamento_homologado():
    memoria = _memoria()
    assert _consumidos(memoria) == {
        "base_original": 100000.0,
        "valor_atualizado": 108000.0,
        "retroativo": 8000.0,
        "evidencias": 1,
    }
    conferencia = _conferencia(memoria)
    assert conferencia["disponivel"] is True
    assert conferencia["executado_atualizado"] == 108000.0
    assert conferencia["potencial_restante_atualizado"] == 54000.0
    assert conferencia["valor_total_atualizado"] == 162000.0
    assert memoria["ajustes_execucao_resumo"]["status"] == "SEM AJUSTE"
    assert memoria["ajustes_execucao_resumo"]["glosa"] is None


def test_1b_bloco_ausente_e_bloco_vazio_produzem_o_mesmo_resultado():
    """XLS antigo (sem o bloco) == XLS novo com o bloco em branco."""
    sem_bloco = _montar_memoria_por_ciclo(_leitura(), {}, [], {})
    bloco_vazio = _montar_memoria_por_ciclo(_leitura(ajustes={}), {}, [], {})
    assert sem_bloco["ciclos"] == bloco_vazio["ciclos"]
    assert (
        sem_bloco["conferencias_metodologicas"]
        == bloco_vazio["conferencias_metodologicas"]
    )
    assert sem_bloco["vta"] == bloco_vazio["vta"]


# ---------------------------------------------------------------------------
# 2/3. Valor pago e glosa sao duas formas de dizer a mesma coisa
# ---------------------------------------------------------------------------

def test_2_valor_pago_menor():
    memoria = _memoria(ajustes={"C1": _ajuste("Valor pago", 90000.0)})
    registro = memoria["ajustes_execucao_por_ciclo"]["C1"]
    assert registro["status"] == "AJUSTE APLICADO"
    assert registro["valor_calculado_execucao"] == 100000.0
    assert registro["valor_pago_considerado"] == 90000.0
    assert registro["glosa"] == 10000.0
    assert registro["valor_pago_atualizado"] == 97200.0
    assert registro["retroativo"] == 7200.0


def test_3_glosa_informada_equivale_ao_valor_pago():
    """O resultado economico e o MESMO; so o que o fiscal digitou difere."""
    por_pago = _memoria(ajustes={"C1": _ajuste("Valor pago", 90000.0)})
    por_glosa = _memoria(ajustes={"C1": _ajuste("Glosa", 10000.0)})
    economicos = (
        "status", "valor_calculado_execucao", "valor_pago_considerado",
        "glosa", "fator", "valor_pago_atualizado", "retroativo",
    )
    registro_pago = por_pago["ajustes_execucao_por_ciclo"]["C1"]
    registro_glosa = por_glosa["ajustes_execucao_por_ciclo"]["C1"]
    assert {c: registro_glosa[c] for c in economicos} == {
        c: registro_pago[c] for c in economicos
    }
    # A entrada crua fica preservada como o fiscal informou — nao normalizada.
    assert (registro_pago["tipo_ajuste"], registro_pago["valor_informado"]) == (
        "Valor pago", 90000.0
    )
    assert (registro_glosa["tipo_ajuste"], registro_glosa["valor_informado"]) == (
        "Glosa", 10000.0
    )
    # E toda a cadeia derivada e identica.
    assert por_glosa["ciclos"] == por_pago["ciclos"]
    assert (
        por_glosa["conferencias_metodologicas"]
        == por_pago["conferencias_metodologicas"]
    )
    for campo in ("glosa", "valor_pago_considerado", "valor_pago_atualizado"):
        assert (
            por_glosa["ajustes_execucao_resumo"][campo]
            == por_pago["ajustes_execucao_resumo"][campo]
        )


def test_3b_retroativo_nunca_reajusta_a_parcela_glosada():
    """100.000 x 1,08 - 10.000 = 98.000 seria ERRADO; o correto e 97.200."""
    registro = _memoria(ajustes={"C1": _ajuste("Glosa", 10000.0)})[
        "ajustes_execucao_por_ciclo"
    ]["C1"]
    assert registro["valor_pago_atualizado"] == 97200.0
    assert registro["valor_pago_atualizado"] != 98000.0


# ---------------------------------------------------------------------------
# 4. Zero e valor valido; ausencia nao e zero
# ---------------------------------------------------------------------------

def test_4_valor_pago_zero_e_ajuste_valido():
    memoria = _memoria(ajustes={"C1": _ajuste("Valor pago", 0)})
    registro = memoria["ajustes_execucao_por_ciclo"]["C1"]
    assert registro["status"] == "AJUSTE APLICADO"
    assert registro["valor_pago_considerado"] == 0.0
    assert registro["glosa"] == 100000.0
    assert registro["valor_pago_atualizado"] == 0.0
    assert registro["retroativo"] == 0.0
    assert _consumidos(memoria)["valor_atualizado"] == 0.0


def test_4b_vazio_nao_e_zero():
    """Sem tipo e sem valor o ciclo nao tem ajuste — nao vira glosa integral."""
    memoria = _memoria(ajustes={})
    registro = memoria["ajustes_execucao_por_ciclo"]["C1"]
    assert registro["informado"] is False
    assert registro["status"] == "SEM AJUSTE"
    assert registro["glosa"] is None
    assert _consumidos(memoria)["valor_atualizado"] == 108000.0


# ---------------------------------------------------------------------------
# 5/6 + demais validacoes — invalido nunca vira zero
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    ("bruto", "esperado"),
    [
        (_ajuste("Valor pago", 105000.0), "REVISAR: VALOR PAGO MAIOR QUE O CALCULADO"),
        (_ajuste("Glosa", 110000.0), "REVISAR: GLOSA MAIOR QUE O CALCULADO"),
        (_ajuste("Valor pago", -1.0), "REVISAR: VALOR PAGO NEGATIVO"),
        (_ajuste("Glosa", -1.0), "REVISAR: GLOSA NEGATIVA"),
        (_ajuste("Valor pago", None), "REVISAR: TIPO DE AJUSTE SEM VALOR"),
        (_ajuste("", 90000.0), "REVISAR: VALOR INFORMADO SEM TIPO DE AJUSTE"),
        (_ajuste("Desconto", 90000.0), "REVISAR: TIPO DE AJUSTE INVALIDO"),
        (_ajuste("Glosa", "dez mil"), "REVISAR: VALOR INFORMADO NAO NUMERICO"),
    ],
)
def test_5_6_entradas_invalidas_viram_revisar_e_fecham_o_metodo(bruto, esperado):
    memoria = _memoria(ajustes={"C1": bruto})
    registro = memoria["ajustes_execucao_por_ciclo"]["C1"]
    assert registro["status"] == esperado
    # Nunca vira zero e nunca e ignorado em silencio: fecha o metodo inteiro.
    assert registro["valor_pago_considerado"] is None
    assert registro["glosa"] is None
    conferencia = _conferencia(memoria)
    assert conferencia["disponivel"] is False
    assert conferencia["executado_atualizado"] is None
    assert conferencia["valor_total_atualizado"] is None
    assert memoria["ajustes_execucao_resumo"]["status"] == "REVISAR"


def test_5b_ajuste_em_ciclo_sem_execucao_calculada_e_revisar():
    """C3 tem fator novo valido, mas nenhum consumo informado."""
    memoria = _montar_memoria_por_ciclo(
        _leitura(
            ajustes={"C3": _ajuste("Glosa", 100.0)},
            fatores={"C0": 1.0, "C1": FATOR_C1, "C3": 1.1},
            percentuais={"C1": 0.08, "C3": 0.02},
        ),
        {}, [], {},
    )
    registro = memoria["ajustes_execucao_por_ciclo"]["C3"]
    assert registro["fator"] == 1.02  # o fator novo existe...
    # ...mas nao ha execucao naquele ciclo para comparar com a glosa.
    assert registro["status"] == "REVISAR: SEM EXECUCAO CALCULADA NO CICLO"


# ---------------------------------------------------------------------------
# MULTICICLO — o reajuste anterior JA formalizado nao pode ser reaplicado
# ---------------------------------------------------------------------------

# C1 formalizado (8%, FORA da apuracao) + C2 em apuracao (5%):
#   F (acumulado)   = 1,08 x 1,05 = 1,134
#   D (fator novo)  = 1,05
#   F / D           = 1,08   <- ja embutido no que foi pago
#   F - 1   = 0,134  != F - F/D = 0,054   (o cenario so prova se forem !=)
FATOR_ACUM_C2 = 1.134
FATOR_NOVO_C2 = 1.05

_ITEM_C2 = {
    "item": "I1",
    "qtd_contratada": 1500,
    "vu_original": 100.0,
    "qtd_total": 1000,
    "check": "OK",
    "consumos": {
        "C0": {"qtd": None, "valor": None},
        "C1": {"qtd": None, "valor": None},
        "C2": {"qtd": 1000, "valor": 113400.0},  # 1000 x 100 x 1,134
        "C3": {"qtd": None, "valor": None},
        "C4": {"qtd": None, "valor": None},
    },
}


def _leitura_multiciclo(ajustes=None):
    return _leitura(
        itens=[_ITEM_C2],
        ajustes=ajustes,
        fatores={"C0": 1.0, "C1": FATOR_C1, "C2": FATOR_ACUM_C2},
        percentuais={"C1": 0.08, "C2": 0.05},
        computar={"C1": "Nao", "C2": "Sim"},  # C1 ja formalizado
    )


def test_multiciclo_o_cenario_de_fato_distingue_as_duas_formulas():
    """Guarda do proprio teste: se F-1 == F-F/D, ele nao provaria nada."""
    assert round(FATOR_ACUM_C2 - 1, 3) == 0.134
    assert round(FATOR_ACUM_C2 - FATOR_ACUM_C2 / FATOR_NOVO_C2, 3) == 0.054


def test_multiciclo_valor_calculado_esta_na_base_vigente_anterior():
    memoria = _montar_memoria_por_ciclo(_leitura_multiciclo(), {}, [], {})
    registro = memoria["ajustes_execucao_por_ciclo"]["C2"]
    # 1000 x 100 x (1,134 / 1,05) = 108.000 — e nao 100.000.
    assert registro["valor_calculado_execucao"] == 108000.0
    # O fator publicado e o do reajuste NOVO, nao o acumulado.
    assert registro["fator"] == FATOR_NOVO_C2


def test_multiciclo_nao_reaplica_o_reajuste_anterior():
    memoria = _montar_memoria_por_ciclo(
        _leitura_multiciclo({"C2": _ajuste("Glosa", 10000.0)}), {}, [], {}
    )
    registro = memoria["ajustes_execucao_por_ciclo"]["C2"]
    assert registro["status"] == "AJUSTE APLICADO"
    assert registro["valor_calculado_execucao"] == 108000.0
    assert registro["valor_pago_considerado"] == 98000.0
    assert registro["glosa"] == 10000.0
    assert registro["valor_pago_atualizado"] == 102900.0   # 98.000 x 1,05
    assert registro["retroativo"] == 4900.0                # 98.000 x 0,05
    # O erro que este teste existe para impedir: 98.000 x 1,134 = 111.132,
    # com retroativo 13.132 — cobraria os 8% de C1 uma segunda vez.
    assert registro["valor_pago_atualizado"] != 111132.0
    assert registro["retroativo"] != 13132.0


def test_multiciclo_valor_pago_e_glosa_continuam_equivalentes():
    por_glosa = _montar_memoria_por_ciclo(
        _leitura_multiciclo({"C2": _ajuste("Glosa", 10000.0)}), {}, [], {}
    )["ajustes_execucao_por_ciclo"]["C2"]
    por_pago = _montar_memoria_por_ciclo(
        _leitura_multiciclo({"C2": _ajuste("Valor pago", 98000.0)}), {}, [], {}
    )["ajustes_execucao_por_ciclo"]["C2"]
    for campo in (
        "valor_calculado_execucao", "valor_pago_considerado", "glosa",
        "fator", "valor_pago_atualizado", "retroativo",
    ):
        assert por_glosa[campo] == por_pago[campo], campo


def test_multiciclo_retroativo_converge_com_a_formula_homologada_do_xls():
    """AF (bloco) == D12 (memoria) == base_pago x (D - 1).

    A formula homologada e `base x (F - F/D)`, que e identicamente
    `(base x F/D) x (D - 1)`. Com o ajuste, o primeiro fator vira o valor pago
    considerado — e as duas superficies publicam o MESMO retroativo.
    """
    memoria = _montar_memoria_por_ciclo(
        _leitura_multiciclo({"C2": _ajuste("Glosa", 10000.0)}), {}, [], {}
    )
    registro = memoria["ajustes_execucao_por_ciclo"]["C2"]
    homologada = round(
        registro["valor_pago_considerado"] * (FATOR_NOVO_C2 - 1), 2
    )
    assert registro["retroativo"] == homologada == 4900.0
    # A cadeia canonica publica o mesmo numero — uma so medida de retroativo.
    assert _consumidos(memoria, "C2")["retroativo"] == 4900.0
    assert memoria["ajustes_execucao_resumo"]["retroativo"] == 4900.0


def test_multiciclo_sem_ajuste_nao_muda_nada():
    base = _montar_memoria_por_ciclo(_leitura_multiciclo(), {}, [], {})
    bloco = _consumidos(base, "C2")
    assert bloco["base_original"] == 100000.0
    assert bloco["valor_atualizado"] == 113400.0
    assert base["ajustes_execucao_resumo"]["status"] == "SEM AJUSTE"


def test_multiciclo_execucao_atualizada_nao_muda_de_base():
    """Sem glosa, valor_pago_considerado x D volta a ser base x F."""
    sem = _consumidos(
        _montar_memoria_por_ciclo(_leitura_multiciclo(), {}, [], {}), "C2"
    )["valor_atualizado"]
    com = _consumidos(
        _montar_memoria_por_ciclo(
            _leitura_multiciclo({"C2": _ajuste("Glosa", 0)}), {}, [], {}
        ),
        "C2",
    )["valor_atualizado"]
    assert sem == 113400.0
    assert com == 113400.0  # glosa zero nao pode deslocar a base monetaria


def test_ajuste_em_ciclo_fora_da_apuracao_e_fail_closed():
    """C1 esta formalizado: nao ha reajuste novo para incidir sobre o pago.

    O ciclo TEM execucao (senao o gate de execucao dispararia antes), mas
    COMPUTAR_NESTA_APURACAO=Nao — entao nao existe fator novo e o ajuste nao
    pode ser aplicado. Fail-closed, sem inventar reajuste.
    """
    item = dict(_ITEM_C2)
    item["qtd_contratada"] = 2500
    item["qtd_total"] = 2000
    item["consumos"] = dict(_ITEM_C2["consumos"])
    item["consumos"]["C1"] = {"qtd": 1000, "valor": 108000.0}
    leitura = _leitura(
        itens=[item],
        ajustes={"C1": _ajuste("Glosa", 10.0)},
        fatores={"C0": 1.0, "C1": FATOR_C1, "C2": FATOR_ACUM_C2},
        percentuais={"C1": 0.08, "C2": 0.05},
        computar={"C1": "Nao", "C2": "Sim"},
    )
    memoria = _montar_memoria_por_ciclo(leitura, {}, [], {})
    registro = memoria["ajustes_execucao_por_ciclo"]["C1"]
    assert registro["valor_calculado_execucao"] is None
    assert registro["status"] == "REVISAR: CICLO FORA DA APURACAO OU SEM FATOR NOVO"
    assert _conferencia(memoria)["disponivel"] is False


# ---------------------------------------------------------------------------
# 7. Dois ciclos — so o ajustado muda
# ---------------------------------------------------------------------------

def test_7_dois_ciclos_apenas_o_ajustado_muda():
    fatores = {"C0": 1.0, "C1": FATOR_C1, "C2": FATOR_C2}
    itens = [_item_dois_ciclos()]
    base = _montar_memoria_por_ciclo(
        _leitura(itens=itens, fatores=fatores), {}, [], {}
    )
    ajustada = _montar_memoria_por_ciclo(
        _leitura(
            itens=itens,
            fatores=fatores,
            ajustes={"C1": _ajuste("Glosa", 10000.0)},
        ),
        {}, [], {},
    )
    # C2 nao foi tocado.
    assert _consumidos(ajustada, "C2") == _consumidos(base, "C2")
    assert ajustada["ajustes_execucao_por_ciclo"]["C2"]["status"] == "SEM AJUSTE"
    # C1 mudou exatamente pelo ajuste.
    assert _consumidos(base, "C1")["valor_atualizado"] == 108000.0
    assert _consumidos(ajustada, "C1")["valor_atualizado"] == 97200.0


# ---------------------------------------------------------------------------
# 8/9. Retroativo e VTA
# ---------------------------------------------------------------------------

def test_8_retroativo_do_ciclo_usa_o_valor_pago_considerado():
    bloco = _consumidos(_memoria(ajustes={"C1": _ajuste("Glosa", 10000.0)}))
    assert bloco["base_original"] == 90000.0
    assert bloco["valor_atualizado"] == 97200.0
    assert bloco["retroativo"] == 7200.0
    # A evidencia continua sendo a mesma execucao fisica.
    assert bloco["evidencias"] == 1


def test_9_vta_incorpora_a_execucao_considerada_sem_dupla_contagem():
    base = _memoria()
    ajustada = _memoria(ajustes={"C1": _ajuste("Glosa", 10000.0)})
    assert _conferencia(base)["valor_total_atualizado"] == 162000.0
    conferencia = _conferencia(ajustada)
    assert conferencia["executado_atualizado"] == 97200.0
    assert conferencia["potencial_restante_atualizado"] == 54000.0
    # 97.200 + 54.000 — o retroativo NAO entra outra vez.
    assert conferencia["valor_total_atualizado"] == 151200.0
    assert ajustada["vta"]["valor_total_atualizado"] == 151200.0


def test_9b_resumo_canonico_e_a_unica_fonte_dos_agregados():
    resumo = _memoria(ajustes={"C1": _ajuste("Glosa", 10000.0)})[
        "ajustes_execucao_resumo"
    ]
    assert resumo["status"] == "AJUSTE APLICADO"
    assert resumo["ciclos_ajustados"] == ["C1"]
    assert resumo["valor_calculado_execucao"] == 100000.0
    assert resumo["glosa"] == 10000.0
    assert resumo["valor_pago_considerado"] == 90000.0
    assert resumo["valor_pago_atualizado"] == 97200.0
    assert resumo["retroativo"] == 7200.0


# ---------------------------------------------------------------------------
# 10/11. Quantidades e remanescente fisico intactos
# ---------------------------------------------------------------------------

def test_10_quantidades_consumidas_inalteradas():
    itens = [dict(_ITEM)]
    _montar_memoria_por_ciclo(
        _leitura(itens=itens, ajustes={"C1": _ajuste("Glosa", 10000.0)}),
        {}, [], {},
    )
    assert itens[0]["qtd_total"] == 1000
    assert itens[0]["qtd_contratada"] == 1500
    assert itens[0]["consumos"]["C1"]["qtd"] == 1000


def test_11_remanescente_fisico_identico_com_e_sem_glosa():
    base = _conferencia(_memoria())["potencial_restante_atualizado"]
    for informado in (0, 10000.0, 100000.0):
        ajustada = _memoria(ajustes={"C1": _ajuste("Glosa", informado)})
        assert (
            _conferencia(ajustada)["potencial_restante_atualizado"] == base
        ), f"glosa de {informado} mexeu no remanescente"


# ---------------------------------------------------------------------------
# 12/13. Financeiro e PC preservados
# ---------------------------------------------------------------------------

_PARCELA_FINANCEIRO = {
    "fonte_parcela": "Financeiro",
    "identificador": "financeiro:2026-01",
    "ciclo": "C1",
    "valor": 5000.0,
    "valor_atualizado": 5400.0,
}
_PC = {
    "elegivel_retroativo_pc": True,
    "ciclo_calculado": "C1",
    "efeito_financeiro_pc": "Sim",
    "valor_pago": 3000.0,
}


def test_12_13_financeiro_e_pc_intactos_sob_glosa():
    sem = _montar_memoria_por_ciclo(
        _leitura(financeiro=[_PARCELA_FINANCEIRO]), {}, [_PC], {}
    )
    com = _montar_memoria_por_ciclo(
        _leitura(
            financeiro=[_PARCELA_FINANCEIRO],
            ajustes={"C1": _ajuste("Glosa", 10000.0)},
        ),
        {}, [_PC], {},
    )
    for metodo in ("financeiro", "pc"):
        assert (
            _ciclo(com, "C1")["retroativo"][metodo]
            == _ciclo(sem, "C1")["retroativo"][metodo]
        ), metodo
        assert _conferencia(com, metodo) == _conferencia(sem, metodo), metodo
    assert com["controle_pcs"] == sem["controle_pcs"]


# ---------------------------------------------------------------------------
# Leitor — compatibilidade e os dois unicos campos manuais
# ---------------------------------------------------------------------------

def _planilha(cabecalhos, linhas):
    wb = Workbook()
    ws = wb.active
    for coluna, titulo in enumerate(cabecalhos, start=1):
        ws.cell(1, coluna).value = titulo
    for offset, valores in enumerate(linhas, start=2):
        for coluna, valor in enumerate(valores, start=1):
            ws.cell(offset, coluna).value = valor
    return ws


def test_leitor_xls_antigo_sem_o_bloco_devolve_ausencia_de_ajuste():
    ws = _planilha(["ITEM", "QTD_CONTRATADA"], [["I1", 10]])
    assert _ler_ajustes_execucao_consumidos(
        ws, _mapear_colunas_por_cabecalho(ws)
    ) == {}


def test_leitor_captura_apenas_os_dois_campos_manuais():
    ws = _planilha(
        ["AJUSTE_CICLO", "AJUSTE_TIPO", "AJUSTE_VALOR_INFORMADO"],
        [
            ["C0", None, None],       # vazio real: nao entra
            ["C1", "Glosa", 10000.0],
            ["C2", "Valor pago", 0],  # zero explicito: entra
        ],
    )
    lido = _ler_ajustes_execucao_consumidos(ws, _mapear_colunas_por_cabecalho(ws))
    assert set(lido) == {"C1", "C2"}
    assert lido["C1"]["tipo"] == "Glosa"
    assert lido["C1"]["valor_bruto"] == 10000.0
    assert lido["C2"]["tipo"] == "Valor pago"
    assert lido["C2"]["valor_vazio"] is False


def test_leitor_preserva_entrada_invalida_em_vez_de_descartar():
    ws = _planilha(
        ["AJUSTE_CICLO", "AJUSTE_TIPO", "AJUSTE_VALOR_INFORMADO"],
        [["C1", "Desconto", "dez mil"]],
    )
    lido = _ler_ajustes_execucao_consumidos(ws, _mapear_colunas_por_cabecalho(ws))
    assert lido["C1"]["tipo"] is None
    assert lido["C1"]["tipo_bruto"] == "Desconto"
    assert lido["C1"]["valor_bruto"] == "dez mil"


# ---------------------------------------------------------------------------
# Template oficial — estrutura do bloco e travas das formulas reescritas
# ---------------------------------------------------------------------------

_CABECALHOS_BLOCO = {
    "X": "AJUSTE_CICLO",
    "Y": "AJUSTE_VALOR_CALCULADO",
    "Z": "AJUSTE_TIPO",
    "AA": "AJUSTE_VALOR_INFORMADO",
    "AB": "AJUSTE_VALOR_PAGO_CONSIDERADO",
    "AC": "AJUSTE_GLOSA",
    "AD": "AJUSTE_FATOR_NOVO",
    "AE": "AJUSTE_VALOR_PAGO_ATUALIZADO",
    "AF": "AJUSTE_RETROATIVO",
    "AG": "AJUSTE_STATUS",
}


@pytest.fixture(scope="module")
def workbook():
    wb = load_workbook(TEMPLATE)
    yield wb
    wb.close()


def test_xls_bloco_lateral_nao_invade_a_grade_item_a_item(workbook):
    ws = workbook["itens_Consumidos"]
    for coluna, titulo in _CABECALHOS_BLOCO.items():
        assert ws[f"{coluna}1"].value == titulo
    # W permanece vazia como separador da grade por item (A:V).
    assert all(ws[f"W{linha}"].value in (None, "") for linha in range(1, 201))
    for offset, ciclo in enumerate(("C0", "C1", "C2", "C3", "C4")):
        assert ws.cell(2 + offset, 24).value == ciclo


def test_xls_campos_manuais_ficam_vazios_e_com_lista_de_opcoes(workbook):
    ws = workbook["itens_Consumidos"]
    for linha in range(2, 7):
        assert ws[f"Z{linha}"].value in (None, "")
        assert ws[f"AA{linha}"].value in (None, "")
    listas = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    assert any("Z2:Z6" in str(dv.sqref) for dv in listas)


def test_xls_base_da_glosa_nunca_e_valor_cons(workbook):
    """VALOR_CONS_Cn ja embute o reajuste em apuracao e nunca serve de base.

    O que entra na comparacao e QTD x VU_ORIGINAL convertido para a base
    monetaria vigente anterior pelo fator F/D — a conversao em si esta
    coberta por test_xls_valor_calculado_esta_na_base_vigente_anterior.
    Aqui so se prova que o SUMPRODUCT parte de VU_ORIGINAL, nao de VALOR_CONS.
    """
    ws = workbook["itens_Consumidos"]
    for coluna_qtd, linha in (("E", 2), ("G", 3), ("I", 4), ("K", 5), ("M", 6)):
        formula = ws[f"Y{linha}"].value
        assert f"SUMPRODUCT(${coluna_qtd}$2:${coluna_qtd}$200,$C$2:$C$200)" in formula


def test_xls_f20_preserva_o_ramo_sem_ajuste_de_cada_ciclo(workbook):
    f20 = workbook["MEMORIA_RESULTADOS"]["F20"].value
    for coluna in ("F", "H", "J", "L", "N"):
        assert f"SUM(itens_Consumidos!${coluna}$2:${coluna}$200)" in f20
    for linha in range(2, 7):
        assert f'itens_Consumidos!$AG${linha}="AJUSTE APLICADO"' in f20
    # Fail-closed: qualquer REVISAR fecha a medida, nunca a transforma em zero.
    assert 'COUNTIF(itens_Consumidos!$AG$2:$AG$6,"REVISAR*")>0' in f20


def test_xls_retroativo_itens_encaixa_o_ajuste_na_regra_homologada(workbook):
    """O ramo sem ajuste fica intacto; o ajustado e a MESMA conta.

        base x (F - F/D)  ==  (base x F/D) x (D - 1)

    O ramo ajustado troca `base x F/D` pelo valor pago considerado (AB), que
    esta na mesma base monetaria — e aplica so `(D - 1)`, o reajuste NOVO.
    Aplicar `(F - F/D)` sobre AB reajustaria duas vezes os ciclos anteriores.
    """
    mem = workbook["MEMORIA_RESULTADOS"]
    for ciclo in range(1, 5):
        formula = mem[f"D{10 + ciclo}"].value
        fator = f"parametros!$F{ciclo + 2}"
        apuracao = f"parametros!$D{11 + ciclo}"
        # Ramo legado: a expressao homologada, uma unica vez e intacta.
        assert formula.count(f"({fator}-{fator}/{apuracao})") == 1
        # Ramo ajustado: valor pago considerado x (fator novo - 1).
        assert (
            f"ROUND(itens_Consumidos!$AB${ciclo + 2}*({apuracao}-1),2)"
            in formula
        )
        # E o acumulado NUNCA multiplica o valor pago.
        assert f"$AB${ciclo + 2}*({fator}" not in formula
        assert 'COUNTIF(itens_Consumidos!$AG$2:$AG$6,"REVISAR*")>0' in formula


def test_xls_valor_calculado_esta_na_base_vigente_anterior(workbook):
    """Y = QTD x VU x (F / D): a base em que o valor foi efetivamente pago."""
    ws = workbook["itens_Consumidos"]
    for ciclo, (qtd, linha) in enumerate(
        (("E", 2), ("G", 3), ("I", 4), ("K", 5), ("M", 6))
    ):
        formula = ws[f"Y{linha}"].value
        fator = f"parametros!$F{ciclo + 2}"
        assert f"SUMPRODUCT(${qtd}$2:${qtd}$200,$C$2:$C$200)*{fator}/$AD{linha}" in formula
    # E o fator publicado e o do reajuste NOVO, nao o acumulado da coluna U.
    for linha in range(2, 7):
        assert ws[f"AD{linha}"].value == (
            f"=IF(ISNUMBER(parametros!$D{9 + linha}),parametros!$D{9 + linha},\"\")"
        )
        assert "$U$" not in ws[f"AD{linha}"].value


def test_xls_b26_e_remanescente_do_ramo_itens_intactos(workbook):
    """A trava do enunciado: B26/Itens = F20 + D35 + ajustes, sem somar B21."""
    mem = workbook["MEMORIA_RESULTADOS"]
    ramo_itens = (
        'IF(OR($F$20="",D35="",AND(B24<>"",NOT(ISNUMBER(B24)))),"",'
        'ROUND($F$20+D35+IF(ISNUMBER(B24),B24,0)+IF(ISNUMBER($N$263),$N$263,0),2))'
    )
    assert ramo_itens in mem["B26"].value
    assert mem["D35"].value == (
        '=IF($B$4="PCs",$T$23,IF($B$4="Financeiro",D32,IF($B$4="Itens",D33,"")))'
    )
    # C33/D33 (remanescente derivado dos itens) nao foram tocados.
    assert "itens_Consumidos!$V$2:$V$200" in mem["C33"].value
    assert "AJUSTE" not in (mem["C33"].value or "")
    assert "AJUSTE" not in (mem["D33"].value or "")


def test_xls_resultados_nao_foi_tocada(workbook):
    """A aba RESULTADOS nao tem linha visivel livre — e nao recebeu nada.

    Todas as suas linhas vazias entre 1 e 87 sao ancoras de leiaute de
    frentes anteriores: 8/14/23/32/39/52 sao separadores brancos geridos pela
    Etapa 50.3, 31/40/51 sao linhas OCULTAS de 7pt e 78 fecha a tabela 8
    (test_resultados_final_1.py::test_tabela_8_padronizada exige A78 vazia).
    Abaixo da 87 esta a camada que o rollback da UX2 removeu.

    Nenhuma delas foi povoada por esta frente — a aba continua exibindo o
    ajuste sozinha, porque F20 alimenta o executado apurado, o retroativo e
    o VTA que ela ja mostra.
    """
    res = workbook["RESULTADOS"]
    assert res.max_row == 87
    assert not [
        celula.coordinate
        for linha in res.iter_rows(min_row=88, max_row=200)
        for celula in linha
        if celula.value is not None
    ]
    for linha in (8, 14, 23, 31, 32, 39, 40, 51, 52, 78):
        for coluna in "ABCDEFGHIJ":
            valor = str(res[f"{coluna}{linha}"].value or "")
            assert "itens_Consumidos!$A" not in valor
            assert "$T$71" not in valor, f"{coluna}{linha} virou faixa de glosa"
    assert res["A78"].value is None
    for linha in (31, 40, 51):
        assert res.row_dimensions[linha].hidden is True


def test_xls_medidas_do_ajuste_vivem_na_memoria_auditavel(workbook):
    """As seis medidas do ajuste sao publicadas em MEMORIA_RESULTADOS."""
    mem = workbook["MEMORIA_RESULTADOS"]
    assert mem["S69"].value == (
        "AJUSTES DA EXECUCAO - VALOR PAGO / GLOSA (metodo Itens)"
    )
    rotulos = [mem[f"S{linha}"].value for linha in range(70, 76)]
    assert rotulos == [
        "Valor calculado da execucao (ciclos ajustados)",
        "Glosa total",
        "Valor pago considerado",
        "Valor pago atualizado",
        "Retroativo do valor pago",
        "Status dos ajustes",
    ]
    for linha in range(70, 75):
        assert mem[f"T{linha}"].value.startswith(
            '=IF($T$75<>"AJUSTE APLICADO","",'
        ), linha
    assert 'COUNTIF(itens_Consumidos!$AG$2:$AG$6,"REVISAR*")>0' in mem["T75"].value
