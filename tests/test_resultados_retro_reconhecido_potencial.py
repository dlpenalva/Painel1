"""Coleta 12.1 — RESULTADOS: card RETRO. RECONHECIDO + POTENCIAL e situacao.

Somente apresentacao. Dois blocos:

* ESTRUTURA (CI): ordem dos cards, fontes do card novo (as MESMAS dos cards
  RETROATIVO RECONHECIDO e RETROATIVO POTENCIAL), soma sem nome definido e sem
  leitores, faixa da situacao com tres estados e formatacao condicional que so
  le a propria aba.
* EXCEL REAL (RUN_EXCEL_INTEGRATION=1): o caso de uso da 11.8/12.0 (PCs C0..C2,
  um PC nao pago), com PC_PAGO e a lista de PCs variados, gera por calculo real
  os casos reconhecido/potencial > 0 ou = 0, o bloqueio REVISE e os tres
  estados da situacao; VTA, reconhecido e potencial sao comparados com o
  template da Coleta 12.0 (A/B) e o arquivo abre, salva e reabre sem reparo.
"""
from __future__ import annotations

import gc
import io
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_resultados_retro_reconhecido_potencial as novo  # noqa: E402
from _resultados_abas import ABA_RESULTADOS  # noqa: E402

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
# Merge do PR #188 (ultima Coleta 12.0): base do A/B no Excel real.
COMMIT_BASE_120 = "7c124052384a92b013b35bfe2d7026f4ed3e780e"

_CACHE: dict[str, object] = {}


def _template():
    if "template" not in _CACHE:
        _CACHE["template"] = load_workbook(TEMPLATE)
    return _CACHE["template"]


def _gerado():
    if "gerado" not in _CACHE:
        from _coleta_oficial import obter_coleta_oficial_bytes

        _CACHE["gerado"] = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    return _CACHE["gerado"]


# =========================================================================== #
# ESTRUTURA
# =========================================================================== #
def test_cards_na_ordem_pedida():
    for wb in (_template(), _gerado()):
        ws = wb[ABA_RESULTADOS]
        assert [ws[f"{c}8"].value for c in "BCDEFG"] == [
            "VTA OFICIAL",
            "RETROATIVO RECONHECIDO",
            "RETROATIVO POTENCIAL",
            "RETRO. RECONHECIDO + POTENCIAL",
            "SALDO REMANESCENTE",
            "ÍNDICE ACUMULADO",
        ]


def test_cards_existentes_mantem_as_mesmas_fontes():
    ws = _template()[ABA_RESULTADOS]
    assert ws["B9"].value == "=IF(ISNUMBER(VTA_FINAL),VTA_FINAL,MEMORIA_RESULTADOS!$AG$2)"
    assert ws["C9"].value == (
        "=IF(ISNUMBER(RESULTADOS_DETALHE!$D$22),RESULTADOS_DETALHE!$D$22,"
        "MEMORIA_RESULTADOS!$AG$2)"
    )
    assert ws["D9"].value == (
        '=IF(NOT(METODO_RETROATIVO="PCs"),MEMORIA_RESULTADOS!$AG$2,'
        "IF(ISNUMBER(RETROATIVO_POTENCIAL_VTA),RETROATIVO_POTENCIAL_VTA,"
        "MEMORIA_RESULTADOS!$AG$2))"
    )
    assert "SALDO_REMANESCENTE_ATUAL" in ws["F9"].value
    assert ws["G9"].value == "=F6"


def test_card_novo_espelha_a_soma_da_camada_tecnica():
    ws = _template()[ABA_RESULTADOS]
    assert ws["E9"].value == (
        "=IF(ISNUMBER(MEMORIA_RESULTADOS!$T$66),MEMORIA_RESULTADOS!$T$66,"
        "MEMORIA_RESULTADOS!$AG$2)"
    )
    mem = _template()["MEMORIA_RESULTADOS"]
    assert mem["T66"].value == novo.FORMULAS_MEMORIA["T66"]
    assert mem["AG65"].value == "Soma do retroativo reconhecido com a parcela potencial."


def test_soma_usa_somente_as_fontes_dos_dois_cards():
    """reconhecido + potencial: nenhuma outra aba, nenhum recalculo."""
    t66 = _template()["MEMORIA_RESULTADOS"]["T66"].value
    assert set(re.findall(r"[A-Za-z_]+!\$?[A-Z]+\$?\d+", t66)) == {
        "RESULTADOS_DETALHE!$D$22"}
    nomes = set(re.findall(r"\b[A-Z][A-Z_]{4,}\b", re.sub(r'"[^"]*"', "", t66)))
    assert nomes - {"RESULTADOS_DETALHE", "ISNUMBER", "ROUND"} == {
        "METODO_RETROATIVO", "RETROATIVO_POTENCIAL_VTA"}
    assert "ROUND(RESULTADOS_DETALHE!$D$22+RETROATIVO_POTENCIAL_VTA,2)" in t66
    for aba in ("itens_PC", "financeiro", "itens_Consumidos", "itens_Remanesc",
                "CICLO_EM_EXECUCAO", "aditivos", "historico_VU"):
        assert aba not in t66


def test_soma_nao_tem_nome_nem_leitores():
    """A soma-resumo nao vira parcela: nenhum nome definido, nenhuma formula
    alem do card E9 e nenhum leitor Python apontam para T66."""
    wb = _template()
    for nome, definicao in wb.defined_names.items():
        assert "$T$66" not in definicao.attr_text, nome
    leitores = [
        (ws.title, c.coordinate)
        for ws in wb.worksheets for linha in ws.iter_rows() for c in linha
        if isinstance(c.value, str) and c.value.startswith("=") and (
            "$T$66" in c.value
            or (ws.title == "MEMORIA_RESULTADOS" and re.search(r"\$?T\$?66\b", c.value)))
    ]
    assert leitores == [(ABA_RESULTADOS, "E9")]
    for arquivo in RAIZ.glob("*.py"):
        texto = arquivo.read_text(encoding="utf-8", errors="ignore")
        assert not re.search(r"""["']T66["']""", texto), arquivo.name


def test_faixa_da_situacao_tem_tres_estados():
    ws = _template()[ABA_RESULTADOS]
    assert ws["B12"].value == "SITUAÇÃO DA APURAÇÃO"
    assert ws["C12"].value == (
        '=IF(LEFT(RESULTADOS_DETALHE!$A$7,4)="APUR",RESULTADOS_DETALHE!$A$7,'
        "MEMORIA_RESULTADOS!$AG$67)"
    )
    assert ws["C13"].value == "=MEMORIA_RESULTADOS!$T$67"
    mem = _template()["MEMORIA_RESULTADOS"]
    assert mem["AG67"].value == "REVISE A APURAÇÃO"
    assert mem["AG68"].value == (
        "Existem pendências que precisam ser corrigidas antes da conclusão.")
    assert mem["AG69"].value + "R$ 119,20" + mem["AG70"].value == (
        "Nenhuma pendência impeditiva identificada. Há R$ 119,20 de retroativo "
        "potencial, sujeito à confirmação."
    )
    t67 = mem["T67"].value
    # A regra de pendencia continua sendo a do detalhe (A7: $J$5 = 0).
    assert 'LEFT(RESULTADOS_DETALHE!$A$7,4)<>"APUR"' in t67
    assert "FIXED(RETROATIVO_POTENCIAL_VTA,2)" in t67 and "TEXT(" not in t67
    assert t67.count("RESULTADOS_DETALHE!$B$7") == 4
    assert {"C12:G12", "C13:G13"} <= {str(m) for m in ws.merged_cells.ranges}


def test_cores_da_situacao_so_leem_a_propria_aba():
    for wb in (_template(), _gerado()):
        ws = wb[ABA_RESULTADOS]
        regras = {
            regra.formula[0]: regra.dxf.fill.bgColor.rgb
            for faixa in ws.conditional_formatting if str(faixa.sqref) == "B12:G13"
            for regra in faixa.rules
        }
        assert regras == {
            'LEFT($C$12,4)="REVI"': "FFFFC7CE",
            'AND(LEFT($C$12,4)="APUR",N($D$9)>0)': "FFEEF3CF",
            'AND(LEFT($C$12,4)="APUR",N($D$9)<=0)': "FFC6EFCE",
        }
        assert "G8:G10" not in {str(f.sqref) for f in ws.conditional_formatting}


def test_descritor_bate_com_template_e_coleta_gerada():
    esperado = novo.formulas_executivo()
    for wb in (_template(), _gerado()):
        ws = wb[ABA_RESULTADOS]
        formulas = {
            c.coordinate: c.value for linha in ws.iter_rows() for c in linha
            if isinstance(c.value, str) and c.value.startswith("=")
        }
        assert formulas == esperado
    novo.validar_ascii_e_parenteses(esperado)
    novo.validar_ascii_e_parenteses(
        {k: v for k, v in novo.FORMULAS_MEMORIA.items() if v.startswith("=")})


def test_formatos_cor_e_area_de_impressao():
    ws = _template()[ABA_RESULTADOS]
    for endereco in ("B9", "C9", "D9", "E9", "F9"):
        assert "R$" in ws[endereco].number_format, endereco
    assert ws["G9"].number_format.startswith("0.00%")
    assert ws["E9"].fill.fgColor.rgb == "FFEDE7F6"      # lilas muito claro
    assert str(ws.print_area).replace("'", "").endswith("RESULTADOS!$B$1:$G$74")


# =========================================================================== #
# EXCEL REAL
# =========================================================================== #
excel_real = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar o Excel COM",
)

VERDE, AMARELO_ESVERDEADO, ROSA = "C6EFCE", "EEF3CF", "FFC7CE"


def _cor(rgb: str) -> int:
    return int(rgb[4:6], 16) * 65536 + int(rgb[2:4], 16) * 256 + int(rgb[0:2], 16)


def _preencher(livro, metodo: str, pcs, financeiro: bool) -> None:
    """Caso de uso 11.8/12.0 (`test_coleta_120_excel_real._preencher`) com a
    lista de PCs variavel: (numero, data, valor, pago)."""
    from pywintypes import Time

    from test_coleta_118_excel_real import CORTE

    livro.Worksheets("CONTROLE").Range("B1").Value = metodo
    livro.Worksheets("CONTROLE").Range("B3").Value = Time(CORTE)
    itens = livro.Worksheets("itens_Remanesc")
    itens.Range("A2").Value = "1.1"
    itens.Range("B2").Value = 1000
    itens.Range("C2").Value = 10
    itens.Range("E2").Value = 950
    aba = livro.Worksheets("itens_PC")
    for linha, (numero, data_pc, valor, pago) in enumerate(pcs, start=2):
        aba.Range(f"A{linha}").Value = numero
        aba.Range(f"B{linha}").Value = Time(data_pc)
        aba.Range(f"D{linha}").Value = valor
        aba.Range(f"G{linha}").Value = pago
    if financeiro:
        fin = livro.Worksheets("financeiro")
        for linha in range(2, 74):
            competencia = fin.Range(f"A{linha}").Value
            if hasattr(competencia, "year") and competencia.replace(tzinfo=None) <= CORTE:
                fin.Range(f"C{linha}").Value = 50
    ciclo = livro.Worksheets("CICLO_EM_EXECUCAO")
    ciclo.Range("D5").Value2 = livro.Worksheets("CONTROLE").Range("B3").Value2
    livro.Application.CalculateFullRebuild()
    ciclo.Range("C13").Value = 800


def _pcs(pagos: dict[str, str] | None = None, *, so=None):
    """PCs do fixture 11.8 (PC-C0, PC-C1, PC-C1-NAO, PC-C2) com PC_PAGO trocado."""
    from test_coleta_118_excel_real import PCS

    pagos = pagos or {}
    return tuple(
        (numero, data_pc, valor, pagos.get(numero, pago))
        for numero, data_pc, valor, pago in PCS
        if so is None or numero in so
    )


def _rodar(tmp_path: Path, nome: str, metodo: str, *, pcs=(), financeiro=False,
           template: Path | None = None, monkeypatch=None) -> dict:
    import pythoncom
    import win32com.client

    import _coleta_oficial
    from test_coleta_118_excel_real import _dados, _tentar

    if template is not None:
        monkeypatch.setattr(_coleta_oficial, "TEMPLATE_COLETA_OFICIAL", template)
    caminho = tmp_path / f"Coleta_{nome}.xlsx"
    try:
        caminho.write_bytes(_coleta_oficial.gerar_coleta_oficial_preenchida(_dados()))
    finally:
        if template is not None:
            monkeypatch.undo()
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    livro = None

    def ler():
        mem = livro.Worksheets("MEMORIA_RESULTADOS")
        det = livro.Worksheets("RESULTADOS_DETALHE")
        exe = livro.Worksheets("RESULTADOS")
        r = {c: mem.Range(c).Value2 for c in ("B16", "B26", "T39", "T40", "T66", "T67")}
        r.update({c: det.Range(c).Value2 for c in ("D22", "J5", "A7", "B7")})
        r.update({f"exe_{c}": exe.Range(c).Value2 for c in (
            "B9", "C9", "D9", "E9", "F9", "G9", "E10", "C12", "C13")})
        r["cor_faixa"] = int(exe.Range("B12").DisplayFormat.Interior.Color)
        r["valores"] = {
            aba: livro.Worksheets(aba).Range("A1:AG290").Value2
            for aba in ("MEMORIA_RESULTADOS", "RESULTADOS_DETALHE")
        }
        return r

    try:
        livro = _tentar(pythoncom, lambda: excel.Workbooks.Open(
            str(caminho.resolve()), UpdateLinks=0, ReadOnly=False, CorruptLoad=0))
        _tentar(pythoncom, lambda: _preencher(livro, metodo, pcs, financeiro))
        _tentar(pythoncom, excel.CalculateFullRebuild)
        resultado = _tentar(pythoncom, ler)
        _tentar(pythoncom, livro.Save)
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = None
    finally:
        if livro is not None:
            try:
                livro.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()
    return resultado


def _pc():
    from test_coleta_118_excel_real import PC

    return PC


@excel_real
def test_1_reconhecido_e_potencial_positivos(tmp_path):
    r = _rodar(tmp_path, "rec_pot", _pc(), pcs=_pcs())
    assert r["exe_C9"] > 0 and r["exe_D9"] > 0
    assert r["exe_E9"] == round(r["exe_C9"] + r["exe_D9"], 2)
    assert r["exe_E10"] == "Soma do retroativo reconhecido com a parcela potencial."


@excel_real
def test_2_reconhecido_positivo_e_potencial_zero(tmp_path):
    r = _rodar(tmp_path, "so_rec", _pc(), pcs=_pcs({"PC-C1-NAO": "Sim"}))
    assert r["exe_C9"] > 0 and r["exe_D9"] == 0
    assert r["exe_E9"] == r["exe_C9"]


@excel_real
def test_3_reconhecido_zero_e_potencial_positivo(tmp_path):
    r = _rodar(tmp_path, "so_pot", _pc(), pcs=_pcs({"PC-C1": "Nao"}))
    assert r["exe_C9"] == 0 and r["exe_D9"] > 0
    assert r["exe_E9"] == r["exe_D9"]


@excel_real
def test_4_reconhecido_e_potencial_zero(tmp_path):
    # So os PCs do C0 (sem reajuste) e do C2 (precluso): nada a reconhecer.
    r = _rodar(tmp_path, "zeros", _pc(), pcs=_pcs(so=("PC-C0", "PC-C2")))
    assert r["exe_C9"] == 0 and r["exe_D9"] == 0
    assert r["exe_E9"] == 0          # R$ 0,00, mesma semantica dos dois cards


@excel_real
def test_5_e_8_bloqueio_revise_nao_inventa_valor(tmp_path):
    # Metodo PCs sem PCs (com financeiro): bloqueio de coerencia da 12.0.
    r = _rodar(tmp_path, "bloqueio", _pc(), financeiro=True)
    assert r["D22"] in (None, "") and r["T66"] in (None, "")
    assert r["exe_C9"] == "—" and r["exe_E9"] == "—"
    assert r["exe_C12"] == "REVISE A APURAÇÃO"
    assert r["exe_C13"] == (
        "Existem pendências que precisam ser corrigidas antes da conclusão. " + r["B7"])
    assert r["cor_faixa"] == _cor(ROSA)


@excel_real
def test_6_e_7_situacao_concluida_sem_e_com_potencial(tmp_path):
    sem_pot = _rodar(tmp_path, "sit_sem", _pc(), pcs=_pcs({"PC-C1-NAO": "Sim"}))
    com_pot = _rodar(tmp_path, "sit_pot", _pc(), pcs=_pcs())
    assert sem_pot["J5"] == com_pot["J5"] == 0
    assert sem_pot["exe_C12"] == com_pot["exe_C12"] == "APURAÇÃO CONCLUÍDA"
    assert sem_pot["exe_C13"] == (
        "Nenhuma pendência registrada pelas validações desta apuração.")
    assert sem_pot["cor_faixa"] == _cor(VERDE)
    valor = f"{com_pot['exe_D9']:.2f}".replace(".", ",")
    assert com_pot["exe_C13"] == (
        f"Nenhuma pendência impeditiva identificada. Há R$ {valor} de retroativo "
        "potencial, sujeito à confirmação."
    )
    assert com_pot["cor_faixa"] == _cor(AMARELO_ESVERDEADO) != _cor(VERDE)


@excel_real
def test_9_10_11_vta_reconhecido_e_potencial_identicos_a_coleta_120(tmp_path, monkeypatch):
    """A/B: mesmo caso no template 12.0 (merge do PR #188) e no 12.1."""
    from test_coleta_118_excel_real import FINANCEIRO

    base = tmp_path / "template_120.xlsx"
    base.write_bytes(subprocess.run(
        ["git", "show", f"{COMMIT_BASE_120}:templates/COLETA_REAJUSTE_OFICIAL.xlsx"],
        cwd=RAIZ, capture_output=True, check=True).stdout)
    casos = {
        "pcs": (_pc(), {"pcs": _pcs()}),
        "pcs_so_rec": (_pc(), {"pcs": _pcs({"PC-C1-NAO": "Sim"})}),
        "fin": (FINANCEIRO, {"financeiro": True}),
        "bloqueio": (_pc(), {"financeiro": True}),
    }
    novas = {(66, 19), (66, 20), (67, 19), (67, 20)} | {
        (linha, col) for linha in range(65, 71) for col in (32, 33)}
    for nome, (metodo, opcoes) in casos.items():
        antes = _rodar(tmp_path, f"base_{nome}", metodo, template=base,
                       monkeypatch=monkeypatch, **opcoes)
        depois = _rodar(tmp_path, f"novo_{nome}", metodo, **opcoes)
        for chave in ("B26", "T40", "T39", "B16", "D22", "J5", "A7", "B7"):
            assert depois[chave] == antes[chave], (nome, chave)
        # Camada tecnica inteira identica, fora das celulas novas.
        for i, linha in enumerate(antes["valores"]["MEMORIA_RESULTADOS"], start=1):
            for j, valor in enumerate(linha, start=1):
                if (i, j) not in novas:
                    assert depois["valores"]["MEMORIA_RESULTADOS"][i - 1][j - 1] == valor, (
                        nome, i, j)
        assert depois["valores"]["RESULTADOS_DETALHE"] == antes["valores"]["RESULTADOS_DETALHE"]
        # Os cards preservados mostram exatamente o que mostravam.
        assert depois["exe_B9"] == antes["exe_B9"]                  # VTA
        assert depois["exe_C9"] == antes["exe_C9"]                  # reconhecido
        assert depois["exe_D9"] == antes["exe_D9"]                  # potencial
        assert depois["exe_F9"] == antes["exe_E9"]                  # saldo: E -> F
        assert depois["exe_G9"] == antes["exe_F9"]                  # indice: F -> G
        if nome == "fin":
            # Metodo sem parcela potencial: o total coincide com o reconhecido.
            assert depois["exe_E9"] == depois["exe_C9"]
            assert depois["exe_E10"] == (
                "Neste método, não há parcela potencial; o total coincide com o "
                "retroativo reconhecido."
            )


@excel_real
def test_12_abre_salva_e_reabre_no_excel_sem_reparo(tmp_path):
    from _coleta_oficial import obter_coleta_oficial_bytes
    from validar_visual_resultados_v3 import validar

    arquivo = tmp_path / "coleta_branco_121.xlsx"
    arquivo.write_bytes(obter_coleta_oficial_bytes())
    assert validar(arquivo, tmp_path / "saida", estresse=True) == []
