"""Coleta 11.7 — historico_VU e CICLO_EM_EXECUCAO seguem a base economica do VU.

Estrutura (CI). A prova economica no Excel real fica em
tests/test_coleta_117_excel_real.py (RUN_EXCEL_INTEGRATION=1).
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from _ciclo_em_execucao import _formula_vu
from _coleta_oficial import (
    TEMPLATE_COLETA_OFICIAL,
    _COLUNAS_HISTORICO_VU,
    _LINHAS_HISTORICO_VU,
    _formula_historico_vu_base_economica,
    _formula_historico_vu_nascimento,
    _formula_com_gate_vta,
    _formula_status_aditivo,
    _garantir_gate_vta_alertas_aditivos,
    _garantir_historico_vu_base_economica,
    _gate_vta_canonico,
    obter_coleta_oficial_bytes,
)

_VU_C0_11_6 = (
    '=IF(itens_Remanesc!A2="","",IF(posicao_contratual!$Y2="","",'
    'IF(posicao_contratual!$Y2>0,"",itens_Remanesc!C2)))'
)


@pytest.fixture(scope="module")
def wb_template():
    return load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)


@pytest.fixture(scope="module")
def wb_coleta_117():
    return load_workbook(BytesIO(obter_coleta_oficial_bytes()), data_only=False)


def _formulas_historico(ws) -> dict[str, str]:
    return {
        f"{coluna}{linha}": ws[f"{coluna}{linha}"].value
        for coluna, _ in _COLUNAS_HISTORICO_VU
        for linha in _LINHAS_HISTORICO_VU
    }


def test_template_oficial_ja_carrega_a_coleta_117(wb_template):
    ws = wb_template["historico_VU"]
    for coluna, indice in _COLUNAS_HISTORICO_VU:
        for linha in _LINHAS_HISTORICO_VU:
            assert ws[f"{coluna}{linha}"].value == (
                _formula_historico_vu_base_economica(linha, indice)
            ), f"historico_VU!{coluna}{linha}"
    # VU_C0 nao muda: so existe para quem ja existia em C0 (base C0).
    assert ws["C2"].value == _VU_C0_11_6


def test_formula_le_a_base_economica_e_respeita_o_nascimento():
    formula = _formula_historico_vu_base_economica(3, 3)
    assert formula.isascii()
    assert formula.count("(") == formula.count(")")
    assert len(formula) < 8192
    assert "aditivos!$O$2:$O$200" in formula
    assert "MATCH($A3,aditivos!$A$2:$A$200,0)" in formula
    # Antes do nascimento o item nao existe; base nunca posterior a ele.
    assert 'IF(posicao_contratual!$Y3>3,""' in formula
    assert "MIN(INDEX(aditivos!$O$2:$O$200" in formula
    # Ciclo de nascimento: ultimo fator conhecido (mesmo carregamento de aditivos!I).
    assert "INDEX($L$2:$L$6,MIN(3+1,COUNT($L$2:$L$6)))" in formula
    # Ciclos seguintes: fator do proprio ciclo exigido (historico nao inventado).
    assert "NOT(ISNUMBER($L$5))" in formula
    # Nxxx com mais de uma inclusao: vazio, sem cair no nascimento.
    assert (
        'IF(OR(NOT(ISNUMBER(itens_Remanesc!C3)),AND(posicao_contratual!$Y3>0,'
        'COUNTIFS(aditivos!$A$2:$A$200,$A3,aditivos!$D$2:$D$200,"*novo*")>1)),""'
    ) in formula


def test_aditivos_alerta_nxxx_com_mais_de_uma_inclusao(wb_template):
    formula = _formula_status_aditivo(2)
    assert formula.isascii()
    assert formula.count("(") == formula.count(")")
    assert (
        'IF(AND(ISNUMBER(SEARCH("NOVO",D2)),COUNTIFS($A$2:$A$200,A2,'
        '$D$2:$D$200,"*novo*")>1),"ALERTA: NOVO_ITEM_COM_MAIS_DE_UMA_INCLUSAO",'
    ) in formula
    ws = wb_template["aditivos"]
    assert all(ws[f"M{linha}"].value == _formula_status_aditivo(linha) for linha in range(2, 201))


def test_geracao_entrega_o_template_sem_reescrita(wb_template, wb_coleta_117):
    assert _formulas_historico(wb_coleta_117["historico_VU"]) == (
        _formulas_historico(wb_template["historico_VU"])
    )


def test_migracao_runtime_converte_a_forma_11_6_e_e_idempotente():
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    ws = wb["historico_VU"]
    for coluna, indice in _COLUNAS_HISTORICO_VU:
        for linha in _LINHAS_HISTORICO_VU:
            ws[f"{coluna}{linha}"].value = _formula_historico_vu_nascimento(linha, indice)
    estranha = '=IF(A9="","",1)'
    ws["F9"].value = estranha

    _garantir_historico_vu_base_economica(wb)
    assert ws["F9"].value == estranha  # estrutura nao reconhecida: intacta
    assert ws["F8"].value == _formula_historico_vu_base_economica(8, 3)
    assert ws["D2"].value == _formula_historico_vu_base_economica(2, 1)
    assert ws["G200"].value == _formula_historico_vu_base_economica(200, 4)

    ws["F9"].value = _formula_historico_vu_base_economica(9, 3)
    antes = _formulas_historico(ws)
    _garantir_historico_vu_base_economica(wb)
    assert _formulas_historico(ws) == antes


def test_migracao_nao_age_sem_base_economica_canonica_em_aditivos():
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    ws = wb["historico_VU"]
    ws["F3"].value = _formula_historico_vu_nascimento(3, 3)
    wb["aditivos"]["O3"].value = None  # aditivos!O fora da forma canonica
    _garantir_historico_vu_base_economica(wb)
    assert ws["F3"].value == _formula_historico_vu_nascimento(3, 3)


def test_gate_do_vta_na_origem_canonica(wb_template, wb_coleta_117):
    for wb in (wb_template, wb_coleta_117):
        mem = wb["MEMORIA_RESULTADOS"]
        assert _gate_vta_canonico(mem)
        assert mem["T48"].value == '=COUNTIF(aditivos!$M$2:$M$200,"ALERTA:*")'
        # B26: gate logo apos a governanca B24/B25 (inicio homologado mantido).
        assert mem["B26"].value.startswith('=IF(AND(B24<>"",B25<>""),"",IF($T$48>0,"",')
        assert "IF(ISNUMBER(B25),B25," in mem["B26"].value
        assert mem["T40"].value.startswith('=IF($T$48>0,"",')
        for celula in ("B26", "T40"):
            formula = mem[celula].value
            assert formula.isascii() and formula.count("(") == formula.count(")")
    # Os nomes publicados continuam apontando para a origem com gate.
    nomes = wb_template.defined_names
    assert nomes["VTA_FINAL"].attr_text == "MEMORIA_RESULTADOS!$B$26"
    assert nomes["VTA_SEM_POTENCIAL"].attr_text == "MEMORIA_RESULTADOS!$T$40"


def test_migracao_do_gate_e_idempotente():
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    mem = wb["MEMORIA_RESULTADOS"]
    gate = 'IF($T$48>0,"",'
    for celula in ("B26", "T40"):
        com_gate = mem[celula].value
        mem[celula].value = com_gate.replace(gate, "", 1)[:-1]
        assert _formula_com_gate_vta(celula, mem[celula].value) == com_gate
    mem["S48"].value = mem["T48"].value = None

    _garantir_gate_vta_alertas_aditivos(wb)
    assert _gate_vta_canonico(mem)
    antes = {c: mem[c].value for c in ("B26", "T40", "S48", "T48")}
    _garantir_gate_vta_alertas_aditivos(wb)
    assert {c: mem[c].value for c in antes} == antes
    assert _formula_com_gate_vta("B26", antes["B26"]) == antes["B26"]


def test_fallback_da_ciclo_em_execucao_usa_a_mesma_base(wb_coleta_117):
    formula = _formula_vu(14, 3)
    assert formula.isascii()
    assert formula.count("(") == formula.count(")")
    assert (
        "MIN(INDEX(aditivos!$O$2:$O$200,"
        "MATCH(itens_Remanesc!$A3,aditivos!$A$2:$A$200,0)),posicao_contratual!$Y3)"
    ) in formula
    # Nunca reajusta antes do nascimento; historico_VU continua prevalecendo.
    assert "posicao_contratual!$Y3>(MATCH($C$3" in formula
    assert (
        'AND(posicao_contratual!$Y3>0,COUNTIFS(aditivos!$A$2:$A$200,'
        'itens_Remanesc!$A3,aditivos!$D$2:$D$200,"*novo*")>1)),""'
    ) in formula
    assert formula.startswith('=IF(A14="","",IF(ISNUMBER(IF($C$3="C0",historico_VU!$C3')
    assert wb_coleta_117["CICLO_EM_EXECUCAO"]["E14"].value == formula
