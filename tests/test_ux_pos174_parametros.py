"""AJUSTES-XLS-UX pos-PR #174 — Frente 3 — parametros: aviso dinamico de percentuais HISTORICOS ausentes
(A8) e nota curta da MEMORIA DO FATOR (A17). Formulas E/F intactas.

Testes de Excel real rodam com RUN_EXCEL_INTEGRATION=1.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_ajustes_xls_ux_pos174 as ux  # noqa: E402
from _coleta_oficial import (  # noqa: E402
    gerar_coleta_oficial_preenchida,
    obter_coleta_oficial_bytes,
)
from _memoria_calculo import (  # noqa: E402
    FILL_FRONTEIRA_IST,
    escrever_memoria_calculo,
    ler_memoria_calculo,
)

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
COLUNAS_JR = ("J", "K", "L", "M", "N", "O", "P", "Q", "R")

com = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1",
)


@pytest.fixture(scope="module")
def template():
    wb = load_workbook(TEMPLATE)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def coleta_branca():
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    yield wb
    wb.close()


def _rgb(cell) -> str | None:
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


# ------------------------------------------------------------------ payloads
PERCENTUAIS = {1: 0.0512, 2: 0.0374, 3: 0.0289, 4: 0.0315}


def _payload(numeros: list[int]) -> dict:
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2021, 1, 1),
        "data_corte": date(2021 + max(numeros), 12, 31),
        "ciclos": [
            {
                "ciclo": f"C{n}",
                "data_inicio": date(2021 + n, 1, 1),
                "percentual": PERCENTUAIS[n],
                "inicio_efeito_financeiro": date(2021 + n, 1, 1),
                "situacao_aplicada": "✅ TEMPESTIVO",
                "data_pedido": date(2021 + n, 1, 10),
            }
            for n in numeros
        ],
    }


# =================================================================== FRENTE 3
def test_f3_formulas_economicas_de_parametros_intactas(template):
    ws = template["parametros"]
    for linha in range(3, 7):
        anterior = linha - 1
        assert ws[f"F{linha}"].value == (
            f'=IF(E{linha}="","",IF(NOT(ISNUMBER(E{linha})),"",'
            f'IF(NOT(ISNUMBER(F{anterior})),"",F{anterior}*(1+E{linha}))))'
        )
    assert ws["F2"].value == "=1"
    regras = [
        (str(faixa.sqref), regra.formula)
        for faixa in ws.conditional_formatting for regra in faixa.rules
    ]
    assert ("E3:E6", ['AND($A3="Nao",$C3<>"",$E3="")']) in regras


def test_f3_aviso_historico_no_template(template):
    ws = template["parametros"]
    formula = ws[ux.CELULA_AVISO].value
    assert formula == ux.formula_aviso_historico()
    ux.validar_ascii_e_parenteses({"A8": formula})
    # Historico = ciclos ANTERIORES ao vigente canonico (CONTROLE!B2).
    assert "CONTROLE!$B$2" in formula
    assert "$E$6" not in formula  # C4 nunca e historico
    assert ux.FAIXA_AVISO in {str(m) for m in ws.merged_cells.ranges}
    assert any(str(f.sqref) == ux.FAIXA_AVISO for f in ws.conditional_formatting)
    mem = template["MEMORIA_RESULTADOS"]
    for linha, chave, texto in ux.TEXTOS_HISTORICO:
        assert mem[f"AF{linha}"].value == chave
        assert mem[f"AG{linha}"].value == texto


def test_f3_memoria_do_fator_mantem_semantica(template):
    ws = template["parametros"]
    assert ws["A9"].value == "MEMORIA DO FATOR APLICAVEL A APURACAO"
    assert ws[ux.CELULA_NOTA_FATOR].value == ux.NOTA_FATOR
    # So os ciclos COMPUTADOS: C/D/E/F seguem condicionados a B (=A da linha).
    for linha, origem in zip(range(12, 16), range(3, 7)):
        assert ws[f"B{linha}"].value == f"=A{origem}"
        assert ws[f"C{linha}"].value == f'=IF(B{linha}="Sim",E{origem},"")'


def test_f3_aviso_nao_e_lido_como_ciclo():
    from _leitor_masterfile_v10 import ler_masterfile_v10

    conteudo = gerar_coleta_oficial_preenchida(_payload([3]))
    lido = ler_masterfile_v10(conteudo)
    parametros = lido["parametros_v10"]
    assert list(parametros["por_ciclo"]) == ["C0", "C1", "C2", "C3", "C4"]
    # (O alerta preexistente da linha 10 — cabecalho "COMPUTA?" — e da base.)
    assert not [
        a for a in parametros["alertas"]
        if "em linha 8:" in a or "em linha 17:" in a
    ]


def _avaliar_aviso(caminhos_e_edicoes):
    """Abre cada Coleta no Excel real, aplica edicoes em parametros!E e le A8."""
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    leituras = []
    try:
        for caminho, edicoes in caminhos_e_edicoes:
            wb = excel.Workbooks.Open(str(caminho))
            try:
                ws = wb.Worksheets("parametros")
                excel.CalculateFull()
                estado = [ws.Range("A8").Text]
                for celula, valor in edicoes:
                    ws.Range(celula).Value = valor
                    excel.CalculateFull()
                    estado.append(ws.Range("A8").Text)
                leituras.append(estado)
            finally:
                wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()
    return leituras


@com
def test_f3_aviso_avaliado_no_excel(tmp_path):
    cenarios = {
        "c1": ([1], [], [""]),
        "c2": ([2], [("E3", 0.05)],
               ["HISTÓRICO INCOMPLETO — informe o percentual de C1.", ""]),
        "c3": ([3], [("E3", 0.05), ("E4", 0.04)],
               ["HISTÓRICO INCOMPLETO — informe os percentuais de C1 e C2.",
                "HISTÓRICO INCOMPLETO — informe o percentual de C2.", ""]),
        "c4": ([4], [("E4", 0.04), ("E3", 0.05), ("E5", 0.03)],
               ["HISTÓRICO INCOMPLETO — informe os percentuais de C1, C2 e C3.",
                "HISTÓRICO INCOMPLETO — informe os percentuais de C1 e C3.",
                "HISTÓRICO INCOMPLETO — informe o percentual de C3.", ""]),
        "c1_c3": ([1, 2, 3], [], [""]),
    }
    entradas, esperados = [], []
    for nome, (numeros, edicoes, esperado) in cenarios.items():
        caminho = tmp_path / f"{nome}.xlsx"
        caminho.write_bytes(gerar_coleta_oficial_preenchida(_payload(numeros)))
        entradas.append((caminho, edicoes))
        esperados.append(esperado)
    assert _avaliar_aviso(entradas) == esperados



def test_f3_datas_do_pedido_cabem_sem_cerquilhas(template):
    """U (DATA_PEDIDO) e V (PROXIMA_DATA_REAJUSTE) recebem dd/mm/aaaa."""
    ws = template["parametros"]
    for col in ux.COLUNAS_DATAS_PEDIDO:
        assert ws.column_dimensions[col].width >= ux.LARGURA_DATAS_PEDIDO - 0.7
