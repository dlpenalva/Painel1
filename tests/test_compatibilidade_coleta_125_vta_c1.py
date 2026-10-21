from __future__ import annotations

import os
import shutil
import time
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from _compatibilidade_valores import (
    AVISO_COMPATIBILIDADE_125,
    _compatibilidade_125_vta_c1,
    aviso_compatibilidade_coleta_125,
)
from _contexto_coleta import ContextoColeta
from _coleta_reajuste import ler_coleta_reajuste
from _coleta_oficial import (
    _FORMULA_W50_126,
    _FORMULA_W67_126,
    _FORMULAS_RETRO_FINANCEIRO_126,
    _ROTULO_COBERTURA_B6_126,
)

TOL = 0.005
ROOT = Path(__file__).resolve().parents[1]
FIXTURES_COMPAT = Path(__file__).resolve().parent / "fixtures" / "coletas_compat"
# Valores sinteticos: nenhum dado real de fiscalizacao entra no repositorio.
VTA_SINTETICO = 620_123.45


def _workbook_125_minimo() -> Workbook:
    wb = Workbook()
    wb.active.title = "CONTROLE"
    for nome in (
        "parametros", "financeiro", "CICLO_EM_EXECUCAO",
        "MEMORIA_RESULTADOS", "RESULTADOS_DETALHE",
    ):
        wb.create_sheet(nome)
    wb["CONTROLE"]["B1"] = "Financeiro (Mensalidade)"
    par = wb["parametros"]
    par["A2"], par["I2"] = "Nao", date(2024, 8, 1)
    par["A3"], par["I3"] = "Sim", date(2025, 9, 15)
    fin = wb["financeiro"]
    dados = [
        (date(2025, 7, 1), "c0", 100_000.0, 100_000.0, 0.0, "Nao"),
        (date(2025, 8, 1), "c1", 10_000.0, 10_000.0, 0.0, "Nao"),
        (date(2025, 9, 1), "c1", 10_000.0, 10_000.0, 0.0, "Nao"),
    ]
    for linha, (competencia, ciclo, pago, atualizado, delta, efeito) in enumerate(dados, 2):
        for coluna, valor in zip((1, 2, 3, 5, 6, 7), (
            competencia, ciclo, pago, atualizado, delta, efeito
        )):
            fin.cell(linha, coluna).value = valor
    mem = wb["MEMORIA_RESULTADOS"]
    for celula, valor in {
        "D4": 0.005, "B20": 600_000.0, "B22": 20_123.45,
        "D20": 120_000.0, "D35": 500_123.45, "T25": VTA_SINTETICO,
        "N263": 0.0, "T48": 0.0, "T58": 0.0, "W46": 1.0, "B44": 0.0,
        "W53": 500_123.45, "W54": 0.0, "W61": 100_000.0,
        "W66": 100_000.0, "W67": 100_000.0, "W48": 600_123.45,
        "W50": 600_123.45, "W51": 0.0, "W52": "RECONCILIADO",
    }.items():
        mem[celula] = valor
    cee = wb["CICLO_EM_EXECUCAO"]
    cee["D5"], cee["G13"] = date(2025, 9, 15), 500_123.45
    det = wb["RESULTADOS_DETALHE"]
    det["B17"] = 20_000.0
    det["H8"], det["H14"], det["B3"] = "REVISE", "REVISE", "REVISE"
    det["H24"], det["H33"] = "VALIDADO", "VALIDADO"
    return wb


def test_compatibilidade_125_recompoe_so_derivados_e_reconcilia():
    wb = _workbook_125_minimo()
    entradas = {
        (aba, celula): wb[aba][celula].value
        for aba, celula in (
            ("CONTROLE", "B1"), ("parametros", "A3"),
            ("financeiro", "C3"), ("financeiro", "C4"),
            ("CICLO_EM_EXECUCAO", "G13"),
        )
    }
    auditoria = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    mem = wb["MEMORIA_RESULTADOS"]
    assert auditoria["aplicada"] is True
    assert mem["B11"].value == 0.0
    assert mem["B16"].value == 0.0
    assert mem["B26"].value == pytest.approx(VTA_SINTETICO, abs=TOL)
    assert mem["W48"].value == pytest.approx(VTA_SINTETICO, abs=TOL)
    assert mem["W50"].value == pytest.approx(VTA_SINTETICO, abs=TOL)
    assert mem["W51"].value == 0.0 and mem["W52"].value == "RECONCILIADO"
    assert wb["RESULTADOS_DETALHE"]["B3"].value == "VALIDADO"
    assert entradas == {(a, c): wb[a][c].value for a, c in entradas}


def test_compatibilidade_125_falha_fechado_se_t25_nao_confirma_vta():
    wb = _workbook_125_minimo()
    wb["MEMORIA_RESULTADOS"]["T25"] = 600_000.0
    auditoria = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert auditoria["aplicada"] is False
    assert "T25 divergente" in auditoria["motivo"]
    assert wb["MEMORIA_RESULTADOS"]["B26"].value is None


@pytest.mark.parametrize("efeito", [None, "", "Talvez", "S"])
def test_compatibilidade_125_falha_fechado_com_efeito_ausente_ou_invalido(efeito):
    # C1 em execucao: competencia paga sem EFEITO_FINANCEIRO valido nao
    # comprova retroativo zero; o hotfix nao e aplicado.
    wb = _workbook_125_minimo()
    wb["financeiro"]["G4"] = efeito
    auditoria = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert auditoria["aplicada"] is False
    assert "EFEITO_FINANCEIRO ausente ou invalido em C1" in auditoria["motivo"]
    assert wb["MEMORIA_RESULTADOS"]["B11"].value is None
    assert wb["MEMORIA_RESULTADOS"]["B26"].value is None
    assert aviso_compatibilidade_coleta_125(auditoria) is None


def test_compatibilidade_125_falha_fechado_com_linha_incompleta_no_b44():
    wb = _workbook_125_minimo()
    wb["MEMORIA_RESULTADOS"]["B44"] = 1
    auditoria = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert auditoria["aplicada"] is False
    assert "linhas financeiras incompletas" in auditoria["motivo"]
    assert wb["RESULTADOS_DETALHE"]["H14"].value == "REVISE"


def test_compatibilidade_125_aceita_nao_com_acento_e_sim_com_delta():
    wb = _workbook_125_minimo()
    wb["financeiro"]["G3"] = "Não"
    wb["financeiro"]["G2"] = "Sim"  # C0: delta 0,00 numerico (fator 1)
    auditoria = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert auditoria["aplicada"] is True
    assert wb["MEMORIA_RESULTADOS"]["B11"].value == 0.0


def test_compatibilidade_125_falha_fechado_com_sim_sem_delta():
    wb = _workbook_125_minimo()
    wb["financeiro"]["G3"] = "Sim"
    wb["financeiro"]["F3"] = None
    auditoria = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert auditoria["aplicada"] is False
    assert "sem delta financeiro calculado em C1" in auditoria["motivo"]


def test_aviso_125_so_quando_hotfix_efetivamente_aplicado():
    wb = _workbook_125_minimo()
    aplicada = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert aviso_compatibilidade_coleta_125(aplicada) == AVISO_COMPATIBILIDADE_125
    assert AVISO_COMPATIBILIDADE_125 == (
        "Coleta 12.5: correção de compatibilidade aplicada automaticamente ao "
        "processamento. O arquivo original não foi alterado. Recomenda-se "
        "utilizar a Coleta 12.6 nas próximas apurações."
    )
    wb = _workbook_125_minimo()
    wb["MEMORIA_RESULTADOS"]["T25"] = 600_000.0
    recusada = _compatibilidade_125_vta_c1(
        wb, {"marcador_publico": "12.5", "codigo": "COLETA_11"}
    )
    assert aviso_compatibilidade_coleta_125(recusada) is None
    # Compatibilidade legada aplicada nao e o hotfix 12.5: sem aviso.
    assert aviso_compatibilidade_coleta_125(
        {"aplicada": True, "blocos_adaptados": ["financeiro"]}
    ) is None
    assert aviso_compatibilidade_coleta_125(None) is None


@pytest.mark.parametrize("fixture", ["pre11_l1_financeiro", "coleta_11_financeiro"])
def test_upload_legado_ou_atual_nao_exibe_aviso_125(fixture):
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    conteudo = (FIXTURES_COMPAT / f"{fixture}.xlsx").read_bytes()
    resultado, _ = processar_coleta_oficial_runtime(conteudo)
    assert resultado["aviso_compatibilidade_coleta_125"] is None


def test_pagina_valor_global_exibe_aviso_125_condicional_e_discreto():
    fonte = (ROOT / "pages" / "03_Valor_Global.py").read_text(encoding="utf-8")
    assert 'if resultado.get("aviso_compatibilidade_coleta_125"):' in fonte
    assert "st.caption(f\"ℹ️ {resultado['aviso_compatibilidade_coleta_125']}\")" in fonte


real_excel = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1"
    or not os.environ.get("COLETA_125_REAL_XLSX"),
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1 e COLETA_125_REAL_XLSX",
)


def _entradas(wb) -> dict[tuple[str, str], object]:
    abas = (
        "CONTROLE", "parametros", "financeiro", "itens_Remanesc",
        "itens_Consumidos", "itens_PC", "aditivos", "CICLO_EM_EXECUCAO",
    )
    ignoradas = {("CONTROLE", "B16"), ("CONTROLE", "B17")}
    return {
        (aba, celula.coordinate): celula.value
        for aba in abas
        for linha in wb[aba].iter_rows()
        for celula in linha
        if hasattr(celula, "coordinate")
        and (aba, celula.coordinate) not in ignoradas
        and not (isinstance(celula.value, str) and celula.value.startswith("="))
    }


@real_excel
def test_coleta_125_real_compatibilizada_e_mesmos_dados_na_126(tmp_path):
    origem = Path(os.environ["COLETA_125_REAL_XLSX"])
    destino = tmp_path / "coleta_12_6.xlsx"
    shutil.copy2(origem, destino)
    original = load_workbook(origem, data_only=False, read_only=True)
    assert original["CONTROLE"]["B16"].value == "12.5"
    assert original["MEMORIA_RESULTADOS"]["B26"].value.startswith("=IF(")
    original_valores = load_workbook(origem, data_only=True, read_only=True)
    assert original_valores["MEMORIA_RESULTADOS"]["B26"].value is None
    # O VTA esperado vem do proprio arquivo (T25), sem fixar valor real aqui.
    vta_esperado = original_valores["MEMORIA_RESULTADOS"]["T25"].value
    assert isinstance(vta_esperado, (int, float)) and vta_esperado > 0
    diagnostico = ler_coleta_reajuste(origem.read_bytes())
    assert diagnostico["status_base"] == "APTO_PARA_ANALISE"
    assert diagnostico["pronto_para_consolidar"] is True
    assert diagnostico["metadados"]["status_resultados"]["geral"] == "VALIDADO"
    assert diagnostico["metadados"]["status_resultados"]["valores"]["vta_oficial"] == pytest.approx(
        vta_esperado, abs=TOL
    )
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    resultado, _ = processar_coleta_oficial_runtime(origem.read_bytes())
    assert resultado["aviso_compatibilidade_coleta_125"] == AVISO_COMPATIBILIDADE_125
    with ContextoColeta(origem.read_bytes()) as contexto:
        mem = contexto.workbook_valores["MEMORIA_RESULTADOS"]
        assert mem["B26"].value == pytest.approx(vta_esperado, abs=TOL)
        assert mem["W48"].value == mem["W50"].value

    pythoncom = pytest.importorskip("pythoncom")
    client = pytest.importorskip("win32com.client")
    pythoncom.CoInitialize()
    excel = client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    book = excel.Workbooks.Open(str(destino), UpdateLinks=0, ReadOnly=False)
    try:
        mem = book.Worksheets("MEMORIA_RESULTADOS")
        for celula, formula in _FORMULAS_RETRO_FINANCEIRO_126.items():
            mem.Range(celula).Formula = formula
        mem.Range("W50").Formula = _FORMULA_W50_126
        mem.Range("W67").Formula = _FORMULA_W67_126
        book.Worksheets("cobertura_temporal").Range("A6").Value = _ROTULO_COBERTURA_B6_126
        controle = book.Worksheets("CONTROLE")
        controle.Unprotect()
        controle.Range("B16").Value = "12.6"
        controle.Range("B17").Value = "12.13"
        controle.Protect(DrawingObjects=True, Contents=True, Scenarios=True)
        book.Worksheets("MEMORIA_RESULTADOS").Visible = 0
        book.Worksheets("RESULTADOS_DETALHE").Visible = 0
        book.Worksheets("RESULTADOS").Visible = -1
        for _ in range(8):
            try:
                excel.CalculateFullRebuild()
                break
            except Exception:
                pythoncom.PumpWaitingMessages()
                time.sleep(0.5)
        book.Save()
    finally:
        book.Close(SaveChanges=False)
        try:
            excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()

    novo = load_workbook(destino, data_only=True, read_only=True)
    mem = novo["MEMORIA_RESULTADOS"]
    assert mem["B11"].value == 0.0 and mem["B16"].value == 0.0
    assert mem["B26"].value == pytest.approx(vta_esperado, abs=TOL)
    assert mem["W48"].value == pytest.approx(vta_esperado, abs=TOL)
    assert mem["W50"].value == pytest.approx(vta_esperado, abs=TOL)
    assert mem["W51"].value == 0.0 and mem["W52"].value == "RECONCILIADO"
    assert novo["RESULTADOS_DETALHE"]["B3"].value == "VALIDADO"
    assert novo["MEMORIA_RESULTADOS"].sheet_state == "hidden"
    assert novo["RESULTADOS_DETALHE"].sheet_state == "hidden"
    assert novo["RESULTADOS"].sheet_state == "visible"
    convertido = load_workbook(destino, data_only=False, read_only=True)
    assert _entradas(original) == _entradas(convertido)
