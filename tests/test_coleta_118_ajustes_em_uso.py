"""Coleta 11.8 — ajustes da Coleta em uso (estrutura do template e do runtime).

* itens_PC!K aceita "Nao"/"Não" em PC_PAGO_A_CONTRATADA; itens_PC!K/L: PC de
  ciclo PRECLUSO sem INICIO_EFEITO_FINANCEIRO fica com L="Nao" (sem alerta);
* VTA (Financeiro/PCs) com a posicao atual da CICLO_EM_EXECUCAO quando falta a
  fotografia do ciclo vigente (MEMORIA_RESULTADOS!T49:T53);
* ajustes visuais de parametros, posicao_referencia, Nxxx e cobertura_temporal.

O comportamento numerico no Excel real esta em test_coleta_118_excel_real.py.
"""
from __future__ import annotations

import io
import re
from datetime import date

from openpyxl import load_workbook

import _coleta_oficial as co
from _efeitos_financeiros_pc import efeito_financeiro_pc
from _memoria_calculo import FORMATO_FATOR_ACUMULADO, FORMATO_VARIACAO_FINAL


def _template():
    return load_workbook(co.TEMPLATE_COLETA_OFICIAL, data_only=False)


def _gerada():
    return load_workbook(io.BytesIO(co.obter_coleta_oficial_bytes()), data_only=False)


def _regras(ws):
    return [
        (str(cf.sqref), regra)
        for cf in ws.conditional_formatting
        for regra in cf.rules
    ]


def _formulas(wb):
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str) and celula.value.startswith("="):
                    yield ws.title, celula.coordinate, celula.value


# --------------------------------------------------------------- itens_PC
def test_itens_pc_k_l_na_forma_11_8_em_toda_a_grade():
    ws = _template()["itens_PC"]
    for linha in range(2, 5002):
        assert ws[f"K{linha}"].value == co._formula_check_pc(linha)
        assert ws[f"L{linha}"].value == co._formula_efeito_pc(linha)


def test_check_pc_aceita_nao_com_e_sem_acento_e_mantem_o_contrato():
    k = co._formula_check_pc(2)
    # Contrato da 27B preservado; o "Não" (Ã = UNICHAR 195) vira "NAO".
    assert 'AND(G2<>"Sim",G2<>"Nao",G2<>"")' in k
    assert 'SUBSTITUTE(UPPER(G2),_xlfn.UNICHAR(195),"A")<>"NAO"' in k
    assert "INICIO_EFEITO ausente: PC" in k
    assert 'ISERROR(SEARCH("PRECLUSO",' in k
    assert k.count("(") == k.count(")")


def test_efeito_pc_precluso_sem_inicio_e_nao_e_ativo_sem_inicio_segue_vazio():
    l_ = co._formula_efeito_pc(2)
    assert 'IF(ISNUMBER(SEARCH("PRECLUSO",' in l_ and '"Nao",""),' in l_
    assert "B2>=" in l_
    assert l_.count("(") == l_.count(")")


def test_python_alinhado_ao_xls_no_pc_de_ciclo_precluso():
    precluso = {"computar_nesta_apuracao": "Sim", "situacao": "❌ PRECLUSO | SEM PEDIDO NESTE CICLO"}
    ativo = {"computar_nesta_apuracao": "Sim", "situacao": "✅ TEMPESTIVO"}
    assert efeito_financeiro_pc(date(2025, 3, 1), "C2", precluso) == "Nao"
    assert efeito_financeiro_pc(date(2025, 3, 1), "C2", ativo) is None
    com_inicio = {**precluso, "inicio_efeito_financeiro": date(2025, 1, 1)}
    assert efeito_financeiro_pc(date(2025, 3, 1), "C2", com_inicio) == "Sim"


# ------------------------------------------------------- VTA posicao atual
def test_memoria_resultados_t49_t53_e_celulas_do_vta_na_forma_11_8():
    tpl = _template()
    mem = tpl["MEMORIA_RESULTADOS"]
    for celula, (rotulo, formula) in co._CELULAS_POSICAO_ATUAL_VTA.items():
        assert mem[f"S{celula[1:]}"].value == rotulo
        assert mem[celula].value == formula
        assert formula.isascii() and formula.count("(") == formula.count(")")
    for (aba, celula), trocas in co._TRECHOS_POSICAO_ATUAL_VTA.items():
        valor = tpl[aba][celula].value
        assert all(troca in valor for _, troca in trocas), f"{aba}!{celula}"
        assert valor.count("(") == valor.count(")")


def test_sem_posicao_atual_as_formulas_reproduzem_a_11_7():
    """T49=0: cada troca so acrescenta um ramo que devolve a expressao 11.7."""
    mem = _template()["MEMORIA_RESULTADOS"]
    assert mem["T23"].value == "=IF($T$49=1,$T$51,SUM($Y$2:$Y$201))"
    assert mem["T50"].value.startswith("=IF($T$49<>1,0,")
    assert "AND($T$26>0,$T$49<>1)" in mem["T25"].value
    assert "ROUND($T$21+$T$22+$T$50+$T$23+$T$39,2)" in mem["T25"].value
    assert mem["T40"].value.startswith('=IF($T$48>0,"",')  # gate 11.7 intacto
    assert '=IF(AND(B24<>"",B25<>""),"",IF($T$48>0,"",' in mem["B26"].value
    assert 'IF($T$49=1,$T$50,IFERROR(SUM(INDIRECT("CICLO_EM_EXECUCAO!F13:F211")),0))' in mem["W50"].value
    assert "OR($T$49=1,AND(ISNUMBER($W$51),ROUND($W$51,2)=0))" in mem["E26"].value


def test_t49_so_liga_com_residual_ausente_posicao_completa_e_metodo_fin_pc():
    t49 = co._CELULAS_POSICAO_ATUAL_VTA["T49"][1]
    assert '$T$26=0,$W$49<>1),0' in t49
    assert 'IF($B$4="PCs",1,' in t49 and '$B$4="Financeiro"' in t49
    # Itens Consumidos nunca usa a posicao atual (ja e execucao + saldo).
    assert '"Itens"' not in t49


def test_resultados_detalhe_bloco_do_ciclo_atual_sem_referencia_circular():
    rd = _template()["RESULTADOS_DETALHE"]
    assert rd["B36"].value.startswith("=IF(MEMORIA_RESULTADOS!$T$49=1,MEMORIA_RESULTADOS!$T$50,")
    assert rd["B37"].value.startswith("=IF(MEMORIA_RESULTADOS!$T$49=1,")
    assert "$B$38" not in rd["B37"].value  # B38 le B37: sem ciclo
    assert rd["B38"].value.startswith("=IF(MEMORIA_RESULTADOS!$T$49=1,MEMORIA_RESULTADOS!$T$51,")


def test_ciclo_em_execucao_a9_nao_exige_referencia_no_ciclo_atual():
    ws = _gerada()["CICLO_EM_EXECUCAO"]
    assert "$P$13:$P$211" not in ws["A9"].value
    assert 'COUNTIF($K$13:$K$211,"INCOMPLETO:*")>0),""' in ws["A9"].value


# ------------------------------------------------------------- visuais
def test_parametros_formatos_tempestivo_e_campo_de_entrada():
    ws = _template()["parametros"]
    assert {ws[f"P{r}"].number_format for r in range(2, 81)} == {"0.0000"}
    assert {ws[f"Q{r}"].number_format for r in range(2, 81)} == {"0.00%"}
    assert not any(ws[f"{c}{r}"].font.b for c in "ABCDE" for r in range(3, 7))
    regras = [(faixa, r) for faixa, r in _regras(ws) if faixa == "A2:E6"]
    assert len(regras) == 1
    assert regras[0][1].formula == [
        'AND(ISNUMBER(SEARCH("TEMPESTIVO",$G2)),ISERROR(SEARCH("INTEMPESTIVO",$G2)))'
    ]
    assert regras[0][1].dxf.font.b
    # Alerta vermelho de E3:E6 continua antes (maior prioridade).
    alerta = [r for faixa, r in _regras(ws) if faixa == "E3:E6"][0]
    assert alerta.priority < regras[0][1].priority
    assert {ws[f"G{r}"].fill.fgColor.rgb for r in range(12, 16)} == {"FFFFF2CC"}
    assert FORMATO_FATOR_ACUMULADO == "0.0000" and FORMATO_VARIACAO_FINAL == "0.00%"


def test_posicao_referencia_colunas_ajustadas():
    ws = _template()["posicao_referencia"]
    assert ws.column_dimensions["F"].width > 50
    assert ws.column_dimensions["H"].width > 42
    assert ws.column_dimensions["I"].width > 72


def test_nxxx_em_verde_nas_abas_itemizadas_com_menor_prioridade():
    tpl = _template()
    esperado = co._FORMULA_DESTAQUE_NOVOS_ITENS
    for aba, faixa in co._FAIXAS_DESTAQUE_NOVOS_ITENS_118.items():
        regras = _regras(tpl[aba])
        verdes = [(f, r) for f, r in regras if r.formula == [esperado]]
        assert [f for f, _ in verdes] == [faixa], aba
        regra = verdes[0][1]
        assert regra.dxf.font.color.rgb.endswith(co._COR_FONTE_NOVOS_ITENS)
        assert regra.priority == max(r.priority for _, r in regras), aba
    cee = _gerada()["CICLO_EM_EXECUCAO"]
    verdes = [f for f, r in _regras(cee) if r.formula == [co.formula_destaque_novo_item(13)]]
    assert verdes == ["A13:I211"]


def test_cobertura_temporal_b8_data_real_e_legenda_removida_sem_dependentes():
    tpl = _template()
    ws = tpl["cobertura_temporal"]
    assert ws["B8"].value == co._FORMULA_DATA_POSICAO_COBERTURA
    assert "CONTROLE!B3" not in ws["B8"].value
    assert ws["C8"].value == co._AJUDA_DATA_POSICAO_COBERTURA
    assert ws.column_dimensions["B"].width >= 50
    assert all(ws[f"B{r}"].alignment.wrap_text for r in range(2, 24))
    assert ws.max_row <= 24
    assert not ws.merged_cells.ranges
    referencias = re.compile(r"cobertura_temporal!\$?[A-Z]+\$?(\d+)")
    linhas = {
        int(m) for _, _, formula in _formulas(tpl) for m in referencias.findall(formula)
    }
    assert max(linhas) < 25


def test_geracao_nao_conserta_os_pontos_da_11_8():
    tpl, gerada = _template(), _gerada()
    pontos = [("MEMORIA_RESULTADOS", c) for c in co._CELULAS_POSICAO_ATUAL_VTA]
    pontos += list(co._TRECHOS_POSICAO_ATUAL_VTA)
    pontos += [("itens_PC", "K2"), ("itens_PC", "L2"), ("itens_PC", "K5001"),
               ("cobertura_temporal", "B8")]
    assert [p for p in pontos if tpl[p[0]][p[1]].value != gerada[p[0]][p[1]].value] == []


def test_dropdown_pc_pago_segue_so_sim_nao_e_leitura_normaliza_o_acento():
    from _leitor_masterfile_v10 import _pc_pago_canonico

    tpl = _template()
    dvs = [(str(dv.sqref), dv.formula1) for dv in tpl["itens_PC"].data_validations.dataValidation]
    assert dvs == [("G2:G5001", "OPCOES_SIM_NAO")]
    assert [tpl["parametros"][c].value for c in ("T2", "T3")] == ["Sim", "Nao"]
    assert [_pc_pago_canonico(v) for v in ("Não", "NÃO", "nao", "Nao", "sim", "", None, "x")] == [
        "Nao", "Nao", "Nao", "Nao", "Sim", "", None, "x",
    ]
