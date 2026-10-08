"""Coleta 11.5 — historico necessario em parametros!A8 + orientacao de itens
novos por aditivo (mensagens de entrada em itens_Remanesc e aditivos).

Cenario do falso alerta (espelho do PR #177): somente C1 apurado, data de
corte 08/10/2026 (ciclo em execucao C4). C2/C3/C4 nao sao historico da
apuracao atual e nunca podem ser pedidos. Os testes rapidos conferem template
e Coleta gerada; o bloco Excel COM (RUN_EXCEL_INTEGRATION=1) recalcula os
cenarios no Excel real.
"""
from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import tools.aplicar_ajustes_xls_ux_pos174 as ux
from tests._fabrica_coleta import bytes_coleta_oficial, workbook_coleta_oficial
from tests.test_coleta_114_aditivos_ciclo_execucao import PAYLOAD_C1
from tools.aplicar_coleta_115_historico_orientacao import (
    FORMULA_CF_HISTORICO,
    ORIENTACOES,
    validar_orientacoes,
)

RAIZ = Path(__file__).resolve().parents[1]
TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
PERCENTUAIS = {1: 0.0512, 2: 0.0374, 3: 0.0289}


def _payload(numeros: list[int], corte: date) -> dict:
    """Analise dos ciclos `numeros` (data-base 01/2021, C4 a partir de 01/2025)."""
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2021, 1, 1),
        "data_corte": corte,
        "ciclos": [
            {
                "ciclo": f"C{n}",
                "data_inicio": date(2021 + n, 1, 1),
                "percentual": PERCENTUAIS[n],
                "inicio_efeito_financeiro": date(2021 + n, 1, 1),
                "situacao_aplicada": "✅ TEMPESTIVO",
                "data_pedido": date(2021 + n, 1, 10),
            }
            for n in numeros
        ],
    }


def _prompts(wb) -> dict[tuple[str, str], tuple]:
    return {
        (aba, str(dv.sqref)): (
            dv.type, dv.formula1, dv.promptTitle,
            (dv.prompt or "").replace("_x000a_", "\n"), dv.showInputMessage,
        )
        for aba in ("itens_Remanesc", "aditivos")
        for dv in wb[aba].data_validations.dataValidation
    }


@pytest.fixture(scope="module")
def template():
    wb = load_workbook(TEMPLATE)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def gerado():
    return workbook_coleta_oficial(PAYLOAD_C1)


# --- historico -------------------------------------------------------------

@pytest.mark.parametrize("fonte", ["template", "gerado"])
def test_aviso_historico_nao_depende_do_ciclo_em_execucao(fonte, request):
    ws = request.getfixturevalue(fonte)["parametros"]
    formula = ws[ux.CELULA_AVISO].value
    assert formula == ux.formula_aviso_historico()
    assert "CONTROLE!" not in formula  # nunca o ciclo em execucao (B2)
    # Cn so e pedido se ha ciclo computado depois dele; C4 nunca e historico.
    assert 'AND(COUNTIF($A$4:$A$6,"Sim")>0,NOT(ISNUMBER($E$3)))' in formula
    assert 'AND(COUNTIF($A$5:$A$6,"Sim")>0,NOT(ISNUMBER($E$4)))' in formula
    assert 'AND(COUNTIF($A$6:$A$6,"Sim")>0,NOT(ISNUMBER($E$5)))' in formula
    assert "$E$6" not in formula
    ux.validar_ascii_e_parenteses({"A8": formula})


@pytest.mark.parametrize("fonte", ["template", "gerado"])
def test_destaque_de_percentual_historico_usa_a_mesma_regra(fonte, request):
    ws = request.getfixturevalue(fonte)["parametros"]
    regras = [(str(f.sqref), r.formula) for f in ws.conditional_formatting for r in f.rules]
    assert ("E3:E6", [FORMULA_CF_HISTORICO]) in regras


def test_ciclo_em_execucao_preservado(gerado):
    from tools.aplicar_coleta_114_aditivos_ciclo_execucao import FORMULA_CICLO_VIGENTE

    assert gerado["CONTROLE"]["B2"].value == FORMULA_CICLO_VIGENTE
    assert gerado["CICLO_EM_EXECUCAO"]["C3"].value == "=CONTROLE!$B$2"


# --- orientacao ------------------------------------------------------------

def test_textos_cabem_nos_limites_do_excel():
    validar_orientacoes()


@pytest.mark.parametrize("fonte", ["template", "gerado"])
def test_mensagens_de_entrada_presentes(fonte, request):
    prompts = _prompts(request.getfixturevalue(fonte))
    for aba, faixa, titulo, mensagem in ORIENTACOES:
        _tipo, _lista, titulo_lido, mensagem_lida, mostra = prompts[(aba, faixa)]
        assert (titulo_lido, mensagem_lida, mostra) == (titulo, mensagem, True)


@pytest.mark.parametrize("fonte", ["template", "gerado"])
def test_orientacao_nao_restringe_valores_e_preserva_dropdown(fonte, request):
    prompts = _prompts(request.getfixturevalue(fonte))
    for chave in (("itens_Remanesc", "A2:A200"), ("itens_Remanesc", "B2:B200"),
                  ("itens_Remanesc", "C2:C200"), ("aditivos", "A2:A200"),
                  ("aditivos", "E2:E200")):
        assert prompts[chave][:2] == (None, None)  # "qualquer valor": N001 nao e regra
    assert prompts[("aditivos", "D2:D200")][:2] == (
        "list", '"Acrescimo,Acréscimo - novo item,Supressao"')


def test_conteudo_da_orientacao():
    texto = {(aba, faixa): (titulo, msg) for aba, faixa, titulo, msg in ORIENTACOES}
    titulo, item = texto[("itens_Remanesc", "A2:A200")]
    assert titulo == "ITEM NOVO POR ADITIVO"
    assert "N001 = código do item | 0 = quantidade-base original." in item
    assert "PRIMEIRO" in item and "nasce na aba aditivos" in item
    assert "0 aparece sozinho" in texto[("itens_Remanesc", "B2:B200")][1]
    assert "cadastre-o ANTES em itens_Remanesc" in texto[("aditivos", "A2:A200")][1]
    tipo = texto[("aditivos", "D2:D200")][1]
    assert '"Acréscimo - novo item" somente no aditivo em que o item nasce' in tipo
    assert 'Novo aumento desse item depois: use apenas "Acrescimo"' in tipo
    qtd = texto[("aditivos", "E2:E200")][1]
    assert "1.851, nunca 3.702" in qtd and "2.051" in qtd


# --- Excel real ------------------------------------------------------------

ROSA = 0xCEC7FF  # FFC7CE em BGR: destaque de percentual historico ausente


@pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar o Excel COM",
)
def test_cenarios_de_historico_no_excel(tmp_path):
    import pythoncom
    import win32com.client as win32

    corte_c4 = datetime(2025, 6, 30)
    cenarios = [
        # (nome, payload, corte, edicoes, avisos esperados apos cada passo,
        #  E com destaque rosa no estado inicial)
        ("A_c1_corte_c4", PAYLOAD_C1, datetime(2026, 10, 8), [], [""], set()),
        ("B_c1_corte_c1", PAYLOAD_C1, datetime(2024, 2, 29), [], [""], set()),
        ("C_D_c3_corte_c4", _payload([3], date(2025, 6, 30)), corte_c4,
         [("E3", 0.05), ("E4", 0.04)],
         ["HISTÓRICO INCOMPLETO — informe os percentuais de C1 e C2.",
          "HISTÓRICO INCOMPLETO — informe o percentual de C2.", ""],
         {"E3", "E4"}),
        ("C1_C3_sem_C2", _payload([1, 3], date(2025, 6, 30)), corte_c4, [],
         ["HISTÓRICO INCOMPLETO — informe o percentual de C2."], {"E4"}),
    ]
    pythoncom.CoInitialize()
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        for nome, payload, corte, edicoes, esperados, rosas in cenarios:
            caminho = tmp_path / f"{nome}.xlsx"
            caminho.write_bytes(bytes_coleta_oficial(payload))
            wb = xl.Workbooks.Open(str(caminho))
            try:
                wb.Worksheets("CONTROLE").Range("B3").Value = corte
                xl.CalculateFull()
                ws = wb.Worksheets("parametros")
                if nome != "B_c1_corte_c1":
                    assert wb.Worksheets("CONTROLE").Range("B2").Value == "C4", nome
                destacadas = {
                    f"E{r}" for r in range(3, 7)
                    if int(ws.Range(f"E{r}").DisplayFormat.Interior.Color) == ROSA
                }
                assert destacadas == rosas, nome
                lidos = [ws.Range("A8").Text]
                for celula, valor in edicoes:
                    ws.Range(celula).Value = valor
                    xl.CalculateFull()
                    lidos.append(ws.Range("A8").Text)
                assert lidos == esperados, nome
                if nome == "A_c1_corte_c4":
                    msg = wb.Worksheets("aditivos").Range("D2").Validation.InputMessage
                    assert "\n" in msg and "_x000a_" not in msg
            finally:
                wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
        pythoncom.CoUninitialize()
