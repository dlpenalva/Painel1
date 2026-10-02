"""Paridade estrutural: MODELOS em branco x documentos processados.

Regra de negocio: o modelo em branco e o MESMO documento do Cl8us, com
placeholders no lugar dos dados. Diferencas de conteudo (valores, datas, nomes,
redacao assertiva x instrutiva) sao esperadas; diferencas de ESTRUTURA por
desatualizacao nao sao.

O teste reduz cada DOCX a uma sequencia de MARCOS estruturais — propriedades de
secao, titulos, itens numerados, considerandos, titulos de quadro, cabecalhos e
numero de colunas das tabelas, quebra de pagina do anexo e assinaturas — e
compara o modelo com o processado. So o que pode variar legitimamente fica de
fora da comparacao, e cada excecao e declarada e fixada aqui:

* valores, datas, nomes, referencias e placeholders (o texto das celulas e dos
  paragrafos comuns nao entra nos marcos);
* a REGIAO METODO-DEPENDENTE (secao 2 do Termo; quadros da secao 3 do
  Despacho), que tem redacao propria por metodo canonico e e fixada por
  testes posicionais proprios;
* marcos OPCIONAIS no processado (so existem quando ha dado), que o modelo
  sempre exibe.

Se alguem acrescentar um quadro, mudar um cabecalho, reordenar uma secao ou
criar um anexo no documento processado e esquecer o modelo, estes testes falham.
"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph

from _sumario_executivo import TEXTO_PERCENTUAL_FATOR
from _templates_documentos import (
    CABECALHO_QUADRO1_TERMO,
    CABECALHO_QUADRO2_FINANCEIRO_TERMO,
    CABECALHO_QUADRO3_TERMO,
    CABECALHO_QUADRO4_TERMO,
    TITULO_QUADRO1_TERMO,
    TITULO_QUADRO2_FINANCEIRO_TERMO,
    TITULO_QUADRO3_TERMO,
    TITULO_QUADRO4_TERMO,
    gerar_despacho_saneador,
    gerar_modelo_branco_despacho,
    gerar_modelo_branco_termo,
    gerar_termo_apostila,
)
from test_templates_documentos import (
    CAMPOS_SANEADOR,
    CAMPOS_TERMO,
    _leitura_consumidos_canonica,
    _leitura_financeiro,
    _leitura_pc_com_potencial,
    _leitura_pc_sem_potencial,
)

ROOT = Path(__file__).resolve().parents[1]
PAGE14 = (ROOT / "pages" / "14_Central_Modelos_Ferramentas.py").read_text(
    encoding="utf-8"
)

TITULO_TERMO_2 = "2. Da apuração financeira dos valores retroativos"
TITULO_TERMO_3 = "3. Da composição do Valor Total Atualizado do Contrato"
TITULO_DESPACHO_3 = "3. RESULTADO"
TITULO_DESPACHO_4 = "4. DOCUMENTOS E VERIFICAÇÕES"

CAMPOS_TERMO_COMPLETOS = dict(
    CAMPOS_TERMO,
    deliberacao_institucional="Deliberação institucional de teste",
    instrumentos_posteriores="Termo Aditivo de teste",
)

METODOS = {
    "pc_sem_potencial": _leitura_pc_sem_potencial,
    "pc_com_potencial": _leitura_pc_com_potencial,
    "financeiro": _leitura_financeiro,
    "consumidos": _leitura_consumidos_canonica,
}


# ---------------------------------------------------------------------------
# Representacao normalizada do DOCX
# ---------------------------------------------------------------------------

def _norm(texto: str) -> str:
    return re.sub(r"\s+", " ", texto or "").strip()


def _corpo(doc):
    for elemento in doc.element.body.iterchildren():
        if elemento.tag.endswith("}p"):
            yield Paragraph(elemento, doc)
        elif elemento.tag.endswith("}tbl"):
            yield Table(elemento, doc)


def _marcos_paragrafo(par: Paragraph, texto: str) -> list[tuple]:
    runs = [r for r in par.runs if r.text.strip()]
    todos_negrito = bool(runs) and all(r.bold for r in runs)
    primeiro_negrito = bool(runs) and bool(runs[0].bold)

    if texto.startswith(("Quadro", "Tabela")) and len(texto) < 140:
        return [("QUADRO", texto.replace(" (continuação)", ""))]
    if texto == "TELECOMUNICAÇÕES BRASILEIRAS S.A. - TELEBRAS":
        return [("ASSINATURA",)]
    if texto.startswith("Brasília/DF,"):
        return [("LOCAL_DATA",)]
    if todos_negrito:
        # O modelo se identifica como "MODELO PADRAO"; o resto do titulo e igual.
        return [("TITULO", texto.replace(" - MODELO PADRÃO", ""))]
    item = re.match(r"^(\d+\.\d+)\. ", texto)
    if item and not primeiro_negrito:
        return [("ITEM", item.group(1))]
    considerando = re.match(r"^(\d+)\. ", texto)
    if (
        considerando and primeiro_negrito
        and par.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    ):
        return [("CONSIDERANDO", int(considerando.group(1)))]
    if texto.startswith("Assunto:"):
        return [("ASSUNTO",)]
    if texto.startswith("Referência(s):"):
        return [("REFERENCIA",)]
    if texto.startswith("Contrato nº"):
        return [("CONTRATO_LINHA",)]
    if texto.startswith("Acumulado dos ciclos:"):
        return [("ACUMULADO",)]
    if texto.startswith("Para o cálculo, o percentual de reajuste de cada ciclo"):
        return [("REGRA_DUAS_CASAS",)]
    if texto.startswith("SITUAÇÃO DOS VALORES RETROATIVOS"):
        return [("BOX_RETROATIVOS",)]
    return []


def _marcos(docx_bytes: bytes) -> list[tuple]:
    doc = Document(BytesIO(docx_bytes))
    marcos: list[tuple] = []
    for secao in doc.sections:
        marcos.append((
            "SECAO", int(secao.orientation),
            round(secao.page_width.cm, 2), round(secao.page_height.cm, 2),
            round(secao.left_margin.cm, 2), round(secao.right_margin.cm, 2),
            round(secao.top_margin.cm, 2), round(secao.bottom_margin.cm, 2),
        ))
    dentro_anexo = False
    for bloco in _corpo(doc):
        if isinstance(bloco, Table):
            cabecalho = tuple(_norm(c.text) for c in bloco.rows[0].cells)
            if dentro_anexo:
                # Ciclos do anexo variam: so o formato do cabecalho e fixo.
                assert cabecalho[0] == "Item"
                assert all(re.fullmatch(r"VU_C\d+", c) for c in cabecalho[1:])
                marcos.append(("TABELA_ANEXO_VU", cabecalho[0]))
            else:
                marcos.append(("TABELA", len(bloco.columns), cabecalho))
            continue
        if 'w:br w:type="page"' in bloco._p.xml:
            marcos.append(("QUEBRA_DE_PAGINA",))
        texto = _norm(bloco.text)
        if not texto:
            continue
        novos = _marcos_paragrafo(bloco, texto)
        if novos and novos[0][0] == "TITULO" and novos[0][1].startswith("ANEXO 1"):
            dentro_anexo = True
        marcos.extend(novos)
    return marcos


def _sem_regiao(marcos: list[tuple], inicio: tuple, fim: tuple, *,
                manter: tuple[str, ...] = ()) -> list[tuple]:
    """Remove a regiao metodo-dependente entre dois titulos (exclusive)."""
    saida: list[tuple] = []
    dentro = False
    for marco in marcos:
        if marco == fim:
            dentro = False
        if not dentro or marco[0] in manter:
            saida.append(marco)
        if marco == inicio:
            dentro = True
    return saida


def _nucleo_termo(marcos: list[tuple]) -> list[tuple]:
    return _sem_regiao(
        marcos, ("TITULO", TITULO_TERMO_2), ("TITULO", TITULO_TERMO_3)
    )


def _nucleo_despacho(marcos: list[tuple]) -> list[tuple]:
    # Na secao 3 do Despacho os quadros variam por metodo; o titulo do Quadro 2
    # e fixo e permanece na comparacao.
    return _sem_regiao(
        marcos, ("TITULO", TITULO_DESPACHO_3), ("TITULO", TITULO_DESPACHO_4),
        manter=("QUADRO",),
    )


def _regiao(marcos: list[tuple], inicio: tuple, fim: tuple) -> list[tuple]:
    dentro = False
    saida: list[tuple] = []
    for marco in marcos:
        if marco == fim:
            dentro = False
        if dentro:
            saida.append(marco)
        if marco == inicio:
            dentro = True
    return saida


def _opcional_termo(marco: tuple) -> bool:
    """Marcos que o processado so emite quando ha dado correspondente."""
    if marco[0] == "CONSIDERANDO":
        return marco[1] >= 12
    if marco[0] == "ITEM":
        secao, subitem = (int(parte) for parte in marco[1].split("."))
        return (secao == 1 and subitem >= 3) or (secao == 3 and subitem >= 2)
    if marco[0] == "QUADRO":
        return marco[1] == TITULO_QUADRO4_TERMO or marco[1].startswith("Tabela 1")
    if marco[0] == "TABELA":
        return marco[2] == tuple(CABECALHO_QUADRO4_TERMO)
    # Anexo: o processado so o emite quando ha historico de valores unitarios.
    return marco[0] in {"QUEBRA_DE_PAGINA", "TABELA_ANEXO_VU"} or (
        marco[0] == "TITULO" and marco[1].startswith("ANEXO 1")
    )


def _eh_subsequencia(sub: list[tuple], seq: list[tuple]) -> bool:
    it = iter(seq)
    return all(any(item == candidato for candidato in it) for item in sub)


def _primeira_diferenca(a: list[tuple], b: list[tuple]) -> str:
    for i, (x, y) in enumerate(zip(a, b)):
        if x != y:
            return f"posicao {i}: modelo={x!r} processado={y!r}"
    return f"tamanhos diferentes: modelo={len(a)} processado={len(b)}"


@pytest.fixture(scope="module")
def modelo_despacho() -> list[tuple]:
    return _marcos(gerar_modelo_branco_despacho())


@pytest.fixture(scope="module")
def modelo_termo() -> list[tuple]:
    return _marcos(gerar_modelo_branco_termo())


def _despacho_processado(nome: str) -> list[tuple]:
    return _marcos(gerar_despacho_saneador(
        METODOS[nome](), campos_manuais=dict(CAMPOS_SANEADOR)
    ))


def _termo_processado(nome: str, campos: dict | None = None) -> list[tuple]:
    return _marcos(gerar_termo_apostila(
        METODOS[nome](), campos_manuais=dict(campos or CAMPOS_TERMO)
    ))


# ---------------------------------------------------------------------------
# 1. Paridade estrutural exata (processado com todos os blocos opcionais)
# ---------------------------------------------------------------------------

def test_despacho_modelo_tem_a_mesma_estrutura_do_processado(modelo_despacho):
    modelo = _nucleo_despacho(modelo_despacho)
    processado = _nucleo_despacho(_despacho_processado("financeiro"))
    assert modelo == processado, _primeira_diferenca(modelo, processado)


def test_termo_modelo_tem_a_mesma_estrutura_do_processado(modelo_termo):
    modelo = _nucleo_termo(modelo_termo)
    processado = _nucleo_termo(
        _termo_processado("financeiro", CAMPOS_TERMO_COMPLETOS)
    )
    assert modelo == processado, _primeira_diferenca(modelo, processado)


# ---------------------------------------------------------------------------
# 2. Todo marco do processado, em qualquer metodo, existe no modelo, na ordem
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("metodo", sorted(METODOS))
def test_despacho_processado_nao_tem_estrutura_ausente_do_modelo(
    metodo, modelo_despacho
):
    modelo = _nucleo_despacho(modelo_despacho)
    processado = _nucleo_despacho(_despacho_processado(metodo))
    assert _eh_subsequencia(processado, modelo), (
        f"{metodo}: estrutura do processado ausente do modelo — "
        + _primeira_diferenca(modelo, processado)
    )


@pytest.mark.parametrize("metodo", sorted(METODOS))
def test_termo_processado_nao_tem_estrutura_ausente_do_modelo(
    metodo, modelo_termo
):
    modelo = _nucleo_termo(modelo_termo)
    processado = _nucleo_termo(_termo_processado(metodo))
    obrigatorio = [m for m in processado if not _opcional_termo(m)]
    assert _eh_subsequencia(obrigatorio, modelo), (
        f"{metodo}: estrutura do processado ausente do modelo — "
        + _primeira_diferenca(modelo, processado)
    )


def test_termo_modelo_exibe_todos_os_marcos_opcionais(modelo_termo):
    """O que o processado emite so com dado, o modelo mostra sempre."""
    assert ("CONSIDERANDO", 12) in modelo_termo
    assert ("CONSIDERANDO", 13) in modelo_termo
    for item in ("1.2", "1.3", "3.2", "5.1", "5.2"):
        assert ("ITEM", item) in modelo_termo, item
    assert ("QUADRO", TITULO_QUADRO4_TERMO) in modelo_termo
    assert (
        "TABELA", len(CABECALHO_QUADRO4_TERMO), tuple(CABECALHO_QUADRO4_TERMO)
    ) in modelo_termo
    assert ("QUEBRA_DE_PAGINA",) in modelo_termo
    assert ("TABELA_ANEXO_VU", "Item") in modelo_termo


# ---------------------------------------------------------------------------
# 3. Regiao metodo-dependente: fixada por teste posicional
# ---------------------------------------------------------------------------

def _secao2_termo(marcos: list[tuple]) -> list[tuple]:
    return _regiao(
        marcos, ("TITULO", TITULO_TERMO_2), ("TITULO", TITULO_TERMO_3)
    )


def _secao3_despacho(marcos: list[tuple]) -> list[tuple]:
    return _regiao(
        marcos, ("TITULO", TITULO_DESPACHO_3), ("TITULO", TITULO_DESPACHO_4)
    )


CABECALHO_Q2_FIN = tuple(CABECALHO_QUADRO2_FINANCEIRO_TERMO)
SECAO2_FINANCEIRO = [
    ("ITEM", "2.1"),
    ("QUADRO", TITULO_QUADRO2_FINANCEIRO_TERMO),
    ("TABELA", 4, CABECALHO_Q2_FIN),
]


def test_termo_secao2_modelo_e_igual_a_do_metodo_financeiro(modelo_termo):
    """O modelo adota o formato Financeiro; titulo e cabecalho tem fonte unica."""
    assert _secao2_termo(modelo_termo) == SECAO2_FINANCEIRO
    assert _secao2_termo(_termo_processado("financeiro")) == SECAO2_FINANCEIRO


def test_termo_secao2_pc_vigente():
    esperado = [
        ("ITEM", "2.1"), ("ITEM", "2.2"), ("ITEM", "2.3"),
        ("QUADRO", "Quadro 2 — Execução reconhecida e retroativo por ciclo"),
        ("TABELA", 4, (
            "Ciclo", "Pedidos de Compra reconhecidos / valor original",
            "Pedidos de Compra reconhecidos / valor atualizado",
            "Retroativo reconhecido",
        )),
    ]
    assert _secao2_termo(_termo_processado("pc_sem_potencial")) == esperado
    com_potencial = _secao2_termo(_termo_processado("pc_com_potencial"))
    assert com_potencial[:len(esperado)] == esperado
    assert ("QUADRO", "Quadro complementar — Situação dos valores retroativos") \
        in com_potencial
    assert ("ITEM", "2.4") in com_potencial


def test_termo_secao2_consumidos_vigente():
    assert _secao2_termo(_termo_processado("consumidos")) == [
        ("ITEM", "2.1"), ("ITEM", "2.2"), ("ITEM", "2.3"),
    ]


def test_despacho_secao3_financeiro_e_modelo_tem_o_mesmo_quadro_2(modelo_despacho):
    esperado = [
        ("QUADRO", "Quadro 2 - Síntese financeira"),
        ("TABELA", 2, ("Resultado", "Valor")),
    ]
    assert _secao3_despacho(modelo_despacho) == esperado
    assert _secao3_despacho(_despacho_processado("financeiro")) == esperado
    assert _secao3_despacho(_despacho_processado("consumidos")) == esperado


def test_despacho_secao3_pc_vigente():
    marcos = _secao3_despacho(_despacho_processado("pc_sem_potencial"))
    assert marcos == [
        ("QUADRO", "Quadro 2 - Síntese financeira"),
        ("TABELA", 4, (
            "Ciclo", "Valor original", "Valor atualizado",
            "Retroativo reconhecido",
        )),
        ("TABELA", 2, ("Resultado", "Valor")),
        ("BOX_RETROATIVOS",),
    ]


# ---------------------------------------------------------------------------
# 4. Regressao: o que estava ultrapassado e a estrutura vigente esperada
# ---------------------------------------------------------------------------

def _texto_ordenado(docx_bytes: bytes) -> list[str]:
    doc = Document(BytesIO(docx_bytes))
    saida = []
    for bloco in _corpo(doc):
        if isinstance(bloco, Table):
            saida.append("TABELA:" + " | ".join(
                _norm(c.text) for c in bloco.rows[0].cells))
            for linha in bloco.rows[1:]:
                saida.append("LINHA:" + " | ".join(_norm(c.text) for c in linha.cells))
        elif _norm(bloco.text):
            saida.append(_norm(bloco.text))
    return saida


def test_despacho_modelo_traz_a_regra_das_duas_casas_apos_o_acumulado():
    linhas = _texto_ordenado(gerar_modelo_branco_despacho())
    indice = next(i for i, t in enumerate(linhas) if t.startswith("Acumulado dos ciclos:"))
    assert linhas[indice + 1] == TEXTO_PERCENTUAL_FATOR
    assert "por exemplo" not in linhas[indice + 1]


def test_termo_modelo_quadro_1_tem_linha_de_acumulado_e_itens_1_2_e_1_3():
    linhas = _texto_ordenado(gerar_modelo_branco_termo())
    q1 = linhas.index(TITULO_QUADRO1_TERMO)
    assert linhas[q1 + 1] == "TABELA:" + " | ".join(CABECALHO_QUADRO1_TERMO)
    ultima = linhas[q1 + 3]
    assert ultima.startswith("LINHA:")
    celulas = [c.strip() for c in ultima[len("LINHA:"):].split(" | ")]
    assert celulas[1] == "Acumulado"
    assert celulas[3] == "Conforme composição dos ciclos"
    assert linhas[q1 + 4] == "1.2. " + TEXTO_PERCENTUAL_FATOR
    assert linhas[q1 + 5].startswith("1.3. Havendo competências não alcançadas")


def test_termo_modelo_quadro_3_tem_parcelas_e_total_como_no_processado():
    linhas = _texto_ordenado(gerar_modelo_branco_termo())
    q3 = linhas.index(TITULO_QUADRO3_TERMO)
    assert linhas[q3 + 1] == "TABELA:" + " | ".join(CABECALHO_QUADRO3_TERMO)
    assert linhas[q3 + 2].startswith("LINHA:A | ")
    assert linhas[q3 + 3].startswith("LINHA:B | ")
    # A composicao tem tamanho variavel: C cobre a terceira parcela (ex.:
    # retroativo potencial) e o 3.1 manda acrescentar linhas quando preciso.
    assert linhas[q3 + 4].startswith("LINHA:C | ")
    assert "quando houver" in linhas[q3 + 4]
    assert linhas[q3 + 5].startswith(
        "LINHA:Total (A + B + C) | Valor Total Atualizado do Contrato | "
    )
    assert linhas[q3 + 6].startswith("3.2. Havendo retroativo já incorporado")
    item_31 = next(t for t in linhas if t.startswith("3.1."))
    assert "uma linha para cada parcela da composição" in item_31


def test_modelo_avisa_que_a_secao_dependente_do_metodo_deve_ser_adaptada():
    termo = "\n".join(_texto_ordenado(gerar_modelo_branco_termo()))
    despacho = "\n".join(_texto_ordenado(gerar_modelo_branco_despacho()))
    aviso = "(Financeiro, Pedidos de Compra ou Itens Consumidos)"
    assert "Esta seção deverá ser adaptada ao método de apuração adotado " + aviso in termo
    assert "O quadro deverá ser adaptado ao método de apuração adotado " + aviso in despacho


def test_termo_modelo_aditivos_usa_quadro_4_como_o_processado():
    linhas = _texto_ordenado(gerar_modelo_branco_termo())
    i51 = next(i for i, t in enumerate(linhas) if t.startswith("5.1."))
    assert linhas[i51 + 1] == TITULO_QUADRO4_TERMO
    assert linhas[i51 + 2] == "TABELA:" + " | ".join(CABECALHO_QUADRO4_TERMO)
    assert linhas[i51 + 4].startswith("5.2.")


def test_termo_processado_mantem_quadro_1_com_acumulado_e_item_1_2():
    linhas = _texto_ordenado(gerar_termo_apostila(
        _leitura_financeiro(), campos_manuais=dict(CAMPOS_TERMO)
    ))
    q1 = linhas.index(TITULO_QUADRO1_TERMO)
    assert linhas[q1 + 1] == "TABELA:" + " | ".join(CABECALHO_QUADRO1_TERMO)
    assert "Acumulado" in linhas[q1 + 3]
    assert linhas[q1 + 4] == "1.2. " + TEXTO_PERCENTUAL_FATOR


# ---------------------------------------------------------------------------
# 5. Seguranca do modelo: a estrutura nova nao afirma fato
# ---------------------------------------------------------------------------

AFIRMACOES_PROIBIDAS = [
    "Em razão da data do pedido, os efeitos financeiros do reajuste",
    "já está incorporado ao valor da execução atualizada considerado na "
    "composição do Quadro 3. Por essa razão",
    "Percentual acumulado apurado",
    "Foram consideradas na apuração as alterações contratuais",
    "Não foram identificados aditivos",
    "R$ 0,00",
]


@pytest.mark.parametrize("gerar", [
    gerar_modelo_branco_despacho, gerar_modelo_branco_termo,
])
def test_modelo_nao_afirma_fato_da_estrutura_nova(gerar):
    texto = "\n".join(_texto_ordenado(gerar()))
    achou = [f for f in AFIRMACOES_PROIBIDAS if f in texto]
    assert not achou, achou


FRASE_SEM_PENDENCIAS = "Não existem pendências nesta data."
TEXTO_PENDENCIAS_MODELO = (
    "Registrar abaixo as pendências identificadas na análise, quando houver. "
    "Caso não existam pendências, registrar expressamente essa condição."
)


def _secao6_despacho_texto(docx_bytes: bytes) -> list[str]:
    linhas = _texto_ordenado(docx_bytes)
    inicio = linhas.index("6. PENDÊNCIAS")
    fim = linhas.index("7. CONCLUSÃO")
    return linhas[inicio + 1:fim]


def test_modelo_do_despacho_nao_afirma_existencia_nem_ausencia_de_pendencias():
    modelo = gerar_modelo_branco_despacho()
    texto = "\n".join(_texto_ordenado(modelo))
    # Nunca a frase de ausencia, em nenhuma parte do modelo.
    assert FRASE_SEM_PENDENCIAS not in texto
    assert "Não existem pendências" not in texto
    # A secao existe, com orientacao neutra seguida do campo a preencher.
    assert _secao6_despacho_texto(modelo) == [
        TEXTO_PENDENCIAS_MODELO, "[PREENCHER, EM CASO DE PENDÊNCIA]",
    ]
    # O placeholder segue destacado em amarelo.
    runs = [
        r for p in Document(BytesIO(modelo)).paragraphs for r in p.runs
        if r.text == "[PREENCHER, EM CASO DE PENDÊNCIA]"
    ]
    assert len(runs) == 1 and runs[0]._element.rPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight"
    ) is not None


def test_secao_pendencias_do_modelo_mantem_a_posicao_do_processado(modelo_despacho):
    processado = _despacho_processado("pc_sem_potencial")
    titulos = [m[1] for m in modelo_despacho if m[0] == "TITULO"]
    assert titulos.index("6. PENDÊNCIAS") == titulos.index("5. CONTROLE DA ADEQUAÇÃO ORÇAMENTÁRIA") + 1
    assert titulos.index("7. CONCLUSÃO") == titulos.index("6. PENDÊNCIAS") + 1
    assert [m for m in modelo_despacho if m[0] == "TITULO"] == [
        m for m in processado if m[0] == "TITULO"
    ]


def test_despacho_processado_sem_pendencias_continua_podendo_afirmar_ausencia():
    """A frase e legitima quando o estado real a justifica."""
    processado = gerar_despacho_saneador(
        _leitura_pc_sem_potencial(), campos_manuais=dict(CAMPOS_SANEADOR)
    )
    assert _secao6_despacho_texto(processado) == [FRASE_SEM_PENDENCIAS]


def test_despacho_processado_com_pendencia_nao_usa_a_frase_nem_a_orientacao_do_modelo():
    processado = gerar_despacho_saneador(
        _leitura_pc_com_potencial(), campos_manuais=dict(CAMPOS_SANEADOR)
    )
    secao = "\n".join(_secao6_despacho_texto(processado))
    assert FRASE_SEM_PENDENCIAS not in secao
    assert TEXTO_PENDENCIAS_MODELO not in secao
    assert "PENDÊNCIA TÉCNICA" in secao or "PROVIDÊNCIA DA ÁREA GESTORA" in secao


@pytest.mark.parametrize("gerar", [
    gerar_modelo_branco_despacho, gerar_modelo_branco_termo,
])
def test_modelo_nao_carrega_id_de_apuracao(gerar):
    doc = Document(BytesIO(gerar()))
    xml = doc.element.xml + "".join(
        s.footer._element.xml + s.header._element.xml for s in doc.sections
    )
    assert "APUR-" not in xml


# ---------------------------------------------------------------------------
# 6. Central de Modelos: pagina simples, sem logica documental
# ---------------------------------------------------------------------------

def test_central_descreve_os_modelos_como_mesma_estrutura_dos_gerados():
    fonte = re.sub(r'"\s*\n\s*"', "", PAGE14)
    assert (
        "Modelo atualizado com a mesma estrutura utilizada no Despacho "
        "Saneador gerado pelo Cl8us." in fonte
    )
    assert (
        "Modelo atualizado com a mesma estrutura utilizada no Termo de "
        "Apostila gerado pelo Cl8us." in fonte
    )


def test_central_qualifica_a_regiao_que_depende_do_metodo():
    """A paridade vale para a estrutura; o formato por metodo e declarado."""
    fonte = re.sub(r'"\s*\n\s*"', "", PAGE14)
    assert "A seção 2 segue o formato da apuração financeira e deve ser adaptada ao método adotado." in fonte
    assert "O quadro de resultado financeiro segue o formato da apuração financeira e deve ser adaptado ao método adotado." in fonte


def test_central_apenas_chama_os_wrappers_dos_modelos():
    assert "gerar_modelo_branco_despacho()" in PAGE14
    assert "gerar_modelo_branco_termo()" in PAGE14
    for proibido in ("from docx", "import docx", "_ta_", "_ds_",
                     "gerar_despacho_saneador", "gerar_termo_apostila",
                     "modo_modelo_em_branco"):
        assert proibido not in PAGE14, proibido
