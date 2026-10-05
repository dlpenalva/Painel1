"""AJUSTES-XLS-UX pos-PR #174 — Frente 2 — RESULTADOS_DETALHE oculta (hidden, nunca veryHidden); D31 sem
hiperlink; nomes, leitores e ajustes manuais intactos.

Testes de Excel real rodam com RUN_EXCEL_INTEGRATION=1.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_ajustes_xls_ux_pos174 as ux  # noqa: E402
from _coleta_oficial import (  # noqa: E402
    gerar_coleta_oficial_preenchida,
    obter_coleta_oficial_bytes,
)
from _memoria_calculo import (  # noqa: E402
    FILL_FRONTEIRA_IST,
    LEGENDA_FRONTEIRA_IST,
    LINHA_LEGENDA_FRONTEIRA,
    escrever_memoria_calculo,
    ler_memoria_calculo,
)

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
COLUNAS_JR = ("J", "K", "L", "M", "N", "O", "P", "Q", "R")

com = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1",
)


@pytest.fixture(scope="module")
def template():
    wb = load_workbook(TEMPLATE)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def coleta_branca():
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    yield wb
    wb.close()


def _rgb(cell) -> str | None:
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


# ------------------------------------------------------------------ payloads
PERCENTUAIS = {1: 0.0512, 2: 0.0374, 3: 0.0289, 4: 0.0315}


def _payload(numeros: list[int]) -> dict:
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2021, 1, 1),
        "data_corte": date(2021 + max(numeros), 12, 31),
        "ciclos": [
            {
                "ciclo": f"C{n}",
                "data_inicio": date(2021 + n, 1, 1),
                "percentual": PERCENTUAIS[n],
                "inicio_efeito_financeiro": date(2021 + n, 1, 1),
                "situacao_aplicada": "✅ TEMPESTIVO",
                "data_pedido": date(2021 + n, 1, 10),
            }
            for n in numeros
        ],
    }


# =================================================================== FRENTE 2
def test_f2_detalhe_oculto_normal_e_executiva_visivel(template, coleta_branca):
    for wb in (template, coleta_branca):
        assert wb["RESULTADOS_DETALHE"].sheet_state == "hidden"
        assert wb["RESULTADOS"].sheet_state == "visible"
        assert wb["MEMORIA_RESULTADOS"].sheet_state == "hidden"


def test_f2_coleta_preenchida_mantem_detalhe_oculto():
    wb = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(_payload([1, 2]))))
    assert wb["RESULTADOS_DETALHE"].sheet_state == "hidden"
    assert wb["RESULTADOS"].sheet_state == "visible"
    vista = wb.views[0]
    assert wb.worksheets[vista.activeTab].sheet_state == "visible"


def test_f2_d31_texto_normal_sem_hiperlink(template, coleta_branca):
    for wb in (template, coleta_branca):
        ws = wb["RESULTADOS"]
        assert ws["D31"].hyperlink is None
        assert ws["D31"].value == ux.TEXTO_AJUSTES_D31
        assert not ws["D31"].font.u
        assert ws["B3"].value == ux.SUBTITULO_B3
    # C31 (situacao dos ajustes) segue espelhando a camada tecnica.
    assert "RESULTADOS_DETALHE!$C$43:$G$50" in template["RESULTADOS"]["C31"].value


def test_f2_nomes_e_ajustes_manuais_preservados(template):
    destinos = {n: d.attr_text for n, d in template.defined_names.items()}
    assert destinos["STATUS_RESULTADOS"] == "RESULTADOS_DETALHE!$B$3"
    assert destinos["OPCOES_APLICAR_MANUAL"] == "RESULTADOS_DETALHE!$J$2:$J$3"
    det = template["RESULTADOS_DETALHE"]
    for linha in range(43, 51):
        for col in "CDEFG":
            assert det[f"{col}{linha}"].value is None  # entradas livres
    dvs = {str(dv.sqref): dv.formula1 for dv in det.data_validations.dataValidation}
    assert dvs.get("G43:G50") == "OPCOES_APLICAR_MANUAL"
    mem = template["MEMORIA_RESULTADOS"]
    refs = [
        c.value for linha in mem.iter_rows() for c in linha
        if isinstance(c.value, str) and "RESULTADOS_DETALHE!" in c.value
        and "43" in c.value
    ]
    assert refs, "MEMORIA_RESULTADOS deixou de ler os ajustes manuais"


def test_f2_leitor_encontra_aba_tecnica_oculta(coleta_branca):
    from _coleta_reajuste import _validar_resultados_integra
    from _resultados_abas import aba_resultados_tecnica

    assert aba_resultados_tecnica(coleta_branca) == "RESULTADOS_DETALHE"
    assert _validar_resultados_integra(coleta_branca, "teste")["visivel"]


def test_f2_very_hidden_continua_rejeitado():
    from _coleta_reajuste import _validar_resultados_integra

    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    wb["RESULTADOS_DETALHE"].sheet_state = "veryHidden"
    with pytest.raises(ValueError, match="RESULTADOS_DETALHE"):
        _validar_resultados_integra(wb, "teste")


def test_f2_ajuste_manual_reexibido_chega_ao_leitor():
    """Apos reexibir, C43:G50 seguem editaveis e o arquivo segue legivel."""
    from _leitor_masterfile_v10 import ler_masterfile_v10

    wb = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(_payload([1]))))
    det = wb["RESULTADOS_DETALHE"]
    det.sheet_state = "visible"  # Reexibir
    det["C43"] = 1234.56
    det["D43"] = "Ajuste de teste"
    det["E43"] = "Fiscal"
    det["F43"] = date(2025, 1, 10)
    det["G43"] = "Sim"
    saida = io.BytesIO()
    wb.save(saida)
    lido = ler_masterfile_v10(saida.getvalue())
    assert "RESULTADOS_DETALHE" not in (lido.get("abas_ausentes") or [])
