"""CONSUMO-GLOSA-DV-1 — o dropdown AJUSTE_TIPO sobrevive a geracao da Coleta.

O template guarda a lista de `itens_Consumidos!Z2:Z6` dentro de um
`mc:AlternateContent`/`x12ac:list`. O openpyxl nao entende essa extensao e
descarta a lista JA NA LEITURA, entao a Coleta gerada saia com uma validacao
`type="list"` sem `formula1` — dropdown vazio para o fiscal.

Estes testes protegem a SAIDA GERADA, nao o template: e a saida que chega ao
usuario, e era exatamente ai que a regressao passava despercebida.
"""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

from _coleta_oficial import obter_coleta_oficial_bytes  # noqa: E402
from _leitor_masterfile_v10 import OPCOES_AJUSTE_TIPO  # noqa: E402

FAIXA = "Z2:Z6"


@pytest.fixture(scope="module")
def coleta_gerada():
    """A Coleta pela MESMA cadeia que a aplicacao usa no download."""
    wb = load_workbook(BytesIO(obter_coleta_oficial_bytes()))
    yield wb
    wb.close()


def _validacao_ajuste_tipo(wb):
    ws = wb["itens_Consumidos"]
    return [
        dv for dv in ws.data_validations.dataValidation
        if str(dv.sqref) == FAIXA
    ]


def _opcoes(dv):
    """Le a lista de forma semantica, sem depender do XML incidental."""
    return [
        opcao.strip()
        for opcao in str(dv.formula1 or "").strip('"').split(",")
        if opcao.strip()
    ]


def test_coleta_gerada_tem_dropdown_com_as_duas_opcoes(coleta_gerada):
    """A regressao que este teste existe para impedir: DV presente, lista vazia."""
    validacoes = _validacao_ajuste_tipo(coleta_gerada)
    assert len(validacoes) == 1, f"esperava 1 validacao em {FAIXA}"
    dv = validacoes[0]
    assert dv.type == "list"
    assert dv.formula1, (
        "dataValidation existe mas a lista esta vazia — o dropdown de "
        "AJUSTE_TIPO nao chegaria ao fiscal"
    )
    assert _opcoes(dv) == ["Valor pago", "Glosa"]


def test_opcoes_do_dropdown_sao_as_que_o_leitor_aceita(coleta_gerada):
    """O que se oferece e o que o motor reconhece nao podem divergir."""
    dv = _validacao_ajuste_tipo(coleta_gerada)[0]
    assert _opcoes(dv) == list(OPCOES_AJUSTE_TIPO)


def test_nenhuma_lista_do_workbook_gerado_fica_sem_fonte(coleta_gerada):
    """Guarda ampla e barata: nenhuma outra validacao pode perder a lista."""
    vazias = [
        (aba, str(dv.sqref))
        for aba in coleta_gerada.sheetnames
        for dv in coleta_gerada[aba].data_validations.dataValidation
        if dv.type == "list" and not dv.formula1
    ]
    assert vazias == [], f"listas sem fonte na Coleta gerada: {vazias}"


def test_a_faixa_e_a_grade_ao_lado_nao_foram_alteradas(coleta_gerada):
    """A correcao e so a ajuda de entrada: nada de layout, formula ou dado."""
    ws = coleta_gerada["itens_Consumidos"]
    # Cabecalhos do bloco intactos.
    assert ws["Z1"].value == "AJUSTE_TIPO"
    assert ws["AA1"].value == "AJUSTE_VALOR_INFORMADO"
    # Os dois campos manuais continuam vazios na Coleta em branco.
    for linha in range(2, 7):
        assert ws[f"Z{linha}"].value in (None, "")
        assert ws[f"AA{linha}"].value in (None, "")
    # E as colunas calculadas do bloco seguem sendo formulas.
    for coluna in ("Y", "AB", "AC", "AD", "AE", "AF", "AG"):
        valor = ws[f"{coluna}3"].value
        assert isinstance(valor, str) and valor.startswith("="), coluna


def test_template_oficial_permanece_intacto():
    """A correcao vive na geracao; o template binario homologado nao muda.

    No template a lista continua no `x12ac:list` que o openpyxl nao le — por
    isso `formula1` vem vazia aqui. Isso e o estado ESPERADO: o arquivo abre
    normalmente no Excel, e quem conserta a lista e o gerador.
    """
    caminho = (
        Path(__file__).resolve().parents[1]
        / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
    )
    wb = load_workbook(caminho)
    try:
        validacoes = [
            dv for dv in wb["itens_Consumidos"].data_validations.dataValidation
            if str(dv.sqref) == FAIXA
        ]
        assert len(validacoes) == 1
        assert validacoes[0].type == "list"
    finally:
        wb.close()
