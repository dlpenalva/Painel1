from __future__ import annotations

from datetime import date
from io import BytesIO
import zipfile

import pytest
from openpyxl import load_workbook

from _ciclo_em_execucao import calcular_posicao_ciclo_por_data
from _coleta_oficial import (
    TEMPLATE_COLETA_OFICIAL,
    _mapa_xml_planilhas,
    obter_coleta_oficial_bytes,
)


@pytest.fixture(scope="module")
def wb_coleta_116():
    return load_workbook(BytesIO(obter_coleta_oficial_bytes()), data_only=False)


def test_motor_usa_referencia_anterior_e_recorta_movimentos_por_data():
    resultado = calcular_posicao_ciclo_por_data(
        ciclo="C3",
        data_inicio=date(2026, 1, 1),
        data_fim=date(2026, 12, 31),
        data_posicao=date(2026, 8, 31),
        itens=[{
            "item": "1.1",
            "remanescente_referencia": 1000,
            "ciclo_referencia": "C1",
            "data_referencia": date(2024, 2, 29),
            "remanescente_atual": 650,
            "vu_atualizado": 10,
        }],
        movimentos=[
            {"item": "1.1", "data_efeito": date(2024, 2, 20), "delta": 80},
            {"item": "1.1", "data_efeito": date(2025, 5, 10), "delta": 200},
            {"item": "1.1", "data_efeito": date(2026, 9, 1), "delta": 300},
        ],
    )
    item = resultado["itens"][0]
    assert item["ciclo_referencia"] == "C1"
    assert item["data_referencia"] == date(2024, 2, 29)
    assert item["alteracoes_liquidas_periodo"] == 200
    assert item["quantidade_consumida"] == 550


def test_motor_sem_referencia_aceita_posicao_atual_mas_nao_inventa_consumo():
    resultado = calcular_posicao_ciclo_por_data(
        ciclo="C3",
        data_inicio=date(2026, 1, 1),
        data_fim=date(2026, 12, 31),
        data_posicao=date(2026, 8, 31),
        itens=[{
            "item": "1.1",
            "remanescente_referencia": None,
            "data_referencia": None,
            "remanescente_atual": 0,
            "vu_atualizado": 10,
        }],
    )
    item = resultado["itens"][0]
    assert item["ciclo_referencia"] is None
    assert item["data_referencia"] is None
    assert item["remanescente_atual"] == 0
    assert item["quantidade_consumida"] is None
    assert item["valor_remanescente_atualizado"] == 0
    assert item["consumo_mensuravel"] is False
    assert "CONSUMO_NAO_CALCULAVEL_SEM_REFERENCIA" in item["alertas"]
    assert resultado["valido"] is False


def test_motor_item_novo_nascido_depois_da_referencia_parte_de_zero():
    resultado = calcular_posicao_ciclo_por_data(
        ciclo="C3",
        data_inicio=date(2026, 1, 1),
        data_fim=date(2026, 12, 31),
        data_posicao=date(2026, 8, 31),
        itens=[{
            "item": "N001",
            "novo_item": True,
            "remanescente_referencia": 999,
            "ciclo_referencia": "C1",
            "data_referencia": date(2024, 2, 29),
            "remanescente_atual": 60,
            "vu_atualizado": 10,
        }],
        movimentos=[
            {"item": "N001", "data_efeito": date(2026, 2, 1), "delta": 100},
        ],
    )
    item = resultado["itens"][0]
    assert item["remanescente_inicio"] == 0
    assert item["alteracoes_liquidas_periodo"] == 100
    assert item["quantidade_consumida"] == 40


def test_ciclo_em_execucao_tem_fallback_por_item_e_rastreabilidade(wb_coleta_116):
    ws = wb_coleta_116["CICLO_EM_EXECUCAO"]
    assert ws["B12"].value == (
        "QTD REMANESCENTE NA ÚLTIMA REFERÊNCIA CONHECIDA (AUTO)"
    )
    assert ws["D12"].value == "QTD CONSUMIDA DESDE A REFERÊNCIA (AUTO)"
    assert ws["P12"].value == "CICLO_REFERENCIA_FISICA"
    assert ws["Q12"].value == "DATA_REFERENCIA_FISICA"
    assert ws.column_dimensions["P"].hidden
    assert ws.column_dimensions["Q"].hidden

    # Recuo (so quando o ciclo em execucao nao resolve): do mais recente ao C0.
    recuo = str(ws["P13"].value).split(",$C$3,", 1)[1]
    assert recuo.index("posicao_contratual!$R2") < recuo.index(
        "posicao_contratual!$N2"
    ) < recuo.index("posicao_contratual!$J2") < recuo.index(
        "posicao_contratual!$F2"
    )
    assert "CHOOSE(MATCH(P13," in str(ws["B13"].value)
    assert 'aditivos!$B$2:$B$200,">="&(INT(Q13)+1)' in str(ws["I13"].value)
    # Coleta 11.8: A9 e a posicao medida qualquer que seja a referencia (P);
    # o VTA nao soma o consumo desde referencia anterior (MEMORIA!T49/T50).
    assert "$P$13:$P$211" not in str(ws["A9"].value)

    dv = next(
        d for d in ws.data_validations.dataValidation
        if str(d.sqref) == "C13:C211"
    )
    assert "NOT(ISNUMBER(B13))" in str(dv.formula1)
    assert "IFERROR(C13<=ROUND(B13+I13,2),FALSE)" in str(dv.formula1)


def test_aditivos_separa_nascimento_fisico_da_base_economica(wb_coleta_116):
    ws = wb_coleta_116["aditivos"]
    assert ws["N1"].value == "ÚLTIMO REAJUSTE JÁ INCORPORADO AO VU"
    assert ws["O1"].value == "INDICE_BASE_ECONOMICA_VU"
    assert ws.column_dimensions["O"].hidden

    validacao = next(
        d for d in ws.data_validations.dataValidation
        if str(d.sqref) == "N2:N200"
    )
    assert validacao.formula1 == '"C0,C1,C2,C3,C4"'

    fator = str(ws["I2"].value)
    assert "$O2" in fator
    assert "posicao_contratual!$Y$2:$Y$200" not in fator
    assert "/INDEX(parametros!$F$2:$F$6,($O2)+1)" in fator
    assert 'UPPER(H2)="SIM"' in str(ws["J2"].value)
    assert "BASE_VU_OBRIGATORIA" in str(ws["M2"].value)
    assert "BASE_VU_NAO_APLICAVEL" in str(ws["M2"].value)
    assert "BASE_VU_INVALIDA_OU_POSTERIOR_AO_FATOR_ALVO" in str(ws["M2"].value)


def test_novos_itens_recebem_fonte_verde_sem_preenchimento(wb_coleta_116):
    ws = wb_coleta_116["itens_Remanesc"]
    regras = [
        regra
        for cf in ws.conditional_formatting
        if str(cf.sqref) == "A2:AC200"
        for regra in cf.rules
    ]
    assert len(regras) == 1
    regra = regras[0]
    assert 'LEN(TRIM($A2))=4' in str(regra.formula[0])
    assert regra.dxf.font.color.rgb == "FF006100"
    assert regra.dxf.fill is None
    prioridades_erros = [
        r.priority
        for cf in ws.conditional_formatting
        if str(cf.sqref) != "A2:AC200"
        for r in cf.rules
    ]
    assert regra.priority > max(prioridades_erros)


def test_geracao_preserva_extensao_condicional_x14_do_template():
    def partes_x14(conteudo: bytes) -> dict[str, int]:
        with zipfile.ZipFile(BytesIO(conteudo)) as pacote:
            mapa = _mapa_xml_planilhas(pacote)
            return {
                aba: pacote.read(nome).count(b"x14:")
                for aba, nome in mapa.items()
                if pacote.read(nome).count(b"x14:")
            }

    origem = partes_x14(TEMPLATE_COLETA_OFICIAL.read_bytes())
    gerado = partes_x14(obter_coleta_oficial_bytes())
    assert origem
    assert gerado == origem


def test_template_oficial_ja_carrega_a_coleta_116():
    """PR3B-1: a estrutura 11.6 mora no template; a geracao nao a conserta."""
    import _coleta_oficial as co

    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    ws = wb["aditivos"]
    assert co._aditivos_base_economica_canonica(ws)
    assert ws["O1"].value == co._CABECALHO_INDICE_BASE_VU
    assert ws.column_dimensions["O"].hidden
    validacao = next(
        d for d in ws.data_validations.dataValidation
        if str(d.sqref) == co._FAIXA_BASE_ECONOMICA_VU
    )
    assert (validacao.type, validacao.formula1) == ("list", '"C0,C1,C2,C3,C4"')
    regras = [
        regra
        for cf in wb["itens_Remanesc"].conditional_formatting
        if str(cf.sqref) == co._FAIXA_DESTAQUE_NOVOS_ITENS
        for regra in cf.rules
    ]
    assert [r.formula[0] for r in regras] == [co._FORMULA_DESTAQUE_NOVOS_ITENS]
