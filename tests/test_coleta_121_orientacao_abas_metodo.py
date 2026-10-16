"""Coleta 12.1 — orientacao visual no topo das abas de metodo.

Somente apresentacao. Dois blocos:

* ESTRUTURA (CI): estado e faixa de texto na linha 1 de financeiro,
  itens_Consumidos e itens_PC, lendo os MESMOS indicadores da coerencia
  metodo x fonte da Coleta 12.0 (MEMORIA_RESULTADOS!B4, T54:T57); cores por
  formatacao condicional que so le a propria aba; cabecalhos, grade, filtros,
  validacoes, congelamento e regras existentes preservados — no template e na
  Coleta gerada.
* EXCEL REAL (RUN_EXCEL_INTEGRATION=1): os quatro estados por calculo real,
  troca de metodo sem regenerar o arquivo, bloqueio/prevalencia da 12.0 e
  resultados financeiros identicos ao template do Pacote 1 (A/B).
"""
from __future__ import annotations

import gc
import io
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_coleta_orientacao_abas_metodo as orient  # noqa: E402
from _coleta_oficial import _CELULAS_COLETA_120, COLUNAS_ITENS_PC_OFICIAL  # noqa: E402

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
# Pacote 1 (Coleta 12.1, card RETRO. RECONHECIDO + POTENCIAL): base do A/B.
COMMIT_BASE_PACOTE_1 = "0fd36d80e79c2fcc30aacb92fc2a762529a3b2da"

_CACHE: dict[str, object] = {}


def _template():
    if "template" not in _CACHE:
        _CACHE["template"] = load_workbook(TEMPLATE)
    return _CACHE["template"]


def _gerado():
    if "gerado" not in _CACHE:
        from _coleta_oficial import obter_coleta_oficial_bytes

        _CACHE["gerado"] = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    return _CACHE["gerado"]


def _livros():
    return (_template(), _gerado())


# =========================================================================== #
# ESTRUTURA
# =========================================================================== #
def test_posicoes_no_topo_de_cada_aba():
    assert {aba: dados[2:] for aba, dados in orient.ABAS.items()} == {
        "financeiro": ("A1:G1", "H1:K1", "L1"),
        "itens_Consumidos": ("A1:AG1", "AH1:AK1", "AL1"),
        "itens_PC": ("A1:L1", "AE1:AH1", "AI1"),
    }
    for wb in _livros():
        for aba, (*_r, faixa, estado) in orient.ABAS.items():
            ws = wb[aba]
            assert faixa in {str(m) for m in ws.merged_cells.ranges}, aba
            assert ws[estado].number_format == ";;;", aba
            assert ws[faixa.split(":")[0]].alignment.wrap_text, aba


def test_descritor_bate_com_template_e_coleta_gerada():
    esperado = orient.formulas()
    orient.validar_ascii_e_parenteses()
    for wb in _livros():
        for (aba, celula), formula in esperado.items():
            assert wb[aba][celula].value == formula, (aba, celula)


def test_estado_reusa_os_indicadores_da_coerencia_120():
    """Metodo = MEMORIA_RESULTADOS!B4; dados = T54/T55/T56; T57 numerico =
    ha metodo. Nenhuma contagem nova de 'fonte preenchida'."""
    contagem = {"financeiro": "T54", "itens_PC": "T55", "itens_Consumidos": "T56"}
    codigo = {"financeiro": "Financeiro", "itens_PC": "PCs", "itens_Consumidos": "Itens"}
    mem = _template()["MEMORIA_RESULTADOS"]
    for celula in ("T54", "T55", "T56", "T57"):
        assert mem[celula].value == _CELULAS_COLETA_120[celula][1]
    for aba, (*_r, estado) in orient.ABAS.items():
        formula = _template()[aba][estado].value
        refs = set(re.findall(r"MEMORIA_RESULTADOS!\$[A-Z]+\$\d+", formula))
        fonte = f"MEMORIA_RESULTADOS!${contagem[aba][0]}${contagem[aba][1:]}"
        assert refs == {"MEMORIA_RESULTADOS!$B$4", "MEMORIA_RESULTADOS!$T$57", fonte}
        assert formula == (
            f'=IF(MEMORIA_RESULTADOS!$B$4="{codigo[aba]}",1,'
            f"IF(ISNUMBER(MEMORIA_RESULTADOS!$T$57),IF({fonte}>0,3,2),4))"
        )


def test_faixa_exibe_o_metodo_real_de_controle_e_os_quatro_textos():
    for aba, (*_r, faixa, estado) in orient.ABAS.items():
        formula = _template()[aba][faixa.split(":")[0]].value
        assert formula.startswith(f"=CHOOSE(${estado[:-1]}${estado[-1]},")
        assert formula.count("CONTROLE!$B$1") == 3      # ativo, outro, outro com dados
        assert formula.count("CHAR(10)") == 4
    mem = _template()["MEMORIA_RESULTADOS"]
    textos = {mem[f"AF{linha}"].value: mem[f"AG{linha}"].value for linha in range(71, 80)}
    assert textos == dict(orient.TEXTOS)
    assert textos["ABA_ATIVO_TIT"] == "MÉTODO ATIVO — "
    assert textos["ABA_OUTRO_TIT"] == "ESTA ABA NÃO CORRESPONDE AO MÉTODO SELECIONADO"
    assert textos["ABA_DADOS_TIT"] == "HÁ DADOS NESTA ABA, MAS ELA NÃO É O MÉTODO ATIVO"
    assert textos["ABA_SEM_TIT"] == "SELECIONE PRIMEIRO O MÉTODO DE APURAÇÃO"
    assert mem["AF80"].value is None and mem["AG80"].value is None


def test_cores_so_leem_a_celula_de_estado_da_propria_aba():
    esperadas = {
        1: "FFE2EFDA",   # verde claro: metodo ativo
        2: "FFFFF2CC",   # amarelo claro: outra aba, sem dados
        3: "FFF8CBAD",   # laranja/pessego: outra aba COM dados
        4: "FFDDEBF7",   # azul claro: metodo nao selecionado
    }
    for wb in _livros():
        for aba, (*_r, cabecalho, faixa, estado) in orient.ABAS.items():
            ws = wb[aba]
            ref = f"${estado[:-1]}${estado[-1]}"
            regras = {
                regra.formula[0]: (regra.dxf.fill.bgColor.rgb, str(bloco.sqref))
                for bloco in ws.conditional_formatting
                for regra in bloco.rules
                if ref in regra.formula[0]
            }
            assert {f: cor for f, (cor, _s) in regras.items()} == {
                f"{ref}={c}": cor for c, cor in esperadas.items()}, aba
            for _cor, sqref in regras.values():
                # Contiguos o Excel funde (A1:G1 + H1:K1 = A1:K1).
                assert sqref in (f"{cabecalho} {faixa}",
                                 f"{cabecalho.split(':')[0]}:{faixa.split(':')[1]}"), aba
            for bloco in ws.conditional_formatting:
                for regra in bloco.rules:
                    assert "!" not in "".join(regra.formula), (aba, regra.formula)


def test_cabecalhos_e_grade_preservados():
    for wb in _livros():
        fin, ic, pc = wb["financeiro"], wb["itens_Consumidos"], wb["itens_PC"]
        assert [fin.cell(1, c).value for c in range(1, 8)] == [
            "COMPETENCIA", "CICLO", "VALOR_PAGO", "FATOR_APLICAVEL",
            "VALOR_ATUALIZADO", "DELTA", "EFEITO_FINANCEIRO"]
        assert [pc.cell(1, c).value for c in range(1, 8)] == COLUNAS_ITENS_PC_OFICIAL[:7]
        assert pc["L1"].value == "EFEITO_FINANCEIRO_PC"
        assert pc["M1"].value == "TODOS OS PCs CADASTRADOS POR CICLO"
        assert ic["A1"].value == "ITEM" and ic["AG1"].value == "AJUSTE_STATUS"
        assert ic["W1"].value is None
        # Dados continuam a partir da linha 2 (nenhuma linha inserida).
        assert fin["B74"].value == "TOTAL" and fin["C74"].value == "=ROUND(SUM(C2:C73),2)"
        assert pc["AD2"].value.startswith("=IFERROR(IF(UPPER(TRIM(SUBSTITUTE(G2,")
        assert ic["D2"].value == '=IF(OR(A2="",B2="",C2=""),"",ROUND(B2*C2,2))'
        assert ic.freeze_panes == "D2"
        assert pc.auto_filter.ref == "A1:L5001"
        assert {str(d.sqref) for d in fin.data_validations.dataValidation} == {"G2:G73"}
        assert {str(d.sqref) for d in pc.data_validations.dataValidation} == {"G2:G5001"}
        assert {str(d.sqref) for d in ic.data_validations.dataValidation} == {"Z2:Z6"}
        assert {str(m) for m in pc.merged_cells.ranges} >= {"M1:T1", "M31:T31", "M32:T32"}
        assert {str(m) for m in ic.merged_cells.ranges} >= {"X8:AG8", "X13:AG13"}


def test_regras_condicionais_existentes_preservadas():
    for wb in _livros():
        regras = {
            aba: sorted(
                (str(b.sqref), r.formula[0]) for b in wb[aba].conditional_formatting
                for r in b.rules
            )
            for aba in orient.ABAS
        }
        assert ("A2:G73", 'AND($C2<>"",$G2="")') in regras["financeiro"]
        assert ("A2:G73", 'AND($A2<>"",$G2="Nao")') in regras["financeiro"]
        assert ("G2:G5001", '$AD2="INVALIDO"') in regras["itens_PC"]
        assert ("A2:L5001", '$L2="Nao"') in regras["itens_PC"]
        assert ("S3:S9", "N($S3)<>0") in regras["itens_PC"]
        assert ("S12:S18", "N($S12)<>0") in regras["itens_PC"]
        assert len(regras["financeiro"]) == 2 + 4
        assert len(regras["itens_PC"]) == 6 + 4
        assert len(regras["itens_Consumidos"]) == 0 + 4


def test_coleta_gerada_mantem_protecao_e_colunas_tecnicas():
    pc = _gerado()["itens_PC"]
    assert pc.protection.sheet and pc.protection.autoFilter is False
    assert pc["A2"].protection.locked is False and pc["G2"].protection.locked is False
    assert pc["AE1"].protection.locked and pc["AI1"].protection.locked
    assert pc.column_dimensions["V"].hidden and pc.column_dimensions["AD"].hidden
    assert not pc.column_dimensions["AE"].hidden


def test_geometria_cabe_no_gate_de_upload():
    from _seguranca_xlsx import MAX_AREA_TOTAL_WORKBOOK, MAX_COLUNAS_POR_ABA

    wb = _gerado()
    assert wb["itens_PC"].max_column == 35
    assert wb["itens_Consumidos"].max_column == 38
    assert wb["financeiro"].max_column <= 13
    assert max(ws.max_column for ws in wb.worksheets) <= MAX_COLUNAS_POR_ABA
    assert sum(ws.max_row * ws.max_column for ws in wb.worksheets) < MAX_AREA_TOTAL_WORKBOOK


def test_leitor_mapeia_as_colunas_oficiais_como_antes():
    from _leitor_masterfile_v10 import _mapear_colunas_por_cabecalho

    wb = _gerado()
    fin = _mapear_colunas_por_cabecalho(wb["financeiro"])
    assert (fin["competencia"], fin["valor_pago"], fin["efeito_financeiro"]) == (1, 3, 7)
    pc = _mapear_colunas_por_cabecalho(wb["itens_PC"])
    assert (pc["numero_pc"], pc["valor_pc"], pc["pc_pago_a_contratada"]) == (1, 4, 7)
    ic = _mapear_colunas_por_cabecalho(wb["itens_Consumidos"])
    assert (ic["item"], ic["qtd_cons_c0"], ic["ajuste_status"]) == (1, 5, 33)
    assert "ciclo_pc" not in ic and "numero_pc" not in ic


# =========================================================================== #
# EXCEL REAL
# =========================================================================== #
excel_real = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar o Excel COM",
)

SEL = "SELECIONE AQUI"
FIN = "Financeiro (Mensalidade)"
PC = "PC (Pedidos de Compra)"
IT = "Itens Consumidos"
TITULOS = {
    1: "MÉTODO ATIVO — ",
    2: "ESTA ABA NÃO CORRESPONDE AO MÉTODO SELECIONADO\n",
    3: "HÁ DADOS NESTA ABA, MAS ELA NÃO É O MÉTODO ATIVO\n",
    4: "SELECIONE PRIMEIRO O MÉTODO DE APURAÇÃO\n",
}


def _cor(rgb: str) -> int:
    return int(rgb[4:6], 16) * 65536 + int(rgb[2:4], 16) * 256 + int(rgb[0:2], 16)


def _preencher(livro, metodo: str, *, pcs=False, fin=False, itens=False) -> None:
    """Caso de uso 11.8/12.0 (itens_Remanesc 1.1, corte 31/08/2025)."""
    from pywintypes import Time

    from test_coleta_118_excel_real import CORTE, PCS

    livro.Worksheets("CONTROLE").Range("B1").Value = metodo
    livro.Worksheets("CONTROLE").Range("B3").Value = Time(CORTE)
    rem = livro.Worksheets("itens_Remanesc")
    for celula, valor in (("A2", "1.1"), ("B2", 1000), ("C2", 10), ("E2", 950)):
        rem.Range(celula).Value = valor
    if pcs:
        aba = livro.Worksheets("itens_PC")
        for linha, (numero, data_pc, valor, pago) in enumerate(PCS, start=2):
            aba.Range(f"A{linha}").Value = numero
            aba.Range(f"B{linha}").Value = Time(data_pc)
            aba.Range(f"D{linha}").Value = valor
            aba.Range(f"G{linha}").Value = pago
    if fin:
        aba = livro.Worksheets("financeiro")
        for linha in range(2, 74):
            competencia = aba.Range(f"A{linha}").Value
            if hasattr(competencia, "year") and competencia.replace(tzinfo=None) <= CORTE:
                aba.Range(f"C{linha}").Value = 50
    if itens:
        aba = livro.Worksheets("itens_Consumidos")
        for celula, valor in (("A2", "1.1"), ("B2", 1000), ("C2", 10), ("E2", 30),
                              ("G2", 40), ("I2", 20)):
            aba.Range(celula).Value = valor
    ciclo = livro.Worksheets("CICLO_EM_EXECUCAO")
    ciclo.Range("D5").Value2 = livro.Worksheets("CONTROLE").Range("B3").Value2
    livro.Application.CalculateFullRebuild()
    ciclo.Range("C13").Value = 800


def _ler(livro) -> dict:
    mem = livro.Worksheets("MEMORIA_RESULTADOS")
    det = livro.Worksheets("RESULTADOS_DETALHE")
    exe = livro.Worksheets("RESULTADOS")
    r = {"abas": {}}
    for aba, (*_x, faixa, estado) in orient.ABAS.items():
        ws = livro.Worksheets(aba)
        faixa = faixa.split(":")[0]           # celula ancora da mescla
        r["abas"][aba] = {
            "estado": ws.Range(estado).Value2,
            "visivel": ws.Range(estado).Text,
            "texto": ws.Range(faixa).Value2,
            "cor_cabecalho": int(ws.Range("A1").DisplayFormat.Interior.Color),
            "cor_faixa": int(ws.Range(faixa).DisplayFormat.Interior.Color),
        }
    r.update({c: mem.Range(c).Value2 for c in ("B4", "B26", "T40", "T57", "T58", "T66")})
    r.update({c: det.Range(c).Value2 for c in ("D22", "J5", "A7", "B3", "B7")})
    r.update({f"exe_{c}": exe.Range(c).Value2 for c in ("B9", "C9", "D9", "E9", "F9", "C12")})
    r["C1"] = livro.Worksheets("CONTROLE").Range("C1").Value2
    r["valores"] = {
        aba: livro.Worksheets(aba).Range("A1:AG290").Value2
        for aba in ("MEMORIA_RESULTADOS", "RESULTADOS_DETALHE", "RESULTADOS")
    }
    return r


def _rodar(tmp_path: Path, nome: str, passos, *, template: Path | None = None,
           monkeypatch=None) -> list[dict]:
    """passos: [(metodo, opcoes de preenchimento ou None = so trocar B1)]."""
    import pythoncom
    import win32com.client

    import _coleta_oficial
    from test_coleta_118_excel_real import _dados, _tentar

    if template is not None:
        monkeypatch.setattr(_coleta_oficial, "TEMPLATE_COLETA_OFICIAL", template)
    caminho = tmp_path / f"Coleta_orientacao_{nome}.xlsx"
    try:
        caminho.write_bytes(_coleta_oficial.gerar_coleta_oficial_preenchida(_dados()))
    finally:
        if template is not None:
            monkeypatch.undo()
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    livro = None
    leituras = []
    try:
        livro = _tentar(pythoncom, lambda: excel.Workbooks.Open(
            str(caminho.resolve()), UpdateLinks=0, ReadOnly=False, CorruptLoad=0))
        for metodo, opcoes in passos:
            if opcoes is None:
                _tentar(pythoncom, lambda: setattr(
                    livro.Worksheets("CONTROLE").Range("B1"), "Value", metodo))
            else:
                _tentar(pythoncom, lambda: _preencher(livro, metodo, **opcoes))
            _tentar(pythoncom, excel.CalculateFullRebuild)
            leituras.append(_tentar(pythoncom, lambda: _ler(livro)))
        _tentar(pythoncom, livro.Save)
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = None
    finally:
        if livro is not None:
            try:
                livro.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()
    return leituras


def _confere(leitura: dict, esperado: dict[str, int], metodo: str) -> None:
    for aba, estado in esperado.items():
        a = leitura["abas"][aba]
        assert a["estado"] == estado, (aba, a)
        assert a["visivel"] == "", aba                      # celula de estado invisivel
        assert a["texto"].startswith(TITULOS[estado]), (aba, a["texto"])
        if estado in (1, 2, 3):
            assert metodo in a["texto"], (aba, a["texto"])   # nome real de CONTROLE!B1
        fundo = orient.CORES[estado][0]
        assert a["cor_cabecalho"] == a["cor_faixa"] == _cor(fundo), aba


@excel_real
def test_A_nenhum_metodo_selecionado(tmp_path):
    vazio = _rodar(tmp_path, "sem_metodo", [(SEL, {})])[0]
    com_dados = _rodar(tmp_path, "sem_metodo_dados",
                       [(SEL, dict(pcs=True, fin=True, itens=True))])[0]
    for leitura in (vazio, com_dados):
        _confere(leitura, {"financeiro": 4, "itens_Consumidos": 4, "itens_PC": 4}, SEL)
    assert vazio["abas"]["financeiro"]["texto"] == (
        "SELECIONE PRIMEIRO O MÉTODO DE APURAÇÃO\nSelecione o método na aba "
        "CONTROLE. Depois, preencha a aba correspondente.")


@excel_real
def test_B_C_D_metodo_ativo_verde_e_demais_amarelas(tmp_path):
    casos = {
        FIN: (dict(fin=True), "financeiro"),
        IT: (dict(itens=True), "itens_Consumidos"),
        PC: (dict(pcs=True), "itens_PC"),
    }
    for metodo, (opcoes, ativa) in casos.items():
        leitura = _rodar(tmp_path, ativa, [(metodo, opcoes)])[0]
        esperado = {aba: (1 if aba == ativa else 2) for aba in orient.ABAS}
        _confere(leitura, esperado, metodo)
        assert leitura["abas"][ativa]["texto"] == (
            f"MÉTODO ATIVO — {metodo}\nPreencha os dados desta aba para a apuração.")
        outra = next(aba for aba in orient.ABAS if aba != ativa)
        assert leitura["abas"][outra]["texto"] == (
            "ESTA ABA NÃO CORRESPONDE AO MÉTODO SELECIONADO\nMétodo atual: "
            f"{metodo}. Os dados desta aba não serão utilizados como fonte "
            "principal desta apuração.")


@excel_real
def test_E_aba_nao_ativa_com_dados_fica_laranja(tmp_path):
    leitura = _rodar(tmp_path, "it_com_tudo", [(IT, dict(fin=True, pcs=True, itens=True))])[0]
    _confere(leitura, {"itens_Consumidos": 1, "financeiro": 3, "itens_PC": 3}, IT)
    assert leitura["abas"]["itens_PC"]["texto"] == (
        "HÁ DADOS NESTA ABA, MAS ELA NÃO É O MÉTODO ATIVO\nMétodo atual: "
        f"{IT}. Confirme em CONTROLE se a escolha do método está correta.")
    assert orient.CORES[3][0] != orient.CORES[2][0]


@excel_real
def test_F_troca_de_metodo_atualiza_sem_regenerar(tmp_path):
    fin_e_pcs = dict(fin=True, pcs=True)
    leituras = _rodar(tmp_path, "troca",
                      [(FIN, fin_e_pcs), (PC, None), (IT, None), (SEL, None)])
    _confere(leituras[0], {"financeiro": 1, "itens_PC": 3, "itens_Consumidos": 2}, FIN)
    _confere(leituras[1], {"itens_PC": 1, "financeiro": 3, "itens_Consumidos": 2}, PC)
    _confere(leituras[2], {"itens_Consumidos": 1, "financeiro": 3, "itens_PC": 3}, IT)
    _confere(leituras[3], {"financeiro": 4, "itens_PC": 4, "itens_Consumidos": 4}, SEL)


@excel_real
def test_G_H_bloqueio_e_prevalencia_da_120_inalterados(tmp_path):
    # G: metodo PCs sem PCs (com financeiro) -> bloqueio da 12.0.
    bloqueio = _rodar(tmp_path, "bloqueio", [(PC, dict(fin=True))])[0]
    assert bloqueio["T58"] == 1
    assert bloqueio["D22"] in (None, "") and bloqueio["B26"] in (None, "")
    assert bloqueio["C1"].startswith("REVISE O MÉTODO: aba itens_PC")
    _confere(bloqueio, {"itens_PC": 1, "financeiro": 3, "itens_Consumidos": 2}, PC)
    # H: duas fontes preenchidas -> aviso da 12.0; o metodo selecionado prevalece.
    duas = _rodar(tmp_path, "duas_fontes", [(PC, dict(fin=True, pcs=True))])[0]
    assert duas["T58"] == 2 and isinstance(duas["D22"], float)
    assert duas["C1"].startswith("CONFERIR: também há dados em financeiro")
    _confere(duas, {"itens_PC": 1, "financeiro": 3, "itens_Consumidos": 2}, PC)


@excel_real
def test_I_resultados_financeiros_identicos_ao_pacote_1(tmp_path, monkeypatch):
    """A/B: mesmos dados no template do Pacote 1 e no atual."""
    base = tmp_path / "template_pacote_1.xlsx"
    base.write_bytes(subprocess.run(
        ["git", "show", f"{COMMIT_BASE_PACOTE_1}:templates/COLETA_REAJUSTE_OFICIAL.xlsx"],
        cwd=RAIZ, capture_output=True, check=True).stdout)
    casos = {
        "pcs": [(PC, dict(pcs=True))],
        "fin": [(FIN, dict(fin=True))],
        "itens": [(IT, dict(itens=True))],
        "bloqueio": [(PC, dict(fin=True))],
        "duas": [(FIN, dict(fin=True, pcs=True))],
    }
    novas = {(linha, col) for linha in range(71, 80) for col in (32, 33)}   # AF:AG
    for nome, passos in casos.items():
        antes = _rodar(tmp_path, f"base_{nome}", passos, template=base,
                       monkeypatch=monkeypatch)[0]
        depois = _rodar(tmp_path, f"novo_{nome}", passos)[0]
        for chave in ("B26", "T40", "T58", "T66", "D22", "J5", "A7", "B3", "B7", "C1"):
            assert depois[chave] == antes[chave], (nome, chave)
        for chave in ("exe_B9", "exe_C9", "exe_D9", "exe_E9", "exe_F9", "exe_C12"):
            assert depois[chave] == antes[chave], (nome, chave)
        for i, linha in enumerate(antes["valores"]["MEMORIA_RESULTADOS"], start=1):
            for j, valor in enumerate(linha, start=1):
                if (i, j) not in novas:
                    assert depois["valores"]["MEMORIA_RESULTADOS"][i - 1][j - 1] == valor, (
                        nome, i, j)
        assert depois["valores"]["RESULTADOS_DETALHE"] == antes["valores"]["RESULTADOS_DETALHE"]
        assert depois["valores"]["RESULTADOS"] == antes["valores"]["RESULTADOS"]


@excel_real
def test_K_abre_recalcula_salva_e_reabre_sem_reparo(tmp_path):
    from _coleta_oficial import obter_coleta_oficial_bytes
    from validar_visual_resultados_v3 import validar

    arquivo = tmp_path / "coleta_branco_orientacao.xlsx"
    arquivo.write_bytes(obter_coleta_oficial_bytes())
    assert validar(arquivo, tmp_path / "saida", estresse=True) == []
