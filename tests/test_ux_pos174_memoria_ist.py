"""AJUSTES-XLS-UX pos-PR #174 — Frente 4 — memoria IST (parametros!J:R): destaque da competencia de
fronteira entre ciclos + legenda fora do bloco reservado. Valores identicos.

Testes de Excel real rodam com RUN_EXCEL_INTEGRATION=1.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_ajustes_xls_ux_pos174 as ux  # noqa: E402
from _coleta_oficial import (  # noqa: E402
    gerar_coleta_oficial_preenchida,
    obter_coleta_oficial_bytes,
)
from _memoria_calculo import (  # noqa: E402
    FILL_FRONTEIRA_IST,
    COLUNA_EXPLICACAO_FRONTEIRA,
    EXPLICACAO_FRONTEIRA_IST,
    LARGURA_EXPLICACAO_FRONTEIRA,
    escrever_memoria_calculo,
    ler_memoria_calculo,
)

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
COLUNAS_JR = ("J", "K", "L", "M", "N", "O", "P", "Q", "R")

com = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1",
)


@pytest.fixture(scope="module")
def template():
    wb = load_workbook(TEMPLATE)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def coleta_branca():
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    yield wb
    wb.close()


def _rgb(cell) -> str | None:
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


# ------------------------------------------------------------------ payloads
PERCENTUAIS = {1: 0.0512, 2: 0.0374, 3: 0.0289, 4: 0.0315}


def _payload(numeros: list[int]) -> dict:
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2021, 1, 1),
        "data_corte": date(2021 + max(numeros), 12, 31),
        "ciclos": [
            {
                "ciclo": f"C{n}",
                "data_inicio": date(2021 + n, 1, 1),
                "percentual": PERCENTUAIS[n],
                "inicio_efeito_financeiro": date(2021 + n, 1, 1),
                "situacao_aplicada": "✅ TEMPESTIVO",
                "data_pedido": date(2021 + n, 1, 10),
            }
            for n in numeros
        ],
    }


# =================================================================== FRENTE 4
def _ist(base_comp: str, base: float, final_comp: str, final: float):
    return [
        {"tipo": "INDICE", "ordem": 1, "competencia": base_comp, "valor_indice": base},
        {"tipo": "INDICE", "ordem": 2, "competencia": final_comp, "valor_indice": final},
        {"tipo": "RESULTADO", "ordem": 3, "fator_acumulado": final / base,
         "variacao_final": round(final / base - 1, 4), "metodo_fonte": "IST"},
    ]


def _ws_memoria():
    wb = Workbook()
    ws = wb.active
    for col, cab in zip(COLUNAS_JR, ("CICLO", "TIPO_REGISTRO", "ORDEM", "COMPETENCIA",
                                     "VALOR_INDICE", "FATOR_MENSAL", "FATOR_ACUMULADO",
                                     "VARIACAO_FINAL", "METODO_FONTE")):
        ws[f"{col}1"] = cab
    return ws


def _conteudo_jr(ws):
    return [
        tuple(ws[f"{col}{linha}"].value for col in COLUNAS_JR)
        for linha in range(2, 81)
    ]


def _destacadas(ws):
    return [
        linha for linha in range(2, 81)
        if _rgb(ws[f"J{linha}"]) == FILL_FRONTEIRA_IST
    ]


def _explicacoes(ws):
    return {
        linha: ws[f"{COLUNA_EXPLICACAO_FRONTEIRA}{linha}"].value
        for linha in range(2, 81)
        if ws[f"{COLUNA_EXPLICACAO_FRONTEIRA}{linha}"].value is not None
    }


CICLOS_FRONTEIRA = {
    "C3": {"memoria_calculo": _ist("2022-10-01", 100.0, "2023-10-01", 104.0)},
    "C4": {"memoria_calculo": _ist("2023-10-01", 104.0, "2024-10-01", 107.0)},
}


def test_f4_fronteira_destaca_somente_base_do_segundo_ciclo():
    ws = _ws_memoria()
    escrever_memoria_calculo(ws, CICLOS_FRONTEIRA)
    assert ws["J5"].value == "C4" and ws["M5"].value.month == 10
    assert _destacadas(ws) == [5]
    assert all(_rgb(ws[f"{col}5"]) == FILL_FRONTEIRA_IST for col in COLUNAS_JR)
    assert _explicacoes(ws) == {5: EXPLICACAO_FRONTEIRA_IST}
    assert _rgb(ws[f"{COLUNA_EXPLICACAO_FRONTEIRA}5"]) == FILL_FRONTEIRA_IST
    assert ws.column_dimensions[COLUNA_EXPLICACAO_FRONTEIRA].width >= LARGURA_EXPLICACAO_FRONTEIRA
    assert ws["K82"].value is None  # sem legenda distante da linha laranja


def test_f4_conteudo_jr_identico_com_e_sem_destaque(monkeypatch):
    import _memoria_calculo as mc

    com_destaque = _ws_memoria()
    escrever_memoria_calculo(com_destaque, CICLOS_FRONTEIRA)
    monkeypatch.setattr(mc, "_linhas_fronteira_ist", lambda planos: set())
    sem_destaque = _ws_memoria()
    escrever_memoria_calculo(sem_destaque, CICLOS_FRONTEIRA)
    assert _conteudo_jr(com_destaque) == _conteudo_jr(sem_destaque)
    assert ler_memoria_calculo(com_destaque) == ler_memoria_calculo(sem_destaque)
    assert [r["tipo"] for r in ler_memoria_calculo(com_destaque)["C4"]] == [
        "INDICE", "INDICE", "RESULTADO"]


def test_f4_ciclo_unico_sem_destaque():
    ws = _ws_memoria()
    escrever_memoria_calculo(ws, {"C2": {"memoria_calculo": _ist(
        "2022-10-01", 100.0, "2023-10-01", 104.0)}})
    assert _destacadas(ws) == []
    assert _explicacoes(ws) == {}


def test_f4_competencias_diferentes_sem_destaque():
    ws = _ws_memoria()
    escrever_memoria_calculo(ws, {
        "C1": {"memoria_calculo": _ist("2021-10-01", 100.0, "2022-10-01", 103.0)},
        "C2": {"memoria_calculo": _ist("2022-11-01", 103.5, "2023-11-01", 106.0)},
    })
    assert _destacadas(ws) == []
    assert _explicacoes(ws) == {}


def test_f4_so_a_linha_que_abre_o_ciclo_e_destacada():
    ws = _ws_memoria()
    escrever_memoria_calculo(ws, {
        "C1": {"memoria_calculo": _ist("2021-10-01", 100.0, "2022-10-01", 103.0)},
        "C2": {"memoria_calculo": _ist("2021-10-01", 100.0, "2022-10-01", 103.0)},
    })
    assert _destacadas(ws) == [5]


def test_f4_coleta_gerada_destaca_fronteira():
    payload = _payload([1, 2])
    payload["ciclos"][0]["memoria_calculo"] = _ist("2021-10-01", 101.0, "2022-10-01", 102.0)
    payload["ciclos"][1]["memoria_calculo"] = _ist("2022-10-01", 102.0, "2023-10-01", 103.0)
    ws = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(payload)))["parametros"]
    assert _destacadas(ws) == [5]
    assert _explicacoes(ws) == {5: EXPLICACAO_FRONTEIRA_IST}



def _mes(competencias: list[str], taxa: float):
    registros = [
        {"tipo": "MES", "ordem": i, "competencia": c, "valor_indice": taxa}
        for i, c in enumerate(competencias, start=1)
    ]
    registros.append({"tipo": "RESULTADO", "ordem": len(registros) + 1,
                      "fator_acumulado": 1.01, "variacao_final": 0.01,
                      "metodo_fonte": "ICTI"})
    return registros


def test_f4_memoria_mes_nunca_recebe_explicacao_do_ist():
    """ICTI/SGS gravam MES (produtorio de taxas): mesmo com competencia
    repetida entre ciclos, nao ha fronteira de numero-indice."""
    ws = _ws_memoria()
    escrever_memoria_calculo(ws, {
        "C1": {"memoria_calculo": _mes(["2022-09-01", "2022-10-01"], 0.004)},
        "C2": {"memoria_calculo": _mes(["2022-10-01", "2022-11-01"], 0.005)},
    })
    assert _destacadas(ws) == []
    assert _explicacoes(ws) == {}


def test_f4_explicacao_nao_entra_na_leitura_da_memoria():
    """R (METODO_FONTE) segue vazia nas linhas INDICE; o leitor nao ve S."""
    ws = _ws_memoria()
    escrever_memoria_calculo(ws, CICLOS_FRONTEIRA)
    lido = ler_memoria_calculo(ws)
    assert all(r["metodo_fonte"] is None for r in lido["C4"] if r["tipo"] == "INDICE")
    assert EXPLICACAO_FRONTEIRA_IST not in str(lido)
