"""Coleta 12.4 — precisao financeira canonica + SALDO REMANESCENTE — EVOLUCAO.

Precisao: so agregacoes monetarias ganham ROUND(...,2); percentual, fator,
quantidades e contagens seguem intocados. O web normaliza o zero negativo na
exibicao (nunca "-R$ 0,00"). Saldo: a nova tabela da RESULTADOS so espelha
fontes canonicas e mostra "—" quando a posicao nao existe ou esta incompleta.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_coleta_124_precisao_saldo as c124  # noqa: E402

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
_CACHE: dict[str, object] = {}


def _template():
    if "template" not in _CACHE:
        _CACHE["template"] = load_workbook(TEMPLATE)
    return _CACHE["template"]


# ------------------------------------------------------------------ precisao
def test_agregacoes_monetarias_fecham_em_duas_casas():
    wb = _template()
    for aba, celulas in c124.ROUND_CELULAS_INTEIRAS.items():
        for celula in celulas:
            formula = wb[aba][celula].value
            assert formula.startswith("=ROUND(") and formula.endswith(",2)"), (aba, celula)
    for (aba, celula), trechos in c124.ROUND_TRECHOS.items():
        formula = wb[aba][celula].value
        for _, troca in trechos:
            assert formula.count(troca) == 1, (aba, celula)
    consumidos = wb["itens_Consumidos"]
    for linha in c124.LINHAS_CONSUMIDOS:
        assert consumidos[f"P{linha}"].value == c124.formula_consumidos_p(linha)[1]


def test_percentual_fator_quantidades_e_contagens_nao_sao_arredondados():
    wb = _template()
    par = wb["parametros"]
    for linha in range(3, 7):
        assert "ROUND" not in par[f"F{linha}"].value          # fator acumulado
    assert par["D12"].value == '=IF(C12="","",1+C12)'          # fator do ciclo
    pc = wb["itens_PC"]
    for linha in range(3, 10):
        assert "ROUND" not in pc[f"N{linha}"].value            # qtd. de PCs
        assert "ROUND" not in pc[f"T{linha}"].value            # qtd. com alerta
    assert wb["itens_Consumidos"]["O2"].value == '=IF(A2="","",SUM(E2,G2,I2,K2,M2))'


def test_coleta_gerada_nao_reembrulha_b85():
    from _coleta_oficial import obter_coleta_oficial_bytes

    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    b85 = wb["RESULTADOS_DETALHE"]["B85"].value
    assert b85.count('ROUND(SUM(INDIRECT("CICLO_EM_EXECUCAO!G13:G211")),2)') == 1
    assert "$W$49=1,SUM(INDIRECT(" not in b85
    assert b85.count('=IF($B$5="PCs",') == 1


@pytest.mark.parametrize("modulo, funcao", [
    ("_reajuste_utils", "_formatar_moeda_br"),
    ("_resultado_consolidado", "_moeda_br"),
    ("_painel_executivo", "_brl"),
])
def test_web_nunca_exibe_zero_negativo(modulo, funcao):
    formatar = getattr(__import__(modulo), funcao)
    assert formatar(-0.000000022) == "R$ 0,00"
    assert formatar(-0.004) == "R$ 0,00"
    assert formatar(-0.01) == "R$ -0,01"
    assert formatar(38322075.53000001) == "R$ 38.322.075,53"


# ------------------------------------------------------------------- saldo
def test_tabela_saldo_espelha_fontes_canonicas():
    ws = _template()["RESULTADOS"]
    for celula, texto in c124.rotulos_saldo().items():
        assert ws[celula].value == texto
    for celula, formula in c124.formulas_saldo().items():
        assert ws[celula].value == formula
    assert ws["B76"].value == "SALDO REMANESCENTE — EVOLUÇÃO"
    assert [ws[f"B{linha}"].value for linha in range(78, 84)] == [
        "Abertura C0", "Abertura C1", "Abertura C2", "Abertura C3", "Abertura C4",
        "Última posição informada",
    ]
    assert ws.print_area == "'RESULTADOS'!$B$1:$G$83"
    c124.validar_ascii_e_parenteses()


def test_ausencia_de_posicao_nunca_vira_zero():
    formulas = c124.formulas_saldo()
    for linha in range(78, 84):
        assert formulas[f"D{linha}"].endswith(f"{c124.TRACO})")   # fallback e "—"
        assert ",0)" not in formulas[f"D{linha}"]
    # abertura so vale com a fotografia completa (W41:W45), nunca a soma parcial
    for n in range(5):
        assert f"MEMORIA_RESULTADOS!$W${41 + n}=1" in formulas[f"D{78 + n}"]
    mem = _template()["MEMORIA_RESULTADOS"]
    # ultima posicao: a escolha fica na camada de apoio; a executiva so espelha
    for linha, (rotulo, formula) in c124.APOIO_ULTIMA.items():
        assert mem[f"V{linha}"].value == rotulo
        assert mem[f"W{linha}"].value == formula
    assert '$W$49=1' in mem["W69"].value and "$W$46+1" in mem["W69"].value
    assert formulas["D83"] == "=IF(ISNUMBER(MEMORIA_RESULTADOS!$W$69),MEMORIA_RESULTADOS!$W$69,MEMORIA_RESULTADOS!$AG$2)"
    for formula in formulas.values():
        assert "INDIRECT(" not in formula and "+" not in formula
    assert mem["AG2"].value == "—"
    assert [mem[f"AG{linha}"].value for linha in (81, 82, 83)] == [
        "Posição de abertura", "Não informado", "Posição atual"]
    assert mem["AF80"].value is None and mem["AG80"].value is None


def test_tabela_saldo_usa_o_estilo_existente_da_resultados():
    ws = _template()["RESULTADOS"]

    def assinatura(celula):
        c = ws[celula]
        return (c.fill.fgColor.rgb, c.font.sz, c.font.b,
                c.font.color.rgb if c.font.color else None)

    assert assinatura("B76") == assinatura("B44")    # titulo azul-escuro
    assert assinatura("B77") == assinatura("B45")    # cabecalho azul-claro
    assert assinatura("D78") == assinatura("C46")    # moeda
    assert assinatura("B83") == assinatura("B52")    # linha final no padrao TOTAL
    assert ws["D79"].number_format == ws["C46"].number_format
    assert ws["C79"].number_format == ws["C38"].number_format      # dd/mm/aaaa
    assert ws["C83"].number_format == ws["C38"].number_format


# ------------------------------------------------- Excel real (opt-in)
def _dados_c1_c4() -> dict:
    from _memoria_calculo import normalizar_memoria_calculo
    from _reajuste_utils import fator_acumulado_oficial, resolver_tratamento_variacao_negativa

    brutos = (0.03813785, 0.04496, 0.021349, 0.050371)
    ciclos, fator = [], 1.0
    for i, bruto in enumerate(brutos, start=1):
        ano = 2020 + i
        r = resolver_tratamento_variacao_negativa(bruto)
        memoria = {"indice": "IST", "variacao": bruto, "meses": [],
                   "indice_inicial": 100.0, "indice_final": 100 * (1 + bruto)}
        fator *= r["fator"]
        ciclos.append({
            "ciclo": f"C{i}", "data_base": f"01/08/{ano}", "periodo_inicio": f"01/08/{ano}",
            "periodo_fim": f"31/07/{ano + 1}", "financeiro_inicio": f"01/08/{ano + 1}",
            "percentual_indice": bruto, "percentual_aplicado": r["percentual_aplicado"],
            "variacao": r["percentual_aplicado"], "fator": r["fator"], "fator_acumulado": fator,
            "situacao": "TEMPESTIVO", "situacao_aplicada": "TEMPESTIVO",
            "memoria_calculo": normalizar_memoria_calculo(
                memoria, r["fator"], r["percentual_aplicado"]),
        })
    acumulado = fator_acumulado_oficial(brutos)
    return {"origem": "Reajuste Represado", "indice": "IST", "data_base_original": "01/08/2021",
            "fator": acumulado, "fator_acumulado": acumulado, "variacao": acumulado - 1,
            "variacao_acumulada": acumulado - 1, "ciclos": ciclos}


def _com(funcao, *args, **kwargs):
    """Repete a chamada enquanto o Excel recusa por estar ocupado."""
    import time

    for _ in range(120):
        try:
            return funcao(*args, **kwargs)
        except Exception as exc:  # pragma: no cover - somente Excel COM
            if getattr(exc, "hresult", None) != -2147418111:
                raise
            time.sleep(0.5)
    return funcao(*args, **kwargs)


@pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar Excel COM",
)
def test_excel_real_sem_residuo_e_ausencia_como_traco(tmp_path):
    client = pytest.importorskip("win32com.client")
    pythoncom = pytest.importorskip("pythoncom")
    from _coleta_oficial import gerar_coleta_oficial_preenchida

    wb = load_workbook(io.BytesIO(gerar_coleta_oficial_preenchida(_dados_c1_c4())))
    rem = wb["itens_Remanesc"]
    for i in range(30):
        linha = 2 + i
        rem[f"A{linha}"], rem[f"B{linha}"], rem[f"C{linha}"] = f"{i + 1}.1", 900, 17.31 + i
        for col, qtd in zip("EGIK", (800, 700, 600, 500)):
            rem[f"{col}{linha}"] = qtd - i
    rem["I4"] = None                                   # abertura C3 incompleta
    pc = wb["itens_PC"]
    for i in range(400):
        pc[f"A{2 + i}"] = f"PC-{i + 1:04d}"
        pc[f"B{2 + i}"] = datetime(2021 + i // 100, 9 + (i % 4), 1 + i % 27)
        pc[f"D{2 + i}"] = round(100.1 + i * 617.37, 2)
        pc[f"G{2 + i}"] = "Sim" if i % 3 else "Nao"
    wb["CONTROLE"]["B1"] = "PCs"
    caminho = tmp_path / "c124.xlsx"
    wb.save(caminho)

    pythoncom.CoInitialize()
    app = client.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    try:
        book = _com(app.Workbooks.Open, str(caminho), UpdateLinks=0, ReadOnly=True,
                    CorruptLoad=0)
        try:
            _com(app.CalculateFullRebuild)

            def v(aba, celula):   # Value2: moeda volta como float (Value daria Decimal)
                return _com(lambda: book.Worksheets(aba).Range(celula).Value2)

            for aba, celulas in c124.ROUND_CELULAS_INTEIRAS.items():
                for celula in celulas:
                    valor = v(aba, celula)
                    if isinstance(valor, float):
                        assert valor == round(valor, 2), (aba, celula, valor)
                        assert not (valor < 0 and round(valor, 2) == 0), (aba, celula)
            assert v("parametros", "E3") == 0.0381 and v("parametros", "F3") == 1.0381
            assert v("RESULTADOS", "D81") == "—"                   # C3 incompleta
            assert v("RESULTADOS", "E81") == "Não informado"
            assert isinstance(v("RESULTADOS", "D80"), float)       # C2 completa
            assert v("RESULTADOS", "E80") == "Posição de abertura"
            assert v("RESULTADOS", "D83") == v("RESULTADOS_DETALHE", "C30")
            assert v("RESULTADOS", "E83") == "Posição atual"
        finally:
            _com(book.Close, False)
    finally:
        app.Quit()
        pythoncom.CoUninitialize()
