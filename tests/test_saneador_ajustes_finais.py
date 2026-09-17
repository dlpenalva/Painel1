"""Ajustes finais do DESPACHO SANEADOR (titulo, acumulado, TEMPESTIVO*,
quadros, capitulo 3, Quadro 3 e conclusao).

Todo o conteudo verificado aqui e APRESENTACAO: o gerador consome dados
canonicos ja apurados (rotulos de ciclo, percentual acumulado, inicio dos
efeitos financeiros, proxima data de reajuste) e nao recalcula nada. Os
testes provam justamente isso — inclusive que o acumulado exibido e o
canonico, diferente da soma dos percentuais dos ciclos.
"""
from __future__ import annotations

import sys
import zipfile
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from _templates_documentos import (  # noqa: E402
    _extrair_dados,
    gerar_despacho_saneador,
    gerar_modelo_branco_despacho,
    gerar_termo_apostila,
)
from test_sumario_executivo import (  # noqa: E402
    leitura_multiciclo_pc,
    leitura_simples_financeiro,
)
from test_templates_documentos import CAMPOS_SANEADOR, CAMPOS_TERMO  # noqa: E402

SITUACAO_TEMPESTIVO_ASTERISCO = "✅ TEMPESTIVO*"


def _texto(conteudo: bytes) -> str:
    doc = Document(BytesIO(conteudo))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.extend(celula.text for celula in linha.cells)
    return "\n".join(partes)


def _paragrafos_de_quadro(conteudo: bytes) -> list:
    doc = Document(BytesIO(conteudo))
    return [p for p in doc.paragraphs if p.text.strip().startswith("Quadro ")]


def _leitura_com_tempestivo_asterisco():
    """C1 tempestivo com efeito retardado — apresentado como TEMPESTIVO*."""
    leitura = leitura_simples_financeiro()
    leitura["parametros_v10"]["por_ciclo"]["C1"]["situacao"] = (
        SITUACAO_TEMPESTIVO_ASTERISCO
    )
    return leitura


# ---------------------------------------------------------------------------
# 1 — Titulo principal
# ---------------------------------------------------------------------------

def test_titulo_principal_com_um_unico_ciclo():
    texto = _texto(gerar_despacho_saneador(
        leitura_simples_financeiro(), campos_manuais=CAMPOS_SANEADOR
    ))
    assert (
        "Assunto: Saneamento para formalização de atualização contratual - "
        "reajuste do Ciclo C1 (Empresa XPTO S.A.)"
    ) in texto
    assert "reajustes dos Ciclos" not in texto


def test_titulo_principal_com_varios_ciclos_na_ordem_canonica():
    leitura = leitura_multiciclo_pc()
    ciclos = [
        c["ciclo"] for c in _extrair_dados(leitura, None)["ciclos_computados"]
    ]
    assert ciclos == ["C1", "C2"], "fixture deve ter mais de um ciclo computado"
    texto = _texto(gerar_despacho_saneador(leitura, campos_manuais=CAMPOS_SANEADOR))
    assert (
        "Assunto: Saneamento para formalização de atualização contratual - "
        "reajustes dos Ciclos C1, C2 (Empresa XPTO S.A.)"
    ) in texto


def test_titulo_principal_sem_dado_usa_marcador_e_nao_inventa_ciclo():
    texto = _texto(gerar_modelo_branco_despacho())
    assert (
        "Saneamento para formalização de atualização contratual - "
        "reajuste(s) do(s) Ciclo(s) [PREENCHER: Ciclos de reajuste] "
        "([PREENCHER: Nome da empresa contratada])"
    ) in texto


# ---------------------------------------------------------------------------
# 2 e 3 — Acumulado imediatamente abaixo do Quadro 1, sem soma no gerador
# ---------------------------------------------------------------------------

FRASE_ACUMULADO = (
    ", apurado pela composição dos fatores de reajuste de cada ciclo, e não "
    "pela soma dos respectivos percentuais."
)


def test_frase_do_acumulado_vem_imediatamente_apos_o_quadro_1():
    doc = Document(BytesIO(gerar_despacho_saneador(
        leitura_multiciclo_pc(), campos_manuais=CAMPOS_SANEADOR
    )))
    corpo = list(doc.element.body)
    indice_primeira_tabela = next(
        i for i, elemento in enumerate(corpo) if elemento.tag == qn("w:tbl")
    )
    seguinte = corpo[indice_primeira_tabela + 1]
    assert seguinte.tag == qn("w:p")
    texto_seguinte = "".join(no.text or "" for no in seguinte.iter(qn("w:t")))
    assert texto_seguinte.startswith("Acumulado dos ciclos: ")
    assert texto_seguinte.endswith(FRASE_ACUMULADO)


def test_acumulado_usa_o_percentual_canonico_e_nao_a_soma_dos_ciclos():
    leitura = leitura_multiciclo_pc()
    dados = _extrair_dados(leitura, None)
    canonico = dados["var_acumulada"]
    soma = sum(c["percentual_reajuste"] for c in dados["ciclos_computados"])
    assert round(canonico, 6) != round(soma, 6), (
        "fixture nao discrimina composicao x soma"
    )
    texto = _texto(gerar_despacho_saneador(leitura, campos_manuais=CAMPOS_SANEADOR))
    assert f"Acumulado dos ciclos: 5,99%{FRASE_ACUMULADO}" in texto
    assert "Acumulado dos ciclos: 5,90%" not in texto


def test_acumulado_em_negrito_somente_ate_o_percentual():
    doc = Document(BytesIO(gerar_despacho_saneador(
        leitura_multiciclo_pc(), campos_manuais=CAMPOS_SANEADOR
    )))
    p = next(
        p for p in doc.paragraphs if p.text.startswith("Acumulado dos ciclos: ")
    )
    negrito = [r.text for r in p.runs if r.bold]
    normal = [r.text for r in p.runs if not r.bold]
    assert negrito == ["Acumulado dos ciclos: 5,99%"]
    assert normal == [FRASE_ACUMULADO]


def test_acumulado_ausente_vira_marcador_e_nunca_zero():
    texto = _texto(gerar_modelo_branco_despacho())
    assert (
        "Acumulado dos ciclos: [PREENCHER: Percentual acumulado dos ciclos]"
        + FRASE_ACUMULADO
    ) in texto
    assert "Acumulado dos ciclos: 0,00%" not in texto


# ---------------------------------------------------------------------------
# 4 e 5 — TEMPESTIVO*
# ---------------------------------------------------------------------------

FRASE_ATUAL_SEM_ASTERISCO = (
    "Em razão da data do pedido, os efeitos financeiros do reajuste deste "
    "ciclo iniciam-se em 04/2025, não alcançando as competências de "
    "fevereiro e março de 2025."
)


def test_sem_tempestivo_asterisco_preserva_a_redacao_atual():
    texto = _texto(gerar_despacho_saneador(
        leitura_simples_financeiro(), campos_manuais=CAMPOS_SANEADOR
    ))
    assert FRASE_ATUAL_SEM_ASTERISCO in texto
    assert "TEMPESTIVO*" not in texto


def test_com_tempestivo_asterisco_usa_a_nova_redacao_dinamica():
    texto = _texto(gerar_despacho_saneador(
        _leitura_com_tempestivo_asterisco(), campos_manuais=CAMPOS_SANEADOR
    ))
    assert (
        "Em razão da data do pedido, os efeitos financeiros do reajuste "
        "referente ao ciclo C1 iniciam-se em 04/2025, não alcançando as "
        "competências de fevereiro e março de 2025. Por essa razão, a "
        "situação indicada como TEMPESTIVO* é acompanhada de asterisco, a fim "
        "de sinalizar essa particularidade quanto ao termo inicial dos "
        "efeitos financeiros."
    ) in texto
    assert FRASE_ATUAL_SEM_ASTERISCO not in texto


def test_titulo_do_capitulo_2_permanece_inalterado():
    for leitura in (leitura_simples_financeiro(), _leitura_com_tempestivo_asterisco()):
        texto = _texto(gerar_despacho_saneador(leitura, campos_manuais=CAMPOS_SANEADOR))
        assert "2. PEDIDO E PARÂMETROS DA ANÁLISE" in texto


def test_termo_de_apostila_nao_recebe_a_redacao_do_asterisco():
    texto = _texto(gerar_termo_apostila(
        _leitura_com_tempestivo_asterisco(), campos_manuais=CAMPOS_TERMO
    ))
    assert FRASE_ATUAL_SEM_ASTERISCO in texto
    assert "é acompanhada de asterisco" not in texto


# ---------------------------------------------------------------------------
# 6 e 7 — Titulos dos quadros
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("leitura", [leitura_simples_financeiro, leitura_multiciclo_pc])
def test_titulos_dos_quadros_do_saneador_sao_esquerda_italico_sem_negrito(leitura):
    paragrafos = _paragrafos_de_quadro(
        gerar_despacho_saneador(leitura(), campos_manuais=CAMPOS_SANEADOR)
    )
    assert [p.text for p in paragrafos] == [
        "Quadro 1 - Síntese da análise",
        "Quadro 2 - Síntese financeira",
        "Quadro 3 - Documentos e verificações",
    ]
    for p in paragrafos:
        assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT, p.text
        for run in p.runs:
            assert run.italic is True, p.text
            assert run.bold is False, p.text


def test_quadros_do_termo_de_apostila_ficam_como_estavam():
    paragrafos = _paragrafos_de_quadro(
        gerar_termo_apostila(leitura_multiciclo_pc(), campos_manuais=CAMPOS_TERMO)
    )
    assert paragrafos, "o Termo continua declarando os seus quadros"
    for p in paragrafos:
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER, p.text


# ---------------------------------------------------------------------------
# 8 e 9 — Capitulo 3 e Quadro 3
# ---------------------------------------------------------------------------

def test_capitulo_3_tem_o_titulo_exato():
    doc = Document(BytesIO(gerar_despacho_saneador(
        leitura_multiciclo_pc(), campos_manuais=CAMPOS_SANEADOR
    )))
    titulos = [p.text for p in doc.paragraphs if p.text.startswith("3.")]
    assert titulos == ["3. RESULTADO"]


def test_quadro_3_ganha_a_concordancia_da_area_gestora_sem_perder_linhas():
    doc = Document(BytesIO(gerar_despacho_saneador(
        leitura_multiciclo_pc(), campos_manuais=CAMPOS_SANEADOR
    )))
    quadro3 = doc.tables[-1]
    rotulos = [linha.cells[0].text for linha in quadro3.rows]
    assert rotulos == [
        "Documento ou verificação",
        "Memória de cálculo",
        "Adequação orçamentária",
        "Regularidade da contratada",
        "Concordância da contratada",
        "Concordância da área gestora",
        "Garantia contratual",
    ]
    linha_gestora = next(
        linha for linha in quadro3.rows
        if linha.cells[0].text == "Concordância da área gestora"
    )
    # Nunca preenchida automaticamente: sem fonte canonica, fica como marcador.
    assert linha_gestora.cells[1].text == (
        "[PREENCHER: Referencia da concordancia da area gestora] / "
        "[PREENCHER: Situacao da concordancia da area gestora]"
    )
    assert all(
        parte.strip().startswith("[PREENCHER:")
        for parte in linha_gestora.cells[1].text.split(" / ")
    )
    for afirmacao in ("de acordo", "validado", "concordou", "manifestou"):
        assert afirmacao not in linha_gestora.cells[1].text.lower()


# ---------------------------------------------------------------------------
# 10 — Conclusao
# ---------------------------------------------------------------------------

def test_conclusao_declara_ciclo_inicio_e_proxima_data_canonicos():
    leitura = leitura_multiciclo_pc()
    ciclo = leitura["parametros_v10"]["por_ciclo"]["C2"]
    ciclo["inicio_efeito_financeiro"] = date(2026, 3, 1)
    ciclo["proxima_data_reajuste"] = date(2027, 3, 26)
    texto = _texto(gerar_despacho_saneador(leitura, campos_manuais=CAMPOS_SANEADOR))
    assert (
        "Considerando que, no âmbito da presente atualização contratual, os "
        "efeitos financeiros relativos ao ciclo C2 são reconhecidos a partir "
        "de 03/2026, registra-se que o próximo ciclo de reajuste contratual "
        "estará apto a partir de 26/03/2027, observados os termos e a "
        "periodicidade previstos no contrato, desde que o instrumento "
        "jurídico ainda esteja em vigência."
    ) in texto
    # A proxima data e a canonica do ciclo, nao "inicio dos efeitos + 12 meses".
    assert "01/03/2027" not in texto


def test_conclusao_preserva_a_regra_do_segundo_paragrafo():
    frase = (
        "Após a complementação e conferência das informações documentais "
        "indicadas, deverá ser avaliado o prosseguimento da instrução para "
        "formalização."
    )
    completo = _texto(gerar_despacho_saneador(
        leitura_multiciclo_pc(), campos_manuais=CAMPOS_SANEADOR
    ))
    assert frase not in completo
    pendente = _texto(gerar_despacho_saneador(
        leitura_multiciclo_pc(),
        campos_manuais={
            k: v for k, v in CAMPOS_SANEADOR.items()
            if k != "adequacao_orcamentaria_ref"
        },
    ))
    assert frase in pendente


# ---------------------------------------------------------------------------
# 11 — DOCX real
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("leitura", [
    leitura_simples_financeiro, leitura_multiciclo_pc,
    _leitura_com_tempestivo_asterisco,
])
def test_docx_gerado_e_integro_e_reabre_sem_reparo(leitura, tmp_path):
    conteudo = gerar_despacho_saneador(leitura(), campos_manuais=CAMPOS_SANEADOR)
    arquivo = tmp_path / "saneador.docx"
    arquivo.write_bytes(conteudo)
    with zipfile.ZipFile(arquivo) as pacote:
        assert pacote.testzip() is None
        nomes = set(pacote.namelist())
    assert {"[Content_Types].xml", "word/document.xml"} <= nomes
    redocumento = Document(str(arquivo))
    assert redocumento.paragraphs[0].text == "DESPACHO SANEADOR"
    assert redocumento.tables


def test_word_real_abre_o_docx_sem_reparo(tmp_path):
    """Abertura no Word real (pulado onde o Word/pywin32 nao existir)."""
    client = pytest.importorskip("win32com.client")
    arquivo = tmp_path / "saneador_word.docx"
    arquivo.write_bytes(gerar_despacho_saneador(
        leitura_multiciclo_pc(), campos_manuais=CAMPOS_SANEADOR
    ))
    try:
        word = client.DispatchEx("Word.Application")
    except Exception as excecao:  # pragma: no cover - ambiente sem Word
        pytest.skip(f"Word indisponivel: {excecao}")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        documento = word.Documents.Open(
            str(arquivo), ConfirmConversions=False, ReadOnly=True,
            AddToRecentFiles=False,
        )
        try:
            assert documento.Paragraphs.Count > 0
            assert documento.Tables.Count == 3
        finally:
            documento.Close(False)
    finally:
        word.Quit()
