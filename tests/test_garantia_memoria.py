"""Regressão da memória XLSX da Garantia Contratual."""
from __future__ import annotations

import ast
import inspect
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import unittest
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from _garantia_calculo import (
    COLUNA_EVENTO_DATA,
    COLUNA_EVENTO_GARANTIA,
    COLUNA_EVENTO_TIPO,
    COLUNA_EVENTO_VALIDADE,
    COLUNA_EVENTO_VALOR,
    COLUNA_EVENTO_VIGENCIA,
    DIAS_VALIDADE_MINIMA,
    TIPO_ADITIVO,
    TIPO_OUTRO,
    TIPO_PRORROGACAO,
    TIPO_REAJUSTE,
    TIPO_REPACTUACAO,
    analisar_garantia,
    calcular_situacao_atual,
    normalizar_eventos,
)
from _garantia_memoria import NOME_ABA, TRACO, gerar_memoria_garantia_xlsx


ROOT = Path(__file__).resolve().parents[1]
FUSO_BRASILIA = ZoneInfo("America/Sao_Paulo")
GERADO_EM = datetime(2026, 10, 1, 14, 30, tzinfo=FUSO_BRASILIA)


def _registro(tipo, valor=None, vigencia=None, data=None, garantia=None, validade=None):
    return {
        COLUNA_EVENTO_TIPO: tipo,
        COLUNA_EVENTO_DATA: data,
        COLUNA_EVENTO_VALOR: valor,
        COLUNA_EVENTO_VIGENCIA: vigencia,
        COLUNA_EVENTO_GARANTIA: garantia,
        COLUNA_EVENTO_VALIDADE: validade,
    }


def _canonico(*, valor="1.000.000,00", percentual="5", vigencia=date(2026, 12, 31),
              garantia=None, validade=None, registros=()):
    eventos, avisos, pendencias = normalizar_eventos(registros)
    assert not avisos and not pendencias
    situacao = calcular_situacao_atual(
        valor, percentual, vigencia, eventos, garantia, validade
    )
    analise = analisar_garantia(
        situacao["valor_atual"], situacao["percentual"], situacao["vigencia_atual"],
        situacao["garantia_apresentada"], situacao["validade_apresentada"],
    )
    return situacao, analise


def _workbook(situacao, analise):
    conteudo = gerar_memoria_garantia_xlsx(
        situacao,
        analise,
        versao_cl8us="11.2",
        gerado_em=GERADO_EM,
        dias_validade_minima=DIAS_VALIDADE_MINIMA,
    )
    return load_workbook(BytesIO(conteudo), data_only=False)


def _valor_por_rotulo(ws, rotulo):
    for linha in ws.iter_rows():
        for celula in linha:
            if celula.value == rotulo:
                return ws.cell(celula.row, 5).value
    raise AssertionError(f"rótulo não encontrado: {rotulo}")


def _linha_por_texto(ws, texto):
    for linha in ws.iter_rows():
        if any(celula.value == texto for celula in linha):
            return linha[0].row
    raise AssertionError(f"texto não encontrado: {texto}")


def _como_data(valor):
    return valor.date() if isinstance(valor, datetime) else valor


class IntegridadeMemoriaTests(unittest.TestCase):
    def test_uma_unica_aba_sem_formulas_e_com_configuracao_de_impressao(self):
        situacao, analise = _canonico()
        wb = _workbook(situacao, analise)
        self.assertEqual(wb.sheetnames, [NOME_ABA])
        ws = wb[NOME_ABA]
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertTrue(ws.print_area)
        self.assertFalse(any(
            isinstance(celula.value, str) and celula.value.startswith("=")
            for linha in ws.iter_rows() for celula in linha
        ))

    def test_horario_geracao_e_convertido_para_brasilia_sem_deslocamento_silencioso(self):
        situacao, analise = _canonico()
        conteudo = gerar_memoria_garantia_xlsx(
            situacao,
            analise,
            versao_cl8us="11.2",
            gerado_em=datetime(2026, 10, 2, 0, 30, tzinfo=timezone.utc),
            dias_validade_minima=DIAS_VALIDADE_MINIMA,
        )
        ws = load_workbook(BytesIO(conteudo), data_only=False)[NOME_ABA]
        self.assertEqual(_valor_por_rotulo(ws, "Data e hora de geração"), "01/10/2026 21:30")

        with self.assertRaisesRegex(ValueError, "fuso horário explícito"):
            gerar_memoria_garantia_xlsx(
                situacao,
                analise,
                versao_cl8us="11.2",
                gerado_em=datetime(2026, 10, 2, 0, 30),
                dias_validade_minima=DIAS_VALIDADE_MINIMA,
            )

    def test_historico_longo_mantem_impressoes_sem_cabecalho_global(self):
        registros = [
            _registro(
                TIPO_REAJUSTE,
                f"{1_000_000 + indice * 10_000:,}".replace(",", ".") + ",00",
                data=date(2027 + (indice - 1) // 12, (indice - 1) % 12 + 1, 15),
            )
            for indice in range(1, 21)
        ]
        situacao, analise = _canonico(registros=registros)
        wb = _workbook(situacao, analise)
        ws = wb[NOME_ABA]

        self.assertGreater(len(situacao["linha_do_tempo"]), 18)
        self.assertIsNone(ws.print_title_rows)
        self.assertNotIn("_xlnm.Print_Titles", wb.defined_names)
        self.assertLess(
            _linha_por_texto(ws, "SITUAÇÃO ATUAL"),
            _linha_por_texto(ws, "CONCLUSÃO"),
        )
        self.assertIn(f"$I${ws.max_row}", str(ws.print_area))

    def test_exportador_nao_importa_streamlit_nem_motor_e_nao_recalcula_regra(self):
        import _garantia_memoria as modulo

        fonte = inspect.getsource(modulo)
        arvore = ast.parse(fonte)
        imports = {
            alias.name
            for no in ast.walk(arvore)
            if isinstance(no, (ast.Import, ast.ImportFrom))
            for alias in no.names
        }
        self.assertNotIn("streamlit", imports)
        self.assertNotIn("_garantia_calculo", imports)
        for proibido in (
            "calcular_situacao_atual(", "analisar_garantia(",
            "normalizar_eventos(", "calcular_garantia_necessaria(",
            "calcular_validade_minima(", "ROUND_HALF_UP", "timedelta(",
        ):
            self.assertNotIn(proibido, fonte)

    def test_celulas_relevantes_reproduzem_o_resultado_canonico(self):
        situacao, analise = _canonico(
            valor="5.117.331,91",
            vigencia=date(2029, 12, 31),
            garantia="255.866,60",
            validade=date(2031, 3, 23),
            registros=[_registro(
                TIPO_REAJUSTE, "5.240.904,95", vigencia=date(2030, 12, 23),
                data=date(2027, 1, 15),
            )],
        )
        ws = _workbook(situacao, analise)[NOME_ABA]
        esperados = {
            "Valor original do contrato": situacao["valor_original"],
            "Valor atualizado total do contrato": situacao["valor_atual"],
            "Garantia atualmente exigida": analise["garantia_necessaria"],
            "Última garantia apresentada": analise["garantia_apresentada"],
            "Complemento necessário": analise["complemento"],
            "Validade mínima exigida": analise["validade_minima"],
        }
        for rotulo, esperado in esperados.items():
            obtido = _valor_por_rotulo(ws, rotulo)
            if isinstance(esperado, Decimal):
                self.assertEqual(Decimal(str(obtido)), esperado, rotulo)
            else:
                self.assertEqual(_como_data(obtido), esperado, rotulo)
        self.assertEqual(
            _como_data(_valor_por_rotulo(ws, "Término da vigência original")),
            situacao["vigencia_original"],
        )
        self.assertEqual(
            Decimal(str(_valor_por_rotulo(ws, "Percentual da garantia"))),
            situacao["percentual"] / Decimal("100"),
        )

    def test_evolucao_corresponde_exatamente_a_linha_do_tempo_canonica(self):
        registros = [
            _registro(TIPO_REAJUSTE, "1.100.000,00", data=date(2027, 1, 10)),
            _registro(TIPO_PRORROGACAO, vigencia=date(2027, 12, 31),
                      garantia="60.000,00", validade=date(2028, 4, 1)),
            _registro(TIPO_ADITIVO, "1.250.000,00", data=date(2027, 8, 1)),
        ]
        situacao, analise = _canonico(
            garantia="50.000,00", validade=date(2027, 3, 31), registros=registros
        )
        ws = _workbook(situacao, analise)[NOME_ABA]
        cabecalho = _linha_por_texto(ws, "EVOLUÇÃO DO CONTRATO E DA GARANTIA") + 1
        for indice, etapa in enumerate(situacao["linha_do_tempo"], start=1):
            linha = cabecalho + 1 + indice
            self.assertEqual(ws.cell(linha, 1).value, etapa["numero"])
            self.assertEqual(ws.cell(linha, 2).value, etapa["tipo"])
            self.assertEqual(_como_data(ws.cell(linha, 3).value), etapa["data"] or TRACO)
            self.assertEqual(Decimal(str(ws.cell(linha, 4).value)), etapa["valor"])
            self.assertEqual(Decimal(str(ws.cell(linha, 5).value)), etapa["variacao"])
            self.assertEqual(Decimal(str(ws.cell(linha, 6).value)), etapa["garantia_exigida"])
            self.assertEqual(_como_data(ws.cell(linha, 7).value), etapa["vigencia"])
            apresentada = ws.cell(linha, 8).value
            if etapa["garantia_apresentada"] is None:
                self.assertEqual(apresentada, TRACO)
            else:
                self.assertEqual(Decimal(str(apresentada)), etapa["garantia_apresentada"])
            self.assertEqual(
                _como_data(ws.cell(linha, 9).value), etapa["validade_apresentada"] or TRACO
            )

    def test_ausencia_de_garantia_permanece_travessao_e_nunca_zero(self):
        situacao, analise = _canonico()
        ws = _workbook(situacao, analise)[NOME_ABA]
        self.assertEqual(_valor_por_rotulo(ws, "Garantia apresentada na assinatura"), TRACO)
        self.assertEqual(_valor_por_rotulo(ws, "Última garantia apresentada"), TRACO)
        linha_memoria = _linha_por_texto(ws, "Complemento necessário")
        textos = [ws.cell(linha, 3).value for linha in range(linha_memoria, ws.max_row + 1)]
        self.assertTrue(any("Não foi informada garantia apresentada" in str(t) for t in textos))


class CenariosMemoriaTests(unittest.TestCase):
    def test_os_doze_cenarios_geram_memoria_coerente(self):
        minimo_2027 = date(2027, 3, 31)
        cenarios = [
            ("sem alterações", {}, 0),
            ("um reajuste", {"registros": [_registro(TIPO_REAJUSTE, "1.100.000,00")]}, 1),
            ("reajuste e prorrogação", {"registros": [
                _registro(TIPO_REAJUSTE, "1.100.000,00"),
                _registro(TIPO_PRORROGACAO, vigencia=date(2027, 12, 31)),
            ]}, 2),
            ("múltiplos eventos", {"registros": [
                _registro(TIPO_REAJUSTE, "1.100.000,00"),
                _registro(TIPO_REPACTUACAO, "1.150.000,00"),
                _registro(TIPO_ADITIVO, "1.250.000,00"),
                _registro(TIPO_OUTRO, garantia="70.000,00", validade=date(2028, 1, 1)),
            ]}, 4),
            ("garantia suficiente", {"garantia": "50.000,00", "validade": minimo_2027}, 0),
            ("complemento necessário", {"garantia": "40.000,00", "validade": minimo_2027}, 0),
            ("validade insuficiente", {"garantia": "50.000,00", "validade": date(2027, 1, 1)}, 0),
            ("valor e validade insuficientes", {"garantia": "40.000,00", "validade": date(2027, 1, 1)}, 0),
            ("garantia superior", {"garantia": "80.000,00", "validade": minimo_2027}, 0),
            ("ausência de garantia", {}, 0),
            ("fotografia herdada", {"garantia": "50.000,00", "validade": minimo_2027,
                "registros": [_registro(TIPO_REAJUSTE, "1.100.000,00")]}, 1),
            ("nova fotografia substitui anterior", {"garantia": "50.000,00", "validade": minimo_2027,
                "registros": [_registro(TIPO_REAJUSTE, "1.100.000,00", garantia="60.000,00",
                                         validade=date(2027, 5, 1))]}, 1),
        ]
        for nome, kwargs, quantidade in cenarios:
            with self.subTest(nome=nome):
                situacao, analise = _canonico(**kwargs)
                ws = _workbook(situacao, analise)[NOME_ABA]
                self.assertEqual(situacao["quantidade_eventos"], quantidade)
                self.assertEqual(_valor_por_rotulo(ws, "Diagnóstico"), analise["diagnostico"])
                self.assertEqual(
                    Decimal(str(_valor_por_rotulo(ws, "Complemento necessário"))),
                    analise["complemento"],
                )

    def test_cenario_visual_solicitado(self):
        situacao, analise = _canonico(
            valor="5.117.331,91",
            garantia="255.866,60",
            validade=date(2031, 3, 23),
            vigencia=date(2030, 12, 23),
            registros=[_registro(TIPO_REAJUSTE, "5.240.904,95")],
        )
        self.assertEqual(situacao["valor_atual"], Decimal("5240904.95"))
        self.assertEqual(analise["garantia_necessaria"], Decimal("262045.25"))
        self.assertEqual(analise["complemento"], Decimal("6178.65"))
        self.assertEqual(analise["validade_minima"], date(2031, 3, 23))
        ws = _workbook(situacao, analise)[NOME_ABA]
        self.assertEqual(Decimal(str(_valor_por_rotulo(ws, "Garantia atualmente exigida"))),
                         Decimal("262045.25"))


class IntegracaoPaginaTests(unittest.TestCase):
    def test_box_fica_depois_da_situacao_atual_e_antes_do_resultado(self):
        pagina = (ROOT / "pages" / "05_Garantia.py").read_text(encoding="utf-8")
        pos_resumo = pagina.index("render_resumo_situacao_atual(situacao, analise)",
                                  pagina.index("# Interface"))
        pos_memoria = pagina.index('st.subheader("Memória de cálculo")')
        pos_resultado = pagina.index('st.subheader("Resultado da análise")')
        self.assertLess(pos_resumo, pos_memoria)
        self.assertLess(pos_memoria, pos_resultado)
        self.assertIn('key="baixar_memoria_calculo_garantia"', pagina)
        self.assertIn('file_name="memoria_calculo_garantia.xlsx"', pagina)
        self.assertIn('FUSO_BRASILIA = ZoneInfo("America/Sao_Paulo")', pagina)
        self.assertIn("gerado_em=datetime.now(FUSO_BRASILIA)", pagina)
        self.assertNotIn("gerado_em=datetime.now()", pagina)


if __name__ == "__main__":
    unittest.main()
