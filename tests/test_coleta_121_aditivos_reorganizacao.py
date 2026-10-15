"""Coleta 12.1 (Pacote 3) — reorganizacao controlada da aba aditivos.

Somente geometria e apresentacao: blocos de leitura, VALOR UNITARIO ATUALIZADO
exposto em J (o mesmo termo que o valor atualizado da alteracao ja usava) e
letras fisicas novas para totais/controles. A regra economica nao muda.

Estrutural (sempre): mapa, formulas canonicas, validacoes, formatacao
condicional, referencias externas e leitores nos dois layouts.
Excel real (RUN_EXCEL_INTEGRATION=1): caso N001-N005 nascendo em C3 + item
existente com acrescimo (Sim) e supressao (Nao), cortes em C3 e C4.
"""
from __future__ import annotations

import gc
import os
import time
import zipfile
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

import _coleta_oficial as co
from _layout_aditivos import (
    CABECALHOS_ADITIVOS_12_1,
    COLUNAS_ADITIVOS_12_1,
    COLUNAS_ADITIVOS_LEGADO,
    colunas_aditivos,
    traduzir_formula_legado_para_12_1,
)

TEMPLATE = co.TEMPLATE_COLETA_OFICIAL


@pytest.fixture(scope="module")
def wb_template():
    wb = load_workbook(TEMPLATE, data_only=False)
    yield wb
    wb.close()


def _validacoes(ws) -> dict[str, object]:
    return {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}


# ------------------------------------------------------------------ estrutura
def test_mapa_de_colunas_e_cabecalhos(wb_template):
    ws = wb_template["aditivos"]
    assert {c: ws[f"{c}1"].value for c in CABECALHOS_ADITIVOS_12_1} == CABECALHOS_ADITIVOS_12_1
    assert colunas_aditivos(ws) == COLUNAS_ADITIVOS_12_1
    assert ws.column_dimensions["P"].hidden  # INDICE_BASE_ECONOMICA_VU segue tecnica
    assert ws.freeze_panes == "A2"
    # Linha 1 sem a orientacao longa: cabe cabecalhos de ate 3 linhas.
    assert 50 <= ws.row_dimensions[1].height <= 60
    assert ws.max_column == 16


def test_formulas_canonicas_e_vu_atualizado_e_a_origem_do_valor(wb_template):
    ws = wb_template["aditivos"]
    assert co._aditivos_base_economica_canonica(ws)
    for linha in (2, 99, 200):
        assert ws[f"J{linha}"].value == (
            f'=IF(F{linha}="","",IF(UPPER(H{linha})="SIM",'
            f'IF(ISNUMBER(I{linha}),ROUND(F{linha}*I{linha},2),""),F{linha}))'
        )
        # Valor atualizado = DELTA x VU atualizado: uma unica origem do VU.
        assert ws[f"L{linha}"].value == (
            f'=IF(OR(N{linha}="",J{linha}=""),"",ROUND(N{linha}*J{linha},2))'
        )
        assert ws[f"K{linha}"].value == (
            f'=IF(OR(E{linha}="",F{linha}=""),"",ROUND(E{linha}*F{linha},2))'
        )
        assert "ROUND(ABS(E" in ws[f"N{linha}"].value
        assert ws[f"O{linha}"].value == co._formula_status_aditivo(linha)


def test_vu_atualizado_reproduz_o_termo_da_formula_da_coleta_12_0():
    """A forma 12.0 do valor atualizado e a forma 12.1 de L (com J dentro)
    tem os mesmos ramos e os mesmos arredondamentos."""
    antiga = (
        '=IF(OR(L2="",F2=""),"",IF(UPPER(H2)="SIM",'
        'IF(ISNUMBER(I2),ROUND(L2*ROUND(F2*I2,2),2),""),ROUND(L2*F2,2)))'
    )
    traduzida = traduzir_formula_legado_para_12_1(antiga, aba_da_formula="aditivos")
    vu = co._formula_vu_atualizado_aditivo(2)[1:]
    expandida = co._formula_valor_aditivo(2).replace("J2", f"({vu})")
    for termo in ('UPPER(H2)="SIM"', "ISNUMBER(I2)", "ROUND(F2*I2,2)", 'N2=""'):
        assert termo in traduzida and termo in expandida


def test_dropdowns_e_orientacao_de_novo_item(wb_template):
    dvs = _validacoes(wb_template["aditivos"])
    assert set(dvs) == {"H2:H200 M2:M200", "D2:D200", "A2:A200", "E2:E200", "G2:G200"}
    simnao = dvs["H2:H200 M2:M200"]
    assert (simnao.type, simnao.formula1) == ("list", '"Sim,Nao"')
    assert dvs["D2:D200"].formula1 == '"Acrescimo,Acréscimo - novo item,Supressao"'
    base = dvs["G2:G200"]
    assert (base.type, base.formula1) == ("list", '"C0,C1,C2,C3,C4"')
    assert base.promptTitle == "Base econômica do VU" and base.showErrorMessage
    item = dvs["A2:A200"]
    # Coleta 12.2: titulo NOVO ITEM e texto do cadastro Nxxx (mesma validacao).
    assert item.showInputMessage and item.promptTitle == "NOVO ITEM"
    assert "Cadastre primeiro em itens_Remanesc como N001, N002..." in item.prompt
    assert "Informe o VU_ORIGINAL e, depois, registre aqui a data e a quantidade" in item.prompt
    assert len(item.prompt) <= 255
    assert dvs["E2:E200"].promptTitle == "QUANTIDADE DO ADITIVO"


def test_formatacao_condicional_acompanha_a_geometria(wb_template):
    regras = [
        (str(cf.sqref), r.formula[0])
        for cf in wb_template["aditivos"].conditional_formatting
        for r in cf.rules
    ]
    supressao = 'OR(LEFT(UPPER($D2),6)="SUPRES",LEFT(UPPER($D2),4)="DECR")'
    assert regras.count(("A2:F200 H2:M200", supressao)) == 2
    assert ("A2:N200 P2:P200", co._FORMULA_DESTAQUE_NOVOS_ITENS) in regras
    assert ("G2:G200", co._FORMULA_CF_BASE_VU_NAO_APLICAVEL) in regras
    assert sum(1 for faixa, _ in regras if faixa == "C2:C200") == 5
    assert len(regras) == 9


def test_referencias_externas_seguem_a_geometria(wb_template):
    assert "aditivos!$N$2:$N$200" in wb_template["posicao_contratual"]["D2"].value
    assert "aditivos!$N$2:$N$200" in wb_template["posicao_referencia"]["C2"].value
    assert "INDEX(aditivos!$P$2:$P$200" in wb_template["historico_VU"]["D2"].value
    mem = wb_template["MEMORIA_RESULTADOS"]
    assert mem["T48"].value == '=COUNTIF(aditivos!$O$2:$O$200,"ALERTA:*")'
    assert mem["B45"].value == '=COUNTIF(aditivos!$M$2:$M$200,"Sim")'


def _aba_aditivos(layout: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "aditivos"
    if layout is COLUNAS_ADITIVOS_12_1:
        for coluna, texto in CABECALHOS_ADITIVOS_12_1.items():
            ws[f"{coluna}1"] = texto
    else:
        ws["J1"] = "Valor atualizado da alteracao (QTD x VU atualizado)"
    ws["A2"], ws["B2"], ws["C2"] = "N001", datetime(2026, 2, 1), "C3"
    ws["D2"], ws["E2"], ws["F2"] = "Acréscimo - novo item", 3, 623976.61
    ws[f"{layout['valor_original']}2"] = 1871929.83
    ws[f"{layout['valor_atualizado']}2"] = 1943250.36
    ws[f"{layout['considerado']}2"] = "Sim"
    return wb


@pytest.mark.parametrize(
    "layout", [COLUNAS_ADITIVOS_12_1, COLUNAS_ADITIVOS_LEGADO], ids=["12.1", "ate_12.0"]
)
def test_leitor_do_upload_reconhece_os_dois_layouts(layout):
    from _leitor_masterfile_v10 import _ler_parcelas_sombra_aditivos

    parcelas = _ler_parcelas_sombra_aditivos(_aba_aditivos(layout))
    assert len(parcelas) == 1
    assert parcelas[0]["valor"] == pytest.approx(1943250.36)
    assert parcelas[0]["valor_original"] == pytest.approx(1871929.83)
    assert parcelas[0]["ciclo"] == "C3"


def test_coleta_entregue_mantem_a_forma_12_1():
    wb = load_workbook(BytesIO(co.obter_coleta_oficial_bytes()), data_only=False)
    ws = wb["aditivos"]
    assert colunas_aditivos(ws) == COLUNAS_ADITIVOS_12_1
    assert co._aditivos_base_economica_canonica(ws)
    cee = wb["CICLO_EM_EXECUCAO"]
    assert "aditivos!$N$2:$N$200" in cee["I13"].value
    assert "aditivos!$P$2:$P$200" in cee["E13"].value
    assert ws["O2"].alignment.wrap_text is True
    wb.close()


# ------------------------------------------------------------------ Excel real
excel_real = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar o Excel COM",
)

# Caso real (valores de referencia): N001..N005 nascem em C3 com base C0;
# unico fator conhecido C1 = 3,81% -> fator aplicavel 1,0381.
NOVOS = [  # item, linha itens_Remanesc, VU, quantidade, VU atualizado, total original, total atualizado
    ("N001", 4, 623976.61, 3, 647750.12, 1871929.83, 1943250.36),
    ("N002", 5, 623976.61, 3, 647750.12, 1871929.83, 1943250.36),
    ("N003", 6, 2599.67, 1851, 2698.72, 4811989.17, 4995330.72),
    ("N004", 7, 6164.80, 60, 6399.68, 369888.00, 383980.80),
    ("N005", 8, 150097.41, 24, 155816.12, 3602337.84, 3739586.88),
]


def _dados(corte: datetime) -> dict:
    return {
        "origem": "Coleta 12.1 aditivos",
        "indice": "IST",
        "data_base_original": "01/01/2023",
        "data_corte": corte.date(),
        "ciclos": [{
            "ciclo": "C1", "data_inicio": date(2024, 1, 1), "data_fim": date(2024, 12, 31),
            "data_pedido": date(2024, 1, 1), "financeiro_inicio": date(2024, 1, 1),
            "percentual_aplicado": 0.0381, "objeto_analise_atual": True,
        }],
    }


def _tentar(pythoncom, acao):
    for _ in range(60):
        try:
            return acao()
        except Exception as exc:  # pragma: no cover - somente Excel COM
            codigo = getattr(exc, "hresult", None)
            if codigo is None and getattr(exc, "args", ()):
                codigo = exc.args[0]
            if codigo != -2147418111:
                raise
            pythoncom.PumpWaitingMessages()
            time.sleep(0.3)
    return acao()


def _preencher(livro, corte: datetime) -> None:
    from pywintypes import Time

    livro.Worksheets("CONTROLE").Range("B1").Value = "Financeiro (Mensalidade)"
    livro.Worksheets("CONTROLE").Range("B3").Value = Time(corte)
    livro.Worksheets("financeiro").Range("C14").Value = 600
    itens = livro.Worksheets("itens_Remanesc")
    for linha, item, base, vu, c1, c3 in (
        (2, "1.1", 1000, 10.0, 950, 850),
        (3, "1.2", 500, 20.0, 480, 400),
    ):
        itens.Range(f"A{linha}").Value = item
        itens.Range(f"B{linha}").Value = base
        itens.Range(f"C{linha}").Value = vu
        itens.Range(f"E{linha}").Value = c1
        itens.Range(f"I{linha}").Value = c3
    for item, linha, vu, *_ in NOVOS:
        itens.Range(f"A{linha}").Value = item
        itens.Range(f"C{linha}").Value = vu
    ws = livro.Worksheets("aditivos")
    eventos = [(item, "Acréscimo - novo item", qtd, "Sim", "C0") for item, _l, _v, qtd, *_ in NOVOS]
    eventos += [("1.1", "Acrescimo", 100, "Sim", None), ("1.2", "Supressao", 50, "Nao", None)]
    for linha, (item, tipo, qtd, aplicar, base) in enumerate(eventos, start=2):
        ws.Range(f"A{linha}").Value = item
        ws.Range(f"B{linha}").Value = Time(datetime(2026, 2, 1))
        ws.Range(f"D{linha}").Value = tipo
        ws.Range(f"E{linha}").Value = qtd
        ws.Range(f"H{linha}").Value = aplicar
        ws.Range(f"M{linha}").Value = "Sim"
        if base:
            ws.Range(f"G{linha}").Value = base
    cee = livro.Worksheets("CICLO_EM_EXECUCAO")
    cee.Range("D5").Value2 = livro.Worksheets("CONTROLE").Range("B3").Value2
    livro.Application.CalculateFullRebuild()
    for i, consumo in enumerate((800, 300, 1, 2, 1000, 30, 10)):
        cee.Range(f"C{13 + i}").Value = consumo


def _ler(livro) -> dict:
    ws = livro.Worksheets("aditivos")
    hist = livro.Worksheets("historico_VU")
    pos = livro.Worksheets("posicao_contratual")
    cee = livro.Worksheets("CICLO_EM_EXECUCAO")
    linhas = range(2, 9)
    return {
        "adt": {c: [ws.Range(f"{c}{l}").Value for l in linhas] for c in "GHIJKLNOP"},
        "hist": {l: [hist.Range(f"{c}{l}").Value for c in "CDEFG"] for l in linhas},
        "pos": {l: [pos.Range(f"{c}{l}").Value for c in "EIMQU"] for l in linhas},
        "cee_ciclo": cee.Range("C3").Value,
        "cee_E": {l: cee.Range(f"E{l}").Value for l in range(13, 20)},
        "gate": livro.Worksheets("MEMORIA_RESULTADOS").Range("T48").Value,
        "dv_a2": (ws.Range("A2").Validation.InputTitle, ws.Range("A2").Validation.InputMessage),
        "dv_h2": str(ws.Range("H2").Validation.Formula1).replace(";", ","),
        "dv_m2": str(ws.Range("M2").Validation.Formula1).replace(";", ","),
        "dv_g2": str(ws.Range("G2").Validation.Formula1).replace(";", ","),
        "j_formula": ws.Range("J2").Formula,
    }


@pytest.fixture(scope="module", params=["C3", "C4"])
def excel_caso_real(request, tmp_path_factory):
    import pythoncom
    import win32com.client

    corte = datetime(2026, 8, 31) if request.param == "C3" else datetime(2027, 3, 31)
    caminho = tmp_path_factory.mktemp("adt121") / f"Coleta_121_aditivos_{request.param}.xlsx"
    caminho.write_bytes(co.gerar_coleta_oficial_preenchida(_dados(corte)))
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    livro = None
    try:
        abrir = lambda: excel.Workbooks.Open(  # noqa: E731
            str(caminho.resolve()), UpdateLinks=0, ReadOnly=False, CorruptLoad=0
        )
        livro = _tentar(pythoncom, abrir)
        time.sleep(1.0)
        _tentar(pythoncom, lambda: _preencher(livro, corte))
        _tentar(pythoncom, excel.CalculateFullRebuild)
        resultado = _tentar(pythoncom, lambda: _ler(livro))
        _tentar(pythoncom, livro.Save)
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = _tentar(pythoncom, abrir)
        _tentar(pythoncom, excel.CalculateFullRebuild)
        resultado["reaberto"] = _tentar(pythoncom, lambda: _ler(livro))
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = None
    finally:
        if livro is not None:
            try:
                livro.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()
    with zipfile.ZipFile(caminho) as pacote:
        assert not any(b"repairLoad" in pacote.read(n) for n in pacote.namelist())
    resultado["corte"] = request.param
    return resultado


@excel_real
def test_excel_novos_itens_nxxx_nascendo_em_c3(excel_caso_real):
    r = excel_caso_real
    adt = r["adt"]
    for indice, (_item, linha_itens, _vu, qtd, vu_atual, total_orig, total_atual) in enumerate(NOVOS):
        assert adt["G"][indice] == "C0"                                  # G base economica
        assert adt["I"][indice] == pytest.approx(1.0381)                # fator
        assert adt["J"][indice] == pytest.approx(vu_atual, abs=1e-9)     # H VU atualizado
        assert adt["K"][indice] == pytest.approx(total_orig, abs=1e-9)   # I valor original
        assert adt["L"][indice] == pytest.approx(total_atual, abs=1e-9)  # J valor atualizado
        assert adt["N"][indice] == qtd                                    # K DELTA
        assert adt["O"][indice] == "OK"                                   # L CHECK
        assert adt["P"][indice] == 0
        # Nascimento em C3: sem VU nem quantidade antes; continuidade depois.
        hist = r["hist"][linha_itens]
        assert hist[:3] == ["", "", ""]
        assert hist[3] == pytest.approx(vu_atual, abs=1e-9)  # J = VU economico do historico
        pos = r["pos"][linha_itens]
        assert pos[:3] == [0, 0, 0]
        assert pos[3:] == [qtd, qtd]
        # CICLO_EM_EXECUCAO: C3 pelo historico; C4 pelo fallback com a mesma base.
        assert r["cee_E"][13 + linha_itens - 2] == pytest.approx(vu_atual, abs=1e-9)
    assert r["cee_ciclo"] == r["corte"]
    assert r["gate"] == 0


@excel_real
def test_excel_item_existente_acrescimo_sim_e_supressao_nao(excel_caso_real):
    adt = excel_caso_real["adt"]
    acr, sup = 5, 6
    # A + E: acrescimo com reajuste -> VU atualizado 10,38.
    assert adt["I"][acr] == pytest.approx(1.0381)
    assert adt["J"][acr] == pytest.approx(10.38)
    assert adt["K"][acr] == pytest.approx(1000.0)
    assert adt["L"][acr] == pytest.approx(1038.0)
    assert adt["N"][acr] == 100
    assert adt["O"][acr] == "OK"
    assert adt["G"][acr] is None
    # B + F: supressao sem reajuste -> VU atualizado = VU original.
    assert adt["J"][sup] == pytest.approx(20.0)
    assert adt["K"][sup] == pytest.approx(1000.0)
    assert adt["L"][sup] == pytest.approx(-1000.0)
    assert adt["N"][sup] == -50
    assert adt["O"][sup] == "OK"


@excel_real
def test_excel_reabertura_dropdowns_e_orientacao(excel_caso_real):
    r = excel_caso_real
    reaberto = r["reaberto"]
    for chave in ("adt", "hist", "pos", "cee_E", "cee_ciclo", "gate"):
        assert reaberto[chave] == r[chave], chave
    titulo, mensagem = reaberto["dv_a2"]
    assert titulo == "NOVO ITEM"                                   # Coleta 12.2
    assert "Cadastre primeiro em itens_Remanesc" in mensagem
    assert reaberto["dv_h2"] == "Sim,Nao" and reaberto["dv_m2"] == "Sim,Nao"
    assert reaberto["dv_g2"] == "C0,C1,C2,C3,C4"
    assert reaberto["j_formula"] == co._formula_vu_atualizado_aditivo(2)
