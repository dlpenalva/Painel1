"""Coleta 12.0 — itens_PC com filtro, PC_PAGO normalizado e coerencia do metodo.

* itens_PC: AutoFilter A1:L5001 utilizavel com a aba protegida (ordenar segue
  bloqueado); AD (PC_PAGO_NORMALIZADO, oculta) normaliza G; H/I/J/K e os
  quadros leem AD, nunca a grafia bruta de G;
* MEMORIA_RESULTADOS!T54:T65: fontes preenchidas, coerencia metodo x fonte e
  contagens de grafia; o bloqueio (T58 = 1) esvazia VTA e retroativo e entra
  nas pendencias; nada usa outra fonte automaticamente;
* leitor Python alinhado: so Sim/Nao (sem maiusculas, acento e espacos).

O comportamento numerico no Excel real esta em test_coleta_120_excel_real.py.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

import _coleta_oficial as co
from _leitor_masterfile_v10 import _pc_pago_canonico, ler_masterfile_v10
from test_vta_pc_independente_27b import _workbook_pc


def _template():
    return load_workbook(co.TEMPLATE_COLETA_OFICIAL, data_only=False)


def _gerada():
    return load_workbook(io.BytesIO(co.obter_coleta_oficial_bytes()), data_only=False)


def _formulas(wb):
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str) and celula.value.startswith("="):
                    yield ws.title, celula.coordinate, celula.value


def _ascii_balanceada(formula: str) -> bool:
    return formula.isascii() and formula.count("(") == formula.count(")")


# --------------------------------------------------------------- itens_PC
def test_ad_normaliza_g_em_toda_a_grade_e_h_i_j_k_leem_ad():
    ws = _template()["itens_PC"]
    assert ws["AD1"].value == co.CABECALHO_PC_PAGO_NORMALIZADO
    for linha in range(2, 5002):
        assert ws[f"AD{linha}"].value == co._formula_pc_pago_normalizado(linha)
        assert ws[f"H{linha}"].value == co._formula_retro_pc(linha)
        assert ws[f"I{linha}"].value == co._formula_em_analise_pc(linha)
        assert ws[f"J{linha}"].value == co._formula_potencial_pc(linha)
        assert ws[f"K{linha}"].value == co._formula_check_pc(linha)


def test_normalizacao_ignora_maiusculas_acento_e_espacos_e_nao_inventa_significado():
    ad = co._formula_pc_pago_normalizado(2)
    assert _ascii_balanceada(ad)
    assert "UPPER(TRIM(SUBSTITUTE(G2,CHAR(160)," in ad  # espacos, inclusive nao separavel
    assert '"SIM","Sim"' in ad
    assert "_xlfn.UNICHAR(195)" in ad and "_xlfn.UNICHAR(227)" in ad  # Ã e ã
    assert '"NAO","Nao"' in ad
    assert ad.count('"INVALIDO"') == 2  # qualquer outro texto (e erro) = invalido


def test_formulas_criticas_nao_dependem_da_grafia_de_g():
    for linha in (2, 5001):
        for formula in (co._formula_retro_pc(linha), co._formula_em_analise_pc(linha),
                        co._formula_potencial_pc(linha)):
            assert f'IF(AD{linha}="Sim",' in formula
            assert f'G{linha}="Sim"' not in formula
        k = co._formula_check_pc(linha)
        assert f'IF(AD{linha}="INVALIDO","PC_PAGO_A_CONTRATADA invalido",' in k
        assert "UNICHAR" not in k and f'G{linha}<>"Nao"' not in k
        assert "INICIO_EFEITO ausente: PC" in k and 'ISERROR(SEARCH("PRECLUSO",' in k
        assert _ascii_balanceada(k)


def test_nenhum_criterio_pc_pago_sim_le_a_coluna_g_bruta():
    wb = _template()
    for aba, coord, formula in _formulas(wb):
        # EXACT(G,"Sim") so conta a grafia divergente (MEMORIA_RESULTADOS!T64).
        sem_contagem = formula.replace('EXACT(itens_PC!$G$2:$G$5001,"Sim")', "")
        assert '$G$2:$G$5001,"Sim"' not in sem_contagem, f"{aba}!{coord}"
    for aba, faixa in co._FAIXAS_PC_PAGO_SIM.items():
        for linha in wb[aba][faixa]:
            for celula in linha:
                assert '$AD$2:$AD$5001,"Sim"' in celula.value, f"{aba}!{celula.coordinate}"


def test_d44_continua_testando_so_presenca_em_g():
    d44 = _template()["MEMORIA_RESULTADOS"]["D44"].value
    assert '(itens_PC!$G$2:$G$5001<>"")' in d44
    assert '(itens_PC!$K$2:$K$5001<>"OK")' in d44


def test_avisos_de_grafia_e_destaque_em_g():
    ws = _template()["itens_PC"]
    assert ws["M31"].value == co._FORMULA_AVISO_GRAFIA_PC
    assert ws["M32"].value == co._FORMULA_AVISO_INVALIDO_PC
    assert {"M31:T31", "M32:T32"} <= {str(r) for r in ws.merged_cells.ranges}
    regras = {
        (str(cf.sqref), tuple(regra.formula)): regra.priority
        for cf in ws.conditional_formatting
        for regra in cf.rules
    }
    grafia = regras[("G2:G5001", (co._FORMULA_DESTAQUE_GRAFIA_PC,))]
    invalido = regras[("G2:G5001", (co._FORMULA_DESTAQUE_INVALIDO_PC,))]
    linha = regras[("A2:L5001", ('AND(OR($A2<>"",$B2<>"",$D2<>"",$G2<>""),$L2="")',))]
    assert grafia < linha and invalido < linha  # destaque de G acima das regras de linha
    assert ("A2:L5001", ('$L2="Nao"',)) in regras  # regras existentes preservadas


def test_dropdown_de_g_segue_so_sim_nao():
    wb = _gerada()
    ws = wb["itens_PC"]
    listas = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
    assert listas["G2:G5001"] == "OPCOES_SIM_NAO"
    assert [wb["parametros"]["T2"].value, wb["parametros"]["T3"].value] == ["Sim", "Nao"]
    assert wb.defined_names["OPCOES_SIM_NAO"].attr_text == "parametros!$T$2:$T$3"


def test_filtro_nativo_com_a_aba_protegida_e_ordenacao_bloqueada():
    assert _template()["itens_PC"].auto_filter.ref == co.FAIXA_FILTRO_ITENS_PC == "A1:L5001"
    ws = _gerada()["itens_PC"]
    assert ws.auto_filter.ref == "A1:L5001"
    assert ws.protection.sheet is True
    assert ws.protection.autoFilter is False  # False = filtrar permitido
    assert ws.protection.sort is True  # ordenar moveria formulas protegidas
    assert ws["A1"].protection.locked and ws["H2"].protection.locked
    assert not ws["G2"].protection.locked


def test_coluna_ad_oculta_na_copia_entregue():
    ws = _gerada()["itens_PC"]
    assert ws.column_dimensions["AD"].hidden
    assert ws.column_dimensions["V"].hidden  # V:AC seguem ocultas
    assert ws["AD2"].protection.locked


# ------------------------------------------------------ coerencia do metodo
def test_celulas_de_coerencia_e_textos_na_memoria():
    ws = _template()["MEMORIA_RESULTADOS"]
    for celula, (rotulo, formula) in co._CELULAS_COLETA_120.items():
        assert ws[f"S{celula[1:]}"].value == rotulo
        assert ws[celula].value == formula
        assert _ascii_balanceada(formula)
    for celula, texto in co._TEXTOS_COLETA_120.items():
        assert ws[celula].value == texto
    t58 = co._CELULAS_COLETA_120["T58"][1]
    assert t58 == '=IF(NOT(ISNUMBER($T$57)),0,IF($T$57=0,1,IF($T$54+$T$55+$T$56>$T$57,2,0)))'


def test_mensagens_citam_metodo_fonte_esperada_ausente_outra_fonte_e_acao():
    t61 = co._CELULAS_COLETA_120["T61"][1]
    for trecho in ("CONTROLE!$B$1", "$T$59", "$T$60", "$AG$46", "$AG$47"):
        assert trecho in t61
    textos = co._TEXTOS_COLETA_120
    assert textos["AG43"].startswith("REVISE O MÉTODO DE APURAÇÃO. Método selecionado")
    assert "Fonte esperada: aba" in textos["AG44"]
    assert "Foram encontrados dados em" in textos["AG45"]
    assert "Confirme o método da apuração antes de prosseguir" in textos["AG46"]
    assert "Não há dados informados para o método selecionado" in textos["AG47"]
    assert "O método selecionado prevalece" in textos["AG51"]


def test_bloqueio_esvazia_vta_e_retroativo_sem_trocar_de_fonte():
    wb = _template()
    mem, rd = wb["MEMORIA_RESULTADOS"], wb["RESULTADOS_DETALHE"]
    assert mem["B26"].value.startswith(
        '=IF(AND(B24<>"",B25<>""),"",IF($T$48>0,"",IF($T$58=1,"",')
    assert co._gate_vta_canonico(mem)  # gate de aditivos da 11.7 intacto
    assert mem["T40"].value.startswith('=IF($T$48>0,"",IF($T$58=1,"",')
    assert rd["D22"].value == '=IF(MEMORIA_RESULTADOS!$T$58=1,"",IFERROR(RETRO_OFICIAL,""))'
    assert rd["H14"].value.startswith('=IF(MEMORIA_RESULTADOS!$T$58=1,"REVISE",')
    assert rd["B3"].value.startswith('=IF(OR(MEMORIA_RESULTADOS!$T$58=1,$H$8="REVISE",')
    assert rd["J5"].value.endswith('+(MEMORIA_RESULTADOS!$T$58=1)')
    assert 'IF(MEMORIA_RESULTADOS!$T$58=1,MEMORIA_RESULTADOS!$T$61&" | ","")' in rd["B7"].value
    for aba, celula in co._TRECHOS_COERENCIA_120:
        formula = wb[aba][celula].value
        assert formula.count("(") == formula.count(")"), f"{aba}!{celula}"


def test_aviso_curto_ao_lado_do_metodo_e_b2_da_11_9_intocado():
    controle = _template()["CONTROLE"]
    assert controle["C1"].value == "=MEMORIA_RESULTADOS!$T$62"
    assert controle["C1"].protection.locked
    assert controle["B2"].value.startswith('=IFERROR(IF(NOT(ISNUMBER($B$3)),"",')
    assert not any(str(dv.sqref) == "B2" for dv in controle.data_validations.dataValidation)


def test_transformacoes_sao_idempotentes():
    wb = _template()
    for (aba, celula) in co._TRECHOS_COERENCIA_120:
        atual = wb[aba][celula].value
        assert co._formula_coerencia_120(aba, celula, atual) == atual


# ------------------------------------------------------------ leitor Python
def test_leitor_reconhece_so_sim_e_nao_com_tolerancia_de_grafia():
    assert [_pc_pago_canonico(v) for v in ("SIM", " sim ", "Não", "NÃO", "não", " nao ")] == [
        "Sim", "Sim", "Nao", "Nao", "Nao", "Nao",
    ]
    for invalido in ("S", "N", "1", "0", "x", "Talvez", "Pago"):
        assert _pc_pago_canonico(invalido) not in ("Sim", "Nao")


def test_leitor_nao_da_significado_a_s_ou_n():
    for texto in ("S", "N"):
        leitura = ler_masterfile_v10(_workbook_pc(texto))
        alertas = (leitura.get("itens_pc_v10") or {}).get("alertas") or []
        assert any("PC_PAGO_A_CONTRATADA invalido" in a for a in alertas), texto
