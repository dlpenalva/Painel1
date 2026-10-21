"""Issue #195 — Painel Executivo com o VTA canonico, nunca R$ 0,00 artificial.

O Painel Executivo e o Assistente Fiscal exibiam como "VTA" a referencia
historica do resumo do arquivo; no formato Coleta esse resumo nao existe e
`_to_float(None)` virava R$ 0,00. O VTA exibido passa a ser o canonico ja
apurado (memoria_por_ciclo.vta == resultado_consolidado.vta da pagina Valor
Global); a referencia historica e a conferencia sombra seguem separadas.
Todos os dados sao sinteticos (fixtures versionadas e cenarios gerados aqui).
"""
from __future__ import annotations

import copy
import io
import os
import warnings
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from _ciclo_em_execucao import garantir_aba_ciclo_em_execucao
from _coleta_oficial import obter_coleta_oficial_bytes
from _coleta_reajuste_documentos import processar_coleta_oficial_runtime
from _leitor_masterfile_v10 import ler_masterfile_v10
from _objeto_processo_reajuste import consumidor_do_objeto
from _painel_executivo import (
    montar_painel_executivo,
    vincular_vta_canonico,
    vta_canonico_da_memoria,
)
from _resultado_consolidado import ORIGEM_VTA_CANONICA, ORIGEM_VTA_INDISPONIVEL

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "coletas_compat"
TOL = 0.005
_CACHE: dict[str, tuple] = {}


def _processar(nome: str, conteudo: bytes) -> tuple[dict, dict, dict]:
    if nome not in _CACHE:
        resultado, _ = processar_coleta_oficial_runtime(conteudo)
        leitura = ler_masterfile_v10(conteudo, exigir_modelo_oficial=True)
        _CACHE[nome] = (resultado, leitura, montar_painel_executivo(leitura))
    return _CACHE[nome]


def _fixture(nome: str) -> tuple[dict, dict, dict]:
    return _processar(nome, (FIXTURES / f"{nome}.xlsx").read_bytes())


def _coleta_126_sem_cache() -> bytes:
    """Coleta 12.6 preenchida sem passar pelo Excel: sem cache de formulas."""
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    ctl = wb["CONTROLE"]
    ctl["B1"], ctl["B2"], ctl["B3"] = "Financeiro (Mensalidade)", "C1", date(2025, 10, 15)
    par = wb["parametros"]
    par["A2"], par["B2"], par["C2"], par["D2"], par["I2"] = (
        "Nao", "C0", date(2024, 8, 1), date(2025, 7, 31), date(2024, 8, 1)
    )
    par["A3"], par["B3"], par["C3"], par["D3"], par["E3"] = (
        "Sim", "C1", date(2025, 8, 1), date(2026, 7, 31), 0.05
    )
    par["H3"], par["I3"] = date(2025, 10, 1), date(2025, 10, 15)
    fin = wb["financeiro"]
    for linha, (competencia, valor, efeito) in enumerate(
        [(date(2025, 7, 15), 100_000.0, "Nao"), (date(2025, 8, 15), 10_000.0, "Sim")], 2
    ):
        fin[f"A{linha}"], fin[f"C{linha}"], fin[f"G{linha}"] = competencia, valor, efeito
    rem = wb["itens_Remanesc"]
    rem["A2"], rem["B2"], rem["C2"], rem["E2"] = "ITEM-01", 100.0, 2000.0, 1500.0
    garantir_aba_ciclo_em_execucao(wb)
    cee = wb["CICLO_EM_EXECUCAO"]
    cee["D5"], cee["C13"] = date(2025, 10, 15), 1500.0
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# --------------------------------------------------------------------------- #
# Fonte canonica: valor valido, zero real e indisponivel sao distintos.
# --------------------------------------------------------------------------- #
def test_vta_canonico_distingue_valor_zero_real_e_indisponivel():
    valido = vta_canonico_da_memoria({"valor_total_atualizado": 1234.56, "metodo": "financeiro"})
    assert valido["valor_total_atualizado"] == 1234.56
    assert valido["vta_disponivel"] is True and valido["origem_vta"] == ORIGEM_VTA_CANONICA
    assert valido["motivo_vta_indisponivel"] is None

    zero = vta_canonico_da_memoria({"valor_total_atualizado": 0.0, "metodo": "financeiro"})
    assert zero["valor_total_atualizado"] == 0.0 and zero["vta_disponivel"] is True

    for ausente in (None, {}, {"valor_total_atualizado": None, "motivo": "sem potencial"},
                    {"valor_total_atualizado": ""}, {"valor_total_atualizado": True}):
        saida = vta_canonico_da_memoria(ausente)
        assert saida["valor_total_atualizado"] is None
        assert saida["vta_disponivel"] is False
        assert saida["origem_vta"] == ORIGEM_VTA_INDISPONIVEL
        assert saida["motivo_vta_indisponivel"]
    assert vta_canonico_da_memoria(
        {"valor_total_atualizado": None, "motivo": "sem potencial"}
    )["motivo_vta_indisponivel"] == "sem potencial"


def test_vincular_preserva_zero_real_no_painel_e_no_assistente():
    painel, assistente = {"vta": {"oficial": None}}, {"vta": {"valor": None}}
    vincular_vta_canonico(painel, assistente, {"valor_total_atualizado": 0.0, "metodo": "pc"})
    for alvo in (painel["vta"], assistente["vta"]):
        assert alvo["valor_total_atualizado"] == 0.0 and alvo["vta_disponivel"] is True
        assert alvo["metodo_vta"] == "pc"


# --------------------------------------------------------------------------- #
# Tres metodos + linhagem legada: Painel == Assistente == Valor Global.
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("nome, metodo", [
    ("coleta_11_financeiro", "financeiro"),
    ("coleta_11_pc", "pc"),
    ("coleta_11_consumidos", "consumidos"),
    ("pre11_l1_financeiro", "financeiro"),
])
def test_painel_assistente_e_valor_global_convergem_no_vta(nome, metodo):
    resultado, leitura, painel = _fixture(nome)
    vta_global = resultado["resultado_consolidado"]["vta"]
    memoria = leitura["objeto_processo"]["memoria_por_ciclo"]["vta"]
    assert vta_global is not None and vta_global > 0
    assert memoria["metodo"] == metodo

    pv = painel["vta"]
    assert pv["valor_total_atualizado"] == pytest.approx(vta_global, abs=TOL)
    assert pv["valor_total_atualizado"] == pytest.approx(memoria["valor_total_atualizado"], abs=TOL)
    assert pv["vta_disponivel"] is True and pv["origem_vta"] == ORIGEM_VTA_CANONICA
    assert pv["metodo_vta"] == metodo

    assistente = consumidor_do_objeto(leitura, "assistente_fiscal")["vta"]
    assert assistente["valor_total_atualizado"] == pytest.approx(vta_global, abs=TOL)
    assert assistente["vta_disponivel"] is True


@pytest.mark.parametrize("nome", ["coleta_11_financeiro", "coleta_11_consumidos"])
def test_coleta_sem_resumo_nao_exibe_zero_artificial(nome):
    _, leitura, painel = _fixture(nome)
    assert not (leitura.get("resumo") or {}).get("valor_total_atualizado")
    pv = painel["vta"]
    # Referencia historica e conferencia sombra sem base: indisponiveis, nunca 0,00.
    assert pv["oficial"] is None
    assert pv["metodologia_sombra"] is None and pv["diferenca_sombra"] is None
    assert pv["sombra_aplicavel"] is False and pv["motivo_sombra"]
    assistente = consumidor_do_objeto(leitura, "assistente_fiscal")
    assert assistente["vta"]["valor"] is None
    evidencias = " ".join(assistente["dossie_decisao"].get("evidencias") or [])
    assert "(nao e VTA): R$ 0,00" not in evidencias


def test_pc_mantem_referencia_historica_separada_do_vta_com_potencial():
    resultado, leitura, painel = _fixture("coleta_11_pc")
    pv = painel["vta"]
    memoria = leitura["objeto_processo"]["memoria_por_ciclo"]["vta"]
    # O VTA do PC e a composicao canonica (inclui o retroativo potencial
    # prudencial); a soma de itens_PC do resumo segue como referencia.
    assert "retroativo_potencial_vta" in memoria
    assert pv["valor_total_atualizado"] == pytest.approx(
        resultado["resultado_consolidado"]["vta"], abs=TOL
    )
    assert pv["sombra_aplicavel"] is True
    assert pv["oficial"] == pytest.approx(
        leitura["resumo"]["valor_total_atualizado"], abs=TOL
    )
    assert pv["oficial"] != pytest.approx(pv["valor_total_atualizado"], abs=TOL)


# --------------------------------------------------------------------------- #
# VTA indisponivel e informacao insuficiente: "Indisponivel", bloqueio mantido.
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("nome", ["template_vazio", "coleta_126_sem_cache"])
def test_vta_indisponivel_fica_indisponivel_e_bloqueio_preservado(nome):
    conteudo = (
        (ROOT / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx").read_bytes()
        if nome == "template_vazio" else _coleta_126_sem_cache()
    )
    resultado, leitura, painel = _processar(nome, conteudo)
    assert resultado["resultado_consolidado"]["vta"] is None
    assert resultado["formalizacao_bloqueada"] is True
    pv = painel["vta"]
    assert pv["valor_total_atualizado"] is None and pv["vta_disponivel"] is False
    assert pv["origem_vta"] == ORIGEM_VTA_INDISPONIVEL and pv["motivo_vta_indisponivel"]
    assert pv["oficial"] is None and pv["metodologia_sombra"] is None
    assistente = consumidor_do_objeto(leitura, "assistente_fiscal")["vta"]
    assert assistente["valor_total_atualizado"] is None
    assert assistente["vta_disponivel"] is False


# --------------------------------------------------------------------------- #
# Masterfile antigo: com resumo a conferencia sombra e preservada; sem resumo
# ela fica indisponivel. O VTA exibido segue o canonico nos dois casos.
# --------------------------------------------------------------------------- #
def _leitura_masterfile_antigo(resumo: dict | None) -> dict:
    _, leitura, _ = _fixture("coleta_11_financeiro")
    leitura = copy.deepcopy(leitura)
    leitura["resumo"] = resumo or {}
    return leitura


def test_masterfile_antigo_com_resumo_preserva_conferencia_sombra():
    leitura = _leitura_masterfile_antigo({"valor_total_atualizado": 1_000_000.0})
    marco = (leitura.get("controle") or {}).get("ciclo_vigente") or "C1"
    pv = montar_painel_executivo(leitura, marco=marco)["vta"]
    assert pv["sombra_aplicavel"] is True and pv["motivo_sombra"] is None
    assert pv["oficial"] == pytest.approx(1_000_000.0)
    assert pv["metodologia_sombra"] is not None
    assert pv["diferenca_sombra"] == pytest.approx(pv["metodologia_sombra"] - 1_000_000.0, abs=TOL)
    memoria = leitura["objeto_processo"]["memoria_por_ciclo"]["vta"]
    assert pv["valor_total_atualizado"] == pytest.approx(memoria["valor_total_atualizado"], abs=TOL)


def test_masterfile_antigo_sem_resumo_nao_fabrica_base_zero():
    leitura = _leitura_masterfile_antigo(None)
    marco = (leitura.get("controle") or {}).get("ciclo_vigente") or "C1"
    pv = montar_painel_executivo(leitura, marco=marco)["vta"]
    assert pv["oficial"] is None and pv["diferenca_sombra"] is None
    assert pv["sombra_aplicavel"] is False and pv["motivo_sombra"]


# --------------------------------------------------------------------------- #
# Apresentacao: o rotulo VTA so mostra o canonico; nada de "VTA (oficial)".
# --------------------------------------------------------------------------- #
def test_telas_rotulam_vta_canonico_e_referencia_historica_separadas():
    painel = (ROOT / "_painel_executivo.py").read_text(encoding="utf-8")
    assistente = (ROOT / "_assistente_fiscal.py").read_text(encoding="utf-8")
    assert '"VTA (oficial)"' not in painel
    assert '"Valor Total Atualizado (VTA)"' in painel
    assert "Valor histórico do arquivo (não é VTA)" in painel
    assert "**Valor atualizado do contrato (VTA):**" not in assistente
    assert "**Valor Total Atualizado (VTA):** " in assistente
    assert "Valor histórico do arquivo (não é VTA)" in assistente


# --------------------------------------------------------------------------- #
# Opcional (fora do CI): Coleta 12.5 real compatibilizada pelo PR #194.
# --------------------------------------------------------------------------- #
@pytest.mark.skipif(
    not os.environ.get("COLETA_125_REAL_XLSX"),
    reason="arquivo real: defina COLETA_125_REAL_XLSX",
)
def test_coleta_125_real_painel_converge_com_valor_global():
    origem = Path(os.environ["COLETA_125_REAL_XLSX"])
    resultado, leitura, painel = _processar("real_125", origem.read_bytes())
    t25 = load_workbook(origem, data_only=True, read_only=True)["MEMORIA_RESULTADOS"]["T25"].value
    assert resultado["aviso_compatibilidade_coleta_125"]
    assert resultado["resultado_consolidado"]["vta"] == pytest.approx(t25, abs=TOL)
    assert painel["vta"]["valor_total_atualizado"] == pytest.approx(t25, abs=TOL)
    assert consumidor_do_objeto(leitura, "assistente_fiscal")["vta"][
        "valor_total_atualizado"
    ] == pytest.approx(t25, abs=TOL)
