"""AJUSTES-XLS-UX pos-PR #174 — Frente 1 — itens_Consumidos!X:AG: bloco opcional com identidade visual;
so Z/AA sao entradas; formulas e dropdown intactos.

Testes de Excel real rodam com RUN_EXCEL_INTEGRATION=1.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tools"))

import aplicar_ajustes_xls_ux_pos174 as ux  # noqa: E402
from _coleta_oficial import (  # noqa: E402
    gerar_coleta_oficial_preenchida,
    obter_coleta_oficial_bytes,
)
from _memoria_calculo import (  # noqa: E402
    FILL_FRONTEIRA_IST,
    escrever_memoria_calculo,
    ler_memoria_calculo,
)

TEMPLATE = RAIZ / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
COLUNAS_JR = ("J", "K", "L", "M", "N", "O", "P", "Q", "R")

com = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1",
)


@pytest.fixture(scope="module")
def template():
    wb = load_workbook(TEMPLATE)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def coleta_branca():
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    yield wb
    wb.close()


def _rgb(cell) -> str | None:
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


# ------------------------------------------------------------------ payloads
PERCENTUAIS = {1: 0.0512, 2: 0.0374, 3: 0.0289, 4: 0.0315}


def _payload(numeros: list[int]) -> dict:
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2021, 1, 1),
        "data_corte": date(2021 + max(numeros), 12, 31),
        "ciclos": [
            {
                "ciclo": f"C{n}",
                "data_inicio": date(2021 + n, 1, 1),
                "percentual": PERCENTUAIS[n],
                "inicio_efeito_financeiro": date(2021 + n, 1, 1),
                "situacao_aplicada": "✅ TEMPESTIVO",
                "data_pedido": date(2021 + n, 1, 10),
            }
            for n in numeros
        ],
    }


# =================================================================== FRENTE 1
FORMULAS_AJUSTE = ("Y", "AB", "AC", "AD", "AE", "AF", "AG")


def test_f1_bloco_x_ag_entradas_e_automaticos(template):
    ws = template["itens_Consumidos"]
    for linha in range(2, 7):
        for col in ux.COLUNAS_ENTRADA:
            assert ws[f"{col}{linha}"].value is None
            assert _rgb(ws[f"{col}{linha}"]) == ux.COR_ENTRADA
        for col in FORMULAS_AJUSTE:
            assert str(ws[f"{col}{linha}"].value).startswith("=")
            assert _rgb(ws[f"{col}{linha}"]) == ux.COR_AUTOMATICA
        assert ws[f"X{linha}"].value == f"C{linha - 2}"
    for col in ux.COLUNAS_ENTRADA:
        assert _rgb(ws[f"{col}1"]) == ux.COR_ENTRADA_CAB
    assert ws["Z1"].value == "AJUSTE_TIPO" and ws["AA1"].value == "AJUSTE_VALOR_INFORMADO"


def test_f1_orientacao_legivel(template):
    ws = template["itens_Consumidos"]
    assert ws["X8"].value == ux.TITULO_AJUSTE
    assert ws["X9"].value == ux.ENTRADAS_AJUSTE + ux.TEXTO_X9_BASE
    mescladas = {str(m) for m in ws.merged_cells.ranges}
    for linha in range(8, 14):
        assert f"X{linha}:AG{linha}" in mescladas
        assert ws[f"X{linha}"].alignment.wrap_text
        assert (ws[f"X{linha}"].font.sz or 11) >= 10


def test_f1_formulas_e_dropdown_na_coleta_gerada(coleta_branca, template):
    ws_t, ws_g = template["itens_Consumidos"], coleta_branca["itens_Consumidos"]
    for linha in range(2, 7):
        for col in FORMULAS_AJUSTE:
            assert ws_g[f"{col}{linha}"].value == ws_t[f"{col}{linha}"].value
    dvs = [dv for dv in ws_g.data_validations.dataValidation if str(dv.sqref) == "Z2:Z6"]
    assert len(dvs) == 1 and dvs[0].formula1 == '"Valor pago,Glosa"'

