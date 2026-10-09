from __future__ import annotations

import gc
import os
import time
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from _coleta_oficial import gerar_coleta_oficial_preenchida


pytestmark = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar o Excel COM",
)


def _dados() -> dict:
    return {
        "origem": "Validacao Coleta 11.6",
        "indice": "IST",
        "data_base_original": "01/01/2023",
        "data_corte": date(2026, 8, 31),
        "ciclos": [{
            "ciclo": "C1",
            "data_inicio": date(2024, 1, 1),
            "data_fim": date(2024, 12, 31),
            "data_pedido": date(2024, 1, 1),
            "financeiro_inicio": date(2024, 1, 1),
            "percentual_aplicado": 0.0381,
            "objeto_analise_atual": True,
        }],
    }


def _enderecos_formulas(conteudo: bytes) -> dict[str, set[str]]:
    wb = load_workbook(BytesIO(conteudo), data_only=False, read_only=False)
    try:
        return {
            ws.title: {
                celula.coordinate
                for linha in ws.iter_rows()
                for celula in linha
                if isinstance(celula.value, str) and celula.value.startswith("=")
            }
            for ws in wb.worksheets
        }
    finally:
        wb.close()


def _tentar(pythoncom, acao):
    ultimo = None
    for _ in range(30):
        try:
            return acao()
        except Exception as exc:  # pragma: no cover - somente Excel COM
            ultimo = exc
            codigo = getattr(exc, "hresult", None)
            if codigo is None and getattr(exc, "args", ()):
                codigo = exc.args[0]
            if codigo != -2147418111:
                raise
            pythoncom.PumpWaitingMessages()
            time.sleep(0.2)
    raise ultimo


def _valor_nome(livro, nome: str):
    return livro.Names(nome).RefersToRange.Value2


def _preencher_cenario(livro) -> None:
    from pywintypes import Time

    controle = livro.Worksheets("CONTROLE")
    parametros = livro.Worksheets("parametros")
    itens = livro.Worksheets("itens_Remanesc")
    aditivos = livro.Worksheets("aditivos")
    ciclo = livro.Worksheets("CICLO_EM_EXECUCAO")
    financeiro = livro.Worksheets("financeiro")

    controle.Range("B1").Value = "Financeiro (Mensalidade)"
    controle.Range("B3").Value = Time(datetime(2026, 8, 31))
    parametros.Range("I4").Value = Time(datetime(2025, 1, 1))
    parametros.Range("I5").Value = Time(datetime(2026, 1, 1))
    ciclo.Range("D5").Value = Time(datetime(2026, 8, 31))

    financeiro.Range("A2").Value = Time(datetime(2023, 1, 1))
    financeiro.Range("C2").Value = 1000
    financeiro.Range("G2").Value = "Nao"

    # C3 presente; fallback C2; fallback C1; sem qualquer referencia.
    entradas = (
        (2, "1.1", 1000, 10, 950, 900, 850),
        (3, "2.1", 1000, 10, 950, 800, None),
        (4, "3.1", 1000, 10, 1000, None, None),
        (5, "4.1", None, 10, None, None, None),
        (6, "N001", None, 10, None, None, None),
        (7, "N002", None, 10, None, None, None),
        (8, "5.1", 500, 10, 450, None, None),
        (9, "N003", None, 10, None, None, None),
    )
    for linha, item, base, vu, c1, c2, c3 in entradas:
        itens.Range(f"A{linha}").Value = item
        # B dos Nxxx conserva a formula canonica de base zero automatica.
        if not item.startswith("N") and base is not None:
            itens.Range(f"B{linha}").Value = base
        itens.Range(f"C{linha}").Value = vu
        itens.Range(f"E{linha}").Value = c1
        itens.Range(f"G{linha}").Value = c2
        itens.Range(f"I{linha}").Value = c3

    eventos = (
        (2, "3.1", datetime(2023, 12, 31), "Acrescimo", 80, "Nao", None),
        (3, "3.1", datetime(2025, 5, 10), "Acrescimo", 200, "Nao", None),
        (4, "3.1", datetime(2026, 9, 1), "Acrescimo", 300, "Nao", None),
        (5, "N001", datetime(2026, 2, 1), "Acréscimo - novo item", 100, "Sim", "C0"),
        (6, "N002", datetime(2026, 3, 1), "Acréscimo - novo item", 100, "Sim", "C1"),
        (7, "5.1", datetime(2026, 4, 1), "Acrescimo", 100, "Sim", None),
        (8, "5.1", datetime(2026, 5, 1), "Acrescimo", 50, "Nao", None),
        (9, "N003", datetime(2026, 6, 1), "Acréscimo - novo item", 10, "Sim", "C2"),
    )
    for linha, item, data, tipo, quantidade, aplicar, base_vu in eventos:
        aditivos.Range(f"A{linha}").Value = item
        aditivos.Range(f"B{linha}").Value = Time(data)
        aditivos.Range(f"D{linha}").Value = tipo
        aditivos.Range(f"E{linha}").Value = quantidade
        aditivos.Range(f"H{linha}").Value = aplicar
        aditivos.Range(f"K{linha}").Value = "Sim"
        aditivos.Range(f"N{linha}").Value = base_vu


def _inspecionar_hashes(livro, pythoncom) -> list[str]:
    achados: list[str] = []
    limites = {
        "itens_Remanesc": (12, 29),
        "aditivos": (12, 15),
        "CICLO_EM_EXECUCAO": (30, 18),
        "RESULTADOS": (90, 8),
    }
    for aba, (ultima_linha, ultima_coluna) in limites.items():
        ws = livro.Worksheets(aba)
        for linha in range(1, ultima_linha + 1):
            for coluna in range(1, ultima_coluna + 1):
                texto = str(_tentar(
                    pythoncom,
                    lambda: ws.Cells(linha, coluna).Text,
                ) or "")
                if "####" in texto:
                    achados.append(f"{aba}!{ws.Cells(linha, coluna).Address}")
    return achados


def _contagem_formulas_excel(livro) -> dict[str, int]:
    contagens: dict[str, int] = {}
    for ws in livro.Worksheets:
        try:
            contagens[ws.Name] = int(ws.Cells.SpecialCells(-4123).CountLarge)
        except Exception:
            contagens[ws.Name] = 0
    return contagens


def test_coleta_116_no_excel_real(tmp_path: Path):
    import pythoncom
    import win32com.client

    caminho = tmp_path / "Coleta_11_6_validacao.xlsx"
    original = gerar_coleta_oficial_preenchida(_dados())
    caminho.write_bytes(original)
    formulas_antes = _enderecos_formulas(original)

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    livro = None
    try:
        livro = _tentar(
            pythoncom,
            lambda: excel.Workbooks.Open(
                str(caminho.resolve()), UpdateLinks=0, ReadOnly=False, CorruptLoad=0
            ),
        )
        time.sleep(1.0)
        _tentar(pythoncom, lambda: _preencher_cenario(livro))
        _tentar(pythoncom, excel.CalculateFullRebuild)

        ciclo = livro.Worksheets("CICLO_EM_EXECUCAO")
        aditivos = livro.Worksheets("aditivos")
        itens = livro.Worksheets("itens_Remanesc")
        historico = livro.Worksheets("historico_VU")
        posicao = livro.Worksheets("posicao_contratual")

        nomes_ab = (
            "VTA_FINAL",
            "VTA_SEM_POTENCIAL",
            "RETRO_OFICIAL",
            "RETROATIVO_POTENCIAL_APURADO",
            "RETROATIVO_POTENCIAL_VTA",
        )
        economico_antes = {nome: _valor_nome(livro, nome) for nome in nomes_ab}

        for linha, valor in enumerate((800, 700, 650, 0, 60, 100, 600, 10), start=13):
            ciclo.Range(f"C{linha}").Value = valor
        _tentar(pythoncom, excel.CalculateFullRebuild)
        economico_depois = {nome: _valor_nome(livro, nome) for nome in nomes_ab}

        assert ciclo.Range("C3").Value == "C3"
        assert ciclo.Range("P13").Value == "C3"
        assert ciclo.Range("P14").Value == "C2"
        assert ciclo.Range("P15").Value == "C1"
        assert ciclo.Range("B15").Value == 1000
        assert ciclo.Range("I15").Value == 200
        assert ciclo.Range("D15").Value == 550
        assert ciclo.Range("B16").Value in (None, "")
        assert ciclo.Range("D16").Value in (None, "")
        assert ciclo.Range("C16").Value == 0
        assert "CONSUMO NAO CALCULAVEL" in str(ciclo.Range("K16").Value).upper()
        assert bool(ciclo.Range("C16").Validation.Value)
        assert ciclo.Range("B17").Value == 0
        assert ciclo.Range("I17").Value == 100
        assert ciclo.Range("D17").Value == 40

        assert aditivos.Range("I5").Value == pytest.approx(1.0381)
        assert aditivos.Range("J5").Value / aditivos.Range("L5").Value == pytest.approx(10.38)
        assert aditivos.Range("I6").Value == pytest.approx(1.0)
        assert aditivos.Range("J6").Value / aditivos.Range("L6").Value == pytest.approx(10.0)
        assert aditivos.Range("O7").Value == 0
        assert aditivos.Range("I7").Value == pytest.approx(1.0381)
        assert aditivos.Range("J8").Value == pytest.approx(500.0)
        assert aditivos.Range("I9").Value in (None, "")
        assert "BASE_VU_INVALIDA" in str(aditivos.Range("M9").Value)

        # Nascimento C3: nada e criado em C1/C2, embora o VU de C3 receba C1.
        assert posicao.Range("Y6").Value == 3
        assert historico.Range("D6").Value in (None, "")
        assert historico.Range("E6").Value in (None, "")

        assert str(aditivos.Range("D2").Validation.Formula1).replace(";", ",") == (
            "Acrescimo,Acréscimo - novo item,Supressao"
        )
        assert str(aditivos.Range("H2").Validation.Formula1).replace(";", ",") == "Sim,Nao"
        assert str(aditivos.Range("N2").Validation.Formula1).replace(";", ",") == "C0,C1,C2,C3,C4"

        verde_escuro = 0 + (97 << 8) + (0 << 16)
        assert int(itens.Range("A6").DisplayFormat.Font.Color) == verde_escuro
        assert int(itens.Range("A6").DisplayFormat.Interior.Color) == int(
            itens.Range("A6").Interior.Color
        )
        assert economico_depois == economico_antes
        assert not _inspecionar_hashes(livro, pythoncom)
        formulas_excel_antes = _contagem_formulas_excel(livro)

        _tentar(pythoncom, livro.Save)
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = None
        livro = _tentar(
            pythoncom,
            lambda: excel.Workbooks.Open(
                str(caminho.resolve()), UpdateLinks=0, ReadOnly=False, CorruptLoad=0
            ),
        )
        _tentar(pythoncom, excel.CalculateFullRebuild)
        assert livro.Worksheets("CICLO_EM_EXECUCAO").Range("D15").Value == 550
        assert not _inspecionar_hashes(livro, pythoncom)
        assert _contagem_formulas_excel(livro) == formulas_excel_antes
        _tentar(pythoncom, livro.Save)
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = None
    finally:
        if livro is not None:
            try:
                livro.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()

    final = caminho.read_bytes()
    formulas_depois = _enderecos_formulas(final)
    substituidas_por_entrada = {"B2", "B3", "B4", "B8"}
    assert formulas_antes["itens_Remanesc"] - formulas_depois["itens_Remanesc"] == (
        substituidas_por_entrada
    )
    assert not (formulas_depois["itens_Remanesc"] - formulas_antes["itens_Remanesc"])
    for aba in formulas_antes.keys() - {"itens_Remanesc"}:
        assert formulas_depois[aba] == formulas_antes[aba]
    with zipfile.ZipFile(BytesIO(final)) as pacote:
        assert not any(b"repairLoad" in pacote.read(nome) for nome in pacote.namelist())
        assert any(
            b"x14:conditionalFormattings" in pacote.read(nome)
            for nome in pacote.namelist()
            if nome.startswith("xl/worksheets/")
        )


def test_coleta_116_ciclo_em_execucao_prevalece_sem_data_exata(tmp_path: Path):
    """Regressao: analise em C1 com o contrato em C3 deixa parametros!I5 vazio.

    A fotografia C3 informada pelo fiscal continua prevalecendo (regra 11.5),
    com o aditivo do dia da abertura somado uma unica vez (Regra B) e N001
    nascido no ciclo partindo de zero. A9 permanece calculada e alimenta a
    cadeia oficial (MEMORIA_RESULTADOS!W49 = 1), exatamente como na 11.5.
    """
    import pythoncom
    import win32com.client
    from pywintypes import Time

    caminho = tmp_path / "Coleta_11_6_ciclo_atual.xlsx"
    caminho.write_bytes(gerar_coleta_oficial_preenchida(_dados()))

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    livro = None
    try:
        livro = _tentar(
            pythoncom,
            lambda: excel.Workbooks.Open(
                str(caminho.resolve()), UpdateLinks=0, ReadOnly=False, CorruptLoad=0
            ),
        )
        time.sleep(1.0)

        def preencher():
            livro.Worksheets("CONTROLE").Range("B1").Value = "Financeiro (Mensalidade)"
            livro.Worksheets("CONTROLE").Range("B3").Value = Time(datetime(2026, 8, 31))
            itens = livro.Worksheets("itens_Remanesc")
            for linha, item, base, c1, c3 in (
                (2, "1.1", 1000, 950, 850),
                (3, "N001", None, None, None),
            ):
                itens.Range(f"A{linha}").Value = item
                if base is not None:
                    itens.Range(f"B{linha}").Value = base
                itens.Range(f"C{linha}").Value = 10
                itens.Range(f"E{linha}").Value = c1
                itens.Range(f"I{linha}").Value = c3
            aditivos = livro.Worksheets("aditivos")
            for linha, item, data, tipo, base_vu in (
                (2, "1.1", datetime(2026, 1, 1), "Acrescimo", None),
                (3, "N001", datetime(2026, 2, 1), "Acréscimo - novo item", "C0"),
            ):
                aditivos.Range(f"A{linha}").Value = item
                aditivos.Range(f"B{linha}").Value = Time(data)
                aditivos.Range(f"D{linha}").Value = tipo
                aditivos.Range(f"E{linha}").Value = 100
                aditivos.Range(f"H{linha}").Value = "Nao"
                aditivos.Range(f"K{linha}").Value = "Sim"
                aditivos.Range(f"N{linha}").Value = base_vu
            livro.Worksheets("CICLO_EM_EXECUCAO").Range("D5").Value = Time(
                datetime(2026, 8, 31)
            )

        _tentar(pythoncom, preencher)
        _tentar(pythoncom, excel.CalculateFullRebuild)
        ciclo = livro.Worksheets("CICLO_EM_EXECUCAO")
        ciclo.Range("C13").Value = 900
        ciclo.Range("C14").Value = 60
        _tentar(pythoncom, excel.CalculateFullRebuild)

        parametros = livro.Worksheets("parametros")
        assert parametros.Range("I5").Value in (None, "")
        assert ciclo.Range("C3").Value == "C3"
        assert str(ciclo.Range("A13").Value) == "1.1"
        assert ciclo.Range("P13").Value == "C3"
        assert ciclo.Range("B13").Value == 950  # 850 + aditivo do dia da abertura
        assert ciclo.Range("I13").Value == 0
        assert ciclo.Range("D13").Value == 50
        assert ciclo.Range("K13").Value == "OK"
        assert ciclo.Range("A14").Value == "N001"
        assert ciclo.Range("P14").Value == "C3"
        assert ciclo.Range("B14").Value == 0
        assert ciclo.Range("I14").Value == 100
        assert ciclo.Range("D14").Value == 40
        assert isinstance(ciclo.Range("A9").Value, float)
        assert livro.Worksheets("MEMORIA_RESULTADOS").Range("W49").Value == 1
        _tentar(pythoncom, lambda: livro.Close(SaveChanges=False))
        livro = None
    finally:
        if livro is not None:
            try:
                livro.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()
