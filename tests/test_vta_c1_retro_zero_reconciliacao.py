from __future__ import annotations

import io
import os
import time
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from _ciclo_em_execucao import garantir_aba_ciclo_em_execucao
from _coleta_oficial import obter_coleta_oficial_bytes

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
TOL = 0.005


def test_formulas_distinguem_execucao_de_efeito_financeiro():
    wb = load_workbook(TEMPLATE, data_only=False)
    mem = wb["MEMORIA_RESULTADOS"]
    for ciclo in range(5):
        formula = str(mem[f"B{10 + ciclo}"].value)
        condicao, soma = formula.split('")=0', 1)
        assert f'financeiro!$B$2:$B$73,"c{ciclo}"' in condicao
        assert 'financeiro!$C$2:$C$73,">0' in condicao
        assert "financeiro!$G$2:$G$73" not in condicao
        assert "financeiro!$G$2:$G$73" in soma
    assert "$B$4=\"Financeiro\"" in str(mem["W50"].value)
    assert "INDEX(parametros!$I$2:$I$6,$W$46+1)" in str(mem["W67"].value)


def test_abas_tecnicas_ocultas_em_todos_os_gerados():
    for origem in (TEMPLATE, io.BytesIO(obter_coleta_oficial_bytes())):
        wb = load_workbook(origem, data_only=False)
        assert wb["MEMORIA_RESULTADOS"].sheet_state == "hidden"
        assert wb["RESULTADOS_DETALHE"].sheet_state == "hidden"
        assert wb["RESULTADOS"].sheet_state == "visible"


def _cenario(destino: Path, pagamentos: list[tuple[date, float, str]]) -> Path:
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    ctl = wb["CONTROLE"]
    ctl["B1"], ctl["B2"], ctl["B3"] = (
        "Financeiro (Mensalidade)", "C1", date(2026, 10, 16)
    )
    par = wb["parametros"]
    par["A2"], par["B2"], par["C2"], par["D2"], par["I2"] = (
        "Nao", "C0", date(2025, 8, 1), date(2026, 7, 31), date(2025, 8, 1)
    )
    par["A3"], par["B3"], par["C3"], par["D3"], par["E3"] = (
        "Sim", "C1", date(2026, 8, 1), date(2027, 7, 31), 0.0444
    )
    par["H3"], par["I3"] = date(2026, 10, 1), date(2026, 10, 16)
    fin = wb["financeiro"]
    linhas = [(date(2026, 7, 15), 857_200.0, "Nao"), *pagamentos]
    for linha, (competencia, valor, efeito) in enumerate(linhas, 2):
        fin[f"A{linha}"], fin[f"C{linha}"], fin[f"G{linha}"] = (
            competencia, valor, efeito
        )
    rem = wb["itens_Remanesc"]
    rem["A2"], rem["B2"], rem["C2"], rem["E2"] = (
        "ITEM-01", 3607.2, 2000.0, 3092.88
    )
    garantir_aba_ciclo_em_execucao(wb)
    cee = wb["CICLO_EM_EXECUCAO"]
    cee["D5"], cee["C13"] = date(2026, 10, 16), 3092.88
    wb.save(destino)
    return destino


@pytest.fixture
def excel():
    client = pytest.importorskip("win32com.client")
    pythoncom = pytest.importorskip("pythoncom")
    pythoncom.CoInitialize()
    app = client.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    yield app
    try:
        app.Quit()
    except Exception:
        pass
    pythoncom.CoUninitialize()


def _recalcular(excel) -> None:
    import pythoncom

    ultimo = None
    for _ in range(8):
        try:
            excel.CalculateFullRebuild()
            return
        except Exception as exc:
            ultimo = exc
            pythoncom.PumpWaitingMessages()
            time.sleep(0.5)
    raise ultimo


def _valores(excel, caminho: Path) -> dict[str, object]:
    book = excel.Workbooks.Open(str(caminho), UpdateLinks=0, ReadOnly=False)
    try:
        _recalcular(excel)
        mem = book.Worksheets("MEMORIA_RESULTADOS")
        out = {
            c: mem.Range(c).Value2
            for c in ("B11", "B16", "B26", "T21", "T23", "T25", "W48", "W50", "W51", "W52")
        }
        book.Save()
        return out
    finally:
        book.Close(SaveChanges=False)


com = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1",
)


@com
def test_c1_execucao_antes_dos_efeitos_retro_zero_e_vta(excel, tmp_path):
    arq = _cenario(tmp_path / "c1_zero.xlsx", [
        (date(2026, 8, 15), 85_720.0, "Nao"),
        (date(2026, 9, 15), 85_720.0, "Nao"),
    ])
    v = _valores(excel, arq)
    assert float(v["B11"]) == pytest.approx(0.0, abs=TOL)
    assert float(v["B16"]) == pytest.approx(0.0, abs=TOL)
    assert float(v["T21"]) == pytest.approx(1_028_640.0, abs=TOL)
    assert float(v["T23"]) == pytest.approx(6_460_407.74, abs=TOL)
    assert float(v["T25"]) == pytest.approx(7_489_047.74, abs=TOL)
    assert float(v["B26"]) == pytest.approx(7_489_047.74, abs=TOL)
    assert float(v["W48"]) == pytest.approx(7_489_047.74, abs=TOL)
    assert float(v["W50"]) == pytest.approx(7_489_047.74, abs=TOL)
    assert float(v["W51"]) == pytest.approx(0.0, abs=TOL)
    assert v["W52"] == "RECONCILIADO"


@com
def test_c1_sem_execucao_permanece_ausente(excel, tmp_path):
    v = _valores(excel, _cenario(tmp_path / "c1_ausente.xlsx", []))
    assert v["B11"] in (None, "")


@com
def test_c1_misto_soma_somente_parcela_elegivel(excel, tmp_path):
    arq = _cenario(tmp_path / "c1_misto.xlsx", [
        (date(2026, 9, 15), 85_720.0, "Nao"),
        (date(2026, 10, 15), 85_720.0, "Sim"),
    ])
    v = _valores(excel, arq)
    assert float(v["B11"]) == pytest.approx(3_805.97, abs=TOL)
    assert float(v["B16"]) == pytest.approx(3_805.97, abs=TOL)
