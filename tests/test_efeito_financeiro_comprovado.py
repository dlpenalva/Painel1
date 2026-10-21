"""P2 do PR #194 — EFEITO_FINANCEIRO precisa comprovar o retroativo zero.

Regra: competencia paga com "Nao" produz retroativo legitimamente zero; "Sim"
produz a diferenca aplicavel (DELTA numerico); vazio ou invalido nao comprova
zero. Nesse caso nem o retroativo, nem o VTA, nem a formalizacao sao liberados
— no XLS (B10:B14, B44, B15 -> B26 e status) e no Python (leitor, memoria por
ciclo, resultado consolidado e politica de entrega). Dados sinteticos.
"""
from __future__ import annotations

import io
import os
import time
import warnings
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import _leitor_masterfile_v10 as leitor
from _ciclo_em_execucao import garantir_aba_ciclo_em_execucao
from _coleta_oficial import (
    _FORMULA_B15_126,
    _FORMULA_B15_126_ANTERIOR,
    _FORMULA_B44_126,
    _FORMULA_B44_126_ANTERIOR,
    _FORMULAS_RETRO_FINANCEIRO_126,
    _FORMULAS_RETRO_FINANCEIRO_126_V1,
    _garantir_retro_zero_e_referencias_temporais,
    obter_coleta_oficial_bytes,
)
from _coleta_reajuste_documentos import processar_coleta_oficial_runtime
from _politica_entrega_segura import MENSAGEM_EFEITO_FINANCEIRO_PENDENTE

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "COLETA_REAJUSTE_OFICIAL.xlsx"
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "coletas_compat"
TOL = 0.005
DOCS_FORMALIZADORES = ("sumario_executivo", "despacho_saneador", "termo_apostila")


# --------------------------------------------------------------------------- #
# XLS 12.6: formulas no template e na Coleta gerada; migracao idempotente.
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("origem", ["template", "gerada"])
def test_formulas_exigem_efeito_comprovado(origem):
    fonte = TEMPLATE if origem == "template" else io.BytesIO(obter_coleta_oficial_bytes())
    mem = load_workbook(fonte, data_only=False)["MEMORIA_RESULTADOS"]
    for celula, formula in _FORMULAS_RETRO_FINANCEIRO_126.items():
        assert mem[celula].value == formula
        assert 'financeiro!$G$2:$G$73="Nao"' in formula
        assert "ISNUMBER(financeiro!$F$2:$F$73)" in formula
    assert mem["B15"].value == _FORMULA_B15_126 and "$B$44>0" in _FORMULA_B15_126
    assert mem["B44"].value == _FORMULA_B44_126


def test_migracao_aceita_12_5_e_12_6_v1_e_rejeita_formato_desconhecido():
    for anterior in ("v1", "12.5"):
        wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
        mem = wb["MEMORIA_RESULTADOS"]
        for celula, v1 in _FORMULAS_RETRO_FINANCEIRO_126_V1.items():
            mem[celula].value = v1 if anterior == "v1" else v1.replace(
                'financeiro!$C$2:$C$73,">0")=0',
                'financeiro!$C$2:$C$73,">0",financeiro!$G$2:$G$73,"Sim")=0', 1,
            )
        mem["B15"].value, mem["B44"].value = _FORMULA_B15_126_ANTERIOR, _FORMULA_B44_126_ANTERIOR
        _garantir_retro_zero_e_referencias_temporais(wb)
        _garantir_retro_zero_e_referencias_temporais(wb)  # idempotente
        assert all(mem[c].value == f for c, f in _FORMULAS_RETRO_FINANCEIRO_126.items())
        assert mem["B15"].value == _FORMULA_B15_126 and mem["B44"].value == _FORMULA_B44_126
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    wb["MEMORIA_RESULTADOS"]["B44"].value = "=0"
    with pytest.raises(ValueError, match="B44"):
        _garantir_retro_zero_e_referencias_temporais(wb)


# --------------------------------------------------------------------------- #
# Leitor: cada competencia paga declara se o efeito esta comprovado.
# --------------------------------------------------------------------------- #
def _aba_financeiro(linhas, *, mensal=True) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "financeiro"
    for coluna, nome in enumerate(
        ("COMPETENCIA", "CICLO", "VALOR_PAGO", "FATOR_APLICAVEL",
         "VALOR_ATUALIZADO", "DELTA", "EFEITO_FINANCEIRO"), 1
    ):
        ws.cell(1, coluna).value = nome
    for linha, (ciclo, pago, delta, efeito) in enumerate(linhas, 2):
        ws.cell(linha, 1).value = date(2025, 8, 1) if mensal else "Agregado"
        ws.cell(linha, 2).value, ws.cell(linha, 3).value = ciclo, pago
        ws.cell(linha, 6).value, ws.cell(linha, 7).value = delta, efeito
    return wb


def test_leitor_marca_efeito_comprovado_por_competencia_paga():
    wb = _aba_financeiro([
        ("C1", 12_000.0, 450.0, "Sim"),     # diferenca aplicavel
        ("C1", 12_000.0, None, "Sim"),      # Sim sem DELTA: nao comprova
        ("C1", 12_000.0, 0.0, "Nao"),       # zero legitimo
        ("C1", 12_000.0, 0.0, "Não"),       # zero legitimo (acento)
        ("C1", 12_000.0, None, None),       # vazio
        ("C1", 12_000.0, 0.0, "Talvez"),    # invalido
        ("C0", 100_000.0, 0.0, "nao"),      # caixa indiferente
    ])
    bases = [p for p in leitor._ler_parcelas_sombra_financeiro(wb) if ":base:" in p["identificador"]]
    assert [p["efeito_comprovado"] for p in bases] == [True, False, True, True, False, False, True]
    assert [p["valor"] for p in bases] == [12_000.0] * 6 + [100_000.0]


def test_leitor_nao_aplica_a_regra_ao_layout_agregado_antigo():
    wb = _aba_financeiro([("C1", 12_000.0, None, None)], mensal=False)
    assert all(
        p.get("efeito_comprovado") is None for p in leitor._ler_parcelas_sombra_financeiro(wb)
    )


# --------------------------------------------------------------------------- #
# Cadeia Python completa: integro inalterado; pendente fecha VTA e formalizacao.
# --------------------------------------------------------------------------- #
def _processar(nome: str, monkeypatch=None, *, pendente: bool = False):
    if pendente:
        original = leitor._ler_parcelas_sombra_financeiro

        def _com_pendencia(wb):
            parcelas = original(wb)
            alvo = next(
                p for p in parcelas
                if p.get("efeito_comprovado") is not None and p["ciclo"] == "C1"
            )
            alvo["efeito_comprovado"] = False  # C1 em execucao sem efeito valido
            return parcelas

        monkeypatch.setattr(leitor, "_ler_parcelas_sombra_financeiro", _com_pendencia)
    conteudo = (FIXTURES / f"{nome}.xlsx").read_bytes()
    resultado, diagnostico = processar_coleta_oficial_runtime(conteudo)
    leitura = leitor.ler_masterfile_v10(conteudo, exigir_modelo_oficial=True)
    return resultado, diagnostico, leitura


@pytest.mark.parametrize("nome, vta", [
    ("coleta_11_financeiro", 1_340_973.60),
    ("coleta_11_pc", 299_670.40),
    ("coleta_11_consumidos", 289_514.88),
])
def test_metodos_integros_preservam_vta_e_formalizacao(nome, vta):
    resultado, _, leitura = _processar(nome)
    assert leitura["financeiro_efeito_pendente"] == []
    assert resultado["pendencia_efeito_financeiro"] is False
    assert resultado["resultado_consolidado"]["vta"] == pytest.approx(vta, abs=TOL)
    assert MENSAGEM_EFEITO_FINANCEIRO_PENDENTE not in (resultado.get("bloqueios_formalizacao") or [])


def test_financeiro_integro_preserva_retroativo_e_documentos():
    resultado, _, _ = _processar("coleta_11_financeiro")
    assert resultado["resultado_consolidado"]["retroativo_reconhecido"] == pytest.approx(26_112.0, abs=TOL)
    documentos = resultado["capacidades"]["documentos"]
    assert all(documentos[d]["habilitado"] for d in DOCS_FORMALIZADORES)


def test_financeiro_com_efeito_pendente_nao_libera_retro_vta_nem_formalizacao(monkeypatch):
    resultado, _, leitura = _processar("coleta_11_financeiro", monkeypatch, pendente=True)
    assert leitura["financeiro_efeito_pendente"]
    assert any("EFEITO_FINANCEIRO valido" in a for a in leitura["avisos"])
    memoria = leitura["objeto_processo"]["memoria_por_ciclo"]
    financeiro = next(c for c in memoria["conferencias_metodologicas"] if c["metodo"] == "financeiro")
    assert financeiro["disponivel"] is False
    assert memoria["vta"]["valor_total_atualizado"] is None
    assert memoria["vta"]["metodo"] == "financeiro"  # sem fallback para PC/Consumidos
    assert "EFEITO_FINANCEIRO" in memoria["vta"]["motivo"]

    consolidado = resultado["resultado_consolidado"]
    assert consolidado["vta"] is None
    assert consolidado["retroativo_reconhecido"] is None
    assert resultado["pendencia_efeito_financeiro"] is True
    assert resultado["formalizacao_bloqueada"] is True
    assert MENSAGEM_EFEITO_FINANCEIRO_PENDENTE in resultado["bloqueios_documentais_duros"]
    documentos = resultado["capacidades"]["documentos"]
    assert not any(documentos[d]["habilitado"] for d in DOCS_FORMALIZADORES)


# --------------------------------------------------------------------------- #
# Excel real (opcional): C1 em execucao com Sim, Nao, vazio e invalido.
# --------------------------------------------------------------------------- #
_DATA_FOTO = date(2025, 10, 15)


def _cenario(destino: Path, efeitos: tuple[str | None, str | None]) -> Path:
    wb = load_workbook(io.BytesIO(obter_coleta_oficial_bytes()))
    ctl = wb["CONTROLE"]
    ctl["B1"], ctl["B2"], ctl["B3"] = "Financeiro (Mensalidade)", "C1", _DATA_FOTO
    par = wb["parametros"]
    par["A2"], par["B2"], par["C2"], par["D2"], par["I2"] = (
        "Nao", "C0", date(2024, 8, 1), date(2025, 7, 31), date(2024, 8, 1)
    )
    par["A3"], par["B3"], par["C3"], par["D3"], par["E3"] = (
        "Sim", "C1", date(2025, 8, 1), date(2026, 7, 31), 0.0375
    )
    par["H3"], par["I3"] = date(2025, 10, 1), _DATA_FOTO
    fin = wb["financeiro"]
    linhas = [
        (date(2025, 7, 15), 100_000.0, "Nao"),
        (date(2025, 9, 15), 12_000.0, efeitos[0]),
        (date(2025, 10, 10), 12_000.0, efeitos[1]),
    ]
    for linha, (competencia, valor, efeito) in enumerate(linhas, 2):
        fin[f"A{linha}"], fin[f"C{linha}"], fin[f"G{linha}"] = competencia, valor, efeito
    rem = wb["itens_Remanesc"]
    rem["A2"], rem["B2"], rem["C2"], rem["E2"] = "ITEM-01", 1730.5, 250.0, 1234.5
    garantir_aba_ciclo_em_execucao(wb)
    cee = wb["CICLO_EM_EXECUCAO"]
    cee["D5"], cee["C13"] = _DATA_FOTO, 1234.5
    wb.save(destino)
    return destino


@pytest.fixture
def excel():
    client = pytest.importorskip("win32com.client")
    pythoncom = pytest.importorskip("pythoncom")
    pythoncom.CoInitialize()
    app = client.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    yield app
    try:
        app.Quit()
    except Exception:
        pass
    pythoncom.CoUninitialize()


def _valores(excel, caminho: Path) -> dict[str, object]:
    import pythoncom

    book = excel.Workbooks.Open(str(caminho), UpdateLinks=0, ReadOnly=False)
    try:
        for _ in range(8):
            try:
                excel.CalculateFullRebuild()
                break
            except Exception:
                pythoncom.PumpWaitingMessages()
                time.sleep(0.5)
        mem = book.Worksheets("MEMORIA_RESULTADOS")
        det = book.Worksheets("RESULTADOS_DETALHE")
        saida = {c: mem.Range(c).Value2 for c in ("B11", "B15", "B16", "B26", "B44", "F16")}
        saida.update({f"DET_{c}": det.Range(c).Value2 for c in ("H8", "H14", "B3")})
        return saida
    finally:
        book.Close(SaveChanges=False)


com = pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="Excel real: defina RUN_EXCEL_INTEGRATION=1",
)


@com
@pytest.mark.parametrize("efeitos, retro, vta", [
    (("Nao", "Nao"), 0.0, 444_204.61),        # zero legitimo preservado
    (("Nao", "Sim"), 450.0, 444_654.61),       # diferenca aplicavel
])
def test_excel_c1_em_execucao_com_efeito_valido(excel, tmp_path, efeitos, retro, vta):
    v = _valores(excel, _cenario(tmp_path / "valido.xlsx", efeitos))
    assert float(v["B44"]) == 0
    assert float(v["B11"]) == pytest.approx(retro, abs=TOL)
    assert float(v["B15"]) == pytest.approx(retro, abs=TOL)
    assert float(v["B26"]) == pytest.approx(vta, abs=TOL)
    assert "CALCULADO" in str(v["F16"])
    assert v["DET_H14"] == "VALIDADO"


@com
@pytest.mark.parametrize("efeitos", [("Nao", None), ("Nao", "Talvez"), (None, "Sim")])
def test_excel_c1_em_execucao_sem_efeito_valido_nao_libera(excel, tmp_path, efeitos):
    v = _valores(excel, _cenario(tmp_path / "pendente.xlsx", efeitos))
    assert float(v["B44"]) >= 1
    for celula in ("B11", "B15", "B16", "B26"):
        assert v[celula] in (None, ""), celula
    assert "CALCULADO" not in str(v["F16"])
    assert v["DET_H14"] == "REVISE" and v["DET_H8"] == "REVISE" and v["DET_B3"] == "REVISE"
