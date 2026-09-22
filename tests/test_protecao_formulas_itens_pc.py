# -*- coding: utf-8 -*-
"""Hotfix — protecao das formulas automaticas de itens_PC.

Incidente real (Coleta_Reajuste_C1_ICTI_22-09-2026.xlsx): C2, E2 e F2 da
primeira linha de itens_PC foram sobrescritas por valores ("C0", 1,
2066146,03) e o upload so dizia "Formula estrutural ausente em itens_PC!C2".

Cobertura:
1. a Coleta nova sai com itens_PC protegida, A/B/D/G editaveis e todas as
   formulas preservadas;
2. a protecao nao altera valores, formulas nem estilos visuais;
3. o upload do layout oficial bloqueia qualquer celula automatica
   sobrescrita, com diagnostico claro, sem reparar e sem alterar o arquivo;
4. Coleta normal segue aceita; layout legado incompativel nao recebe a
   nova exigencia;
5. (opt-in RUN_EXCEL_INTEGRATION=1) Excel real: preencher PC em A/B/D/G,
   formulas recusam escrita, valores identicos com e sem protecao, reabre
   sem reparo.

A suite roda no CI rapido (limite de 15 min): cada geracao/upload da Coleta
custa segundos, por isso os cenarios de upload estao agrupados e as regras
puras sao testadas direto nas funcoes.
"""
from __future__ import annotations

import gc
import hashlib
import inspect
import io
import os
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import _coleta_oficial
import _gerador_masterfile
from _capacidade_pcs import CAPACIDADE_PCS, ULTIMA_LINHA_PCS
from _coleta_oficial import (
    COLS_AUTOMATICAS_ITENS_PC,
    COLS_MANUAIS_ITENS_PC,
    TEMPLATE_COLETA_OFICIAL,
    eh_layout_itens_pc_protegivel,
    gerar_coleta_oficial_preenchida,
)
from _coleta_reajuste import (
    _celulas_automaticas_itens_pc_sobrescritas,
    _formulas,
    _mensagem_itens_pc_sobrescritas,
    ler_coleta_reajuste,
)
from tests._fabrica_coleta import bytes_coleta_oficial

MENSAGEM_C2_E2_F2 = (
    "Há células automáticas sobrescritas na aba itens_PC (ex.: C2, E2, F2). "
    "Preencha somente NUMERO_PC, DATA_PC, VALOR_PC e PC_PAGO_A_CONTRATADA. "
    "Regere a Coleta antes do upload."
)

LINHAS_AMOSTRA = (2, 3, 101, 2050, ULTIMA_LINHA_PCS)


def _dados() -> dict:
    return {
        "origem": "Teste protecao itens_PC",
        "indice": "ICTI",
        "data_base_original": "01/02/2023",
        "data_corte": date(2025, 1, 31),
        "ciclos": [{
            "ciclo": "C1",
            "data_inicio": date(2024, 2, 1),
            "data_fim": date(2025, 1, 31),
            "data_pedido": date(2024, 3, 10),
            "financeiro_inicio": date(2024, 4, 18),
            "percentual_aplicado": 0.10,
            "objeto_analise_atual": True,
        }],
    }


def _assinatura_celulas(wb) -> dict:
    """Valor + estilo visual (sem protecao) de todas as celulas."""
    def estilo(c):
        s = getattr(c, "_style", None)
        if s is None:
            return None
        return (s.fontId, s.fillId, s.borderId, s.numFmtId, s.alignmentId)

    return {
        ws.title: {
            c.coordinate: (c.value, estilo(c))
            for row in ws.iter_rows() for c in row
        }
        for ws in wb.worksheets
    }


@pytest.fixture(scope="module")
def bytes_preenchida() -> bytes:
    return bytes_coleta_oficial(_dados())


# --------------------------------------------------------------------------- #
# 1. Coleta nova: protecao + formulas preservadas                              #
# --------------------------------------------------------------------------- #
def test_coleta_nova_protege_itens_pc_e_libera_a_b_d_g(bytes_preenchida):
    ws = load_workbook(io.BytesIO(bytes_preenchida), data_only=False)["itens_PC"]

    assert ws.protection.sheet is True
    assert ws.protection.password is None
    assert ws.protection.selectLockedCells is False
    assert ws.protection.selectUnlockedCells is False
    for linha in LINHAS_AMOSTRA:
        for coluna in COLS_MANUAIS_ITENS_PC:
            assert ws[f"{coluna}{linha}"].protection.locked is False, f"{coluna}{linha}"
        for coluna in COLS_AUTOMATICAS_ITENS_PC:
            assert ws[f"{coluna}{linha}"].protection.locked is True, f"{coluna}{linha}"

    # Toda a grade automatica segue com formula em todas as linhas.
    for coluna in COLS_AUTOMATICAS_ITENS_PC:
        indice = ws[f"{coluna}1"].column
        formulas = sum(
            1
            for (valor,) in ws.iter_rows(
                min_row=2, max_row=ULTIMA_LINHA_PCS,
                min_col=indice, max_col=indice, values_only=True,
            )
            if isinstance(valor, str) and valor.startswith("=")
        )
        assert formulas == CAPACIDADE_PCS, coluna
    # Dropdown de PC_PAGO_A_CONTRATADA intacto.
    assert any(
        f"G2:G{ULTIMA_LINHA_PCS}" in str(dv.sqref)
        for dv in ws.data_validations.dataValidation
    )


def test_os_dois_geradores_aplicam_a_protecao():
    # Coleta em branco e Coleta preenchida passam pelo mesmo guard (regerar
    # a Coleta so para isso custaria segundos no CI rapido).
    assert "_garantir_protecao_formulas_itens_pc(wb)" in inspect.getsource(
        _coleta_oficial.obter_coleta_oficial_bytes
    )
    assert "_garantir_protecao_formulas_itens_pc(wb)" in inspect.getsource(
        _gerador_masterfile.gerar_masterfile_preenchido
    )


def test_protecao_nao_altera_valores_formulas_nem_estilos():
    wb = load_workbook(TEMPLATE_COLETA_OFICIAL, data_only=False)
    antes = _assinatura_celulas(wb)
    protecao_antes = {ws.title: ws.protection.sheet for ws in wb.worksheets}

    _coleta_oficial._garantir_protecao_formulas_itens_pc(wb)

    assert _assinatura_celulas(wb) == antes
    protecao_depois = {ws.title: ws.protection.sheet for ws in wb.worksheets}
    assert protecao_depois.pop("itens_PC") is True
    protecao_antes.pop("itens_PC")
    assert protecao_depois == protecao_antes


def test_layout_legado_nao_e_protegido():
    wb = Workbook()
    ws = wb.active
    ws.title = "itens_PC"
    for coluna, rotulo in zip("ABCDEFG", ("ITEM", "DATA_PC", "CICLO_PC", "VALOR_PC",
                                          "FATOR_ACUMULADO", "VALOR_ATUALIZADO",
                                          "PC_PAGO_A_CONTRATADA")):
        ws[f"{coluna}1"] = rotulo

    assert eh_layout_itens_pc_protegivel(ws) is False
    _coleta_oficial._garantir_protecao_formulas_itens_pc(wb)
    assert ws.protection.sheet is False
    assert ws["A2"].protection.locked is True


# --------------------------------------------------------------------------- #
# 2. Upload: diagnostico de sobrescrita                                        #
# --------------------------------------------------------------------------- #
def _coleta_com_pc(base: bytes, ajustar=None) -> bytes:
    wb = load_workbook(io.BytesIO(base), data_only=False)
    ws = wb["itens_PC"]
    ws["A2"] = "4500012345"
    ws["B2"] = datetime(2024, 6, 10)
    ws["D2"] = 2066146.03
    ws["G2"] = "Sim"
    if ajustar is not None:
        ajustar(ws)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _bloqueios_itens_pc(resultado: dict) -> list[str]:
    return [b for b in resultado["bloqueios_estruturais"] if "itens_PC" in b]


def test_coleta_normal_com_pc_nao_acusa_sobrescrita(bytes_preenchida):
    wb = load_workbook(io.BytesIO(_coleta_com_pc(bytes_preenchida)), data_only=False)
    assert _celulas_automaticas_itens_pc_sobrescritas(wb["itens_PC"], _formulas(wb)) == []


def test_upload_bloqueia_c2_e2_f2_sem_reparar_nem_alterar_arquivo(bytes_preenchida, tmp_path):
    def sobrescrever(ws):
        ws["C2"] = "C0"
        ws["E2"] = 1
        ws["F2"] = 2066146.03

    caminho = tmp_path / "coleta_sobrescrita.xlsx"
    caminho.write_bytes(_coleta_com_pc(bytes_preenchida, sobrescrever))
    assinatura = hashlib.sha256(caminho.read_bytes()).hexdigest()

    resultado = ler_coleta_reajuste(caminho.read_bytes())

    assert resultado["valido"] is False
    # Unico bloqueio: o diagnostico novo (substitui a mensagem generica de C2,
    # sem duplicar). Sem a sobrescrita, esta Coleta seria aceita.
    assert resultado["bloqueios_estruturais"] == [MENSAGEM_C2_E2_F2]
    # Nada reparado: o arquivo fisico segue byte a byte igual.
    assert hashlib.sha256(caminho.read_bytes()).hexdigest() == assinatura


def test_varredura_aponta_sobrescrita_alem_da_primeira_linha(bytes_preenchida):
    # Mesma varredura usada por ler_coleta_reajuste, direto no workbook
    # (o ciclo salvar+upload ja e coberto pelo caso C2/E2/F2 acima).
    wb = load_workbook(io.BytesIO(bytes_preenchida), data_only=False)
    ws = wb["itens_PC"]
    ws["A3"] = "4500099999"      # linha com PC e formula apagada
    ws["B3"] = datetime(2024, 7, 1)
    ws["D3"] = 10.0
    ws["G3"] = "Nao"
    ws["L3"] = None
    ws["F150"] = 123.45          # valor fixo em linha sem PC
    for coluna in COLS_AUTOMATICAS_ITENS_PC:
        ws[f"{coluna}3000"] = None   # linha sem PC apenas limpa: nao acusa
    ws[f"U{ULTIMA_LINHA_PCS}"] = 0

    assert _celulas_automaticas_itens_pc_sobrescritas(ws, _formulas(wb)) == [
        "L3", "F150", f"U{ULTIMA_LINHA_PCS}",
    ]


def test_mensagem_informa_contagem_quando_ha_muitas_celulas():
    celulas = [f"{c}{l}" for l in range(2, 5) for c in ("C", "E", "F")]
    assert _mensagem_itens_pc_sobrescritas(celulas).startswith(
        "Há células automáticas sobrescritas na aba itens_PC "
        "(9 células; ex.: C2, E2, F2, C3, E3, F3)."
    )


def test_varredura_acomoda_grade_menor_de_coleta_anterior():
    # Coleta oficial anterior a 26G: formulas so ate a linha 101.
    wb = Workbook()
    ws = wb.active
    ws.title = "itens_PC"
    for linha in range(2, 102):
        for coluna in COLS_AUTOMATICAS_ITENS_PC:
            ws[f"{coluna}{linha}"] = "=ROW()"
    ws["A150"] = "PC-FORA"  # alem da grade: fora desta regra
    assert _celulas_automaticas_itens_pc_sobrescritas(ws, _formulas(wb)) == []
    ws["E50"] = 1.0
    assert _celulas_automaticas_itens_pc_sobrescritas(ws, _formulas(wb)) == ["E50"]


# --------------------------------------------------------------------------- #
# 3. Excel real (opt-in)                                                       #
# --------------------------------------------------------------------------- #
PCS_EXCEL = (
    (2, "PC-ANTES", datetime(2024, 4, 10), 100.0, "Nao"),
    (3, "PC-EXATO", datetime(2024, 4, 18), 2066146.03, "Nao"),
    (4, "PC-DEPOIS", datetime(2024, 4, 25), 555.55, "Sim"),
)
COLUNAS_LIDAS = ("C", "E", "F", "H", "I", "J", "K", "L", "U")


def _excel_preencher_e_ler(caminho: Path, *, tentar_formulas: bool):
    client = pytest.importorskip("win32com.client")
    pythoncom = pytest.importorskip("pythoncom")
    import pywintypes

    pythoncom.CoInitialize()
    excel = client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    travadas: dict[str, bool] = {}
    try:
        pasta = excel.Workbooks.Open(str(caminho), UpdateLinks=0, ReadOnly=False, CorruptLoad=0)
        assert "repar" not in str(pasta.Name).lower()
        ws = pasta.Worksheets("itens_PC")
        for linha, numero, data_pc, valor, pago in PCS_EXCEL:
            ws.Range(f"A{linha}").Value = numero
            ws.Range(f"B{linha}").Value = data_pc
            ws.Range(f"D{linha}").Value = valor
            ws.Range(f"G{linha}").Value = pago
        if tentar_formulas:
            assert bool(ws.ProtectContents) is True
            for coord in ("C2", "E2", "F2", "H2", "L2", "U2", f"F{ULTIMA_LINHA_PCS}"):
                try:
                    ws.Range(coord).Value = 1
                    travadas[coord] = False
                except pywintypes.com_error:
                    travadas[coord] = True
            # Colagem real (area de transferencia): bloco que atravessa coluna
            # automatica e recusado por inteiro; colar so em A continua livre.
            origem = excel.Workbooks.Add()
            fonte = origem.Worksheets(1)
            fonte.Range("A1:D1").Value = (("PC-X", 45000, "C9", 1.0),)
            for faixa, destino in (("A1:D1", "A5"), ("C1", "C2")):
                fonte.Range(faixa).Copy()
                try:
                    ws.Paste(ws.Range(destino))
                    travadas[f"colar {destino}"] = False
                except pywintypes.com_error:
                    travadas[f"colar {destino}"] = True
            assert ws.Range("A5").Value is None
            assert bool(ws.Range("C2").HasFormula) and bool(ws.Range("C5").HasFormula)
            fonte.Range("A1").Copy()
            ws.Paste(ws.Range("A6"))
            assert ws.Range("A6").Value == "PC-X"
            ws.Range("A6").Value = None
            excel.CutCopyMode = False
            origem.Close(SaveChanges=False)
        excel.CalculateFullRebuild()
        valores = {
            linha: tuple(ws.Range(f"{c}{linha}").Value for c in COLUNAS_LIDAS)
            for linha, *_ in PCS_EXCEL
        }
        resumo = tuple(
            ws.Range(f"{c}{linha}").Value for linha in range(3, 10) for c in "NOPQRS"
        )
        pasta.Save()
        pasta.Close(SaveChanges=False)

        reaberta = excel.Workbooks.Open(str(caminho), UpdateLinks=0, ReadOnly=True, CorruptLoad=0)
        assert "repar" not in str(reaberta.Name).lower()
        ws2 = reaberta.Worksheets("itens_PC")
        protegida_reaberta = bool(ws2.ProtectContents)
        assert ws2.Range("A3").Value == "PC-EXATO"
        reaberta.Close(SaveChanges=False)
    finally:
        excel.Quit()
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()
    return valores, resumo, travadas, protegida_reaberta


@pytest.mark.skipif(
    os.environ.get("RUN_EXCEL_INTEGRATION") != "1",
    reason="defina RUN_EXCEL_INTEGRATION=1 para executar o Excel COM",
)
def test_excel_real_protecao_e_valores_identicos(tmp_path, monkeypatch):
    protegido = tmp_path / "coleta_protegida.xlsx"
    protegido.write_bytes(gerar_coleta_oficial_preenchida(_dados()))
    monkeypatch.setattr(_coleta_oficial, "_garantir_protecao_formulas_itens_pc", lambda wb: None)
    sem_protecao = tmp_path / "coleta_sem_protecao.xlsx"
    sem_protecao.write_bytes(gerar_coleta_oficial_preenchida(_dados()))
    assert load_workbook(sem_protecao)["itens_PC"].protection.sheet is False

    valores_p, resumo_p, travadas, reaberta_protegida = _excel_preencher_e_ler(
        protegido, tentar_formulas=True
    )
    valores_s, resumo_s, _, _ = _excel_preencher_e_ler(sem_protecao, tentar_formulas=False)

    assert travadas and all(travadas.values()), travadas
    assert reaberta_protegida is True
    # PC preenchido normalmente: formulas calcularam o ciclo e o valor atualizado.
    assert valores_p[3][0] == "C1"
    assert valores_p[3][2] == pytest.approx(round(2066146.03 * valores_p[3][1], 2))
    # Nenhum valor financeiro muda com a protecao.
    assert valores_p == valores_s
    assert resumo_p == resumo_s

    # Arquivo salvo pelo Excel continua aceito no upload.
    resultado = ler_coleta_reajuste(protegido.read_bytes())
    assert _bloqueios_itens_pc(resultado) == []
