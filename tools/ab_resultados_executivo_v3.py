# -*- coding: utf-8 -*-
"""Prova A/B da RESULTADOS-EXECUTIVO-V3 (BEFORE = main, AFTER = branch).

Etapas (cada uma roda no contexto do repositorio indicado por --repo, para que
BEFORE use o codigo e o template da main e AFTER os da branch):

  montar      gera os XLS dos cenarios pelo fluxo real de geracao
              (`obter_coleta_oficial_bytes` / `gerar_coleta_oficial_preenchida`)
              e preenche SOMENTE entradas de usuario (fiscal / CONTROLE);
  recalcular  abre cada XLS no Excel real, CalculateFullRebuild, salva;
  fotografar  le TODOS os valores calculados de TODAS as abas (data_only) e a
              cadeia web + documentos (`tests/_baseline_fotografia`);
  comparar    diff BEFORE x AFTER com tolerancia ZERO. A aba RESULTADOS da main
              e comparada com a RESULTADOS_DETALHE da branch; as unicas
              diferencas aceitas sao as areas deliberadamente alteradas.

Uso:
  python tools/ab_resultados_executivo_v3.py montar --repo R --out D
  python tools/ab_resultados_executivo_v3.py recalcular --dir D
  python tools/ab_resultados_executivo_v3.py fotografar --repo R --dir D
  python tools/ab_resultados_executivo_v3.py comparar --antes DA --depois DB
"""
from __future__ import annotations

import argparse
import gc
import io
import json
import sys
import time
from datetime import date, datetime
from pathlib import Path


def _preparar_repo(repo: Path) -> None:
    sys.path.insert(0, str(repo / "tests"))
    sys.path.insert(0, str(repo))


# --------------------------------------------------------------------- montar
def _fiscal_ciclo(wb, data: date, restantes: dict[str, float]) -> None:
    """CICLO_EM_EXECUCAO como o fiscal preenche: D5 (data) e C (quantidade)."""
    ws = wb["CICLO_EM_EXECUCAO"]
    ws["D5"] = data
    origem = wb["itens_Remanesc"]
    for k in range(200):
        codigo = origem.cell(2 + k, 1).value
        if codigo in restantes:
            ws.cell(13 + k, 3).value = restantes[codigo]


def _ajuste_manual(wb) -> None:
    """Retroativo manual oficial (linha 43 da camada tecnica)."""
    try:
        from _resultados_abas import aba_resultados_tecnica
        aba = aba_resultados_tecnica(wb)
    except ImportError:          # main anterior: so existe RESULTADOS
        aba = "RESULTADOS"
    ws = wb[aba]
    ws["C43"] = 1234.56
    ws["D43"] = "Ajuste de homologacao A/B"
    ws["E43"] = "Fiscal"
    ws["F43"] = datetime(2025, 1, 10)
    ws["G43"] = "Sim"


def _cenarios():
    import _baseline_cenarios as bc

    def base(wb, ciclos, metodo, vigente, corte):
        bc._base(wb, ciclos=ciclos, metodo=metodo, ciclo_vigente=vigente,
                 data_corte=corte)

    def fin_normal(wb):
        base(wb, bc._um_ciclo(), "Financeiro (Mensalidade)", "C1", date(2024, 12, 31))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 24, 42_500.00))
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)
        bc._cobertura(wb, financeiro_ate=date(2024, 12, 31))

    def fin_multiciclo(wb):
        base(wb, bc._ate(3), "Financeiro (Mensalidade)", "C3", date(2026, 12, 31))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 48, 42_500.00))
        _fiscal_ciclo(wb, date(2026, 12, 31), bc.RESTANTE_PADRAO)
        bc._cobertura(wb, financeiro_ate=date(2026, 12, 31))

    def fin_sem_efeito(wb):
        ciclos = bc._um_ciclo()
        ciclos[1]["inicio_efeito"] = date(2024, 4, 1)
        base(wb, ciclos, "Financeiro (Mensalidade)", "C1", date(2024, 12, 31))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 24, 42_500.00))
        for linha in range(14, 17):                       # jan..mar/2024
            wb["financeiro"].cell(linha, 7).value = "Nao"
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)
        bc._cobertura(wb, financeiro_ate=date(2024, 12, 31))

    def fin_retro_zero(wb):
        ciclos = bc._um_ciclo()
        ciclos[1]["percentual"] = 0.0
        ciclos[1]["situacao"] = "TEMPESTIVO — VARIAÇÃO NEGATIVA NEUTRALIZADA EM 0,00%"
        base(wb, ciclos, "Financeiro (Mensalidade)", "C1", date(2024, 12, 31))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 24, 42_500.00))
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)

    def fin_estimado(wb):
        base(wb, bc._um_ciclo(), "Financeiro (Mensalidade)", "C1", date(2024, 12, 31))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 18, 42_500.00))
        _fiscal_ciclo(wb, date(2024, 6, 30), bc.RESTANTE_PADRAO)
        bc._cobertura(wb, financeiro_ate=date(2024, 6, 30))

    def fin_posterior(wb):
        base(wb, bc._um_ciclo(), "Financeiro (Mensalidade)", "C1", date(2024, 6, 30))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 24, 42_500.00))
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)

    def fin_sem_ciclo_exec(wb):
        base(wb, bc._um_ciclo(), "Financeiro (Mensalidade)", "C1", date(2024, 12, 31))
        bc._financeiro(wb, bc._competencias(date(2023, 1, 1), 24, 42_500.00))

    def fin_ajustes(wb):
        fin_normal(wb)
        _ajuste_manual(wb)

    def pc(wb):
        base(wb, bc._um_ciclo(), "PC (Pedidos de Compra)", "C1", date(2024, 12, 31))
        bc._itens_pc(wb, [
            {"numero": "PC-2024-001", "data": date(2024, 2, 15), "valor": 180_000.00, "pago": "Sim"},
            {"numero": "PC-2024-002", "data": date(2024, 5, 20), "valor": 96_500.00, "pago": "Sim"},
            {"numero": "PC-2024-003", "data": date(2024, 9, 30), "valor": 145_250.00, "pago": "Nao"},
        ])
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)
        bc._cobertura(wb, pcs_ate=date(2024, 12, 31))

    def pc_sem_efeito(wb):
        ciclos = bc._um_ciclo()
        ciclos[1]["inicio_efeito"] = date(2024, 7, 1)
        base(wb, ciclos, "PC (Pedidos de Compra)", "C1", date(2024, 12, 31))
        bc._itens_pc(wb, [
            {"numero": "PC-2024-001", "data": date(2024, 2, 15), "valor": 180_000.00, "pago": "Sim"},
            {"numero": "PC-2024-002", "data": date(2024, 5, 20), "valor": 96_500.00, "pago": "Sim"},
            {"numero": "PC-2024-004", "data": date(2024, 8, 5), "valor": 61_300.00, "pago": "Nao"},
            {"numero": "PC-2025-003", "data": date(2025, 3, 10), "valor": 145_250.00, "pago": "Nao"},
        ])
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)

    def itens(wb):
        base(wb, bc._um_ciclo(), "Itens Consumidos", "C1", date(2024, 12, 31))
        bc._itens_consumidos(wb, [
            {"codigo": "ITEM-001", "quantidade": 75.0, "valor_unitario": 250.00},
            {"codigo": "ITEM-002", "quantidade": 50.0, "valor_unitario": 1_500.00},
            {"codigo": "ITEM-003", "quantidade": 28.0, "valor_unitario": 3_200.00},
        ])
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)

    def itens_sem_efeito(wb):
        ciclos = bc._um_ciclo()
        ciclos[1]["inicio_efeito"] = date(2024, 4, 1)
        base(wb, ciclos, "Itens Consumidos", "C1", date(2024, 12, 31))
        bc._itens_consumidos(wb, [
            {"codigo": "ITEM-001", "quantidade": 75.0, "valor_unitario": 250.00},
            {"codigo": "ITEM-002", "quantidade": 50.0, "valor_unitario": 1_500.00},
        ])
        _fiscal_ciclo(wb, date(2024, 12, 31), bc.RESTANTE_PADRAO)

    def aditivo(wb):
        fin_normal(wb)
        bc._aditivos(wb, [
            {"identificacao": "TA-01/2024", "data": date(2024, 6, 15),
             "valor_unitario": 4_100.00, "computar": "Sim", "novo_item": "Sim"},
        ])

    return {
        "fin_normal": fin_normal, "fin_multiciclo": fin_multiciclo,
        "fin_sem_efeito": fin_sem_efeito, "fin_retro_zero": fin_retro_zero,
        "fin_estimado": fin_estimado, "fin_posterior": fin_posterior,
        "fin_sem_ciclo_exec": fin_sem_ciclo_exec, "fin_ajustes": fin_ajustes,
        "pc": pc, "pc_sem_efeito": pc_sem_efeito, "itens": itens,
        "itens_sem_efeito": itens_sem_efeito, "aditivo": aditivo,
    }


# Caso de homologacao: PCs, multiciclo, potencial, sem efeito, valores altos.
ITENS_HOMOLOG = [
    {"codigo": "ITEM-001", "quantidade": 1200.0, "valor_unitario": 25_000.00,
     "por_ciclo": (1200.0, 1200.0, 1200.0, 1200.0, 1200.0)},
    {"codigo": "ITEM-002", "quantidade": 800.0, "valor_unitario": 95_000.00,
     "por_ciclo": (800.0, 800.0, 800.0, 800.0, 800.0)},
    {"codigo": "ITEM-003", "quantidade": 400.0, "valor_unitario": 62_000.00,
     "por_ciclo": (400.0, 400.0, 400.0, 400.0, 400.0)},
]
PCS_HOMOLOG = [
    ("PC-2023-001", date(2023, 6, 10), 8_500_000.00, "Sim"),
    ("PC-2024-001", date(2024, 1, 20), 4_200_000.00, "Sim"),
    ("PC-2024-002", date(2024, 2, 25), 3_150_000.00, "Sim"),
    ("PC-2024-003", date(2024, 5, 15), 12_800_000.00, "Sim"),
    ("PC-2024-004", date(2024, 10, 1), 6_400_000.00, "Nao"),
    ("PC-2025-001", date(2025, 3, 12), 15_300_000.00, "Sim"),
    ("PC-2025-002", date(2025, 8, 20), 9_750_000.00, "Nao"),
    ("PC-2026-001", date(2026, 1, 15), 5_600_000.00, "Sim"),
    ("PC-2026-002", date(2026, 4, 22), 11_200_000.00, "Sim"),
    ("PC-2026-003", date(2026, 8, 28), 7_900_000.00, "Nao"),
]


def dados_calculadora_homolog() -> dict:
    """Marcos como a Calculadora entrega (C1..C3 analisados, C3 vigente)."""
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2023, 1, 1),
        "data_corte": date(2026, 9, 30),
        "ciclos": [
            {"ciclo": "C1", "data_inicio": date(2024, 1, 1), "percentual": 0.0512,
             "inicio_efeito_financeiro": date(2024, 3, 1),
             "situacao_aplicada": "✅ TEMPESTIVO* — efeitos a partir de 03/2024",
             "data_pedido": date(2024, 2, 14), "efeito_financeiro_retardado": True},
            {"ciclo": "C2", "percentual": 0.0374,
             "inicio_efeito_financeiro": date(2025, 3, 1),
             "situacao_aplicada": "✅ TEMPESTIVO", "data_pedido": date(2025, 2, 10)},
            {"ciclo": "C3", "percentual": 0.0289,
             "inicio_efeito_financeiro": date(2026, 4, 1),
             "situacao_aplicada": "✅ TEMPESTIVO* — efeitos a partir de 04/2026",
             "data_pedido": date(2026, 3, 18), "efeito_financeiro_retardado": True},
        ],
    }


def preencher_homolog(wb) -> None:
    import _baseline_cenarios as bc

    wb["CONTROLE"]["B1"] = "PC (Pedidos de Compra)"
    bc._itens_remanesc(wb, ITENS_HOMOLOG)
    bc._posicao_referencia(wb, {"ITEM-001": 300.0, "ITEM-002": 200.0, "ITEM-003": 100.0})
    bc._itens_pc(wb, [
        {"numero": n, "data": d, "valor": v, "pago": p} for n, d, v, p in PCS_HOMOLOG
    ])
    _fiscal_ciclo(wb, date(2026, 9, 30),
                  {"ITEM-001": 300.0, "ITEM-002": 200.0, "ITEM-003": 100.0})
    bc._cobertura(wb, pcs_ate=date(2026, 9, 30))


def montar(repo: Path, out: Path, cenarios: list[str] | None) -> None:
    _preparar_repo(repo)
    from openpyxl import load_workbook
    from _coleta_oficial import gerar_coleta_oficial_preenchida, obter_coleta_oficial_bytes

    out.mkdir(parents=True, exist_ok=True)
    todos = _cenarios()
    alvo = cenarios or list(todos) + ["homolog_pc"]
    branco = obter_coleta_oficial_bytes()
    for nome in alvo:
        if nome == "homolog_pc":
            conteudo = gerar_coleta_oficial_preenchida(dados_calculadora_homolog())
            wb = load_workbook(io.BytesIO(conteudo))
            preencher_homolog(wb)
        else:
            wb = load_workbook(io.BytesIO(branco))
            todos[nome](wb)
        wb.save(out / f"{nome}.xlsx")
        print("montado", nome)


# ------------------------------------------------------------------ recalcular
def recalcular(pasta: Path, arquivos: list[Path] | None = None) -> None:
    import pythoncom
    import win32com.client as com

    pythoncom.CoInitialize()
    excel = com.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        for caminho in arquivos or sorted(pasta.glob("*.xlsx")):
            wb = excel.Workbooks.Open(str(caminho.resolve()))
            excel.CalculateFullRebuild()
            wb.Save()
            for _ in range(10):
                try:
                    wb.Close(SaveChanges=False)
                    break
                except Exception:
                    time.sleep(1.0)
            wb = None
            print("recalculado", caminho.name)
    finally:
        for _ in range(10):
            try:
                excel.Quit()
                break
            except Exception:
                time.sleep(1.0)
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()


# ------------------------------------------------------------------ fotografar
def _json(valor):
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return valor


def fotografar(repo: Path, pasta: Path) -> None:
    _preparar_repo(repo)
    from openpyxl import load_workbook
    import _baseline_fotografia as bf
    from _coleta_reajuste_documentos import processar_coleta_oficial_runtime

    for caminho in sorted(pasta.glob("*.xlsx")):
        conteudo = caminho.read_bytes()
        wb = load_workbook(io.BytesIO(conteudo), data_only=True)
        abas = {}
        for ws in wb.worksheets:
            abas[ws.title] = {
                c.coordinate: _json(c.value)
                for linha in ws.iter_rows() for c in linha if c.value is not None
            }
        nomes = {}
        for nome, dn in wb.defined_names.items():
            try:
                aba, ref = list(dn.destinations)[0]
                nomes[nome] = {"destino": f"{aba}!{ref}",
                               "valor": _json(wb[aba][ref.replace("$", "").split(":")[0]].value)}
            except Exception as erro:  # noqa: BLE001
                nomes[nome] = {"erro": str(erro)}
        resultado, diagnostico = processar_coleta_oficial_runtime(conteudo)
        foto = {
            "ordem_abas": wb.sheetnames,
            "abas": abas,
            "nomes": nomes,
            "web": bf.fotografar_web(resultado, diagnostico),
            "documentos": bf.fotografar_documentos(resultado),
        }
        destino = caminho.with_suffix(".foto.json")
        destino.write_text(json.dumps(foto, ensure_ascii=False, indent=1, default=str),
                           encoding="utf-8")
        print("fotografado", caminho.name)


# ------------------------------------------------------------------- comparar
# Diferencas DELIBERADAS (documentadas): titulos da camada tecnica, textos da
# aba executiva em MEMORIA_RESULTADOS!AF:AG, marcador de versao da Coleta e
# data/hora de geracao (=NOW()).
def _permitida(aba_depois: str, coord: str) -> bool:
    col = "".join(ch for ch in coord if ch.isalpha())
    if aba_depois == "RESULTADOS_DETALHE" and coord in ("A1", "A2"):
        return True
    if aba_depois == "MEMORIA_RESULTADOS" and col in ("AF", "AG"):
        return True
    if aba_depois == "CONTROLE" and coord == "B14":
        return True
    return False


def _versao(texto: str) -> str:
    return (texto.replace("RESULTADOS_DETALHE", "RESULTADOS")
            .replace("11.2", "11.x").replace("11.1", "11.x"))


def comparar(antes: Path, depois: Path) -> int:
    falhas = 0
    for foto_a in sorted(antes.glob("*.foto.json")):
        foto_b = depois / foto_a.name
        a = json.loads(foto_a.read_text(encoding="utf-8"))
        b = json.loads(foto_b.read_text(encoding="utf-8"))
        difs: list[str] = []
        for aba, celulas_a in a["abas"].items():
            aba_b = "RESULTADOS_DETALHE" if aba == "RESULTADOS" else aba
            celulas_b = b["abas"].get(aba_b)
            if celulas_b is None:
                difs.append(f"aba ausente no AFTER: {aba_b}")
                continue
            for coord in sorted(set(celulas_a) | set(celulas_b)):
                va, vb = celulas_a.get(coord), celulas_b.get(coord)
                if va != vb and not _permitida(aba_b, coord):
                    if aba == "CONTROLE" and isinstance(va, str) and _versao(va) == _versao(str(vb)):
                        continue
                    difs.append(f"{aba_b}!{coord}: {va!r} -> {vb!r}")
        for nome, info in a["nomes"].items():
            info_b = b["nomes"].get(nome)
            if info_b is None or info.get("valor") != info_b.get("valor"):
                difs.append(f"nome {nome}: {info} -> {info_b}")
            elif _versao(info.get("destino", "")) != _versao(info_b.get("destino", "")):
                difs.append(f"destino {nome}: {info['destino']} -> {info_b['destino']}")
        for camada in ("web", "documentos"):
            ta = _versao(json.dumps(a[camada], ensure_ascii=False, sort_keys=True))
            tb = _versao(json.dumps(b[camada], ensure_ascii=False, sort_keys=True))
            if ta != tb:
                ca, cb = json.loads(ta), json.loads(tb)
                for chave in sorted(set(ca) | set(cb)):
                    if ca.get(chave) != cb.get(chave):
                        difs.append(f"{camada}.{chave}: {str(ca.get(chave))[:300]} -> "
                                    f"{str(cb.get(chave))[:300]}")
        extras = set(b["abas"]) - {("RESULTADOS_DETALHE" if x == "RESULTADOS" else x)
                                   for x in a["abas"]}
        status = "OK" if not difs else f"{len(difs)} DIFERENCA(S)"
        print(f"{foto_a.stem:28s} {status}  (abas novas no AFTER: {sorted(extras)})")
        for linha in difs[:40]:
            print("    ", linha)
        falhas += bool(difs)
    return falhas


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("etapa", choices=("montar", "recalcular", "fotografar", "comparar"))
    p.add_argument("--repo", type=Path)
    p.add_argument("--out", type=Path)
    p.add_argument("--dir", type=Path)
    p.add_argument("--antes", type=Path)
    p.add_argument("--depois", type=Path)
    p.add_argument("--cenarios", nargs="*")
    args = p.parse_args()
    if args.etapa == "montar":
        montar(args.repo.resolve(), args.out, args.cenarios)
    elif args.etapa == "recalcular":
        recalcular(args.dir)
    elif args.etapa == "fotografar":
        fotografar(args.repo.resolve(), args.dir)
    else:
        return 1 if comparar(args.antes, args.depois) else 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
