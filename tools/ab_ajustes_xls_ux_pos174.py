# -*- coding: utf-8 -*-
"""Prova A/B dos AJUSTES-XLS-UX pos-PR #174 (BEFORE = base, AFTER = branch).

Reaproveita o harness da RESULTADOS-EXECUTIVO-V3 (`montar` dos cenarios de
preenchimento, `recalcular` no Excel real e `fotografar` valores + web +
documentos) e acrescenta:

* cenarios gerados pela Calculadora (`gerar_coleta_oficial_preenchida`) que
  exercitam as frentes desta etapa: analise so de C1/C2/C3/C4 (aviso de
  historico) e memoria IST com competencia de fronteira entre ciclos;
* fotografia ESTRUTURAL (formulas, estado das abas, validacoes e nomes);
* comparacao aba a aba com a MESMA aba (sem renomeacao) e allowlist explicita
  das unicas celulas deliberadamente alteradas por estas frentes.

Uso:
  python tools/ab_ajustes_xls_ux_pos174.py montar --repo R --out D
  python tools/ab_ajustes_xls_ux_pos174.py recalcular --dir D
  python tools/ab_ajustes_xls_ux_pos174.py fotografar --repo R --dir D
  python tools/ab_ajustes_xls_ux_pos174.py comparar --antes DA --depois DB
"""
from __future__ import annotations

import argparse
import io
import json
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import ab_resultados_executivo_v3 as v3  # noqa: E402


# -------------------------------------------------------------- cenarios novos
def _ist(competencia_base: date, base: float, competencia_final: date, final: float,
         variacao: float) -> list[dict]:
    return [
        {"tipo": "INDICE", "ordem": 1, "competencia": competencia_base.isoformat(),
         "valor_indice": base},
        {"tipo": "INDICE", "ordem": 2, "competencia": competencia_final.isoformat(),
         "valor_indice": final},
        {"tipo": "RESULTADO", "ordem": 3, "fator_acumulado": 1 + variacao,
         "variacao_final": variacao, "metodo_fonte": "IST (Anatel) [teste A/B]"},
    ]


def _ciclo(numero: int, percentual: float, *, memoria: list[dict] | None = None) -> dict:
    ano = 2021 + numero
    ciclo = {
        "ciclo": f"C{numero}",
        "data_inicio": date(ano, 1, 1),
        "percentual": percentual,
        "inicio_efeito_financeiro": date(ano, 1, 1),
        "situacao_aplicada": "✅ TEMPESTIVO",
        "data_pedido": date(ano, 1, 10),
    }
    if memoria is not None:
        ciclo["memoria_calculo"] = memoria
    return ciclo


PERCENTUAIS = {1: 0.0512, 2: 0.0374, 3: 0.0289, 4: 0.0315}


def _payload(numeros: list[int], *, ist_fronteira: bool = False) -> dict:
    ciclos = []
    for n in numeros:
        memoria = None
        if ist_fronteira:
            memoria = _ist(date(2020 + n, 10, 1), 100.0 + n, date(2021 + n, 10, 1),
                           101.0 + n, PERCENTUAIS[n])
        ciclos.append(_ciclo(n, PERCENTUAIS[n], memoria=memoria))
    ultimo = max(numeros)
    return {
        "indice": "IST (Anatel)",
        "data_base": date(2021, 1, 1),
        "data_corte": date(2021 + ultimo, 12, 31),
        "ciclos": ciclos,
    }


CENARIOS_PAYLOAD = {
    "calc_so_c1": _payload([1]),
    "calc_so_c2": _payload([2]),
    "calc_so_c3": _payload([3]),
    "calc_so_c4": _payload([4]),
    "calc_c1_c3": _payload([1, 2, 3]),
    "calc_ist_fronteira": _payload([1, 2], ist_fronteira=True),
    "calc_ist_ciclo_unico": _payload([2], ist_fronteira=True),
}


def montar(repo: Path, out: Path, cenarios: list[str] | None) -> None:
    v3._preparar_repo(repo)
    from openpyxl import load_workbook
    from _coleta_oficial import gerar_coleta_oficial_preenchida

    out.mkdir(parents=True, exist_ok=True)
    proprios = [c for c in (cenarios or CENARIOS_PAYLOAD) if c in CENARIOS_PAYLOAD]
    herdados = [c for c in (cenarios or []) if c not in CENARIOS_PAYLOAD]
    if herdados or not cenarios:
        v3.montar(repo, out, herdados or None)
    for nome in proprios:
        conteudo = gerar_coleta_oficial_preenchida(CENARIOS_PAYLOAD[nome])
        wb = load_workbook(io.BytesIO(conteudo))
        wb.save(out / f"{nome}.xlsx")
        print("montado", nome)


# -------------------------------------------------------- fotografia estrutural
def fotografar_estrutura(pasta: Path) -> None:
    from openpyxl import load_workbook

    for caminho in sorted(pasta.glob("*.xlsx")):
        wb = load_workbook(caminho)
        foto = {
            "estados": {ws.title: ws.sheet_state for ws in wb.worksheets},
            "formulas": {
                ws.title: {
                    c.coordinate: c.value
                    for linha in ws.iter_rows() for c in linha
                    if isinstance(c.value, str) and c.value.startswith("=")
                }
                for ws in wb.worksheets
            },
            "validacoes": {
                ws.title: sorted(
                    f"{dv.sqref}|{dv.type}|{dv.formula1}"
                    for dv in ws.data_validations.dataValidation
                )
                for ws in wb.worksheets
            },
            "nomes": {n: d.attr_text for n, d in wb.defined_names.items()},
        }
        caminho.with_suffix(".estrutura.json").write_text(
            json.dumps(foto, ensure_ascii=False, indent=1, default=str), encoding="utf-8"
        )
        print("estrutura", caminho.name)


# ------------------------------------------------------------------- comparar
# Unicas celulas deliberadamente alteradas por estas frentes (valor e/ou
# formula). Tudo o mais deve ser IDENTICO, inclusive todas as formulas.
ALLOWLIST = {
    "parametros": {"A8", "A17", "K82"},  # K82 = legenda da fronteira IST
    "MEMORIA_RESULTADOS": {"AF41", "AG41", "AF42", "AG42"},
    "itens_Consumidos": {"X7", "X8", "X9", "X10", "X11", "X12", "X13"},
    "RESULTADOS": {"B3", "D31"},
    "CONTROLE": {"B14"},  # =NOW()
    "cobertura_temporal": {"B4"},  # data/hora automatica de geracao
}
ESTADOS_ALTERADOS = {"RESULTADOS_DETALHE": ("visible", "hidden")}


def _permitida(aba: str, coord: str) -> bool:
    return coord in ALLOWLIST.get(aba, set())


def comparar(antes: Path, depois: Path) -> int:
    falhas = 0
    for foto_a in sorted(antes.glob("*.foto.json")):
        nome = foto_a.name.replace(".foto.json", "")
        foto_b = depois / foto_a.name
        if not foto_b.exists():
            print(f"{nome:28s} SEM FOTO NO AFTER")
            falhas += 1
            continue
        a = json.loads(foto_a.read_text(encoding="utf-8"))
        b = json.loads(foto_b.read_text(encoding="utf-8"))
        difs: list[str] = []
        permitidas: list[str] = []
        if a["ordem_abas"] != b["ordem_abas"]:
            difs.append(f"ordem de abas: {a['ordem_abas']} -> {b['ordem_abas']}")
        for aba in sorted(set(a["abas"]) | set(b["abas"])):
            ca, cb = a["abas"].get(aba, {}), b["abas"].get(aba, {})
            for coord in sorted(set(ca) | set(cb)):
                if ca.get(coord) != cb.get(coord):
                    linha = f"{aba}!{coord}: {ca.get(coord)!r} -> {cb.get(coord)!r}"
                    (permitidas if _permitida(aba, coord) else difs).append(linha)
        if a["nomes"] != b["nomes"]:
            difs.append("nomes definidos (destino/valor) divergentes")
        for camada in ("web", "documentos"):
            ta = json.dumps(a[camada], ensure_ascii=False, sort_keys=True)
            tb = json.dumps(b[camada], ensure_ascii=False, sort_keys=True)
            if ta != tb:
                difs.append(f"camada {camada} divergente")

        est_a = antes / f"{nome}.estrutura.json"
        est_b = depois / f"{nome}.estrutura.json"
        if est_a.exists() and est_b.exists():
            ea = json.loads(est_a.read_text(encoding="utf-8"))
            eb = json.loads(est_b.read_text(encoding="utf-8"))
            for aba in sorted(set(ea["estados"]) | set(eb["estados"])):
                sa, sb = ea["estados"].get(aba), eb["estados"].get(aba)
                if sa != sb:
                    if ESTADOS_ALTERADOS.get(aba) == (sa, sb):
                        permitidas.append(f"estado {aba}: {sa} -> {sb}")
                    else:
                        difs.append(f"estado {aba}: {sa} -> {sb}")
            for aba in sorted(set(ea["formulas"]) | set(eb["formulas"])):
                fa, fb = ea["formulas"].get(aba, {}), eb["formulas"].get(aba, {})
                for coord in sorted(set(fa) | set(fb)):
                    if fa.get(coord) != fb.get(coord):
                        linha = f"formula {aba}!{coord}: {fa.get(coord)!r} -> {fb.get(coord)!r}"
                        (permitidas if _permitida(aba, coord) else difs).append(linha)
            if ea["validacoes"] != eb["validacoes"]:
                difs.append("validacoes de dados divergentes")
            if ea["nomes"] != eb["nomes"]:
                difs.append("destinos de nomes definidos divergentes")
        else:
            difs.append("fotografia estrutural ausente")

        status = "OK" if not difs else f"{len(difs)} DIFERENCA(S) FORA DA ALLOWLIST"
        print(f"{nome:28s} {status}  ({len(permitidas)} alteracao(oes) autorizada(s))")
        for linha in difs[:40]:
            print("    FALHA ", linha[:300])
        for linha in permitidas[:12]:
            print("    ok    ", linha[:200])
        falhas += bool(difs)
    return falhas


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("etapa", choices=("montar", "recalcular", "fotografar", "comparar"))
    p.add_argument("--repo", type=Path)
    p.add_argument("--out", type=Path)
    p.add_argument("--dir", type=Path)
    p.add_argument("--antes", type=Path)
    p.add_argument("--depois", type=Path)
    p.add_argument("--cenarios", nargs="*")
    args = p.parse_args()
    if args.etapa == "montar":
        montar(args.repo.resolve(), args.out, args.cenarios)
    elif args.etapa == "recalcular":
        v3.recalcular(args.dir)
    elif args.etapa == "fotografar":
        v3.fotografar(args.repo.resolve(), args.dir)
        fotografar_estrutura(args.dir)
    else:
        return 1 if comparar(args.antes, args.depois) else 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
