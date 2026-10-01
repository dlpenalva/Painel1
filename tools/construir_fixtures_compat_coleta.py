"""Constroi as fixtures permanentes de compatibilidade retroativa da Coleta.

POR QUE ESTE CONSTRUTOR EXISTE
------------------------------
As tres linhagens que o Cl8us 11.0 precisa processar (COLETA_11, PRE_11_L1 e
PRE_11_L2) so podem ser comparadas com arquivos RECALCULADOS pelo Excel: sem o
cache de formulas o produto, corretamente, recusa-se a inventar VTA e
remanescente. O fluxo real do fiscal e: Coleta gerada -> Excel recalcula ->
upload. Este construtor reproduz exatamente esse fluxo.

O QUE ENTRA (E O QUE NAO ENTRA)
-------------------------------
* Estrutura: template de cada linhagem (11.0 = template atual + marcador; L1 =
  ultimo template 15 abas antes do versionamento; L2 = primeiro template de 11
  abas com NUMERO_PC). Nenhum dado contratual real.
* Entradas: cenarios SINTETICOS e deterministicos de
  ``tests/_baseline_cenarios`` (codigos ITEM-001.., valores redondos).
* O dado demonstrativo que o template L2 traz de fabrica e LIMPO antes do
  preenchimento.
* Metadados do pacote (autor, empresa) sao neutralizados.

USO (precisa de Excel desktop + pywin32; nao roda no CI)
--------------------------------------------------------
    <python com pywin32> tools/construir_fixtures_compat_coleta.py [--saida DIR] [nomes...]

As fixtures sao gravadas em ``tests/fixtures/coletas_compat`` junto do
``MANIFESTO.json`` (SHA-256, linhagem, cenario e percentual).
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import subprocess
import sys
import tempfile
import time
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import load_workbook  # noqa: E402

import _baseline_cenarios as cen  # noqa: E402

# Commits que congelam o ultimo template L1 e o primeiro L2 (templates do
# proprio repositorio; nenhum arquivo externo entra).
COMMIT_L1 = "8e12e72"
COMMIT_L2 = "c2ece81"
CAMINHO_TEMPLATE = "templates/COLETA_REAJUSTE_OFICIAL.xlsx"

PASTA_SAIDA = ROOT / "tests" / "fixtures" / "coletas_compat"

# Metodo por linhagem: o L2 usa os rotulos de dropdown da epoca.
METODOS = {
    "coleta_11": {"financeiro": "Financeiro (Mensalidade)",
                  "pc": "PC (Pedidos de Compra)",
                  "consumidos": "Itens Consumidos"},
    "pre11_l1": {"financeiro": "Financeiro (Mensalidade)",
                 "pc": "PC (Pedidos de Compra)",
                 "consumidos": "Itens Consumidos"},
    "pre11_l2": {"financeiro": "Principal", "pc": "PC", "consumidos": "D"},
}
CHAVE_METODO = {"01_financeiro_normal": "financeiro", "02_pc": "pc",
                "03_itens_consumidos": "consumidos"}

# Percentuais do cenario. "oficial" = duas casas; "bruto" = precisao das
# Coletas anteriores a regra vigente (mesmo percentual ANTES do fechamento).
PERCENTUAL_OFICIAL = 0.0512
PERCENTUAL_BRUTO = 0.05123816


def _template_git(commit: str) -> bytes:
    return subprocess.check_output(
        ["git", "show", f"{commit}:{CAMINHO_TEMPLATE}"], cwd=ROOT
    )


def _bytes_template(linhagem: str) -> bytes:
    if linhagem == "coleta_11":
        from _coleta_oficial import obter_coleta_oficial_bytes  # marcador 11.0

        return obter_coleta_oficial_bytes()
    if linhagem == "pre11_l1":
        return _template_git(COMMIT_L1)
    if linhagem == "pre11_l2":
        return _template_git(COMMIT_L2)
    raise KeyError(linhagem)


# --------------------------------------------------------------------------- #
# Preenchimento por linhagem.
# --------------------------------------------------------------------------- #
def _limpar_demonstrativo_l2(wb) -> None:
    """Remove os dados de exemplo que o template L2 traz de fabrica."""
    def limpar(aba: str, colunas: str, linhas: range) -> None:
        ws = wb[aba]
        for linha in linhas:
            for coluna in colunas:
                celula = ws[f"{coluna}{linha}"]
                if not (isinstance(celula.value, str) and celula.value.startswith("=")):
                    celula.value = None

    limpar("parametros", "ACDEG", range(3, 7))          # mantem C0 (linha 2)
    limpar("financeiro", "ACG", range(2, 74))
    limpar("itens_Remanesc", "ABCEGIKM", range(2, 211))
    limpar("itens_Consumidos", "ABC", range(2, 201))
    limpar("aditivos", "ABCDEFHK", range(2, 201))
    ws = wb["CONTROLE"]
    for coordenada in ("B2", "B3", "B7", "B8"):
        ws[coordenada].value = None


def _preencher_l2(wb, construtor: str) -> None:
    """Aplica o cenario no leiaute enxuto do L2 (sem H/U/CICLO_EM_EXECUCAO)."""
    _limpar_demonstrativo_l2(wb)
    ciclos = cen._ate(1)
    for ciclo in ciclos:
        ciclo.pop("inicio_efeito", None)
        ciclo.pop("data_pedido", None)
    cen._parametros(wb, ciclos)
    cen._controle(wb, metodo=METODOS["pre11_l2"][CHAVE_METODO[construtor]],
                  ciclo_vigente="C1", data_corte=date(2024, 12, 31),
                  indice="IST (Anatel)", data_base=date(2023, 1, 1))
    if construtor == "03_itens_consumidos":
        _itens_remanesc_coerente(wb, cen.ITENS_PADRAO)
    else:
        cen._itens_remanesc(wb, cen.ITENS_PADRAO)
    if construtor == "01_financeiro_normal":
        cen._financeiro(wb, cen._competencias(date(2023, 1, 1), 24, 42_500.00))
        _extras_completo(wb)
    elif construtor == "02_pc":
        cen._itens_pc(wb, [
            {"numero": "PC-2024-001", "data": date(2024, 2, 15),
             "valor": 180_000.00, "pago": "Sim"},
            {"numero": "PC-2024-002", "data": date(2024, 5, 20),
             "valor": 96_500.00, "pago": "Sim"},
            {"numero": "PC-2024-003", "data": date(2024, 9, 30),
             "valor": 145_250.00, "pago": "Nao"},
        ])
    else:
        _itens_consumidos_completo(wb)


# Consumo declarado por item e por ciclo (C0 e C1); total = 75/50/28 unidades.
CONSUMO_PADRAO = {
    "ITEM-001": (20.0, 55.0),
    "ITEM-002": (15.0, 35.0),
    "ITEM-003": (8.0, 20.0),
}


def _itens_consumidos_completo(wb, itens=None) -> None:
    """itens_Consumidos!A/B/C + QTD_CONS_C0 (E) e QTD_CONS_C1 (G).

    B e a quantidade contratada (mesma de itens_Remanesc); o consumo por ciclo
    e a entrada que o fiscal declara. Sem ele o metodo nao tem execucao.
    """
    ws = wb["itens_Consumidos"]
    for indice, item in enumerate(cen.ITENS_PADRAO):
        linha = indice + 2
        ws.cell(linha, 1).value = item["codigo"]
        ws.cell(linha, 2).value = item["quantidade"]
        ws.cell(linha, 3).value = item["valor_unitario"]
        c0, c1 = CONSUMO_PADRAO[item["codigo"]]
        ws.cell(linha, 5).value = c0
        ws.cell(linha, 7).value = c1


def _itens_remanesc_coerente(wb, itens) -> None:
    """itens_Remanesc coerente com o consumo: restante = contratado - consumido.

    Mantem A/B/C e grava o quantitativo restante (45/30/12) nas colunas de
    posicao por ciclo (E/G/I/K/M), como o fiscal faria ao declarar o consumo.
    """
    ws = wb["itens_Remanesc"]
    for indice, item in enumerate(cen.ITENS_PADRAO):
        linha = indice + 2
        ws.cell(linha, 1).value = item["codigo"]
        ws.cell(linha, 2).value = item["quantidade"]
        ws.cell(linha, 3).value = item["valor_unitario"]
        restante = cen.RESTANTE_PADRAO[item["codigo"]]
        for coluna in (5, 7, 9, 11, 13):
            ws.cell(linha, coluna).value = restante


PCS_PADRAO = [
    {"numero": "PC-2024-001", "data": date(2024, 2, 15), "valor": 180_000.00, "pago": "Sim"},
    {"numero": "PC-2024-002", "data": date(2024, 5, 20), "valor": 96_500.00, "pago": "Sim"},
    {"numero": "PC-2024-003", "data": date(2024, 9, 30), "valor": 145_250.00, "pago": "Nao"},
]


def _aditivo_simples(wb) -> None:
    """aditivos!A (item), B (data), D (tipo), E (quantidade), H e K (Sim)."""
    ws = wb["aditivos"]
    ws["A2"] = "ITEM-001"
    ws["B2"] = date(2024, 6, 15)
    ws["D2"] = "Acrescimo"
    ws["E2"] = 10.0
    ws["H2"] = "Sim"
    ws["K2"] = "Sim"


def _extras_completo(wb) -> None:
    if _SEM_EXTRAS:
        return
    """Cenario COMPLETO: alem do metodo, preenche PCs, consumo e um aditivo.

    Serve a um proposito so: o par (antigo x regerada) cobre TODOS os blocos de
    derivados no mesmo arquivo, e a prova de equivalencia vale celula a celula.
    """
    cen._itens_pc(wb, PCS_PADRAO)
    _itens_consumidos_completo(wb)
    _aditivo_simples(wb)


def _preencher_ciclo_em_execucao(wb, *, data: date, linhas) -> None:
    """Preenche a aba itemizada CICLO_EM_EXECUCAO (D5 + C13.. quantitativos).

    O cenario original cria uma aba minima; aqui a aba que o runtime ja gera
    (layout 2_ITEMIZADO, o mesmo dos arquivos reais dos fiscais) e preenchida
    no lugar, para nao duplicar a aba nem exceder o limite de abas.
    """
    ws = wb["CICLO_EM_EXECUCAO"]
    ws["D5"] = data
    for indice, item in enumerate(cen.ITENS_PADRAO):
        ws.cell(13 + indice, 3).value = cen.RESTANTE_PADRAO[item["codigo"]]


_SEM_EXTRAS = False


def _montar_xlsx_entrada(linhagem: str, construtor: str, percentual: float) -> bytes:
    """Monta a entrada; `*_item_unico` usa um unico item (ITEM-001).

    `*_sem_extras` monta o Financeiro SO com financeiro (sem PCs, consumo nem
    aditivo): o caso do L2 Financeiro sem a limitacao do aditivo.

    O QTD_REM_OFICIAL do XLS para Itens Consumidos usa `N(intervalo)` dentro de
    SUMPRODUCT e so enxerga o primeiro item: com varios itens o XLS diverge do
    Python mesmo no modelo 11.0 (defeito pre-existente do template). O cenario
    de item unico e o caso integro desse metodo.
    """
    global _SEM_EXTRAS
    if construtor.endswith("_sem_extras"):
        _SEM_EXTRAS = True
        try:
            return _montar_xlsx_entrada(
                linhagem, construtor[: -len("_sem_extras")], percentual
            )
        finally:
            _SEM_EXTRAS = False
    unico = construtor.endswith("_item_unico")
    if not unico:
        return _montar_xlsx_entrada_base(linhagem, construtor, percentual)
    # Quantidades pequenas (10 contratadas, 2+4 consumidas, 4 restantes): a
    # diferenca de ORDEM DE ARREDONDAMENTO entre o agregado do XLS (soma x fator)
    # e o Python (VU arredondado x quantidade) fica abaixo de 1 centavo.
    itens, restante, consumo = cen.ITENS_PADRAO, cen.RESTANTE_PADRAO, dict(CONSUMO_PADRAO)
    cen.ITENS_PADRAO = [dict(itens[0], quantidade=10.0, por_ciclo=(10.0,) * 5)]
    cen.RESTANTE_PADRAO = {"ITEM-001": 4.0}
    CONSUMO_PADRAO["ITEM-001"] = (2.0, 4.0)
    try:
        return _montar_xlsx_entrada_base(
            linhagem, construtor[: -len("_item_unico")], percentual
        )
    finally:
        cen.ITENS_PADRAO, cen.RESTANTE_PADRAO = itens, restante
        CONSUMO_PADRAO.clear()
        CONSUMO_PADRAO.update(consumo)


def _montar_xlsx_entrada_base(linhagem: str, construtor: str, percentual: float) -> bytes:
    wb = load_workbook(io.BytesIO(_bytes_template(linhagem)), data_only=False)
    if linhagem == "pre11_l2":
        _preencher_l2(wb, construtor)
    else:
        from _ciclo_em_execucao import garantir_aba_ciclo_em_execucao

        if "CICLO_EM_EXECUCAO" not in wb.sheetnames:
            # L1 historico: o runtime da epoca acrescentava a aba ao gerar.
            garantir_aba_ciclo_em_execucao(wb, limpar_entradas=True)
        originais = (cen._ciclo_em_execucao, cen._itens_consumidos,
                     cen._itens_remanesc)
        cen._ciclo_em_execucao = _preencher_ciclo_em_execucao
        cen._itens_consumidos = _itens_consumidos_completo
        if construtor == "03_itens_consumidos":
            cen._itens_remanesc = _itens_remanesc_coerente
        try:
            cen._CONSTRUTORES[construtor](wb)
        finally:
            (cen._ciclo_em_execucao, cen._itens_consumidos,
             cen._itens_remanesc) = originais
        if construtor == "01_financeiro_normal":
            _extras_completo(wb)
        if linhagem == "pre11_l1":
            # L1 e anterior ao marcador: garante a area livre.
            for coordenada in ("A24", "B24", "A25", "B25"):
                wb["CONTROLE"][coordenada].value = None
        wb["CONTROLE"]["B1"] = METODOS[linhagem][CHAVE_METODO[construtor]]
    # percentual do unico ciclo computado (C1 = parametros!E3)
    wb["parametros"]["E3"] = percentual
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# --------------------------------------------------------------------------- #
# Recalculo no Excel real (o mesmo que o fiscal faz antes do upload).
# --------------------------------------------------------------------------- #
def _recalcular_no_excel(origem: Path, destino: Path) -> None:
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(origem))
        try:
            excel.CalculateFull()
            for propriedade in ("Author", "Last Author", "Company", "Manager"):
                try:
                    wb.BuiltinDocumentProperties(propriedade).Value = ""
                except Exception:
                    pass
            if destino.exists():
                destino.unlink()
            wb.SaveAs(str(destino), FileFormat=51)
        finally:
            wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()


def _sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


# Matriz de fixtures. Nome -> (linhagem, cenario, percentual).
FIXTURES = {
    "coleta_11_financeiro": ("coleta_11", "01_financeiro_normal", PERCENTUAL_OFICIAL),
    "coleta_11_pc": ("coleta_11", "02_pc", PERCENTUAL_OFICIAL),
    "coleta_11_consumidos": ("coleta_11", "03_itens_consumidos", PERCENTUAL_OFICIAL),
    "pre11_l1_financeiro": ("pre11_l1", "01_financeiro_normal", PERCENTUAL_BRUTO),
    "pre11_l1_pc": ("pre11_l1", "02_pc", PERCENTUAL_BRUTO),
    "pre11_l1_consumidos": ("pre11_l1", "03_itens_consumidos", PERCENTUAL_BRUTO),
    "pre11_l2_financeiro": ("pre11_l2", "01_financeiro_normal", PERCENTUAL_BRUTO),
    "pre11_l2_pc": ("pre11_l2", "02_pc", PERCENTUAL_BRUTO),
    "pre11_l2_consumidos": ("pre11_l2", "03_itens_consumidos", PERCENTUAL_BRUTO),
    # Regeneracao oficial do mesmo caso (percentual fechado em 2 casas): base
    # da prova de equivalencia "arquivo antigo adaptado == Coleta regerada".
    # Etapa 3/03: Itens Consumidos INTEGRO (item unico) no 11.0 e no L1 antigo.
    "coleta_11_consumidos_item_unico": (
        "coleta_11", "03_itens_consumidos_item_unico", PERCENTUAL_OFICIAL),
    "pre11_l1_consumidos_item_unico": (
        "pre11_l1", "03_itens_consumidos_item_unico", PERCENTUAL_BRUTO),
    "pre11_l2_financeiro_sem_extras": (
        "pre11_l2", "01_financeiro_normal_sem_extras", PERCENTUAL_BRUTO),
    "pre11_l1_financeiro_regerada": ("pre11_l1", "01_financeiro_normal", PERCENTUAL_OFICIAL),
    "pre11_l2_financeiro_regerada": ("pre11_l2", "01_financeiro_normal", PERCENTUAL_OFICIAL),
}


def construir(nome: str, pasta: Path) -> dict:
    linhagem, construtor, percentual = FIXTURES[nome]
    entrada = _montar_xlsx_entrada(linhagem, construtor, percentual)
    with tempfile.TemporaryDirectory(prefix="cl8us_fixture_") as tmp:
        bruto = Path(tmp) / f"{nome}_entrada.xlsx"
        bruto.write_bytes(entrada)
        destino = pasta / f"{nome}.xlsx"
        for tentativa in range(1, 4):
            try:
                _recalcular_no_excel(bruto, destino)
                break
            except Exception as erro:  # RPC_E_CALL_REJECTED e afins: refazer do zero
                print(f"   tentativa {tentativa} falhou: {erro}", flush=True)
                if tentativa == 3:
                    raise
                time.sleep(5)
    return {"arquivo": destino.name, "linhagem_esperada": linhagem,
            "cenario": construtor, "percentual_c1": percentual,
            "sha256": _sha256(destino), "bytes": destino.stat().st_size}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--saida", type=Path, default=PASTA_SAIDA)
    parser.add_argument("nomes", nargs="*", help="subconjunto (padrao: todas)")
    args = parser.parse_args()
    args.saida.mkdir(parents=True, exist_ok=True)
    manifesto_arquivo = args.saida / "MANIFESTO.json"
    manifesto = (
        json.loads(manifesto_arquivo.read_text(encoding="utf-8"))
        if manifesto_arquivo.exists() else {}
    )
    for nome in args.nomes or FIXTURES:
        print(f"construindo {nome} ...", flush=True)
        manifesto[nome] = construir(nome, args.saida)
        print("  ", manifesto[nome]["sha256"][:16], manifesto[nome]["bytes"], "bytes")
    manifesto_arquivo.write_text(
        json.dumps(manifesto, indent=2, ensure_ascii=True, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
