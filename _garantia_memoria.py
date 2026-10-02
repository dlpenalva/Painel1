"""Exportação da memória de cálculo da Garantia Contratual.

Este módulo é deliberadamente de apresentação. Ele recebe ``situacao`` e
``analise`` já apuradas pelo motor canônico e apenas organiza esses resultados
em um XLSX. Não importa Streamlit nem chama qualquer função de cálculo da regra
de negócio.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins


NOME_ABA = "Memória da Garantia"
TRACO = "—"
FUSO_BRASILIA = ZoneInfo("America/Sao_Paulo")

AZUL_ESCURO = "173B5D"
AZUL_TEXTO = "24445F"
AZUL_CLARO = "EAF2F8"
AZUL_BORDA = "C8D9E8"
VERDE_CLARO = "E9F4EE"
VERDE_BORDA = "B9DCC6"
VERDE_TEXTO = "14532D"
AMBAR_CLARO = "FDF6E8"
AMBAR_BORDA = "EBDCB4"
AMBAR_TEXTO = "7A5A12"
CINZA_CLARO = "F6F8FA"
CINZA_BORDA = "DDE4EA"
CINZA_TEXTO = "475569"
BRANCO = "FFFFFF"

FONTE = "Aptos"
FORMATO_MOEDA = 'R$ #,##0.00;[Red]-R$ #,##0.00;R$ 0.00'
FORMATO_PERCENTUAL = "0.00%"
FORMATO_DATA = "dd/mm/yyyy"


def _moeda_texto(valor) -> str:
    """Formata um resultado monetário já apurado, sem fazer cálculo."""
    if valor is None:
        return TRACO
    dec = Decimal(str(valor))
    texto = f"{dec:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {texto}"


def _percentual_texto(valor) -> str:
    if valor is None:
        return TRACO
    dec = Decimal(str(valor))
    return f"{dec:.2f}".replace(".", ",") + "%"


def _data_texto(valor) -> str:
    return valor.strftime("%d/%m/%Y") if isinstance(valor, (date, datetime)) else TRACO


def _numero(valor):
    """Conserva números como números no XLSX; ausência continua ausência."""
    if valor is None:
        return TRACO
    return Decimal(str(valor))


def _horario_de_brasilia(valor: datetime | None) -> datetime:
    """Normaliza um instante consciente de fuso para o horário de Brasília."""
    if valor is None:
        return datetime.now(FUSO_BRASILIA)
    if valor.tzinfo is None or valor.utcoffset() is None:
        raise ValueError("gerado_em deve informar um fuso horário explícito")
    return valor.astimezone(FUSO_BRASILIA)


def _borda(cor=CINZA_BORDA):
    lado = Side(style="thin", color=cor)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _estilizar_intervalo(ws, intervalo, *, fill=None, font=None, border=None,
                         alignment=None):
    for linha in ws[intervalo]:
        for celula in linha:
            if fill is not None:
                celula.fill = fill
            if font is not None:
                celula.font = font
            if border is not None:
                celula.border = border
            if alignment is not None:
                celula.alignment = alignment


def _secao(ws, linha: int, titulo: str) -> int:
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=9)
    celula = ws.cell(linha, 1, titulo)
    celula.fill = PatternFill("solid", fgColor=AZUL_CLARO)
    celula.font = Font(name=FONTE, size=11, bold=True, color=AZUL_ESCURO)
    celula.alignment = Alignment(vertical="center")
    celula.border = _borda(AZUL_BORDA)
    ws.row_dimensions[linha].height = 24
    return linha + 1


def _linha_rotulo_valor(ws, linha: int, rotulo: str, valor, *, formato=None,
                        fill=CINZA_CLARO, cor=CINZA_TEXTO, negrito=False) -> int:
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=4)
    ws.merge_cells(start_row=linha, start_column=5, end_row=linha, end_column=9)
    c_rotulo = ws.cell(linha, 1, rotulo)
    c_valor = ws.cell(linha, 5, valor)
    preenchimento = PatternFill("solid", fgColor=fill)
    cor_borda = {
        CINZA_CLARO: CINZA_BORDA,
        AZUL_CLARO: AZUL_BORDA,
        VERDE_CLARO: VERDE_BORDA,
        AMBAR_CLARO: AMBAR_BORDA,
    }.get(fill, CINZA_BORDA)
    borda = _borda(cor_borda)
    for celula in (c_rotulo, c_valor):
        celula.fill = preenchimento
        celula.border = borda
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    c_rotulo.font = Font(name=FONTE, size=10, color=cor, bold=True)
    c_valor.font = Font(name=FONTE, size=11, color=cor, bold=negrito)
    c_valor.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    if formato and valor != TRACO:
        c_valor.number_format = formato
    ws.row_dimensions[linha].height = 25
    return linha + 1


def _linha_explicacao(ws, linha: int, titulo: str, texto: str) -> int:
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
    ws.merge_cells(start_row=linha, start_column=3, end_row=linha, end_column=9)
    c_titulo = ws.cell(linha, 1, titulo)
    c_texto = ws.cell(linha, 3, texto)
    for celula in (c_titulo, c_texto):
        celula.fill = PatternFill("solid", fgColor=BRANCO)
        celula.border = _borda(CINZA_BORDA)
        celula.alignment = Alignment(vertical="top", wrap_text=True)
    c_titulo.font = Font(name=FONTE, size=10, bold=True, color=AZUL_ESCURO)
    c_texto.font = Font(name=FONTE, size=10, color=CINZA_TEXTO)
    ws.row_dimensions[linha].height = 36
    return linha + 1


def gerar_memoria_garantia_xlsx(
    situacao: dict,
    analise: dict,
    *,
    versao_cl8us: str,
    gerado_em: datetime | None = None,
    dias_validade_minima: int,
) -> bytes:
    """Gera a memória em XLSX exclusivamente a partir de resultados canônicos.

    ``situacao`` deve ser a saída de ``calcular_situacao_atual`` e ``analise``
    deve ser a saída de ``analisar_garantia``. O exportador não completa,
    corrige ou recalcula esses dados.
    """
    gerado_em = _horario_de_brasilia(gerado_em)

    wb = Workbook()
    ws = wb.active
    ws.title = NOME_ABA
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.28, right=0.28, top=0.45, bottom=0.45,
                                  header=0.2, footer=0.2)
    ws.oddFooter.left.text = f"Cl8us {versao_cl8us}"
    ws.oddFooter.right.text = "Página &P de &N"

    wb.properties.title = "Memória de cálculo da garantia contratual"
    wb.properties.subject = "Evolução do contrato e da garantia"
    wb.properties.creator = f"Cl8us {versao_cl8us}"
    wb.properties.description = (
        "Documento gerado com os resultados canônicos da Calculadora de Garantia Contratual."
    )

    larguras = {"A": 6, "B": 19, "C": 13, "D": 18, "E": 16,
                "F": 18, "G": 18, "H": 19, "I": 16}
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura

    ws.row_dimensions[1].height = 8
    ws.merge_cells("A2:I2")
    ws["A2"] = "MEMÓRIA DE CÁLCULO DA GARANTIA CONTRATUAL"
    ws["A2"].font = Font(name=FONTE, size=16, bold=True, color=AZUL_ESCURO)
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 28
    ws.merge_cells("A3:I3")
    ws["A3"] = "Resumo da evolução do contrato e da garantia"
    ws["A3"].font = Font(name=FONTE, size=11, italic=True, color=CINZA_TEXTO)
    ws["A3"].alignment = Alignment(vertical="center")
    ws.row_dimensions[3].height = 20
    ws.merge_cells("A4:I4")
    ws["A4"].fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    ws.row_dimensions[4].height = 3

    linha = 5
    linha = _linha_rotulo_valor(ws, linha, "Documento gerado por", f"Cl8us {versao_cl8us}")
    linha = _linha_rotulo_valor(
        ws, linha, "Data e hora de geração", gerado_em.strftime("%d/%m/%Y %H:%M")
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Percentual da garantia", Decimal(str(situacao["percentual"])) / Decimal("100"),
        formato=FORMATO_PERCENTUAL,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Regra de validade mínima",
        f"{dias_validade_minima} dias corridos após o término da vigência",
    )

    linha += 1
    linha = _secao(ws, linha, "SITUAÇÃO ORIGINAL")
    linha = _linha_rotulo_valor(
        ws, linha, "Valor original do contrato", _numero(situacao["valor_original"]),
        formato=FORMATO_MOEDA,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Percentual da garantia", Decimal(str(situacao["percentual"])) / Decimal("100"),
        formato=FORMATO_PERCENTUAL,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Garantia original exigida", _numero(situacao["garantia_original"]),
        formato=FORMATO_MOEDA,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Término da vigência original", situacao["vigencia_original"], formato=FORMATO_DATA,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Validade mínima original", situacao["validade_minima_original"],
        formato=FORMATO_DATA,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Garantia apresentada na assinatura",
        _numero(situacao["garantia_apresentada_original"]), formato=FORMATO_MOEDA,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Validade da garantia apresentada",
        situacao["validade_apresentada_original"] or TRACO, formato=FORMATO_DATA,
    )

    linha += 1
    linha = _secao(ws, linha, "EVOLUÇÃO DO CONTRATO E DA GARANTIA")
    cabecalhos = [
        "Nº", "Evento", "Data", "Valor do contrato", "Variação",
        "Garantia exigida", "Término da vigência", "Garantia apresentada", "Validade",
    ]
    for coluna, titulo in enumerate(cabecalhos, start=1):
        celula = ws.cell(linha, coluna, titulo)
        celula.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        celula.font = Font(name=FONTE, size=9, bold=True, color=BRANCO)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = _borda(BRANCO)
    ws.row_dimensions[linha].height = 34
    linha += 1

    linhas_evolucao = [
        [
            0,
            "Assinatura",
            TRACO,
            _numero(situacao["valor_original"]),
            TRACO,
            _numero(situacao["garantia_original"]),
            situacao["vigencia_original"],
            _numero(situacao["garantia_apresentada_original"]),
            situacao["validade_apresentada_original"] or TRACO,
        ]
    ]
    for etapa in situacao["linha_do_tempo"]:
        linhas_evolucao.append(
            [
                etapa["numero"],
                etapa["tipo"],
                etapa["data"] or TRACO,
                _numero(etapa["valor"]),
                _numero(etapa["variacao"]),
                _numero(etapa["garantia_exigida"]),
                etapa["vigencia"],
                _numero(etapa["garantia_apresentada"]),
                etapa["validade_apresentada"] or TRACO,
            ]
        )

    primeira_linha_evolucao = linha
    for valores in linhas_evolucao:
        for coluna, valor in enumerate(valores, start=1):
            celula = ws.cell(linha, coluna, valor)
            celula.font = Font(name=FONTE, size=9, color=CINZA_TEXTO)
            celula.alignment = Alignment(
                horizontal="left" if coluna == 2 else "right",
                vertical="center",
                wrap_text=coluna == 2,
            )
            celula.border = Border(bottom=Side(style="thin", color=CINZA_BORDA))
            if coluna in (4, 5, 6, 8) and valor != TRACO:
                celula.number_format = FORMATO_MOEDA
            elif coluna in (3, 7, 9) and valor != TRACO:
                celula.number_format = FORMATO_DATA
        if linha == primeira_linha_evolucao:
            _estilizar_intervalo(
                ws, f"A{linha}:I{linha}", fill=PatternFill("solid", fgColor=CINZA_CLARO)
            )
        ws.row_dimensions[linha].height = 25
        linha += 1

    linha += 1
    linha = _secao(ws, linha, "SITUAÇÃO ATUAL")
    linha = _linha_rotulo_valor(
        ws, linha, "Valor atualizado total do contrato", _numero(situacao["valor_atual"]),
        formato=FORMATO_MOEDA, fill=AZUL_CLARO, cor=AZUL_TEXTO,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Garantia atualmente exigida", _numero(analise["garantia_necessaria"]),
        formato=FORMATO_MOEDA, fill=VERDE_CLARO, cor=VERDE_TEXTO, negrito=True,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Última garantia apresentada", _numero(analise["garantia_apresentada"]),
        formato=FORMATO_MOEDA, fill=AZUL_CLARO, cor=AZUL_TEXTO,
    )
    fill_complemento = AMBAR_CLARO if analise["complemento"] > 0 else VERDE_CLARO
    cor_complemento = AMBAR_TEXTO if analise["complemento"] > 0 else VERDE_TEXTO
    linha = _linha_rotulo_valor(
        ws, linha, "Complemento necessário", _numero(analise["complemento"]),
        formato=FORMATO_MOEDA, fill=fill_complemento, cor=cor_complemento, negrito=True,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Validade mínima exigida", analise["validade_minima"],
        formato=FORMATO_DATA, fill=VERDE_CLARO, cor=VERDE_TEXTO, negrito=True,
    )
    linha = _linha_rotulo_valor(
        ws, linha, "Validade da garantia apresentada",
        analise["validade_apresentada"] or TRACO, formato=FORMATO_DATA,
        fill=AMBAR_CLARO if not analise["validade_suficiente"] else AZUL_CLARO,
        cor=AMBAR_TEXTO if not analise["validade_suficiente"] else AZUL_TEXTO,
    )

    linha += 1
    linha = _secao(ws, linha, "MEMÓRIA DO CÁLCULO")
    expressao_garantia = (
        f"{_moeda_texto(analise['valor_total_contrato'])} × "
        f"{_percentual_texto(analise['percentual'])} = "
        f"{_moeda_texto(analise['garantia_necessaria'])}"
    )
    linha = _linha_explicacao(ws, linha, "Garantia exigida", expressao_garantia)

    if not analise["tem_garantia"]:
        expressao_complemento = (
            "Não foi informada garantia apresentada. O complemento necessário corresponde "
            f"à garantia exigida: {_moeda_texto(analise['complemento'])}."
        )
    elif analise["complemento"] > 0:
        expressao_complemento = (
            f"{_moeda_texto(analise['garantia_necessaria'])} − "
            f"{_moeda_texto(analise['garantia_apresentada'])} = "
            f"{_moeda_texto(analise['complemento'])}"
        )
    else:
        expressao_complemento = (
            f"Garantia apresentada: {_moeda_texto(analise['garantia_apresentada'])}. "
            f"Complemento apurado pelo motor: {_moeda_texto(analise['complemento'])}."
        )
    linha = _linha_explicacao(ws, linha, "Complemento necessário", expressao_complemento)
    expressao_validade = (
        f"{_data_texto(analise['data_fim_vigencia'])} + {dias_validade_minima} dias corridos = "
        f"{_data_texto(analise['validade_minima'])}"
    )
    linha = _linha_explicacao(ws, linha, "Validade mínima", expressao_validade)

    linha += 1
    linha = _secao(ws, linha, "CONCLUSÃO")
    conclusao_regular = analise["diagnostico"] == "GARANTIA REGULAR"
    linha = _linha_rotulo_valor(
        ws, linha, "Diagnóstico", analise["diagnostico"],
        fill=VERDE_CLARO if conclusao_regular else AMBAR_CLARO,
        cor=VERDE_TEXTO if conclusao_regular else AMBAR_TEXTO,
        negrito=True,
    )

    ultima_linha = linha - 1
    ws.print_area = f"A1:I{ultima_linha}"

    for linha_celulas in ws.iter_rows(min_row=1, max_row=ultima_linha, min_col=1, max_col=9):
        for celula in linha_celulas:
            if celula.font.name is None:
                celula.font = Font(name=FONTE, size=10, color=CINZA_TEXTO)
            if celula.alignment.vertical is None:
                celula.alignment = Alignment(vertical="center")

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
